# -*- coding: utf-8 -*-
"""DRONE-AREA-CONTROL-V2-MEGA, блок C -- «дата+время» на экранах вылетов.

Семейство вылетов модуля Дроны -- список `/drones/`, сводка и её книга,
`flights.xlsx`, «Источники», расход раствора и его книга -- читает период
одним разбором (`drone_period.parse` через `_drone_filters_from_args`) и
фильтрует `drone_flights.started_at` одним условием
(`_drone_flight_conditions`): начало включительно, конец ПОЛУОТКРЫТЫЙ,
границы -- 19-символьные строки UTC.

Фикстура -- вылеты у самых границ, записанные В ДВУХ НАПИСАНИЯХ:
26 символов ('YYYY-MM-DD HH:MM:SS.ffffff', так пишет ORM) и 19 символов
('YYYY-MM-DD HH:MM:SS', так пишут stdlib-писатели; здесь -- сырым UPDATE
через sqlite3). Первый тест проверяет, что оба написания действительно
лежат в базе: иначе проверка строковой границы была бы проверкой ни о чём.

Что держится:

  * HTML и Excel одного периода дают одни и те же числа (openpyxl, числа
    сверяются, а не «файл открылся»);
  * минутные границы, UTC+5, интервал через полночь;
  * начало позже конца -- пусто, 200 и видимое сообщение RU/UZ;
  * неверное время -- значение по умолчанию и предупреждение;
  * поля времени с 00:00/23:59 по умолчанию, поля дат байт в байт прежние;
  * ссылки (Excel, детализация, «За всё время», страницы) несут время,
    когда оно не по умолчанию, и не несут, когда по умолчанию;
  * имя книги и ячейки «Период» -- с минутами только при не целых днях.

Отрицательные контроли: прежняя реализация (граница-datetime, закрытый
конец) на ТЕХ ЖЕ данных даёт другой набор вылетов; и полуоткрытая граница,
но переданная как datetime, даёт тот же СЧЁТ при другом НАБОРЕ -- поэтому
здесь сверяются наборы id, а не только количества.

Run:
  python -m unittest tests.test_drone_period_filters -v
"""
import html as htmlmod
import io
import re
import sqlite3
import unittest
from datetime import date, datetime, time, timedelta
from urllib.parse import parse_qs, urlsplit

from tests.harness import app, reset_db, create_admin, login, TEST_DB_PATH
from models import db, DroneFlight, DroneUnit, Organization, User

import drone_period
import drones

NBSP = ' '
UTC_OFFSET = timedelta(hours=5)
DJI_BASE = 770000

# key, local moment (UTC+5), stored spelling length, area, machine number.
# Areas are distinct, so a sum names the set and not only its size.
FIXTURE = (
    ('A', (2026, 6, 15, 0, 0, 0), 19, 1.0, 1),     # first second of the day
    ('B', (2026, 6, 14, 23, 59, 59), 19, 2.0, 1),  # last second before it
    ('C', (2026, 6, 15, 23, 59, 41), 26, 3.0, 2),  # inside the last minute
    ('D', (2026, 6, 16, 0, 0, 0), 26, 4.0, 2),     # next day, ORM spelling
    ('E', (2026, 6, 15, 12, 0, 0), 26, 5.0, 1),
    ('F', (2026, 6, 15, 10, 30, 41), 26, 6.0, 1),  # inside minute 10:30
    ('G', (2026, 6, 15, 10, 31, 0), 19, 7.0, 2),   # first second after it
    ('H', (2026, 6, 15, 8, 0, 0), 19, 8.0, 2),     # start of 08:00
    ('I', (2026, 6, 15, 7, 59, 59), 26, 9.0, 1),   # last second before it
    ('J', (2026, 6, 15, 20, 30, 0), 26, 10.0, 2),
    ('K', (2026, 6, 16, 5, 45, 0), 19, 11.0, 1),
    ('L', (2026, 6, 16, 6, 1, 0), 19, 12.0, 2),
    ('M', (2026, 6, 15, 19, 59, 0), 26, 13.0, 1),
    ('N', (2026, 6, 16, 0, 0, 0), 19, 14.0, 1),    # next day, stdlib spelling
)
AREA = {key: area for key, _m, _l, area, _u in FIXTURE}

# Query strings and the sets they must select -- written out BY HAND from the
# table above, not computed by the code under test.
WHOLE_DAY = 'date_from=2026-06-15&date_to=2026-06-15'
MINUTES = ('date_from=2026-06-15&date_to=2026-06-15'
           '&time_from=08:00&time_to=10:30')
CROSS_DAY = ('date_from=2026-06-15&date_to=2026-06-16'
             '&time_from=20:00&time_to=06:00')
TWO_DAYS = 'date_from=2026-06-15&date_to=2026-06-16'
INVERTED = 'date_from=2026-06-16&date_to=2026-06-15'
INVERTED_TIME = ('date_from=2026-06-15&date_to=2026-06-15'
                 '&time_from=11:00&time_to=09:00')
BAD_TIME = 'date_from=2026-06-15&date_to=2026-06-15&time_from=25:61'

EXPECTED = {
    WHOLE_DAY: set('ACEFGHIJM'),
    MINUTES: set('HF'),
    CROSS_DAY: set('JCDNK'),
    TWO_DAYS: set('ACDEFGHIJKLMN'),
    INVERTED: set(),
    INVERTED_TIME: set(),
    BAD_TIME: set('ACEFGHIJM'),
}

MSG_INVERTED_RU = 'Начало периода позже его конца'
MSG_INVERTED_UZ = 'Давр бошланиши тугашидан кейин'
MSG_BAD_FROM_RU = 'Время «с» указано неверно'


def set_language(user_id, lang):
    """The page language comes from users.language through g.lang."""
    with app.app_context():
        user = db.session.get(User, user_id)
        user.language = lang
        db.session.commit()


def seed_units(numbers):
    org = Organization(name='Period Fixture')
    db.session.add(org)
    db.session.flush()
    ids = {}
    for number in numbers:
        unit = DroneUnit(number=number, organization_id=org.id)
        db.session.add(unit)
        db.session.flush()
        ids[number] = unit.id
    return ids


def seed_fixture():
    """The boundary flights, then half of them rewritten to 19 characters."""
    with app.app_context():
        units = seed_units([1, 2])
        for index, (key, moment, _length, area, number) in enumerate(FIXTURE):
            db.session.add(DroneFlight(
                dji_flight_id=DJI_BASE + index,
                drone_unit_id=units[number],
                nickname_raw='period-%s' % key,
                started_at=datetime(*moment) - UTC_OFFSET,
                work_seconds=600,
                area_ha=area,
                spray_liters=area * 10.0,
                usage_type=0,
                raw_json='{}',
            ))
        db.session.commit()
        db.session.remove()
    # [REASON]: the stdlib spelling is written with stdlib sqlite3, exactly
    # as dji_area/store.py and the other stdlib writers do -- a raw UPDATE
    # with a 19-character string, bypassing the ORM's DateTime processor.
    con = sqlite3.connect(TEST_DB_PATH, timeout=10)
    try:
        for index, (key, moment, length, _area, _n) in enumerate(FIXTURE):
            if length == 19:
                utc = datetime(*moment) - UTC_OFFSET
                con.execute('UPDATE drone_flights SET started_at = ? '
                            'WHERE dji_flight_id = ?',
                            (utc.strftime('%Y-%m-%d %H:%M:%S'),
                             DJI_BASE + index))
        con.commit()
    finally:
        con.close()


def key_of(dji_id):
    return FIXTURE[int(dji_id) - DJI_BASE][0]


def area_of(keys):
    return round(sum(AREA[k] for k in keys), 2)


def stat(page, label):
    """The value of a .vs-stat card, by its exact label."""
    match = re.search(
        r'<div class="vs-stat-label">%s</div>\s*'
        r'<div class="vs-stat-value[^"]*">([^<]*)</div>' % re.escape(label),
        page)
    if not match:
        raise AssertionError('stat card %r not found' % label)
    return match.group(1).replace(NBSP, '').strip()


def hrefs(page, path):
    """Every href on the page whose path is `path`, as parsed query dicts."""
    out = []
    for raw in re.findall(r'href="([^"]*)"', page):
        url = urlsplit(htmlmod.unescape(raw))
        if url.path == path:
            out.append({k: v[0] for k, v in
                        parse_qs(url.query, keep_blank_values=True).items()})
    return out


def sheet_pairs(ws):
    """{label: value} of a two-column «Показатель / Значение» sheet."""
    return {row[0]: row[1] for row in ws.iter_rows(min_row=2,
                                                    values_only=True)}


def load_book(response):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(response.data))


def filename(response):
    return response.headers.get('Content-Disposition', '')


class Base(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        reset_db()
        cls.user_id = create_admin('period_admin')
        seed_fixture()

    def setUp(self):
        set_language(self.user_id, 'ru')
        self.client = app.test_client()
        login(self.client, self.user_id)

    def get(self, path, query=''):
        response = self.client.get(path + ('?' + query if query else ''))
        self.assertEqual(response.status_code, 200, path + '?' + query)
        return response

    def page(self, path, query=''):
        return self.get(path, query).get_data(as_text=True)


# ─── 0: the fixture really holds both spellings ──────────────────────────────

class FixtureShape(Base):

    def test_both_spellings_are_in_the_database(self):
        """Без обоих написаний проверка строковой границы ничего не доказала
        бы: на одних 26-символьных строках datetime-граница тоже верна."""
        con = sqlite3.connect(TEST_DB_PATH)
        try:
            rows = dict(con.execute(
                'SELECT dji_flight_id, length(started_at) '
                'FROM drone_flights').fetchall())
        finally:
            con.close()
        for index, (key, _m, length, _a, _n) in enumerate(FIXTURE):
            self.assertEqual(rows[DJI_BASE + index], length, key)
        self.assertEqual(set(rows.values()), {19, 26})


# ─── 1: the conditions themselves ────────────────────────────────────────────

class Conditions(Base):

    def select(self, conds):
        return {key_of(f.dji_flight_id)
                for f in DroneFlight.query.filter(*conds).all()}

    def parsed(self, query):
        with app.test_request_context('/drones/?' + query):
            from flask import request
            return drones._drone_filters_from_args(request.args,
                                                   default_current_month=False)

    def test_every_expected_set(self):
        for query, expected in EXPECTED.items():
            filters = self.parsed(query)
            with app.app_context():
                got = self.select(drones._drone_flight_conditions(filters))
            self.assertEqual(got, expected, query)

    def test_plain_dict_still_works(self):
        """Тесты и отчёты зовут условие обычным словарём без времени."""
        with app.app_context():
            got = self.select(drones._drone_flight_conditions(
                {'date_from': date(2026, 6, 15), 'date_to': date(2026, 6, 15),
                 'unit_id': None, 'region': ''}))
        self.assertEqual(got, EXPECTED[WHOLE_DAY])

    def test_bounds_are_bound_as_19_character_strings(self):
        filters = self.parsed(MINUTES)
        with app.app_context():
            conds = drones._drone_flight_conditions(filters)
            values = [c.right.value for c in conds]
        self.assertEqual(values, ['2026-06-15 03:00:00',
                                  '2026-06-15 05:31:00'])

    def _old_conds(self, date_from, date_to, time_to=None):
        """The pre-V2 implementation: datetime bounds, CLOSED end."""
        start = datetime.combine(date_from, time(0, 0)) - UTC_OFFSET
        end = (datetime.combine(date_to, time_to or datetime.max.time())
               - UTC_OFFSET)
        return [DroneFlight.started_at >= start, DroneFlight.started_at <= end]

    def test_negative_control_old_bounds_select_a_different_set(self):
        """Прежняя реализация на тех же данных теряет 00:00:00 первого дня
        (A: 19 символов против '...19:00:00.000000')."""
        with app.app_context():
            old = self.select(self._old_conds(date(2026, 6, 15),
                                              date(2026, 6, 15)))
        self.assertEqual(old, EXPECTED[WHOLE_DAY] - {'A'})
        self.assertNotEqual(old, EXPECTED[WHOLE_DAY])

    def test_negative_control_closed_minute_loses_both_edges(self):
        """«С 08:00 по 10:30» закрытой datetime-границей: H (08:00:00,
        19 символов) и F (10:30:41) пропадают оба."""
        with app.app_context():
            start = datetime(2026, 6, 15, 8, 0) - UTC_OFFSET
            end = datetime(2026, 6, 15, 10, 30) - UTC_OFFSET
            old = self.select([DroneFlight.started_at >= start,
                               DroneFlight.started_at <= end])
        self.assertEqual(old, set())
        self.assertEqual(EXPECTED[MINUTES], {'H', 'F'})

    def test_negative_control_datetime_bind_keeps_the_count_not_the_set(self):
        """Полуоткрытая граница, но переданная как datetime: A выпадает,
        N (00:00:00 следующего дня, 19 символов) попадает. Количество
        совпадает -- поэтому сверяется набор, а не только счёт."""
        start, end_excl = drone_period.utc_bounds(self.parsed(WHOLE_DAY))
        with app.app_context():
            wrong = self.select([DroneFlight.started_at >= start,
                                 DroneFlight.started_at < end_excl])
        self.assertEqual(len(wrong), len(EXPECTED[WHOLE_DAY]))
        self.assertEqual(wrong, (EXPECTED[WHOLE_DAY] - {'A'}) | {'N'})
        self.assertNotEqual(wrong, EXPECTED[WHOLE_DAY])


# ─── 2: list page and flights.xlsx ───────────────────────────────────────────

class ListAndFlightsBook(Base):

    def total(self, page):
        match = re.search(r'Всего: (\d+)', page)
        self.assertIsNotNone(match)
        return int(match.group(1))

    def book_keys(self, query):
        response = self.get('/drones/flights.xlsx', query)
        ws = load_book(response).worksheets[0]
        header = [c.value for c in ws[1]]
        dji_col = header.index('DJI id')
        area_col = header.index('Гектары')
        keys, area = set(), 0.0
        for row in ws.iter_rows(min_row=2, values_only=True):
            keys.add(key_of(row[dji_col]))
            area += row[area_col]
        return keys, round(area, 2), response

    def test_list_and_book_agree_with_the_expected_sets(self):
        for query, expected in EXPECTED.items():
            page = self.page('/drones/', query)
            keys, area, _resp = self.book_keys(query)
            self.assertEqual(self.total(page), len(expected), query)
            self.assertEqual(keys, expected, query)
            self.assertAlmostEqual(area, area_of(expected), places=2)

    def test_default_list_is_all_time(self):
        page = self.page('/drones/')
        self.assertEqual(self.total(page), len(FIXTURE))

    def test_time_inputs_default_and_date_inputs_unchanged(self):
        page = self.page('/drones/')
        self.assertIn('<input type="date" name="date_from" id="droneDateFrom" '
                      'class="vs-input" value="">', page)
        self.assertIn('<input type="date" name="date_to" id="droneDateTo" '
                      'class="vs-input" value="">', page)
        self.assertIn('<input type="time" step="60" name="time_from" '
                      'id="droneTimeFrom" class="vs-input" value="00:00">', page)
        self.assertIn('<input type="time" step="60" name="time_to" '
                      'id="droneTimeTo" class="vs-input" value="23:59">', page)
        # Время стоит сразу за своей датой: с -- за «с», по -- за «по».
        self.assertLess(page.index('id="droneDateFrom"'),
                        page.index('id="droneTimeFrom"'))
        self.assertLess(page.index('id="droneTimeFrom"'),
                        page.index('id="droneDateTo"'))
        self.assertLess(page.index('id="droneDateTo"'),
                        page.index('id="droneTimeTo"'))

    def test_chosen_time_is_echoed_into_the_inputs(self):
        page = self.page('/drones/', CROSS_DAY)
        self.assertIn('id="droneTimeFrom" class="vs-input" value="20:00"', page)
        self.assertIn('id="droneTimeTo" class="vs-input" value="06:00"', page)
        self.assertIn('id="droneDateFrom" class="vs-input" value="2026-06-15"',
                      page)

    def test_excel_link_carries_the_time_only_when_not_default(self):
        links = hrefs(self.page('/drones/', CROSS_DAY), '/drones/flights.xlsx')
        self.assertEqual(links, [{'date_from': '2026-06-15',
                                  'date_to': '2026-06-16',
                                  'time_from': '20:00', 'time_to': '06:00'}])
        links = hrefs(self.page('/drones/', WHOLE_DAY), '/drones/flights.xlsx')
        self.assertEqual(links, [{'date_from': '2026-06-15',
                                  'date_to': '2026-06-15'}])

    def test_book_filename(self):
        _k, _a, response = self.book_keys(CROSS_DAY)
        self.assertIn('drones_flights_2026-06-15_2026-06-16_2000-0600.xlsx',
                      filename(response))
        _k, _a, response = self.book_keys(WHOLE_DAY)
        self.assertIn('drones_flights_2026-06-15_2026-06-15.xlsx',
                      filename(response))
        self.assertNotIn('_0000-2359', filename(response))

    def test_over_cap_redirect_keeps_the_time(self):
        saved = drones.DRONE_FLIGHTS_XLSX_CAP
        drones.DRONE_FLIGHTS_XLSX_CAP = 1
        try:
            response = self.client.get('/drones/flights.xlsx?' + MINUTES)
        finally:
            drones.DRONE_FLIGHTS_XLSX_CAP = saved
        self.assertEqual(response.status_code, 302)
        query = parse_qs(urlsplit(response.headers['Location']).query)
        self.assertEqual(query['time_from'], ['08:00'])
        self.assertEqual(query['time_to'], ['10:30'])

    def test_inverted_and_bad_time_messages(self):
        page = self.page('/drones/', INVERTED)
        self.assertIn(MSG_INVERTED_RU, page)
        self.assertEqual(self.total(page), 0)
        page = self.page('/drones/', BAD_TIME)
        self.assertIn(MSG_BAD_FROM_RU, page)
        self.assertIn('id="droneTimeFrom" class="vs-input" value="00:00"', page)
        # Отрицательный контроль: при нормальном периоде сообщений нет.
        page = self.page('/drones/', CROSS_DAY)
        self.assertNotIn(MSG_INVERTED_RU, page)
        self.assertNotIn(MSG_BAD_FROM_RU, page)


# ─── 3: summary page and summary.xlsx ────────────────────────────────────────

class SummaryAndBook(Base):

    def test_html_and_book_agree_for_every_period(self):
        for query, expected in EXPECTED.items():
            page = self.page('/drones/summary', query)
            book = load_book(self.get('/drones/summary.xlsx', query))
            pairs = sheet_pairs(book.worksheets[0])
            self.assertEqual(int(stat(page, 'Вылетов')), len(expected), query)
            self.assertEqual(pairs['Вылетов'], len(expected), query)
            self.assertEqual(stat(page, 'Гектаров'),
                             '%.2f' % area_of(expected), query)
            self.assertAlmostEqual(pairs['Гектаров'], area_of(expected),
                                   places=2, msg=query)

    def test_machine_sheet_matches_the_machine_table(self):
        book = load_book(self.get('/drones/summary.xlsx', CROSS_DAY))
        ws = book['По машинам']
        by_number = {row[0]: (row[1], row[2])
                     for row in ws.iter_rows(min_row=2, values_only=True)}
        unit_of = {key: n for key, _m, _l, _a, n in FIXTURE}
        for number in (1, 2):
            keys = {k for k in EXPECTED[CROSS_DAY] if unit_of[k] == number}
            self.assertEqual(by_number[number][0], len(keys))
            self.assertAlmostEqual(by_number[number][1], area_of(keys), 2)

    def test_period_cells(self):
        pairs = sheet_pairs(load_book(
            self.get('/drones/summary.xlsx', WHOLE_DAY)).worksheets[0])
        self.assertEqual(pairs['Период: с'], '2026-06-15')
        self.assertEqual(pairs['Период: по'], '2026-06-15')
        pairs = sheet_pairs(load_book(
            self.get('/drones/summary.xlsx', MINUTES)).worksheets[0])
        self.assertEqual(pairs['Период: с'], '2026-06-15 08:00')
        self.assertEqual(pairs['Период: по'], '2026-06-15 10:30')
        pairs = sheet_pairs(load_book(self.get(
            '/drones/summary.xlsx', 'date_from=&date_to=')).worksheets[0])
        self.assertEqual(pairs['Период: с'], 'не ограничен')
        self.assertEqual(pairs['Вылетов'], len(FIXTURE))

    def test_period_cell_neutralises_a_formula(self):
        pairs = sheet_pairs(load_book(self.get(
            '/drones/summary.xlsx', 'date_from=%3D1%2B1&date_to=')
        ).worksheets[0])
        self.assertEqual(pairs['Период: с'], "'=1+1")

    def test_filename(self):
        self.assertIn('drones_summary_2026-06-15_2026-06-15_0800-1030.xlsx',
                      filename(self.get('/drones/summary.xlsx', MINUTES)))
        self.assertIn('drones_summary_2026-06-15_2026-06-15.xlsx',
                      filename(self.get('/drones/summary.xlsx', WHOLE_DAY)))

    def test_links_carry_the_time(self):
        page = self.page('/drones/summary', MINUTES)
        want = {'time_from': '08:00', 'time_to': '10:30'}
        excel = hrefs(page, '/drones/summary.xlsx')
        self.assertEqual(excel, [dict(want, date_from='2026-06-15',
                                      date_to='2026-06-15')])
        drill = [l for l in hrefs(page, '/drones/') if 'unit_id' in l]
        # H is machine 2 and F is machine 1: one drill-down row each.
        self.assertEqual(len(drill), 2)
        for link in drill:
            self.assertEqual(link['time_from'], '08:00')
            self.assertEqual(link['time_to'], '10:30')
            self.assertEqual(link['date_from'], '2026-06-15')
        all_time = hrefs(page, '/drones/summary')
        all_time = [l for l in all_time if l.get('date_from') == '']
        self.assertEqual(all_time, [dict(want, date_from='', date_to='')])

    def test_whole_day_links_stay_as_they_were(self):
        page = self.page('/drones/summary', WHOLE_DAY)
        for path in ('/drones/summary.xlsx', '/drones/', '/drones/summary'):
            for link in hrefs(page, path):
                self.assertNotIn('time_from', link, path)
                self.assertNotIn('time_to', link, path)
        self.assertEqual(hrefs(page, '/drones/summary.xlsx'),
                         [{'date_from': '2026-06-15',
                           'date_to': '2026-06-15'}])

    def test_drill_down_count_equals_the_row(self):
        """Ссылка строки машины ведёт на список с тем же набором вылетов."""
        page = self.page('/drones/summary', CROSS_DAY)
        drill = [l for l in hrefs(page, '/drones/') if 'unit_id' in l]
        self.assertEqual(len(drill), 2)
        for link in drill:
            query = '&'.join('%s=%s' % kv for kv in link.items())
            listed = self.page('/drones/', query)
            unit = db_unit_number(int(link['unit_id']))
            unit_of = {key: n for key, _m, _l, _a, n in FIXTURE}
            expected = {k for k in EXPECTED[CROSS_DAY] if unit_of[k] == unit}
            self.assertIn('Всего: %d<' % len(expected), listed)

    def test_default_is_current_month_with_default_time(self):
        page = self.page('/drones/summary')
        self.assertRegex(page, r'<input type="date" name="date_from" '
                               r'id="sumDateFrom" class="vs-input" '
                               r'value="\d{4}-\d{2}-01">')
        self.assertIn('id="sumTimeFrom" class="vs-input" value="00:00"', page)
        self.assertIn('id="sumTimeTo" class="vs-input" value="23:59"', page)

    def test_inverted_is_empty_200_and_explained_in_both_languages(self):
        for query in (INVERTED, INVERTED_TIME):
            page = self.page('/drones/summary', query)
            self.assertEqual(stat(page, 'Вылетов'), '0')
            self.assertIn(MSG_INVERTED_RU, page)
        set_language(self.user_id, 'uz')
        page = self.page('/drones/summary', INVERTED)
        self.assertIn(MSG_INVERTED_UZ, page)
        self.assertNotIn(MSG_INVERTED_RU, page)
        pairs = sheet_pairs(load_book(
            self.get('/drones/summary.xlsx', INVERTED)).worksheets[0])
        self.assertEqual(pairs['Парвозлар'], 0)

    def test_bad_time_uses_the_default_and_warns(self):
        page = self.page('/drones/summary', BAD_TIME)
        self.assertIn(MSG_BAD_FROM_RU, page)
        self.assertEqual(int(stat(page, 'Вылетов')),
                         len(EXPECTED[WHOLE_DAY]))
        self.assertIn('id="sumTimeFrom" class="vs-input" value="00:00"', page)

    def test_labels_in_both_languages(self):
        page = self.page('/drones/summary', WHOLE_DAY)
        self.assertIn('>Время с<', page)
        self.assertIn('>Время по<', page)
        set_language(self.user_id, 'uz')
        page = self.page('/drones/summary', WHOLE_DAY)
        self.assertIn('>Вақт дан<', page)
        self.assertIn('>Вақт гача<', page)
        self.assertNotIn('>Время с<', page)


def db_unit_number(unit_id):
    with app.app_context():
        return db.session.get(DroneUnit, unit_id).number


# ─── 4: spray report and spray.xlsx ──────────────────────────────────────────

class SprayAndBook(Base):

    def test_html_and_book_agree_for_every_period(self):
        for query, expected in EXPECTED.items():
            page = self.page('/drones/reports/spray', query)
            pairs = sheet_pairs(load_book(
                self.get('/drones/reports/spray.xlsx', query)).worksheets[0])
            self.assertEqual(int(stat(page, 'Вылеты')), len(expected), query)
            self.assertEqual(pairs['Вылеты'], len(expected), query)
            self.assertEqual(stat(page, 'Гектары опрыскивания'),
                             '%.2f' % area_of(expected), query)
            self.assertAlmostEqual(pairs['Гектары опрыскивания'],
                                   area_of(expected), places=2, msg=query)

    def test_period_cells_and_filename(self):
        response = self.get('/drones/reports/spray.xlsx', MINUTES)
        pairs = sheet_pairs(load_book(response).worksheets[0])
        self.assertEqual(pairs['Период: с'], '2026-06-15 08:00')
        self.assertEqual(pairs['Период: по'], '2026-06-15 10:30')
        self.assertIn('drone_spray_usage_2026-06-15_2026-06-15_0800-1030.xlsx',
                      filename(response))
        response = self.get('/drones/reports/spray.xlsx', WHOLE_DAY)
        pairs = sheet_pairs(load_book(response).worksheets[0])
        self.assertEqual(pairs['Период: с'], '2026-06-15')
        self.assertIn('drone_spray_usage_2026-06-15_2026-06-15.xlsx',
                      filename(response))

    def test_links_carry_the_time(self):
        page = self.page('/drones/reports/spray', CROSS_DAY)
        excel = hrefs(page, '/drones/reports/spray.xlsx')
        self.assertEqual(excel, [{'date_from': '2026-06-15',
                                  'date_to': '2026-06-16',
                                  'time_from': '20:00', 'time_to': '06:00'}])
        self.assertIn({'time_from': '20:00', 'time_to': '06:00'},
                      hrefs(page, '/drones/reports/spray'))
        page = self.page('/drones/reports/spray', WHOLE_DAY)
        for path in ('/drones/reports/spray', '/drones/reports/spray.xlsx'):
            links = hrefs(page, path)
            self.assertTrue(links, path)
            for link in links:
                self.assertNotIn('time_from', link, path)
                self.assertNotIn('time_to', link, path)

    def test_inputs_and_messages(self):
        page = self.page('/drones/reports/spray', CROSS_DAY)
        self.assertIn('id="sTimeFrom" class="vs-input" value="20:00"', page)
        self.assertIn('id="sTimeTo" class="vs-input" value="06:00"', page)
        page = self.page('/drones/reports/spray', INVERTED_TIME)
        self.assertIn(MSG_INVERTED_RU, page)
        self.assertEqual(stat(page, 'Вылеты'), '0')


# ─── 5: sources ──────────────────────────────────────────────────────────────

class Sources(Base):

    def test_counts_for_every_period(self):
        for query, expected in EXPECTED.items():
            page = self.page('/drones/sources', query)
            self.assertEqual(int(stat(page, 'Вылетов за период')),
                             len(expected), query)
            self.assertEqual(stat(page, 'Гектаров за период'),
                             '%.2f' % area_of(expected), query)

    def test_inputs_links_and_messages(self):
        page = self.page('/drones/sources', MINUTES)
        self.assertIn('id="srcTimeFrom" class="vs-input" value="08:00"', page)
        self.assertIn('id="srcTimeTo" class="vs-input" value="10:30"', page)
        self.assertIn({'date_from': '', 'date_to': '', 'time_from': '08:00',
                       'time_to': '10:30'}, hrefs(page, '/drones/sources'))
        page = self.page('/drones/sources', INVERTED)
        self.assertIn(MSG_INVERTED_RU, page)
        page = self.page('/drones/sources')
        self.assertIn('id="srcTimeFrom" class="vs-input" value="00:00"', page)
        self.assertIn({'date_from': '', 'date_to': ''},
                      hrefs(page, '/drones/sources'))


# ─── 6: coverage keeps whole days ────────────────────────────────────────────

class CoverageWholeDays(Base):

    def test_no_time_field_and_time_is_ignored(self):
        page = self.page('/drones/coverage', 'date_from=2026-06-01'
                         '&date_to=2026-06-30&time_from=08:00&time_to=09:00')
        self.assertNotIn('name="time_from"', page)
        self.assertNotIn('name="time_to"', page)
        with app.test_request_context('/drones/coverage?date_from=2026-06-01'
                                      '&time_from=08:00'):
            from flask import request
            filters = drones._drone_coverage_filters(request.args)
        self.assertFalse(filters['with_time'])
        self.assertTrue(filters['time_is_default'])
        self.assertEqual(drone_period.link_args(filters), {})

    def test_date_semantics_are_unchanged(self):
        from flask import request
        with app.test_request_context('/drones/coverage'):
            absent = drones._drone_coverage_filters(request.args)
        with app.test_request_context('/drones/coverage?date_from=&date_to='):
            empty = drones._drone_coverage_filters(request.args)
        with app.test_request_context('/drones/coverage?date_from=2026-6-1'
                                      '&date_to=junk'):
            loose = drones._drone_coverage_filters(request.args)
        self.assertEqual((absent['date_to'] - absent['date_from']).days,
                         drones.DRONE_COVERAGE_DEFAULT_DAYS)
        self.assertEqual((empty['date_from'], empty['date_to']), (None, None))
        self.assertEqual(loose['date_from_s'], '2026-06-01')
        self.assertEqual(loose['date_to_s'], '')
        self.assertIsNone(loose['date_to'])

    def test_inverted_dates_are_explained(self):
        page = self.page('/drones/coverage', INVERTED)
        self.assertIn(MSG_INVERTED_RU, page)
        page = self.page('/drones/coverage', WHOLE_DAY)
        self.assertNotIn(MSG_INVERTED_RU, page)


# ─── 7: the export helper with a filter dict that has no time ───────────────

class ExportName(unittest.TestCase):

    def name(self, filters):
        from openpyxl import Workbook
        with app.test_request_context('/'):
            response = drones._drone_xlsx_response(Workbook(), 'drone_x',
                                                   filters)
        return response.headers['Content-Disposition']

    def test_filter_dict_without_time_keeps_the_old_name(self):
        """Отчёты площади передают словарь без ключей времени."""
        self.assertIn('drone_x_2026-06-01_2026-06-30.xlsx', self.name(
            {'date_from_s': '2026-06-01', 'date_to_s': '2026-06-30'}))
        self.assertIn('drone_x_all.xlsx', self.name(
            {'date_from_s': '', 'date_to_s': ''}))

    def test_time_without_a_date_bounds_nothing_and_is_not_claimed(self):
        filters = drone_period.parse({'time_from': '08:00'})
        self.assertEqual(drone_period.utc_bounds(filters), (None, None))
        self.assertIn('drone_x_all.xlsx', self.name(filters))


# ─── 8: pagination links keep the time ──────────────────────────────────────

class Pagination(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        reset_db()
        cls.user_id = create_admin('period_pages')
        set_language(cls.user_id, 'ru')
        with app.app_context():
            unit = seed_units([7])[7]
            # 55 flights between 09:00 and 09:54 local, one a minute.
            for minute in range(55):
                db.session.add(DroneFlight(
                    dji_flight_id=880000 + minute, drone_unit_id=unit,
                    nickname_raw='pages', raw_json='{}', area_ha=1.0,
                    started_at=datetime(2026, 7, 1, 9, minute) - UTC_OFFSET))
            db.session.commit()

    def page_links(self, page):
        return [l for l in hrefs(page, '/drones/') if 'page' in l]

    def test_whole_day_pages_carry_no_time(self):
        client = app.test_client()
        login(client, self.user_id)
        page = client.get('/drones/?date_from=2026-07-01&date_to=2026-07-01'
                          ).get_data(as_text=True)
        self.assertIn('Всего: 55', page)
        self.assertEqual(self.page_links(page),
                         [{'page': '2', 'date_from': '2026-07-01',
                           'date_to': '2026-07-01'}])

    def test_next_page_link_carries_the_time(self):
        client = app.test_client()
        login(client, self.user_id)
        # From 09:04 on: 51 flights, two pages.
        query = 'date_from=2026-07-01&date_to=2026-07-01&time_from=09:04'
        page = client.get('/drones/?' + query).get_data(as_text=True)
        self.assertIn('Всего: 51', page)
        links = self.page_links(page)
        self.assertEqual(links, [{'page': '2', 'date_from': '2026-07-01',
                                  'date_to': '2026-07-01',
                                  'time_from': '09:04'}])
        # Following the link lands on the same 51 flights, not on 55.
        second = client.get('/drones/?' + '&'.join(
            '%s=%s' % kv for kv in links[0].items())).get_data(as_text=True)
        self.assertIn('Страница 2 / 2', second)
        self.assertIn('Всего: 51', second)
        # From 09:05 on: exactly one page, no next link.
        page = client.get('/drones/?date_from=2026-07-01&date_to=2026-07-01'
                          '&time_from=09:05').get_data(as_text=True)
        self.assertIn('Всего: 50', page)
        self.assertEqual(self.page_links(page), [])


if __name__ == '__main__':
    unittest.main()
