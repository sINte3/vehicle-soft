# -*- coding: utf-8 -*-
"""DJI-AREA-PRODUCTIONIZATION-001: экран, книга и endpoint манифеста во Flask.

Семантику и слова держит `tests/test_dji_area_control_report_001.py` без
приложения. Здесь -- то, что существует только вместе с Flask:

* право `drones` решает оба маршрута; аноним отчёта не получает;
* экран показывает пять величин, статус полноты, разрез по дронам, цепочку
  A -> B -> C со ссылками и слова «Bridge -- не корректируется»;
* числа экрана равны числам `dji_area.control_report` на тех же строках;
* RAW не тронут: `drone_flights.area_ha` после открытия отчёта тот же;
* `billable_area_m2` остаётся NULL;
* техническая страница area-evidence своего контракта не потеряла;
* endpoint манифеста закрыт токеном, ограничен по окну и ничего не пишет.

Нужен Flask: набор идёт локально и на сервере, в CI его нет (контракт
зависимостей workflow -- stdlib + jinja2 + openpyxl).

ВСЕ ДАННЫЕ СИНТЕТИЧЕСКИЕ.
"""

import io
import json
import sqlite3
import unittest
from datetime import date, datetime

from tests.harness import app, login, TEST_DB_PATH
from tests.test_dji_area_report_001 import (Base as ReportBase, HW6, HW7,
                                            PROVIDER, TOKEN)

from models import (db, DjiAreaCalculation, DroneFlight, User,
                    UserModulePermission, ROLE_OPERATOR)

from dji_area import control_report as dji_control
from dji_area import resolver as dji_resolver

DAY = date(2026, 6, 5)
WINDOW = '?date_from=2026-06-01&date_to=2026-06-30'
A, B, C = 900201, 900202, 900203
PARTIAL_C, PENDING_C, REVIEW_C, RULE_MISS = 900213, 900223, 900233, 900240


class Base(ReportBase):

    def seed(self):
        specs = (
            (A, dji_resolver.RAW_CORROBORATED, dji_resolver.AGG_CERTIFIED,
             90000.0, 90000.0, HW6, {}),
            (B, dji_resolver.RAW_CORROBORATED, dji_resolver.AGG_CERTIFIED,
             6000.0, 6000.0, HW6, {}),
            (C, dji_resolver.COUNTER_FLAT_RAW_OVERSTATED,
             dji_resolver.AGG_CERTIFIED, 90000.0, 0.0, HW6,
             {'candidate': True, 'base': A, 'bridges': [B]}),
            (PARTIAL_C, dji_resolver.PARTIAL_RECORDED_OVERSTATEMENT,
             dji_resolver.AGG_CERTIFIED, 80000.0, 5000.0, HW7,
             {'candidate': True, 'base': 900211, 'bridges': [900212]}),
            (PENDING_C, dji_resolver.RAW_UNVERIFIED,
             dji_resolver.AGG_PROVISIONAL, 70000.0, 70000.0, HW7,
             {'candidate': True, 'base': 900221, 'bridges': [900222],
              'v4': False}),
            (REVIEW_C, dji_resolver.COUNTER_FLAT_RAW_OVERSTATED,
             dji_resolver.AGG_UNRESOLVED, 50000.0, None, HW6,
             {'candidate': True, 'base': 900231, 'bridges': [900232],
              'flags': ['APPLICATION_WITH_FLAT_COUNTER']}),
            (RULE_MISS, dji_resolver.COUNTER_FLAT_RAW_OVERSTATED,
             dji_resolver.AGG_CERTIFIED, 58400.0, 0.0, HW7, {}),
        )
        with app.app_context():
            for n, (fid, status, elig, raw, corr, hw, extra) in enumerate(specs):
                obj = self.calc_object(fid, status, elig, raw_m2=raw,
                                       corrected_m2=corr, hardware_id=hw,
                                       minute=n * 5,
                                       flags=extra.get('flags'))
                obj.structural_candidate = bool(extra.get('candidate'))
                obj.scalar_source_check = True if extra.get('candidate') \
                    else None
                obj.candidate_base_flight_id = extra.get('base')
                obj.bridge_flight_ids_json = json.dumps(
                    extra.get('bridges') or [])
                obj.v4_revision_id = 1 if extra.get('v4', True) else None
                obj.structural_rule_version = \
                    'structural-retained-screen-frozen-1'
                db.session.add(obj)
            db.session.commit()

    def control_page(self, query=WINDOW, language='ru', user_id=None):
        client = self.client_as(user_id=user_id, language=language)
        response = client.get('/drones/area-control' + query)
        self.assertEqual(response.status_code, 200, query)
        return response.get_data(as_text=True)

    def control_book(self, query=WINDOW, language='ru'):
        from openpyxl import load_workbook
        client = self.client_as(language=language)
        response = client.get('/drones/area-control.xlsx' + query)
        self.assertEqual(response.status_code, 200, query)
        return response, load_workbook(io.BytesIO(response.data))


class Access(Base):

    def make_operator(self, username, has_drones):
        with app.app_context():
            user = User(username=username, role=ROLE_OPERATOR,
                        full_name='SYNTHETIC operator', language='ru')
            user.set_password('test-password')
            db.session.add(user)
            db.session.flush()
            if has_drones:
                db.session.add(UserModulePermission(
                    user_id=user.id, module_code='drones', has_access=True))
            db.session.commit()
            return user.id

    def test_the_module_permission_decides_both_routes(self):
        allowed = self.make_operator('ctl-allowed', True)
        denied = self.make_operator('ctl-denied', False)
        for path in ('/drones/area-control', '/drones/area-control.xlsx'):
            client = app.test_client()
            login(client, allowed)
            self.assertEqual(client.get(path + WINDOW).status_code, 200, path)
            client = app.test_client()
            login(client, denied)
            self.assertNotEqual(client.get(path + WINDOW).status_code, 200,
                                path)

    def test_an_anonymous_visitor_is_not_served_the_report(self):
        for path in ('/drones/area-control', '/drones/area-control.xlsx'):
            response = app.test_client().get(path + WINDOW)
            self.assertNotEqual(response.status_code, 200, path)


class Screen(Base):

    def setUp(self):
        super(Screen, self).setUp()
        self.seed()

    def test_the_five_figures_match_the_pure_report(self):
        html = self.control_page()
        # RAW 44.44, исключено 9.00 + 7.50 + 5.84 = 22.34, после 22.10,
        # ждёт V4 7.00, на проверке 5.00.
        for figure in ('44.44', '22.34', '22.10', '7.00', '5.00'):
            self.assertIn(figure, html, figure)
        for label in ('DJI RAW, га', 'Подтверждённо исключено, га',
                      'Площадь после подтверждённых корректировок, га',
                      'Ожидает доказательства / V4, га',
                      'Требует проверки, га'):
            self.assertIn(label, html, label)

    def test_the_period_is_called_open_and_never_final(self):
        html = self.control_page()
        self.assertIn('Есть нерешённые записи; они пока учтены по DJI RAW.',
                      html)
        for word in ('точная площадь', 'Финальная площадь',
                     'финальная площадь'):
            self.assertNotIn(word, html)

    def test_the_chain_is_shown_with_links_and_the_bridge_is_not_excluded(self):
        html = self.control_page()
        for fid in (A, B, C):
            self.assertIn('https://www.djiag.com/record/%d' % fid, html)
        self.assertIn('Bridge — не корректируется', html)
        self.assertIn('Промежуточный ручной участок; не исключается '
                      'автоматически.', html)

    def test_the_reason_is_words_not_an_enum(self):
        html = self.control_page()
        self.assertIn('Повтор площади предыдущей Auto-работы', html)
        self.assertIn('V4 подтвердил только 0.5000 га нового прироста', html)
        self.assertIn('Завышение подтверждено V4 контрольной проверкой', html)
        for enum in ('COUNTER_FLAT_RAW_OVERSTATED', 'PHANTOM_PROVEN',
                     'PARTIAL_RECORDED_OVERSTATEMENT'):
            self.assertNotIn(enum, html)

    def test_the_views_filter_the_register(self):
        confirmed = self.control_page(WINDOW + '&view=confirmed')
        self.assertIn('>%d<' % C, confirmed)
        self.assertNotIn('>%d<' % PENDING_C, confirmed)
        pending = self.control_page(WINDOW + '&view=pending')
        self.assertIn('>%d<' % PENDING_C, pending)
        self.assertNotIn('>%d<' % REVIEW_C, pending)
        review = self.control_page(WINDOW + '&view=review')
        self.assertIn('>%d<' % REVIEW_C, review)
        # Итоги от вкладки не зависят.
        for html in (confirmed, pending, review):
            self.assertIn('44.44', html)

    def test_the_drone_filter_narrows_the_totals(self):
        html = self.control_page(WINDOW + '&unit_id=%d' % self.unit6_id)
        # Машина 6: A 9 + B 0.6 + C 9 + REVIEW 5 = 23.60 га RAW.
        self.assertIn('23.60', html)
        self.assertNotIn('44.44', html)

    def test_the_uzbek_page_is_cyrillic_and_carries_the_same_numbers(self):
        html = self.control_page(language='uz')
        self.assertIn('DJI майдони назорати', html)
        self.assertIn('Bridge — тузатилмайди', html)
        self.assertIn('22.34', html)
        self.assertNotIn('Контроль площади DJI', html)

    def test_nothing_but_ids_and_words_leaves_the_server(self):
        html = self.control_page()
        for marker in (TOKEN, PROVIDER, HW6, HW7, 'bridge_flight_ids_json',
                       'calculation_input_hash', 'anomaly_flags_json',
                       'raw_json'):
            self.assertNotIn(marker, html, marker)

    def test_the_screen_numbers_equal_the_pure_module(self):
        with app.app_context():
            rows = []
            for c in DjiAreaCalculation.query.all():
                rows.append({col.name: getattr(c, col.name)
                             for col in DjiAreaCalculation.__table__.columns})
        total = dji_control.build(rows, 'ru')['total']
        html = self.control_page()
        for key in ('raw_m2', 'excluded_m2', 'after_m2', 'pending_m2',
                    'review_m2'):
            self.assertIn('%.2f' % (total[key] / 10000.0), html, key)


class RawIsImmutable(Base):

    def test_opening_the_report_writes_nothing(self):
        with app.app_context():
            db.session.add(DroneFlight(
                dji_flight_id=C, drone_unit_id=self.unit6_id,
                nickname_raw='SYNTHETIC-NICK',
                started_at=datetime(2026, 6, 5, 3, 0), area_ha=9.0,
                raw_json='{}'))
            db.session.commit()
        self.seed()
        before = self.raw('SELECT dji_flight_id, area_ha FROM drone_flights')
        calc_before = self.raw(
            'SELECT flight_id, raw_area_m2, corrected_recorded_area_m2, '
            'billable_area_m2 FROM dji_area_calculations ORDER BY flight_id')
        self.control_page()
        self.control_book()
        self.assertEqual(
            self.raw('SELECT dji_flight_id, area_ha FROM drone_flights'),
            before)
        self.assertEqual(self.raw(
            'SELECT flight_id, raw_area_m2, corrected_recorded_area_m2, '
            'billable_area_m2 FROM dji_area_calculations ORDER BY flight_id'),
            calc_before)
        # Исправленный фантом C: RAW вылета остался 9.0 га.
        self.assertEqual(dict(before)[C], 9.0)

    def test_billable_area_stays_null(self):
        self.seed()
        self.control_page()
        values = {row[0] for row in self.raw(
            'SELECT billable_area_m2 FROM dji_area_calculations')}
        self.assertEqual(values, {None})


class Workbook(Base):

    def test_the_route_serves_the_four_sheets_with_the_period_in_the_name(self):
        self.seed()
        response, book = self.control_book()
        # V2: прежние четыре листа первыми и в прежнем порядке, новые --
        # после них (по дням, полный реестр, история решений).
        self.assertEqual(book.sheetnames, ['Сводка', 'По_дронам',
                                           'Корректировки',
                                           'Требует_проверки', 'По_дням',
                                           'Реестр', 'История_решений'])
        self.assertIn('drone_area_control_2026-06-01_2026-06-30.xlsx',
                      response.headers['Content-Disposition'])

    def test_the_book_is_always_the_full_register(self):
        self.seed()
        _response, book = self.control_book(WINDOW + '&view=review')
        self.assertEqual(book['Корректировки'].max_row, 1 + 3)


class TechnicalPageIsUntouched(Base):

    def test_the_technical_estimate_still_answers_and_links_back(self):
        self.seed()
        html = self.page(WINDOW)
        self.assertIn('Площадь DJI: техническая оценка', html)
        # Производного итога на технической странице по-прежнему нет.
        self.assertNotIn('22.10', html)

    def test_the_hub_offers_both_reports(self):
        client = self.client_as()
        html = client.get('/drones/reports').get_data(as_text=True)
        self.assertIn('Контроль площади DJI', html)
        self.assertIn('Площадь DJI: техническая оценка', html)


class ManifestEndpoint(Base):

    URL = '/drones/api/area_capture_manifest'

    def setUp(self):
        super(ManifestEndpoint, self).setUp()
        self._token = app.config.get('DRONE_API_TOKEN')
        app.config['DRONE_API_TOKEN'] = TOKEN
        self.addCleanup(app.config.__setitem__, 'DRONE_API_TOKEN',
                        self._token)

    def post(self, payload):
        return app.test_client().post(self.URL, json=payload)

    def test_a_wrong_token_is_refused(self):
        self.assertEqual(self.post({'token': 'nope'}).status_code, 401)
        self.assertEqual(self.post({}).status_code, 401)

    def test_the_right_token_gets_a_bounded_answer(self):
        response = self.post({'token': TOKEN, 'date_from': '2026-06-04',
                              'date_to': '2026-06-05'})
        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertEqual(body['status'], 'ok')
        self.assertEqual(body['date_from'], '2026-06-04')
        self.assertEqual(body['structural_rule_version'],
                         'structural-retained-screen-frozen-1')
        self.assertNotIn(TOKEN, json.dumps(body))

    def test_a_window_longer_than_the_cap_is_a_400(self):
        response = self.post({'token': TOKEN, 'date_from': '2026-01-01',
                              'date_to': '2026-06-05'})
        self.assertEqual(response.status_code, 400)

    def test_a_malformed_date_is_a_400_not_a_default(self):
        response = self.post({'token': TOKEN, 'date_from': 'yesterday'})
        self.assertEqual(response.status_code, 400)

    def test_get_is_not_allowed(self):
        self.assertEqual(app.test_client().get(self.URL).status_code, 405)

    def test_the_endpoint_writes_nothing(self):
        import hashlib

        def digest():
            con = sqlite3.connect(TEST_DB_PATH)
            con.execute('PRAGMA wal_checkpoint(TRUNCATE)')
            con.close()
            with open(TEST_DB_PATH, 'rb') as fh:
                return hashlib.sha256(fh.read()).hexdigest()

        before = digest()
        self.post({'token': TOKEN})
        self.assertEqual(digest(), before)


if __name__ == '__main__':
    unittest.main()
