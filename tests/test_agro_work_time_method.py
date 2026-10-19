# -*- coding: utf-8 -*-
"""Сверка заявок метода «время» -- часы зажигания из gps_daily_engine.

Решение владельца 09.10: моточас -- время с включённым зажиганием по
трекеру. Вердикт предрегистрирован в треке GPS (§0.6, 19.10.2026): работа
была -- не меньше 30 минут за окно; работы не было -- только если все сутки
окна известны и даже вместе с «неизвестно» меньше 30 минут; иначе -- без
вердикта. Объём в часах -- допуск В-2. Пока таблицы часов нет, заявка
остаётся с прежней причиной.

Запуск: python -m unittest tests.test_agro_work_time_method -v
"""

import io
import os
import sqlite3
import sys
import unittest
from contextlib import redirect_stderr, redirect_stdout

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from openpyxl import load_workbook                             # noqa: E402

import migrate_gps_engine_001 as engine_mig                    # noqa: E402
from agro_work import reconcile as rc                          # noqa: E402
from gps import tolerance                                      # noqa: E402
from tests.test_agro_work_reconcile import (D, Fixture, W_GA,  # noqa: E402
                                            W_TIME)
from tools import agro_work_reconcile as tool                  # noqa: E402

MIN = 60
HOUR = 3600


class TimeFixture(Fixture):
    def __init__(self, engine=True):
        super().__init__()
        if engine:
            self.con.execute(engine_mig.CREATE_GPS_DAILY_ENGINE)
            self.con.commit()

    def hours(self, unit, day, on, unknown=0, with_ign=100, km=0.0,
              points=100):
        self.con.execute(
            "INSERT INTO gps_daily_engine (work_date, wialon_id, on_seconds, "
            "off_seconds, unknown_seconds, points_total, points_with_ign, "
            "track_km, method_version, computed_at) VALUES (?, ?, ?, 0, ?, ?, "
            "?, ?, 'ignition-R300-P150-2026-10-19', 'now')",
            ('2026-09-%02d' % day, unit, on, unknown, points, with_ign, km))
        self.con.commit()

    def timed(self, **kwargs):
        kwargs.setdefault('work_type', W_TIME)
        kwargs.setdefault('unit', 'HOUR')
        return self.app(**kwargs)


class Base(unittest.TestCase):
    engine = True

    def setUp(self):
        self.fx = TimeFixture(engine=self.engine)
        self.addCleanup(self.fx.close)


class WithoutTheTable(Base):
    engine = False

    def test_a_time_application_keeps_its_old_reason(self):
        app = self.fx.timed(created=10, completed=11)
        row = self.fx.forward(app)
        self.assertEqual((row['verdict'], row['reason']),
                         (rc.V_NONE, rc.R_METHOD_TIME))
        self.assertIsNone(row['gps_on_s'])


class Verdict(Base):
    def test_thirty_minutes_over_the_window_is_work(self):
        self.fx.hours(1001, 10, 20 * MIN)
        self.fx.hours(1001, 11, 10 * MIN)
        row = self.fx.forward(self.fx.timed(created=10, completed=11))
        self.assertEqual(row['verdict'], rc.V_WORK)
        self.assertEqual(row['gps_on_s'], 30 * MIN)
        self.assertEqual(row['work_days'], [D(10), D(11)])
        self.assertIsNone(row['reason'])

    def test_twenty_nine_minutes_with_every_day_known_is_no_work(self):
        self.fx.hours(1001, 10, 29 * MIN)
        self.fx.hours(1001, 11, 0)
        row = self.fx.forward(self.fx.timed(created=10, completed=11))
        self.assertEqual(row['verdict'], rc.V_NO_WORK)
        self.assertEqual(row['work_days'], [D(10)])

    def test_unknown_that_could_reach_thirty_minutes_withholds_the_verdict(self):
        self.fx.hours(1001, 10, 20 * MIN, unknown=10 * MIN)
        self.fx.hours(1001, 11, 0)
        row = self.fx.forward(self.fx.timed(created=10, completed=11))
        self.assertEqual((row['verdict'], row['reason']),
                         (rc.V_NONE, rc.R_NO_GPS))
        # 9 минут неизвестного до 30 не доводят -- работы не было
        fx = TimeFixture()
        self.addCleanup(fx.close)
        fx.hours(1001, 10, 20 * MIN, unknown=9 * MIN)
        fx.hours(1001, 11, 0)
        self.assertEqual(fx.forward(fx.timed(created=10, completed=11))['verdict'],
                         rc.V_NO_WORK)

    def test_a_day_without_a_row_is_unknown_not_idle(self):
        self.fx.hours(1001, 10, 0)
        row = self.fx.forward(self.fx.timed(created=10, completed=11))
        self.assertEqual((row['verdict'], row['reason'], row['unknown_days']),
                         (rc.V_NONE, rc.R_NO_GPS, 1))

    def test_a_day_whose_points_carried_no_ignition_is_unknown(self):
        self.fx.hours(1001, 10, 0, with_ign=0)
        self.fx.hours(1001, 11, 0)
        row = self.fx.forward(self.fx.timed(created=10, completed=11))
        self.assertEqual((row['verdict'], row['unknown_days']), (rc.V_NONE, 1))

    def test_work_found_wins_over_unknown_days(self):
        self.fx.hours(1001, 10, 45 * MIN)
        row = self.fx.forward(self.fx.timed(created=10, completed=11))
        self.assertEqual((row['verdict'], row['unknown_days']), (rc.V_WORK, 1))

    def test_days_after_today_are_unknown(self):
        self.fx.hours(1001, 28, 0)
        self.fx.hours(1001, 29, 0)
        row = self.fx.forward(self.fx.timed(created=28, completed=29),
                              date_from=28, date_to=29)
        self.assertEqual((row['verdict'], row['unknown_days']), (rc.V_NONE, 1))


    def test_hectare_applications_do_not_change(self):
        self.fx.hours(1001, 10, 5 * HOUR)
        row = self.fx.forward(self.fx.app(created=10, completed=11))
        self.assertEqual((row['verdict'], row['gps_ha']), (rc.V_WORK, 2.0))
        self.assertIsNone(row['gps_on_s'])

    def test_the_summary_counts_time_verdicts(self):
        self.fx.hours(1001, 10, 2 * HOUR)
        self.fx.hours(1001, 11, 0)
        self.fx.timed(created=10, completed=11)
        ctx = self.fx.run()
        _, total = ctx.summary()
        self.assertEqual(total['app_' + rc.V_WORK], 1)
        self.assertEqual(total.get('app_reason_' + rc.R_METHOD_TIME, 0), 0)


class Trucks(Base):
    """Решение владельца 19.10 (agro-work §9.8): у грузовика «работала или
    нет» -- по пробегу (3 км), объём -- моточасами. Машина 14 (T4, объект
    1004) -- `yuk_transport`, план-факт исключает её как непольевую."""

    def test_a_truck_is_judged_by_mileage_not_by_ignition(self):
        self.fx.hours(1004, 13, 5 * HOUR, km=0.4)
        row = self.fx.forward(self.fx.timed(transport='T4', created=13,
                                            completed=13))
        self.assertEqual((row['verdict'], row['gps_km']), (rc.V_NO_WORK, 0.4))
        self.assertEqual(row['gps_on_s'], 5 * HOUR)

    def test_three_kilometres_is_work_and_the_volume_is_in_engine_hours(self):
        self.fx.hours(1004, 13, 2 * HOUR, km=1.5)
        self.fx.hours(1004, 14, 6 * HOUR, km=1.5)
        row = self.fx.forward(self.fx.timed(transport='T4', created=13,
                                            completed=14, volume='8'))
        self.assertEqual(row['verdict'], rc.V_WORK)
        self.assertEqual(row['work_days'], [D(13), D(14)])
        self.assertEqual(row['volume_verdict'], tolerance.VERDICT_OK)
        self.assertEqual(row['volume_deviation'], 0.0)

    def test_two_point_nine_kilometres_is_not_work(self):
        self.fx.hours(1004, 13, 5 * HOUR, km=2.9)
        row = self.fx.forward(self.fx.timed(transport='T4', created=13,
                                            completed=13))
        self.assertEqual(row['verdict'], rc.V_NO_WORK)

    def test_the_volume_needs_hours_known_not_kilometres_known(self):
        # 14-го одна точка: пробег неизвестен, но зажигание у неё есть --
        # часы суток известны, и светофор моточасов ставится.
        self.fx.hours(1004, 13, 2 * HOUR, km=5.0)
        self.fx.hours(1004, 14, 0, km=0.0, points=1, with_ign=1)
        row = self.fx.forward(self.fx.timed(transport='T4', created=13,
                                            completed=14, volume='2'))
        self.assertEqual((row['verdict'], row['unknown_days']), (rc.V_WORK, 1))
        self.assertEqual(row['volume_verdict'], tolerance.VERDICT_OK)

    def test_a_day_with_one_point_is_unknown(self):
        self.fx.hours(1004, 13, 0, km=0.0, points=1)
        row = self.fx.forward(self.fx.timed(transport='T4', created=13,
                                            completed=13))
        self.assertEqual((row['verdict'], row['reason'], row['unknown_days']),
                         (rc.V_NONE, rc.R_NO_GPS, 1))

    def test_a_truck_marked_not_ours_stays_excluded(self):
        self.fx.con.execute('UPDATE vialon_mappings SET skip = 1 WHERE '
                            'wialon_id = 1004')
        self.fx.con.commit()
        self.fx.hours(1004, 13, 5 * HOUR, km=40.0)
        row = self.fx.forward(self.fx.timed(transport='T4', created=13,
                                            completed=13))
        # Строка «не наш» (skip = 1) в связки сверки не входит вовсе -- у
        # машины нет объекта Wialon, и по пробегу её не судят.
        self.assertEqual((row['verdict'], row['reason']),
                         (rc.V_NONE, rc.R_NO_WIALON))

    def test_a_passenger_car_is_not_a_truck(self):
        self.fx.con.execute("UPDATE equipment SET category = 'passenger' "
                            "WHERE id = 14")
        self.fx.con.commit()
        self.fx.hours(1004, 13, 5 * HOUR, km=40.0)
        row = self.fx.forward(self.fx.timed(transport='T4', created=13,
                                            completed=13))
        self.assertEqual((row['verdict'], row['reason']),
                         (rc.V_NONE, rc.R_EXCLUDED))

    def test_a_truck_hectare_application_is_still_excluded(self):
        self.fx.hours(1004, 13, 5 * HOUR, km=40.0)
        row = self.fx.forward(self.fx.app(transport='T4', created=13,
                                          completed=13))
        self.assertEqual(row['reason'], rc.R_EXCLUDED)


class HoursVolume(Base):
    def row(self, on, unknown=0, volume='8', unit='HOUR', days=(10, 11)):
        for day in days:
            self.fx.hours(1001, day, on // len(days), unknown // len(days))
        return self.fx.forward(self.fx.timed(created=10, completed=11,
                                             volume=volume, unit=unit))

    def test_within_tolerance(self):
        row = self.row(int(8.5 * HOUR))
        self.assertEqual(row['volume_verdict'], tolerance.VERDICT_OK)
        self.assertEqual(row['volume_deviation'], 0.5)
        self.assertAlmostEqual(row['volume_share'], 0.5 / 8)

    def test_borderline_and_out(self):
        self.assertEqual(self.row(10 * HOUR)['volume_verdict'],
                         tolerance.VERDICT_WARN)
        fx = TimeFixture()
        self.addCleanup(fx.close)
        fx.hours(1001, 10, 11 * HOUR)
        fx.hours(1001, 11, 0)
        row = fx.forward(fx.timed(created=10, completed=11, volume='8'))
        self.assertEqual(row['volume_verdict'], tolerance.VERDICT_FAIL)

    def test_motor_hours_are_hours_too(self):
        self.assertEqual(self.row(8 * HOUR, unit='MOTOR_HOUR')['volume_verdict'],
                         tolerance.VERDICT_OK)

    def test_a_volume_in_hectares_cannot_be_compared(self):
        row = self.row(8 * HOUR, unit='HECTARE')
        self.assertEqual((row['volume_verdict'], row['volume_reason']),
                         (None, rc.VR_NOT_HOURS))

    def test_unknown_days_block_the_comparison(self):
        self.fx.hours(1001, 10, 8 * HOUR)
        row = self.fx.forward(self.fx.timed(created=10, completed=11,
                                            volume='8'))
        self.assertEqual(row['volume_reason'], rc.VR_DAYS_UNKNOWN)

    def test_gaps_that_would_change_the_light_withhold_it(self):
        row = self.row(int(8.5 * HOUR), unknown=2 * HOUR)
        self.assertEqual((row['volume_verdict'], row['volume_reason']),
                         (None, rc.VR_TIME_UNKNOWN))

    def test_gaps_that_would_not_change_the_light_do_not(self):
        row = self.row(int(8.5 * HOUR), unknown=6 * MIN)
        self.assertEqual(row['volume_verdict'], tolerance.VERDICT_OK)

    def test_an_overlapping_application_withholds_the_light(self):
        self.fx.hours(1001, 10, 4 * HOUR)
        self.fx.hours(1001, 11, 4 * HOUR)
        first = self.fx.timed(created=10, completed=11, volume='8')
        self.fx.timed(created=11, completed=11, volume='2')
        row = self.fx.forward(first)
        self.assertEqual((row['volume_verdict'], row['volume_reason']),
                         (None, rc.VR_OVERLAP))
        self.assertIsNone(row['volume_group'])


class HoursTolerance(unittest.TestCase):
    """Допуск В-2 для часов: max(15 %, 0,5 ч) -- в допуске, до 25 % -- на грани."""

    def test_the_boundaries_are_inclusive_and_the_floor_is_half_an_hour(self):
        cases = ((8, 9.2, tolerance.VERDICT_OK),      # 15 % ровно
                 (8, 9.21, tolerance.VERDICT_WARN),
                 (8, 6.8, tolerance.VERDICT_OK),      # и в меньшую сторону
                 (8, 10.0, tolerance.VERDICT_WARN),   # 25 % ровно
                 (8, 10.01, tolerance.VERDICT_FAIL),
                 (2, 2.5, tolerance.VERDICT_OK),      # пол 0,5 ч больше 15 %
                 (2, 2.51, tolerance.VERDICT_FAIL))   # 25 % от 2 ч меньше пола
        for base, fact, expected in cases:
            self.assertEqual(tolerance.verdict_for_hours(base, fact)[0],
                             expected, (base, fact))

    def test_deviation_and_share(self):
        verdict, deviation, share = tolerance.verdict_for_hours(8, 6.0)
        self.assertEqual((verdict, deviation, share),
                         (tolerance.VERDICT_WARN, -2.0, -0.25))
        self.assertIsNone(tolerance.verdict_for_hours(0, 0.4)[2])


class Book(Base):
    def test_hours_go_to_their_own_columns_never_to_the_hectare_ones(self):
        self.fx.hours(1001, 10, int(4.25 * HOUR), unknown=6 * MIN)
        self.fx.hours(1001, 11, 4 * HOUR)
        self.fx.timed(created=10, completed=11, volume='8')
        self.fx.hours(1002, 13, HOUR)
        self.fx.app(transport='T2', created=13, completed=13, volume='4')
        out_path = os.path.join(os.path.dirname(self.fx.path), 'book.xlsx')
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = tool.main(['--db', self.fx.path, '--from', '2026-09-10',
                              '--to', '2026-09-20', '--out', out_path])
        self.assertEqual(code, 0, err.getvalue())
        with open(out_path, 'rb') as handle:
            sheet = list(load_workbook(handle)['Заявка-работа'].iter_rows(
                values_only=True))
        at = {title: index for index, title in enumerate(sheet[0])}
        rows = {row[0]: row for row in sheet[1:]}

        def cell(number, title):
            return rows[number][at[tool.bi(title)]]
        time_app, ga_app = 'N-001', 'N-002'
        self.assertIsNone(cell(time_app, ('Га по GPS в окне',
                                          'Ойнада GPS бўйича га')))
        self.assertIsNone(cell(time_app, ('Объём в заявке, га',
                                          'Буюртмадаги ҳажм, га')))
        self.assertEqual(cell(time_app, ('Часы зажигания по GPS в окне',
                                         'Ойнада GPS бўйича ўт олдириш '
                                         'соатлари')), 8.25)
        self.assertEqual(cell(time_app, ('Неизвестно (дыры в данных), ч',
                                         'Номаълум (маълумотлардаги '
                                         'узилишлар), соат')), 0.1)
        self.assertEqual(cell(time_app, ('Объём в заявке, ч',
                                         'Буюртмадаги ҳажм, соат')), 8.0)
        self.assertEqual(cell(time_app, ('Объём против GPS',
                                         'Ҳажм GPS га қарши')), 'В допуске')
        self.assertEqual(cell(ga_app, ('Объём в заявке, га',
                                       'Буюртмадаги ҳажм, га')), 4.0)
        self.assertIsNone(cell(ga_app, ('Часы зажигания по GPS в окне',
                                        'Ойнада GPS бўйича ўт олдириш '
                                        'соатлари')))
        truck_km = ('Пробег грузовика по GPS в окне, км',
                    'Ойнада GPS бўйича юк машинасининг босиб ўтган йўли, км')
        self.assertIsNone(cell(time_app, truck_km))

    def test_a_truck_row_carries_its_mileage(self):
        self.fx.hours(1004, 13, 2 * HOUR, km=12.345)
        self.fx.timed(transport='T4', created=13, completed=13, volume='2')
        out_path = os.path.join(os.path.dirname(self.fx.path), 'book.xlsx')
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = tool.main(['--db', self.fx.path, '--from', '2026-09-10',
                              '--to', '2026-09-20', '--out', out_path])
        self.assertEqual(code, 0, err.getvalue())
        with open(out_path, 'rb') as handle:
            sheet = list(load_workbook(handle)['Заявка-работа'].iter_rows(
                values_only=True))
        at = {title: index for index, title in enumerate(sheet[0])}
        row = [r for r in sheet[1:] if r[0] == 'N-001'][0]
        self.assertEqual(row[at[tool.bi(('Пробег грузовика по GPS в окне, км',
                                         'Ойнада GPS бўйича юк машинасининг '
                                         'босиб ўтган йўли, км'))]], 12.345)
        self.assertEqual(row[at[tool.bi(('Часы зажигания по GPS в окне',
                                         'Ойнада GPS бўйича ўт олдириш '
                                         'соатлари'))]], 2.0)


class Labels(unittest.TestCase):
    def test_new_reasons_have_both_languages_and_uzbek_is_cyrillic(self):
        import re
        from agro_work import labels
        for reason in (rc.VR_NOT_HOURS, rc.VR_TIME_UNKNOWN):
            ru, uz = labels.VOLUME_REASONS[reason]
            self.assertTrue(ru and uz)
            self.assertIsNone(re.search(r'[A-Za-z]', uz), uz)
        self.assertEqual(set(labels.VOLUME_REASONS), set(rc.VOLUME_REASONS))


if __name__ == '__main__':
    unittest.main()
