# -*- coding: utf-8 -*-
"""agro-work B2: метод сверки ставит владелец, код его не угадывает.

Держится:
  * после импорта метод пуст у ВСЕХ видов работ, включая HECTARE: единица
    из API методом не становится;
  * книга разметки -- все виды работ, единица, число заявок и выпадающий
    список ровно из четырёх методов плана; второй лист объясняет, что
    сверка делает с каждым;
  * загрузка без --apply не меняет базу ни на байт; с --apply -- одна
    транзакция и строка журнала на каждый вид работы;
  * неизвестное значение ячейки -- отказ всего файла;
  * пустая ячейка снимает метод, отсутствующая строка -- не трогает;
  * повторный импорт размеченный метод не трогает.

Запуск: python -m unittest tests.test_agro_work_methods -v
"""

import hashlib
import io
import os
import sqlite3
import sys
import unittest
from contextlib import redirect_stderr, redirect_stdout

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from openpyxl import load_workbook                          # noqa: E402

from agro_work import importer, methods, store              # noqa: E402
from agro_work.client import Client, build_opener           # noqa: E402
from tests import agro_work_db as dbh                       # noqa: E402
from tests import agro_work_fake as fake_api                # noqa: E402
from tools import agro_work_methods as tool                 # noqa: E402

WT_GA = fake_api.uuid_for(0xD, 1)
WT_HOURS = fake_api.uuid_for(0xD, 2)
WT_TRIPS = fake_api.uuid_for(0xD, 3)


def sha(path):
    with open(path, 'rb') as fh:
        return hashlib.sha256(fh.read()).hexdigest()


class MethodsCase(unittest.TestCase):
    def setUp(self):
        self.server = fake_api.FakeAgroWork()
        self.db = dbh.make_db()
        self.folder = os.path.dirname(self.db)
        self.server.work_types = [
            fake_api.work_type(1, name='Култивация', unit='HECTARE'),
            fake_api.work_type(2, name='трактор вақтбай иш (ПСТ)',
                               unit='HOUR', unit_display='Соат'),
            fake_api.work_type(3, name='Юк ташиш', unit='TON_KM',
                               unit_display='Тонна-км')]
        self.server.transports = [fake_api.transport(1)]
        self.server.applications = [
            fake_api.application(n, work_type=2 if n > 3 else 1,
                                 created='2026-09-%02dT08:00:00+05:00' % n)
            for n in range(1, 6)]
        self.server.histories = {fake_api.uuid_for(0xA, n): fake_api.history(
            created='2026-09-%02dT08:00:00+05:00' % n,
            completed='2026-09-%02dT18:00:00+05:00' % n) for n in range(1, 6)}
        self.import_once()

    def tearDown(self):
        self.server.close()

    def import_once(self):
        client = Client({'login': fake_api.LOGIN, 'password': fake_api.PASSWORD},
                        base_url=self.server.base_url, pause=0,
                        log=lambda text: None, opener=build_opener(proxies={}))
        con = store.connect(self.db)
        try:
            importer.run_import(client, con, 10, log=lambda text: None)
        finally:
            con.close()

    def main(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = tool.main(['--db', self.db] + list(args))
        return code, out.getvalue(), err.getvalue()

    def methods_now(self):
        con = sqlite3.connect(self.db)
        try:
            return dict(con.execute('SELECT id, method FROM agro_work_work_types'))
        finally:
            con.close()

    def exported(self):
        path = os.path.join(self.folder, 'methods.xlsx')
        code, out, err = self.main('--export', path)
        self.assertEqual(code, 0, err)
        return path

    def fill(self, path, values):
        """values -- {id: значение ячейки метода}; None -- пустая ячейка."""
        book = load_workbook(path)
        sheet = book[tool.SHEET]
        header = [c.value for c in sheet[1]]
        id_col = header.index(tool.H_ID) + 1
        method_col = header.index(tool.H_METHOD) + 1
        for row in range(2, sheet.max_row + 1):
            work_type_id = sheet.cell(row=row, column=id_col).value
            if work_type_id in values:
                sheet.cell(row=row, column=method_col).value = values[work_type_id]
        book.save(path)


class NothingIsGuessed(MethodsCase):
    def test_after_import_every_method_is_empty_even_for_hectares(self):
        self.assertEqual(self.methods_now(), {WT_GA: None, WT_HOURS: None,
                                              WT_TRIPS: None})

    def test_a_reimport_never_touches_a_marked_method(self):
        path = self.exported()
        self.fill(path, {WT_HOURS: 'время'})
        self.main('--import', path, '--apply')
        self.server.work_types[1]['name'] = 'трактор вақтбай иш'
        self.import_once()
        self.assertEqual(self.methods_now()[WT_HOURS], 'vremya')


class Workbook(MethodsCase):
    def test_every_work_type_with_its_unit_count_and_a_four_value_list(self):
        path = self.exported()
        book = load_workbook(path)
        sheet = book[tool.SHEET]
        self.assertEqual([c.value for c in sheet[1]], list(tool.HEADERS))
        rows = [[c.value for c in r] for r in sheet.iter_rows(min_row=2)]
        self.assertEqual([r[0] for r in rows], [WT_GA, WT_HOURS, WT_TRIPS])
        self.assertEqual([r[4] for r in rows], [3, 2, 0])
        self.assertEqual([r[2] for r in rows], ['HECTARE', 'HOUR', 'TON_KM'])
        self.assertEqual([r[7] for r in rows], [None, None, None])
        validations = sheet.data_validations.dataValidation
        self.assertEqual(len(validations), 1)
        self.assertEqual(validations[0].formula1,
                         '"гектары,время,рейсы,не сверяется"')
        self.assertIn('H2:H4', str(validations[0].sqref))
        legend = book[tool.LEGEND]
        self.assertEqual([legend.cell(row=r, column=1).value for r in range(2, 6)],
                         ['гектары', 'время', 'рейсы', 'не сверяется'])


class Markup(MethodsCase):
    def test_dry_run_changes_nothing_and_apply_writes_with_a_journal(self):
        path = self.exported()
        self.fill(path, {WT_GA: 'Гектары ', WT_HOURS: 'вақт',
                         WT_TRIPS: 'ne_sveryaetsya'})
        before = sha(self.db)
        code, out, err = self.main('--import', path)
        self.assertEqual(code, 0, err)
        self.assertIn('set 3 | change 0 | clear 0 | same 0', out)
        self.assertEqual(sha(self.db), before)
        code, out, err = self.main('--import', path, '--apply')
        self.assertEqual(code, 0, err)
        self.assertTrue(out.isascii())
        self.assertEqual(self.methods_now(), {WT_GA: 'ga', WT_HOURS: 'vremya',
                                              WT_TRIPS: 'ne_sveryaetsya'})
        con = sqlite3.connect(self.db)
        try:
            journal = con.execute("SELECT source, entity_id, old_value, "
                                  "new_value FROM agro_work_changes WHERE "
                                  "field = 'method' ORDER BY entity_id").fetchall()
            source = con.execute('SELECT DISTINCT method_source FROM '
                                 'agro_work_work_types').fetchall()
        finally:
            con.close()
        self.assertEqual(journal, [('methods', WT_GA, None, 'ga'),
                                   ('methods', WT_HOURS, None, 'vremya'),
                                   ('methods', WT_TRIPS, None, 'ne_sveryaetsya')])
        self.assertEqual(source, [('xlsx:methods.xlsx',)])

    def test_an_unknown_value_refuses_the_whole_file(self):
        path = self.exported()
        self.fill(path, {WT_GA: 'гектары', WT_HOURS: 'моточасы'})
        before = sha(self.db)
        code, out, err = self.main('--import', path, '--apply')
        self.assertEqual(code, 2)
        self.assertIn('is not a method', err)
        self.assertEqual(sha(self.db), before)

    def test_empty_cell_clears_and_a_missing_row_is_left_alone(self):
        path = self.exported()
        self.fill(path, {WT_GA: 'гектары', WT_HOURS: 'время'})
        self.main('--import', path, '--apply')
        # Новая книга: у ГА ячейка пустая, строку ВРЕМЕНИ удалили вовсе.
        path2 = self.exported()
        book = load_workbook(path2)
        sheet = book[tool.SHEET]
        header = [c.value for c in sheet[1]]
        method_col = header.index(tool.H_METHOD) + 1
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=1).value == WT_GA:
                sheet.cell(row=row, column=method_col).value = None
        for row in range(sheet.max_row, 1, -1):
            if sheet.cell(row=row, column=1).value == WT_HOURS:
                sheet.delete_rows(row)
        book.save(path2)
        code, out, err = self.main('--import', path2, '--apply')
        self.assertEqual(code, 0, err)
        self.assertIn('set 0 | change 0 | clear 1 | same 1', out)
        self.assertEqual(self.methods_now(), {WT_GA: None, WT_HOURS: 'vremya',
                                              WT_TRIPS: None})

    def test_renamed_headers_are_refused(self):
        path = self.exported()
        book = load_workbook(path)
        book[tool.SHEET].cell(row=1, column=8).value = 'Метод'
        book.save(path)
        code, _, err = self.main('--import', path)
        self.assertEqual(code, 2)
        self.assertIn('is missing in the first row', err)

    def test_parse_method_accepts_slug_and_both_labels_only(self):
        for value, slug in (('ga', 'ga'), ('ГЕКТАРЫ', 'ga'), ('гектарлар', 'ga'),
                            (' рейсы ', 'reysy'), ('солиштирилмайди',
                                                   'ne_sveryaetsya'),
                            ('', None), (None, None)):
            self.assertEqual(methods.parse_method(value), slug, value)
        for value in ('га', 'hectare', 'моточас', 'HECTARE', '1'):
            with self.assertRaises(methods.BadMethod, msg=value):
                methods.parse_method(value)

    def test_status_without_keys_counts_what_waits_for_markup(self):
        code, out, _ = self.main()
        self.assertEqual(code, 0)
        self.assertIn('not marked      3', out)
        self.assertIn('applications of unmarked work types: 5', out)


if __name__ == '__main__':
    unittest.main()
