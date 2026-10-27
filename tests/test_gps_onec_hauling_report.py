# -*- coding: utf-8 -*-
"""tools/gps_onec_hauling_report.py -- недельная сверка перевозок хлопка по рейсам 1С (GPS-35).

Держится то, что решил владелец 27.10 и что записано в треке (§0.7):
  * сутки сверки -- сутки реестра ПК-17; заявка за сутки вне выгрузок не
    судится;
  * лист «Заявки»: каждые сутки в окне каждой заявки «пахта ташиш» / «тола
    ташиш» -- рейсы её машины в 1С (вывоз, приём), нетто, итог: есть рейс,
    нет рейса, машина не сопоставлена; прочие виды работ на лист не идут;
  * лист «Без заявки»: машино-сутки наших машин с рейсами 1С, которые не
    покрывает ни одна заявка любого вида -- по вердикту покрытия экрана
    «Сверка»: отменённая заявка не покрывает, открытая -- «без вердикта» с
    причиной и номером, машина вне agro-work -- «без вердикта»; машины не из
    справочника не идут;
  * заявки «пахта/тола ташиш» без окна не судятся и считаются по причинам;
  * нетто: вывоз -- «Вес нетто», приём -- «Масса физическая»;
  * заголовки двуязычные; ничего не пишется, кроме книги; в консоль --
    ASCII; Flask не импортируется; неверный ввод -- код 2.

Запуск из корня репозитория:
  python -m unittest tests.test_gps_onec_hauling_report
"""
import contextlib
import hashlib
import io
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from datetime import date

import openpyxl

import tools.gps_onec_hauling_report as report
from tests import agro_work_db as dbh
from tests.onec_xlsx import write_xlsx

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VYVOZ_HEADER = ('Завод', 'Пункт приема', 'Гос номер (из названия)',
                'Время первого взвешивания', 'Вес нетто, кг', 'Вес брутто, кг')
REESTR_HEADER = ('Пункт приема', 'Фермер', 'Транспорт', 'Дата брутто',
                 'Масса брутто, кг', 'Масса физическая,кг')

WORLD = {}


def digest(path):
    with open(path, 'rb') as handle:
        return hashlib.sha256(handle.read()).hexdigest()


def setUpModule():
    root = tempfile.mkdtemp()
    db = dbh.make_db(root)
    con = sqlite3.connect(db)
    try:
        dbh.add_org(con, 1, 'Buxoro')
        dbh.add_org(con, 2, 'Mirzachul')
        # A -- хлопок 09--11.10 (09.10 вне выгрузок), волокно 11.10 и прочий
        # груз 10.10; B -- хлопок 10.10 без рейсов
        # и открытая заявка с 11.10; C -- чизель 10--11.10 (рейсы есть, но
        # заявка не перевозка); D -- заявка только 12.10, вне выгрузок, и
        # отменённая 10.10; рейсы D 10.10 -- без заявки; E -- нет в agro-work
        for eq, plate, org in ((10, '80 101 HA', 1), (11, '80 102 HA', 1),
                               (12, '80 103 HA', 2), (13, '80 609 DA', 1),
                               (14, '80 777 HA', 1), (15, '80 777 HA', 1),
                               (16, '80 104 HA', 1)):
            dbh.add_equipment(con, eq, plate, org_id=org, name='МТЗ-%d' % eq)
        con.execute("INSERT INTO agro_work_import_runs (id, started_at, status, "
                    "tool_version) VALUES (1, 't', 'ok', 'v')")
        for work_type, name in (('w-cotton', 'Пахта ташиш'),
                                ('w-lint', 'Тола ташиш (ички)'),
                                ('w-ga', 'Чизеллаш'),
                                ('w-cargo', 'Юк ташиш'),
                                ('w-pick', 'Пахта териш')):
            con.execute("INSERT INTO agro_work_work_types (id, name, unit, method, "
                        "first_seen_at, last_seen_at) VALUES (?, ?, 'TON', NULL, "
                        "'t', 't')", (work_type, name))
        transports = {'TA': 10, 'TB': 11, 'TC': 12, 'TD': 13, 'TU': None}
        for transport, eq in transports.items():
            con.execute("INSERT INTO agro_work_transports (id, plate_number, "
                        "plate_norm, equipment_id, match_status, first_seen_at, "
                        "last_seen_at) VALUES (?, ?, ?, ?, ?, 't', 't')",
                        (transport, 'P' + transport, 'P' + transport, eq,
                         'auto' if eq else 'none'))
        for app_id, transport, work_type, volume, status, first, last in (
                ('X', 'TA', 'w-cotton', '310', 'COMPLETED', '2026-10-09', '2026-10-11'),
                ('T', 'TA', 'w-lint', '5', 'COMPLETED', '2026-10-11', '2026-10-11'),
                ('G', 'TA', 'w-cargo', '1', 'COMPLETED', '2026-10-10', '2026-10-10'),
                ('Y', 'TB', 'w-cotton', '7', 'COMPLETED', '2026-10-10', '2026-10-10'),
                ('P', 'TB', 'w-pick', '2', 'COMPLETED', '2026-10-10', '2026-10-10'),
                ('O', 'TB', 'w-cotton', '6', 'IN_PROGRESS', '2026-10-11', None),
                ('Z', 'TC', 'w-ga', '3', 'COMPLETED', '2026-10-10', '2026-10-11'),
                ('Q', 'TD', 'w-cotton', '9', 'COMPLETED', '2026-10-12', '2026-10-12'),
                ('K', 'TD', 'w-cotton', '8', 'CANCELLED', '2026-10-10', None),
                ('U', 'TU', 'w-cotton', '4', 'COMPLETED', '2026-10-10', '2026-10-10')):
            con.execute(
                "INSERT INTO agro_work_applications (id, application_number, "
                "transport_id, plate_number, work_type_id, unit, volume, status, "
                "created_at, updated_at, created_day, first_seen_run_id, "
                "last_seen_run_id, history_updated_at, initial_status, "
                "completed_day) VALUES (?, ?, ?, ?, ?, 'TON', ?, ?, ?, "
                "'u', ?, 1, 1, 'u', 'IN_PROGRESS', ?)",
                ('app-' + app_id, 'N-' + app_id, transport, 'P' + transport,
                 work_type, volume, status, first + 'T07:00:00+05:00', first, last))
        con.commit()
    finally:
        con.close()

    vyvoz = [VYVOZ_HEADER,
             ('Z', 'Вобкент ПТМ', '80 101 HA', '2026-10-10 08:00:00', '7000', '12000'),
             ('Z', 'Вобкент ПТМ', '80 101 HA', '2026-10-10 11:30:00', '7 500', '12500'),
             ('Z', 'Вобкент ПТМ', '80 101 HA', '2026-10-11 09:00:00', '8000', '13000'),
             ('Z', 'Харгуш ПТМ', '80 103 HA', '2026-10-10 10:00:00', '6000', '11000'),
             ('Z', 'Харгуш ПТМ', '80 609 DA', '2026-10-10 12:15:00', '5000', '10000'),
             ('Z', 'Вобкент ПТМ', '80 609 DA', '2026-10-10 15:40:00', '5500', '10500'),
             ('Z', 'Вобкент ПТМ', '80 999 ZZ', '2026-10-10 09:00:00', '9000', '14000'),
             ('Z', 'Вобкент ПТМ', '80 777 HA', '2026-10-10 09:00:00', '9000', '14000'),
             ('Z', 'Вобкент ПТМ', '80 104 HA', '2026-10-10 10:30:00', '4000', '9000')]
    reestr = [REESTR_HEADER,
              ('К-Бозор ПТМ', 'farmer one', '80 DA 999', '10.10.2026 08:00:00', '9000', '2000'),
              ('К-Бозор ПТМ', 'farmer two', '80 HA 101', '11.10.2026 13:00:00', '9000', '2123,4'),
              ('Кумушкент ПТМ', 'farmer three', '80 HA 102', '11.10.2026 14:00:00', '9000', '1500'),
              ('Кумушкент ПТМ', 'farmer four', '80 HA 103', '11.10.2026 15:00:00', '9000', '1000')]
    vyvoz_path = os.path.join(root, 'vyvoz.xlsx')
    reestr_path = os.path.join(root, 'reestr.xlsx')
    write_xlsx(vyvoz_path, vyvoz)
    write_xlsx(reestr_path, reestr)
    before = {path: digest(path) for path in (db, vyvoz_path, reestr_path)}
    out = os.path.join(root, 'report.xlsx')
    argv = ['--db', db, '--vyvoz', vyvoz_path, '--reestr', reestr_path, '--out', out]
    code, text = run_main(*argv)
    WORLD.update(root=root, db=db, vyvoz=vyvoz_path, reestr=reestr_path,
                 before=before, out=out, code=code, text=text, argv=argv)


def tearDownModule():
    shutil.rmtree(WORLD['root'], True)


def run_main(*argv):
    lines = []
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        code = report.main([str(arg) for arg in argv], out=lines.append)
    return code, '\n'.join(lines) + err.getvalue()


def sheet(index):
    book = openpyxl.load_workbook(WORLD['out'], read_only=True)
    try:
        return [list(row) for row in book.worksheets[index].iter_rows(values_only=True)]
    finally:
        book.close()


def day(value):
    return value.date() if hasattr(value, 'date') else value


class Book(unittest.TestCase):

    def test_sheets_and_bilingual_headers(self):
        self.assertEqual(WORLD['code'], report.EXIT_OK, WORLD['text'])
        book = openpyxl.load_workbook(WORLD['out'], read_only=True)
        try:
            self.assertEqual(book.sheetnames, ['Сводка — Жамланма', 'Заявки — Буюртмалар',
                                               'Без заявки — Буюртмасиз'])
        finally:
            book.close()
        self.assertEqual(sheet(1)[0][0], '№ заявки / Буюртма №')
        self.assertEqual(sheet(2)[0][7], 'Пункты приёма / Қабул пунктлари')

    def test_application_days(self):
        rows = {(row[0], day(row[1])): row for row in sheet(1)[1:]}
        # X 10.10 и 11.10 (09.10 вне выгрузок), T 11.10, Y 10.10, U 10.10; Q
        # 12.10 вне выгрузок, Z и P -- не перевозка, G -- перевозка не
        # хлопка, у O и K нет окна
        self.assertEqual(sorted(rows), [
            ('N-T', date(2026, 10, 11)), ('N-U', date(2026, 10, 10)),
            ('N-X', date(2026, 10, 10)), ('N-X', date(2026, 10, 11)),
            ('N-Y', date(2026, 10, 10))])
        x10 = rows[('N-X', date(2026, 10, 10))]
        self.assertEqual(x10[2:12], ['МТЗ-10', '80 101 HA', 'Buxoro', 'Пахта ташиш',
                                     310, 'TON', 2, 0, 14.5,
                                     'Есть рейс в 1С / 1С да рейс бор'])
        # 11.10: вывоз 8 т и приём 2,1234 т -- реестр пишет серию впереди
        x11 = rows[('N-X', date(2026, 10, 11))]
        self.assertEqual(x11[8:12], [1, 1, 10.123, 'Есть рейс в 1С / 1С да рейс бор'])
        self.assertEqual(rows[('N-T', date(2026, 10, 11))][8:12],
                         [1, 1, 10.123, 'Есть рейс в 1С / 1С да рейс бор'])
        y = rows[('N-Y', date(2026, 10, 10))]
        self.assertEqual(y[8:12], [0, 0, None, 'Нет рейса в 1С / 1С да рейс йўқ'])
        u = rows[('N-U', date(2026, 10, 10))]
        # машины в справочнике нет: название пустое (пустая ячейка), номер --
        # из заявки agro-work
        self.assertEqual((u[2], u[3], u[11]),
                         (None, 'PTU', 'Машина не сопоставлена / Машина боғланмаган'))

    def test_machine_days_not_covered_by_an_application(self):
        rows = {(row[2], day(row[0])): row for row in sheet(2)[1:]}
        # D 10.10 -- рейсы без заявки: его заявка только 12.10, а заявка 10.10
        # отменена; B 11.10 -- приём, покрыть может только открытая заявка;
        # E 10.10 -- машины нет в agro-work; C 10.10 и 11.10 -- в окне чизеля
        # (второй день окна тоже покрыт); чужой номер и номер двух машин не
        # идут вовсе
        self.assertEqual(sorted(rows), [('80 102 HA', date(2026, 10, 11)),
                                        ('80 104 HA', date(2026, 10, 10)),
                                        ('80 609 DA', date(2026, 10, 10))])
        d = rows[('80 609 DA', date(2026, 10, 10))]
        self.assertEqual(d[1:], ['МТЗ-13', '80 609 DA', 'Buxoro', 2, 0, 10.5,
                                 'Вобкент ПТМ, Харгуш ПТМ', '12:15', '15:40',
                                 'Работа без заявки / Буюртмасиз иш', None, None])
        b = rows[('80 102 HA', date(2026, 10, 11))]
        self.assertEqual(b[4:8], [0, 1, 1.5, 'Кумушкент ПТМ'])
        self.assertEqual(b[10:], ['Без вердикта / Ҳукмсиз',
                                  'Может покрыть открытая заявка — вердикт после закрытия / '
                                  'Очиқ буюртма қоплаши мумкин — ҳукм ёпилгандан кейин',
                                  'N-O'])
        e = rows[('80 104 HA', date(2026, 10, 10))]
        self.assertEqual(e[10:], ['Без вердикта / Ҳукмсиз',
                                  'Машины нет в реестре agro-work (или она не связана) / '
                                  'Машина agro-work реестрида йўқ (ёки боғланмаган)', None])

    def test_summary_and_console(self):
        values = {row[0]: row[1] for row in sheet(0)[1:]}
        self.assertEqual(values['Сутки сверки (сутки реестра ПК-17) / Солиштириш кунлари '
                                '(ПК-17 реестри кунлари)'], '10.10.2026 — 11.10.2026, 2')
        self.assertEqual(values['Из них — наших машин / Шулардан — бизнинг машиналар'], 10)
        self.assertEqual(values['Есть рейс в 1С / 1С да рейс бор'], 3)
        self.assertEqual(values['Нет рейса в 1С / 1С да рейс йўқ'], 1)
        self.assertEqual(values['Машина не сопоставлена / Машина боғланмаган'], 1)
        self.assertEqual(values['Заявок «пахта/тола ташиш» без окна — сутки не судятся / '
                                '«пахта/тола ташиш» ойнасиз буюртмалари — кунлари '
                                'баҳоланмайди'], 2)
        self.assertEqual(values['— Заявка открыта — вердикт после закрытия / '
                                '— Буюртма очиқ — ҳукм ёпилгандан кейин'], 1)
        self.assertEqual(values['— Заявка отменена — работу не покрывает / '
                                '— Буюртма бекор қилинган — ишни қопламайди'], 1)
        self.assertEqual(values['Машино-суток рейсов 1С, не покрытых заявкой / Буюртма '
                                'билан қопланмаган 1С рейслари машина-кунлари'], 3)
        self.assertEqual(values['Работа без заявки / Буюртмасиз иш'], 1)
        self.assertEqual(values['Без вердикта / Ҳукмсиз'], 2)
        text = WORLD['text']
        self.assertIn('trips read: 13 (removal 9, receiving 4) | ours 10 | not ours 2 | '
                      'several machines 1 | no time 0', text)
        self.assertIn('cotton hauling application-days 5 | with a 1C trip 3 | without 1 | '
                      'machine not matched 1 | applications without a window 2 '
                      '(otkryta 1, otmenena 1)', text)
        self.assertIn('machine-days with 1C trips not covered by an application: 3 | '
                      'without an application 1 | no verdict 2 '
                      '(net_v_agro_work 1, otkrytaya_zayavka 1)', text)
        self.assertTrue(text.isascii())

    def test_nothing_is_written_but_the_book(self):
        for path, value in WORLD['before'].items():
            self.assertEqual(digest(path), value, path)
        names = sorted(os.listdir(WORLD['root']))
        self.assertIn('report.xlsx', names)
        for row in sheet(1) + sheet(2):
            self.assertFalse(any('farmer' in str(cell) for cell in row))


class Parts(unittest.TestCase):

    def test_parse_mass(self):
        self.assertEqual(report.parse_mass('7 500'), 7500.0)
        self.assertEqual(report.parse_mass('2000,5'), 2000.5)
        self.assertEqual(report.parse_mass('1\xa0234'), 1234.0)
        self.assertEqual(report.parse_mass(None), 0.0)
        self.assertEqual(report.parse_mass('—'), 0.0)


class Refusals(unittest.TestCase):

    def base(self, **changes):
        args = dict(zip(WORLD['argv'][0::2], WORLD['argv'][1::2]))
        args['--out'] = os.path.join(WORLD['root'], 'refused.xlsx')
        args.update(changes)
        argv = []
        for key, value in args.items():
            argv += [key, value]
        return run_main(*argv)

    def test_bad_input_is_code_two_and_nothing_is_written(self):
        no_mass = os.path.join(WORLD['root'], 'no_mass.xlsx')
        write_xlsx(no_mass, [VYVOZ_HEADER[:4], ('Z', 'p', '80 101 HA', '2026-10-10 08:00:00')])
        no_agro = os.path.join(WORLD['root'], 'no_agro.db')
        con = sqlite3.connect(no_agro)
        con.execute('CREATE TABLE equipment (id INTEGER PRIMARY KEY)')
        con.close()
        for changes in ({'--db': os.path.join(WORLD['root'], 'nope.db')},
                        {'--db': no_agro},
                        {'--vyvoz': os.path.join(WORLD['root'], 'nofile.xlsx')},
                        {'--vyvoz': no_mass},
                        {'--out': os.path.join(WORLD['root'], 'report.csv')},
                        {'--out': os.path.join(WORLD['root'], 'nodir', 'a.xlsx')},
                        {'--out': WORLD['vyvoz']}):
            with self.subTest(changes=changes):
                code, text = self.base(**changes)
                self.assertEqual(code, report.EXIT_BAD_INPUT, text)
                self.assertFalse(os.path.exists(os.path.join(WORLD['root'], 'refused.xlsx')))

    def test_no_files_and_no_dated_trips_are_refused(self):
        code, text = run_main('--db', WORLD['db'], '--out',
                              os.path.join(WORLD['root'], 'x.xlsx'))
        self.assertEqual(code, report.EXIT_BAD_INPUT)
        undated = os.path.join(WORLD['root'], 'undated.xlsx')
        write_xlsx(undated, [REESTR_HEADER, ('p', 'f', '80 HA 101', 'нет', '1', '1')])
        code, text = run_main('--db', WORLD['db'], '--reestr', undated, '--out',
                              os.path.join(WORLD['root'], 'y.xlsx'))
        self.assertEqual(code, report.EXIT_BAD_INPUT, text)
        self.assertIn('no dated trips', text)


class Imports(unittest.TestCase):

    def test_flask_is_not_imported(self):
        code = ('import sys; import tools.gps_onec_hauling_report; '
                'print("flask" in sys.modules, "models" in sys.modules)')
        out = subprocess.run([sys.executable, '-B', '-c', code], cwd=REPO,
                             capture_output=True, text=True, check=True).stdout
        self.assertEqual(out.strip(), 'False False')


if __name__ == '__main__':
    unittest.main()
