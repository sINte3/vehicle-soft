# -*- coding: utf-8 -*-
"""DRONE-AREA-CONTROL-V2-MEGA: хранилище решений и журнала циклов, миграция,
блокировка и чистый слой отчёта -- без Flask.

Что держится здесь (stdlib + openpyxl, идёт в CI):

* миграция DRONE_AREA_CONTROL_V2_001 по четырём путям: чистая база, повтор
  («Already applied»), базы нет (код 2, файл не создан), предусловие не
  пройдено (код 1, полный откат, в реестре пусто); триггеры делают историю
  решений append-only на уровне базы; DDL совпадает с ORM (локально);
* правило решений: кто что может, отказ с кодом, эффективный результат,
  «принять автоматический» теряет силу при смене расчёта;
* писатель решений: цепочка chain_seq/supersedes, отмена -- новая строка,
  правка поверх чужой -- отказ, отказ ничего не пишет, автоматический
  расчёт не меняется ни в одной колонке;
* журнал циклов: одна очередь на ручной прогон, захват, закрытие,
  «мёртвый» прогон, редакция секретов;
* блокировка: межпроцессная, не реентерабельна, отпускается;
* отчёт: эффективные итоги, дерево Дрон -> День -> Вылет складывается в
  родителя, вкладки фильтруют по эффективному состоянию, время UTC+5.

ВСЕ ДАННЫЕ СИНТЕТИЧЕСКИЕ.
"""

import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from datetime import date, datetime, timedelta

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

import migrate_dji_area_evidence_001 as evidence_mig  # noqa: E402
import migrate_drone_area_control_v2_001 as mig  # noqa: E402
import migration_utils  # noqa: E402

import dji_area  # noqa: E402
from dji_area import accounting as acc  # noqa: E402
from dji_area import control_report as cr  # noqa: E402
from dji_area import control_store as cs  # noqa: E402
from dji_area import decisions as dec  # noqa: E402
from dji_area import resolver as rs  # noqa: E402
from dji_area import store  # noqa: E402
from drone_collector import runlock  # noqa: E402

UAT_C = 900714


def _index_shape(path, table):
    """Индексы таблицы как сравнимые описания, без имён автоиндексов.

    Для каждого индекса: (имя или '<auto>', уникален, частичный, столбцы,
    текст условия WHERE). Имя автоиндекса SQLite (sqlite_autoindex_*)
    зависит от порядка ограничений и не несёт смысла -- сравнивается его
    форма, а не имя."""
    shapes = set()
    for _seq, name, unique, _origin, partial in _query(
            path, 'PRAGMA index_list(%s)' % table):
        columns = tuple(r[2] for r in _query(path, 'PRAGMA index_info(%s)'
                                             % name))
        sql = _query(path, "SELECT sql FROM sqlite_master WHERE type='index' "
                           "AND name=?", (name,))
        where = ''
        if sql and sql[0][0] and ' WHERE ' in sql[0][0].upper():
            text = sql[0][0]
            where = ' '.join(text[text.upper().index(' WHERE ') + 7:]
                             .replace('(', ' ').replace(')', ' ')
                             .split()).lower()
        label = '<auto>' if name.startswith('sqlite_autoindex_') else name
        shapes.add((label, bool(unique), bool(partial), columns, where))
    return shapes


def _query(path, sql, params=()):
    con = sqlite3.connect(path)
    try:
        return con.execute(sql, params).fetchall()
    finally:
        con.close()


def _make_db(path, with_calcs=True, with_users=True):
    """База, на которой миграция имеет право бежать."""
    con = sqlite3.connect(path)
    try:
        if with_users:
            con.execute('CREATE TABLE users (id INTEGER PRIMARY KEY, '
                        'username TEXT)')
        con.execute('CREATE TABLE drone_flights (id INTEGER PRIMARY KEY, '
                    'dji_flight_id BIGINT, area_ha FLOAT)')
        con.execute("INSERT INTO drone_flights VALUES (1, %d, 6.0)" % UAT_C)
        if with_calcs:
            for _name, ddl in evidence_mig.TABLES:
                con.execute(ddl)
        con.commit()
    finally:
        con.close()


class MigrationBase(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='area_v2_mig_')
        self.db = os.path.join(self.tmp, 'transport.db')
        self._orig = (mig.DB_PATH, migration_utils.DB_PATH)
        mig.DB_PATH = self.db
        migration_utils.DB_PATH = self.db

    def tearDown(self):
        mig.DB_PATH, migration_utils.DB_PATH = self._orig
        shutil.rmtree(self.tmp, ignore_errors=True)

    def run_migration(self):
        try:
            mig.run()
            return 0
        except SystemExit as exc:
            return exc.code


class MigrationPaths(MigrationBase):

    def test_path_1_clean_database_creates_everything_once(self):
        _make_db(self.db)
        self.assertEqual(self.run_migration(), 0)
        names = {r[0] for r in _query(self.db, 'SELECT name FROM '
                                               'sqlite_master')}
        for table in mig.TABLE_NAMES:
            self.assertIn(table, names)
        for name, _sql in mig.INDEXES + mig.TRIGGERS:
            self.assertIn(name, names)
        expected = mig.expected_columns()
        for table, columns in expected.items():
            present = [r[1] for r in _query(self.db,
                                            'PRAGMA table_info(%s)' % table)]
            self.assertEqual(sorted(present), sorted(columns), table)
        self.assertEqual(_query(self.db, 'SELECT COUNT(*) FROM '
                                         'schema_migrations WHERE name=?',
                                (mig.MIGRATION_ID,))[0][0], 1)
        # drone_flights не тронута.
        self.assertEqual(_query(self.db, 'SELECT area_ha, typeof(area_ha) '
                                         'FROM drone_flights'),
                         [(6.0, 'real')])

    def test_path_2_repeat_says_already_applied(self):
        _make_db(self.db)
        self.assertEqual(self.run_migration(), 0)
        buffer = io.StringIO()
        stdout, sys.stdout = sys.stdout, buffer
        try:
            self.assertEqual(self.run_migration(), 0)
        finally:
            sys.stdout = stdout
        self.assertIn('Already applied', buffer.getvalue())
        self.assertEqual(_query(self.db, 'SELECT COUNT(*) FROM '
                                         'schema_migrations')[0][0], 1)

    def test_path_3_missing_database_is_code_2_and_creates_nothing(self):
        self.assertEqual(self.run_migration(), 2)
        self.assertFalse(os.path.exists(self.db))

    def test_path_4_failed_precondition_rolls_everything_back(self):
        _make_db(self.db, with_calcs=False)
        self.assertEqual(self.run_migration(), 1)
        names = {r[0] for r in _query(self.db, 'SELECT name FROM '
                                               'sqlite_master')}
        for table in mig.TABLE_NAMES:
            self.assertNotIn(table, names)
        self.assertEqual(_query(self.db, 'SELECT COUNT(*) FROM '
                                         'schema_migrations')[0][0], 0)
        # Отрицательный контроль: та же база с предусловием -- проходит.
        con = sqlite3.connect(self.db)
        for _name, ddl in evidence_mig.TABLES:
            con.execute(ddl)
        con.commit()
        con.close()
        self.assertEqual(self.run_migration(), 0)

    def test_a_registry_failure_rolls_back_the_tables(self):
        _make_db(self.db)
        original = mig._record_in_transaction

        def refuse(*_args):
            raise RuntimeError('SYNTHETIC registry failure')
        mig._record_in_transaction = refuse
        try:
            self.assertEqual(self.run_migration(), 1)
        finally:
            mig._record_in_transaction = original
        names = {r[0] for r in _query(self.db, 'SELECT name FROM '
                                               'sqlite_master')}
        self.assertFalse(set(mig.TABLE_NAMES) & names)

    def test_the_triggers_make_decisions_append_only(self):
        _make_db(self.db)
        self.run_migration()
        con = sqlite3.connect(self.db)
        con.execute("INSERT INTO drone_area_decisions (flight_id, chain_seq, "
                    "decision_type, is_override, comment, performed_at, "
                    "decisions_version) VALUES (1, 1, 'KEEP_DJI_RAW', 0, "
                    "'SYNTHETIC', '2026-09-23 10:00:00', 'v')")
        con.commit()
        for sql in ("UPDATE drone_area_decisions SET comment='x'",
                    'DELETE FROM drone_area_decisions'):
            with self.assertRaises(sqlite3.DatabaseError):
                con.execute(sql)
        con.close()

    def test_the_chain_is_unique_per_flight(self):
        _make_db(self.db)
        self.run_migration()
        con = sqlite3.connect(self.db)
        insert = ("INSERT INTO drone_area_decisions (flight_id, chain_seq, "
                  "decision_type, is_override, comment, performed_at, "
                  "decisions_version) VALUES (7, 1, 'KEEP_DJI_RAW', 0, 'c', "
                  "'2026-09-23 10:00:00', 'v')")
        con.execute(insert)
        with self.assertRaises(sqlite3.IntegrityError):
            con.execute(insert)
        con.close()

    def test_one_active_manual_run_at_the_database_level(self):
        _make_db(self.db)
        self.run_migration()
        con = sqlite3.connect(self.db)
        insert = ("INSERT INTO drone_area_cycle_runs (trigger_kind, status, "
                  "active_slot, requested_at) VALUES ('MANUAL', 'QUEUED', 1, "
                  "'2026-09-23 10:00:00')")
        con.execute(insert)
        with self.assertRaises(sqlite3.IntegrityError):
            con.execute(insert)
        # Закрытые строки слот не держат.
        con.execute("INSERT INTO drone_area_cycle_runs (trigger_kind, status, "
                    "active_slot, requested_at) VALUES ('MANUAL', 'SUCCESS', "
                    "NULL, '2026-09-23 09:00:00')")
        con.close()

    def test_migration_ddl_matches_the_orm_models(self):
        try:
            from sqlalchemy import create_engine
            import models as m  # noqa: E402
        except ImportError:  # pragma: no cover - CI has no SQLAlchemy
            self.skipTest('SQLAlchemy/Flask not installed: ORM parity runs '
                          'locally with the application dependencies')
        _make_db(self.db)
        self.run_migration()
        orm_db = os.path.join(self.tmp, 'orm.db')
        engine = create_engine('sqlite:///' + orm_db)
        wanted = [m.DroneAreaDecision, m.DroneAreaCycleRun]
        m.db.metadata.create_all(engine, tables=[w.__table__ for w in wanted]
                                 + [m.User.__table__,
                                    m.Organization.__table__])
        engine.dispose()
        for model in wanted:
            table = model.__tablename__
            migrated = {r[1]: (r[2].upper(), r[3]) for r in _query(
                self.db, 'PRAGMA table_info(%s)' % table)}
            fresh = {r[1]: (r[2].upper(), r[3]) for r in _query(
                orm_db, 'PRAGMA table_info(%s)' % table)}
            self.assertEqual(migrated, fresh, table)
            # Индексы сверяются по форме: уникальность, частичность, столбцы
            # и условие -- не только по префиксу имени. Иначе ORM-база с
            # неуникальным ux_... или без WHERE прошла бы проверку.
            self.assertEqual(_index_shape(self.db, table),
                             _index_shape(orm_db, table), table)
        # Отрицательный контроль: индекс той же таблицы и того же имени, но
        # без частичного условия, различается сравнением.
        broken = os.path.join(self.tmp, 'broken.db')
        shutil.copyfile(self.db, broken)
        con = sqlite3.connect(broken)
        con.execute('DROP INDEX ux_drone_area_cycle_runs_active')
        con.execute('CREATE UNIQUE INDEX ux_drone_area_cycle_runs_active '
                    'ON drone_area_cycle_runs (active_slot)')
        con.commit()
        con.close()
        self.assertNotEqual(_index_shape(broken, 'drone_area_cycle_runs'),
                            _index_shape(orm_db, 'drone_area_cycle_runs'))


# ─── Правило решений (чистое) ─────────────────────────────────────────────

class DecisionRules(unittest.TestCase):

    def check(self, action, cls=dec.AUTO_REVIEW, raw=5000.0, active=None,
              comment='проверено вручную', confirmed=True, override=False,
              expected=None):
        return dec.validate(action, cls, raw, active, comment, confirmed,
                            override, expected)

    def test_a_review_accepts_every_decision(self):
        for action in dec.DECISION_TYPES:
            self.assertIsNone(self.check(action), action)

    def test_normal_records_are_not_decidable(self):
        self.assertEqual(self.check(dec.CONFIRM_FULL_PHANTOM,
                                    cls=dec.AUTO_NORMAL), dec.E_NOT_DECIDABLE)

    def test_a_proven_correction_needs_the_override_confirmation(self):
        self.assertEqual(self.check(dec.KEEP_DJI_RAW, cls=dec.AUTO_PROVEN),
                         dec.E_OVERRIDE_NOT_CONFIRMED)
        self.assertIsNone(self.check(dec.KEEP_DJI_RAW, cls=dec.AUTO_PROVEN,
                                     override=True))

    def test_comment_and_confirmation_are_required(self):
        self.assertEqual(self.check(dec.KEEP_DJI_RAW, comment='  ok '),
                         dec.E_COMMENT_REQUIRED)
        self.assertEqual(self.check(dec.KEEP_DJI_RAW, comment='x' * 2001),
                         dec.E_COMMENT_TOO_LONG)
        self.assertEqual(self.check(dec.KEEP_DJI_RAW, confirmed=False),
                         dec.E_NOT_CONFIRMED)

    def test_stale_forms_and_empty_revokes_are_refused(self):
        active = {'id': 5, 'decision_type': dec.KEEP_DJI_RAW}
        self.assertEqual(self.check(dec.CONFIRM_FULL_PHANTOM, active=active),
                         dec.E_STALE_FORM)
        self.assertIsNone(self.check(dec.CONFIRM_FULL_PHANTOM, active=active,
                                     expected=5))
        self.assertEqual(self.check(dec.KEEP_DJI_RAW, active=active,
                                    expected=5), dec.E_SAME_AS_ACTIVE)
        self.assertEqual(self.check(dec.REVOKE), dec.E_NOTHING_TO_REVOKE)
        self.assertEqual(self.check('DROP_TABLE'), dec.E_UNKNOWN_ACTION)

    def test_no_raw_no_phantom_and_no_keep(self):
        self.assertEqual(self.check(dec.CONFIRM_FULL_PHANTOM, raw=None),
                         dec.E_RAW_MISSING)
        self.assertIsNone(self.check(dec.NEEDS_MORE_EVIDENCE, raw=None))

    def test_the_effective_result_of_each_decision(self):
        eff = dec.effective
        self.assertEqual(eff(dec.AUTO_REVIEW, 5000.0, 5000.0, 0.0,
                             {'decision_type': dec.CONFIRM_FULL_PHANTOM}),
                         (dec.S_ADMIN_PHANTOM, 0.0, 5000.0, True))
        self.assertEqual(eff(dec.AUTO_PROVEN, 8000.0, 5.0, 7995.0,
                             {'decision_type': dec.KEEP_DJI_RAW}),
                         (dec.S_ADMIN_KEPT_RAW, 8000.0, 0.0, True))
        self.assertEqual(eff(dec.AUTO_PROVEN, 8000.0, 5.0, 7995.0,
                             {'decision_type': dec.ACCEPT_AUTO_RESULT}),
                         (dec.S_CORRECTED, 5.0, 7995.0, True))
        self.assertEqual(eff(dec.AUTO_REVIEW, 5000.0, 5000.0, 0.0,
                             {'decision_type': dec.ACCEPT_AUTO_RESULT}),
                         (dec.S_ADMIN_ACCEPTED_AUTO, 5000.0, 0.0, True))
        # «Нужны доказательства» по доказанной корректировке -- RAW, открыто.
        self.assertEqual(eff(dec.AUTO_PROVEN, 8000.0, 5.0, 7995.0,
                             {'decision_type': dec.NEEDS_MORE_EVIDENCE}),
                         (dec.S_ADMIN_NEEDS_EVIDENCE, 8000.0, 0.0, True))
        self.assertIn(dec.S_ADMIN_NEEDS_EVIDENCE, dec.OPEN_STATES)

    def test_a_record_that_turned_normal_can_only_be_revoked(self):
        # Пересчёт перевёл запись в обычные (V4 пришёл и опроверг кандидата):
        # действующее решение снимается, новое не принимается.
        active = {'id': 7, 'decision_type': dec.CONFIRM_FULL_PHANTOM}
        self.assertEqual(dec.allowed_actions(dec.AUTO_NORMAL, 5000.0, active),
                         (dec.REVOKE,))
        self.assertEqual(dec.allowed_actions(dec.AUTO_NORMAL, 5000.0, None),
                         ())
        self.assertIsNone(self.check(dec.REVOKE, cls=dec.AUTO_NORMAL,
                                     active=active, expected=7))
        for action in dec.DECISION_TYPES:
            self.assertEqual(self.check(action, cls=dec.AUTO_NORMAL,
                                        active=active, expected=7),
                             dec.E_NOT_DECIDABLE, action)
        # Без действующего решения отменять нечего и у обычной записи.
        self.assertEqual(self.check(dec.REVOKE, cls=dec.AUTO_NORMAL),
                         dec.E_NOT_DECIDABLE)

    def test_the_same_decision_is_confirmed_again_only_after_a_recalc(self):
        active = {'id': 7, 'decision_type': dec.ACCEPT_AUTO_RESULT}
        self.assertEqual(self.check(dec.ACCEPT_AUTO_RESULT, active=active,
                                    expected=7), dec.E_SAME_AS_ACTIVE)
        self.assertIsNone(dec.validate(
            dec.ACCEPT_AUTO_RESULT, dec.AUTO_REVIEW, 5000.0, active,
            'проверено вручную', True, False, 7, active_stale=True))

    def test_accept_auto_lapses_when_the_calculation_changes(self):
        self.assertEqual(dec.effective(
            dec.AUTO_REVIEW, 5000.0, 5000.0, 0.0,
            {'decision_type': dec.ACCEPT_AUTO_RESULT}, stale=True),
            (dec.S_NEEDS_DECISION, 5000.0, 0.0, False))
        # Полный фантом от автомата не зависит: действует и при смене.
        self.assertEqual(dec.effective(
            dec.AUTO_REVIEW, 5000.0, 5000.0, 0.0,
            {'decision_type': dec.CONFIRM_FULL_PHANTOM}, stale=True)[3],
            True)


# ─── Писатель решений и журнала ───────────────────────────────────────────

def calc_row(flight_id, status, eligibility, raw, corrected=None,
             candidate=False, flags=(), day=date(2026, 9, 21), minute=0,
             input_hash='H1', base=None, bridges=()):
    start = datetime(day.year, day.month, day.day, 3, 0) + timedelta(
        minutes=minute)
    return {
        'flight_id': flight_id, 'provider_account_id': 'SYNTHETIC',
        'area_algorithm_version': dji_area.AREA_ALGORITHM_VERSION,
        'calculation_input_hash': input_hash,
        'start_at_utc': start, 'end_at_utc': start + timedelta(minutes=9),
        'report_timezone': dji_area.REPORT_TIMEZONE,
        'report_start_date': day.isoformat(),
        'raw_area_m2': raw, 'corrected_recorded_area_m2': corrected,
        'area_status': status, 'aggregation_eligibility': eligibility,
        'anomaly_flags_json': json.dumps(list(flags)),
        'structural_candidate': candidate, 'scalar_source_check': candidate,
        'candidate_base_flight_id': base,
        'bridge_flight_ids_json': json.dumps(list(bridges)),
    }


class StoreBase(MigrationBase):

    def setUp(self):
        super(StoreBase, self).setUp()
        _make_db(self.db)
        self.assertEqual(self.run_migration(), 0)
        self.con = store.connect(self.db)
        self.addCleanup(self.con.close)

    def add_calc(self, row):
        store.begin_immediate(self.con)
        result = store.insert_calculation(self.con, row)
        self.con.execute('COMMIT')
        return result

    def calcs(self):
        return self.con.execute('SELECT * FROM dji_area_calculations '
                                'ORDER BY id').fetchall()


class DecisionWriter(StoreBase):

    def setUp(self):
        super(DecisionWriter, self).setUp()
        # Кейс UAT: плоский счётчик при наблюдённом распылении -> REVIEW.
        self.add_calc(calc_row(UAT_C, rs.COUNTER_FLAT_RAW_OVERSTATED,
                               rs.AGG_UNRESOLVED, 60000.0, corrected=0.0,
                               candidate=True,
                               flags=['APPLICATION_WITH_FLAT_COUNTER'],
                               base=UAT_C - 3, bridges=[UAT_C - 2]))

    def record(self, action, comment='визуально проверено в DJI',
               expected=None, confirmed=True, override=False):
        return cs.record_decision(self.con, UAT_C, action, comment, confirmed,
                                  override, expected, 1, 'SYNTHETIC admin')

    def test_the_uat_case_becomes_a_full_phantom_without_touching_the_calc(
            self):
        before = [tuple(r) for r in self.calcs()]
        saved = self.record(dec.CONFIRM_FULL_PHANTOM)
        self.assertEqual(saved['auto_class'], acc.REVIEW)
        self.assertEqual(saved['state'], dec.S_ADMIN_PHANTOM)
        self.assertEqual((saved['effective_accepted_m2'],
                          saved['effective_excluded_m2']), (0.0, 60000.0))
        self.assertEqual(saved['calculation_input_hash'], 'H1')
        self.assertEqual([tuple(r) for r in self.calcs()], before)
        active = cs.active_decisions(self.con, [UAT_C])
        self.assertEqual(active[UAT_C]['decision_type'],
                         dec.CONFIRM_FULL_PHANTOM)

    def test_the_chain_supersedes_and_revokes_append_only(self):
        first = self.record(dec.CONFIRM_FULL_PHANTOM)
        second = self.record(dec.KEEP_DJI_RAW, expected=first['id'])
        third = self.record(dec.REVOKE, expected=second['id'])
        chain = cs.decision_chains(self.con, [UAT_C])[UAT_C]
        self.assertEqual([c['chain_seq'] for c in chain], [1, 2, 3])
        self.assertEqual([c['supersedes_decision_id'] for c in chain],
                         [None, first['id'], second['id']])
        self.assertEqual(cs.active_decisions(self.con, [UAT_C]), {})
        history = cs.history(self.con, [UAT_C])
        self.assertEqual([h['is_current'] for h in history],
                         [False, False, False])
        self.assertEqual(third['decision_type'], dec.REVOKE)

    def test_a_refusal_writes_nothing(self):
        with self.assertRaises(cs.DecisionRefused) as ctx:
            self.record(dec.CONFIRM_FULL_PHANTOM, comment='no')
        self.assertEqual(ctx.exception.code, dec.E_COMMENT_REQUIRED)
        self.assertEqual(self.con.execute(
            'SELECT COUNT(*) FROM drone_area_decisions').fetchone()[0], 0)
        # Транзакция закрыта: следующий писатель не упирается в неё.
        self.record(dec.CONFIRM_FULL_PHANTOM)

    def test_an_edit_over_a_newer_decision_is_refused(self):
        self.record(dec.CONFIRM_FULL_PHANTOM)
        with self.assertRaises(cs.DecisionRefused) as ctx:
            self.record(dec.KEEP_DJI_RAW, expected=None)
        self.assertEqual(ctx.exception.code, dec.E_STALE_FORM)

    def test_a_flight_without_a_current_calculation_is_refused(self):
        with self.assertRaises(cs.DecisionRefused):
            cs.record_decision(self.con, 123, dec.KEEP_DJI_RAW,
                               'синтетическая причина', True, False, None, 1,
                               'x')

    def recalc(self, status=rs.COUNTER_FLAT_RAW_OVERSTATED,
               eligibility=rs.AGG_UNRESOLVED, corrected=0.0, candidate=True,
               input_hash='H2', flags=('APPLICATION_WITH_FLAT_COUNTER',)):
        self.add_calc(calc_row(UAT_C, status, eligibility, 60000.0,
                               corrected=corrected, candidate=candidate,
                               input_hash=input_hash, flags=flags))
        return cs.current_calculation(self.con, UAT_C)

    def view(self):
        active = cs.active_decisions(self.con, [UAT_C]).get(UAT_C)
        return cr.record_view(dict(cs.current_calculation(self.con, UAT_C),
                                   machine_key='M', machine_label='M'),
                              'ru', decision=active)

    def count(self):
        return self.con.execute(
            'SELECT COUNT(*) FROM drone_area_decisions').fetchone()[0]

    def test_a_decision_is_revocable_after_the_record_turned_normal(self):
        saved = self.record(dec.CONFIRM_FULL_PHANTOM)
        self.recalc(rs.RAW_CORROBORATED, rs.AGG_CERTIFIED, corrected=60000.0,
                    candidate=False, input_hash='H3', flags=())
        self.assertEqual(self.view()['accounting_class'], dec.AUTO_NORMAL)
        with self.assertRaises(cs.DecisionRefused) as ctx:
            self.record(dec.KEEP_DJI_RAW, expected=saved['id'])
        self.assertEqual(ctx.exception.code, dec.E_NOT_DECIDABLE)
        self.assertEqual(self.count(), 1)
        revoked = self.record(dec.REVOKE, expected=saved['id'])
        self.assertEqual((revoked['decision_type'], revoked['chain_seq']),
                         (dec.REVOKE, 2))
        self.assertEqual(cs.active_decisions(self.con, [UAT_C]), {})

    def test_a_form_opened_before_a_recalculation_is_refused(self):
        seen = cs.current_calculation(self.con, UAT_C)['id']
        now = self.recalc()
        self.assertNotEqual(seen, now['id'])
        with self.assertRaises(cs.DecisionRefused) as ctx:
            cs.record_decision(self.con, UAT_C, dec.CONFIRM_FULL_PHANTOM,
                               'визуально проверено в DJI', True, False, None,
                               1, 'SYNTHETIC admin', expected_calc_id=seen)
        self.assertEqual(ctx.exception.code, dec.E_STALE_CALC)
        self.assertEqual(self.count(), 0)
        # Отрицательный контроль: форма, открытая на нынешнем расчёте.
        saved = cs.record_decision(self.con, UAT_C, dec.CONFIRM_FULL_PHANTOM,
                                   'визуально проверено в DJI', True, False,
                                   None, 1, 'SYNTHETIC admin',
                                   expected_calc_id=now['id'])
        self.assertEqual(saved['calculation_input_hash'], 'H2')

    def test_a_lapsed_accept_is_confirmed_against_the_new_calculation(self):
        first = self.record(dec.ACCEPT_AUTO_RESULT)
        with self.assertRaises(cs.DecisionRefused) as ctx:
            self.record(dec.ACCEPT_AUTO_RESULT, expected=first['id'])
        self.assertEqual(ctx.exception.code, dec.E_SAME_AS_ACTIVE)
        self.recalc()
        self.assertFalse(self.view()['decision']['applied'])
        second = self.record(dec.ACCEPT_AUTO_RESULT, expected=first['id'])
        self.assertEqual((second['chain_seq'],
                          second['calculation_input_hash']), (2, 'H2'))
        item = self.view()
        self.assertTrue(item['decision']['applied'])
        self.assertFalse(item['decision']['stale'])

    def test_the_decision_follows_a_recalculation_as_stale(self):
        saved = self.record(dec.ACCEPT_AUTO_RESULT)
        # Новый вход -> новая строка расчёта: решение «принять автомат»
        # относилось к прежней и перестаёт действовать.
        self.add_calc(calc_row(UAT_C, rs.COUNTER_FLAT_RAW_OVERSTATED,
                               rs.AGG_UNRESOLVED, 60000.0, corrected=0.0,
                               candidate=True, input_hash='H2',
                               flags=['APPLICATION_WITH_FLAT_COUNTER']))
        current = cs.current_calculation(self.con, UAT_C)
        self.assertTrue(dec.decision_is_stale(saved, current))
        item = cr.record_view(dict(current, machine_key='M',
                                   machine_label='M'), 'ru',
                              decision=cs.active_decisions(
                                  self.con, [UAT_C])[UAT_C])
        self.assertEqual(item['state'], dec.S_NEEDS_DECISION)
        self.assertFalse(item['decision']['applied'])
        self.assertTrue(item['decision']['stale'])


class CycleLedger(StoreBase):

    def test_one_manual_run_at_a_time(self):
        run = cs.enqueue_manual(self.con, 1, 'SYNTHETIC admin')
        self.assertEqual((run['status'], run['active_slot']),
                         (cs.STATUS_QUEUED, 1))
        with self.assertRaises(cs.RunActive) as ctx:
            cs.enqueue_manual(self.con, 2, 'someone else')
        self.assertEqual(ctx.exception.run['id'], run['id'])

    def test_a_scheduled_run_also_blocks_a_manual_request(self):
        cs.start_scheduled(self.con, date(2026, 9, 21), date(2026, 9, 23))
        with self.assertRaises(cs.RunActive):
            cs.enqueue_manual(self.con, 1, 'SYNTHETIC admin')

    def test_claim_step_finish(self):
        cs.enqueue_manual(self.con, 1, 'SYNTHETIC admin')
        run = cs.claim_queued(self.con, pid=4242, host='SYNTHETIC-HOST')
        self.assertEqual(run['status'], cs.STATUS_RUNNING)
        self.assertIsNone(cs.claim_queued(self.con))
        cs.set_window(self.con, run['id'], date(2026, 9, 21),
                      date(2026, 9, 23))
        cs.set_step(self.con, run['id'], cs.STEP_SOURCES)
        cs.finish(self.con, run['id'], cs.STATUS_WARNINGS, exit_code=0,
                  result={'outcome': 'SUCCESS_WITH_WARNINGS'})
        done = cs.get_run(self.con, run['id'])
        self.assertEqual((done['status'], done['active_slot'],
                          done['current_step']),
                         (cs.STATUS_WARNINGS, None, cs.STEP_SOURCES))
        self.assertEqual(done['result']['outcome'], 'SUCCESS_WITH_WARNINGS')
        self.assertEqual(cs.last_success(self.con)['id'], run['id'])
        # Слот свободен -- следующий ручной прогон принимается.
        cs.enqueue_manual(self.con, 1, 'SYNTHETIC admin')
        with self.assertRaises(cs.ControlStoreError):
            cs.finish(self.con, run['id'], cs.STATUS_RUNNING)

    def test_a_dead_run_is_shown_and_then_closed(self):
        cs.enqueue_manual(self.con, 1, 'SYNTHETIC admin')
        run = cs.claim_queued(self.con)
        self.assertEqual(cs.effective_status(run, lock_held=True),
                         cs.STATUS_RUNNING)
        self.assertEqual(cs.effective_status(run, lock_held=False),
                         cs.STATUS_INTERRUPTED)
        self.assertEqual(cs.reconcile(self.con, lock_held=True), [])
        self.assertEqual(cs.reconcile(self.con, lock_held=False),
                         [run['id']])
        self.assertEqual(cs.get_run(self.con, run['id'])['status'],
                         cs.STATUS_INTERRUPTED)

    def test_a_queued_run_nobody_took_becomes_a_failed_launch(self):
        now = datetime(2026, 9, 23, 10, 0)
        run = cs.enqueue_manual(self.con, 1, 'SYNTHETIC admin', now=now)
        later = now + cs.QUEUED_GRACE + timedelta(minutes=1)
        self.assertEqual(cs.effective_status(run, False, now=now),
                         cs.STATUS_QUEUED)
        self.assertEqual(cs.effective_status(run, False, now=later),
                         cs.STATUS_LAUNCH_FAILED)
        # Пока блокировку держит другой цикл, очередь -- живая.
        self.assertEqual(cs.effective_status(run, True, now=later),
                         cs.STATUS_QUEUED)

    def test_secrets_never_reach_the_message(self):
        secret = 'SYNTHETIC-TOKEN-VALUE-987654'
        text = 'failed: token=%s, header "Cookie: %s" and raw %s' % (
            secret, secret, secret)
        clean = cs.redact(text, secrets=[secret])
        self.assertNotIn(secret, clean)
        self.assertIn('***', clean)
        env_clean = cs.redact('x ' + secret,
                              secrets=cs.secret_values(
                                  {'DRONE_API_TOKEN': secret, 'PATH': 'C:/'}))
        self.assertNotIn(secret, env_clean)
        # Отрицательный контроль: без знания секрета голое значение осталось бы.
        self.assertIn(secret, 'x ' + secret)

    def test_the_lock_probe_runs_inside_the_writer_transaction(self):
        cs.enqueue_manual(self.con, 1, 'SYNTHETIC admin')
        run = cs.claim_queued(self.con)
        seen = []

        def held():
            seen.append(self.con.in_transaction)
            return True

        def free():
            seen.append(self.con.in_transaction)
            return False
        self.assertEqual(cs.reconcile(self.con, held), [])
        self.assertEqual(cs.reconcile(self.con, free), [run['id']])
        # Проба -- под транзакцией писателя, не до неё.
        self.assertEqual(seen, [True, True])
        self.assertFalse(self.con.in_transaction)

    def test_run_explanations_are_words_in_both_languages(self):
        result = {'failure': cs.FAILURE_CANDIDATE_EVIDENCE,
                  'evidence_misses': {'candidates': [11, 12],
                                      'controls': [13]}}
        self.assertIn('кандидатов: 2', cs.run_explanation(
            cs.STATUS_FAILED, result, '', 'ru'))
        uz = cs.run_explanation(cs.STATUS_FAILED, result, '', 'uz')
        self.assertIn(': 2', uz)
        self.assertNotIn('CANDIDATE', uz)
        warned = cs.run_explanation(
            cs.STATUS_WARNINGS, {'warnings': [cs.WARNING_CONTROL_EVIDENCE],
                                 'evidence_misses': {'controls': [13]}},
            '', 'ru')
        self.assertTrue(warned)
        self.assertNotIn('CONTROL_EVIDENCE', warned)
        # Незнакомый код -- общая фраза, а не сам код.
        unknown = cs.run_explanation(cs.STATUS_FAILED,
                                     {'failure': 'SOMETHING_NEW'}, '', 'ru')
        self.assertTrue(unknown)
        self.assertNotIn('SOMETHING_NEW', unknown)
        # Успех без предупреждений -- сказать нечего сверх статуса.
        self.assertEqual(cs.run_explanation(cs.STATUS_SUCCESS, {}, '', 'ru'),
                         '')
        for pair in (list(cs.FAILURE_TEXTS.values())
                     + list(cs.WARNING_TEXTS.values())
                     + list(cs.STATUS_TEXTS.values())):
            self.assertTrue(pair[0] and pair[1], pair)
            # Подстановки %(имя)d -- не текст для пользователя.
            uz_text = re.sub(r'%\(\w+\)[ds]', '', pair[1])
            for word in ('V4', 'DJI'):
                uz_text = uz_text.replace(word, '')
            self.assertEqual(re.findall(r'[A-Za-z]{2,}', uz_text), [],
                             pair[1])

    def test_the_run_view_speaks_words(self):
        run = cs.enqueue_manual(self.con, 1, 'SYNTHETIC admin')
        view = cs.run_view(run, 'uz', lock_held=False)
        self.assertEqual(view['status_label'], 'Ишга тушишини кутмоқда')
        self.assertEqual(view['trigger_label'], 'қўлда')


# ─── Блокировка ───────────────────────────────────────────────────────────

class Lock(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix='runlock_')
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self.path = runlock.cycle_lock_path(os.path.join(self.tmp, 'x.db'))

    def test_the_lock_lives_next_to_the_database(self):
        self.assertEqual(os.path.dirname(self.path), self.tmp)
        self.assertEqual(os.path.basename(self.path),
                         runlock.CYCLE_LOCK_NAME)

    def test_exclusive_and_released(self):
        first = runlock.RunLock(self.path, 'first')
        second = runlock.RunLock(self.path, 'second')
        self.assertFalse(runlock.is_held(self.path))
        self.assertTrue(first.acquire())
        self.assertTrue(runlock.is_held(self.path))
        self.assertEqual(runlock.owner(self.path)['purpose'], 'first')
        self.assertFalse(second.acquire(wait_s=0.2, poll_s=0.05))
        first.release()
        self.assertFalse(runlock.is_held(self.path))
        self.assertTrue(second.acquire())
        second.release()

    def test_another_process_sees_the_lock_and_it_dies_with_the_holder(self):
        code = ('import sys, time; sys.path.insert(0, %r); '
                'from drone_collector import runlock; '
                'lock = runlock.RunLock(%r, "child"); '
                'print(lock.acquire(), flush=True); time.sleep(30)'
                % (REPO_ROOT, self.path))
        child = subprocess.Popen([sys.executable, '-c', code],
                                 stdout=subprocess.PIPE, text=True)
        try:
            self.assertEqual(child.stdout.readline().strip(), 'True')
            self.assertTrue(runlock.is_held(self.path))
            self.assertFalse(runlock.RunLock(self.path).acquire())
        finally:
            child.kill()
            child.wait()
            child.stdout.close()
        # Процесс убит -- система отпустила блокировку сама.
        self.assertTrue(runlock.RunLock(self.path).acquire(wait_s=5,
                                                           poll_s=0.1))


# ─── Отчёт: эффективный слой и дерево ─────────────────────────────────────

def report_rows():
    rows = []
    specs = (
        # (flight, drone, day, minute, status, eligibility, raw, corrected, extra)
        (1, 'D1', 1, 0, rs.RAW_CORROBORATED, rs.AGG_CERTIFIED, 900.0, 900.0,
         {}),
        (2, 'D1', 1, 5, rs.RAW_CORROBORATED, rs.AGG_CERTIFIED, 60.0, 60.0, {}),
        (3, 'D1', 1, 10, rs.COUNTER_FLAT_RAW_OVERSTATED, rs.AGG_CERTIFIED,
         900.0, 0.0, {'candidate': True, 'base': 1, 'bridges': [2]}),
        (4, 'D1', 2, 0, rs.COUNTER_FLAT_RAW_OVERSTATED, rs.AGG_UNRESOLVED,
         500.0, 0.0, {'candidate': True, 'flags':
                      ['APPLICATION_WITH_FLAT_COUNTER']}),
        (5, 'D2', 1, 0, rs.RAW_UNVERIFIED, rs.AGG_PROVISIONAL, 700.0, 700.0,
         {'candidate': True}),
        (6, 'D2', 2, 0, rs.RAW_CORROBORATED, rs.AGG_CERTIFIED, 300.0, 300.0,
         {}),
    )
    for fid, drone, day, minute, status, elig, raw, corr, extra in specs:
        row = calc_row(fid, status, elig, raw, corrected=corr,
                       candidate=extra.get('candidate', False),
                       flags=extra.get('flags', ()), minute=minute,
                       day=date(2026, 9, 20 + day), base=extra.get('base'),
                       bridges=extra.get('bridges', ()))
        row['report_start_date'] = date(2026, 9, 20 + day)
        row['machine_key'] = row['machine_label'] = drone
        row['v4_revision_id'] = 1
        rows.append(row)
    return rows


class ReportLayer(unittest.TestCase):

    def test_without_decisions_the_totals_equal_the_automatic_ones(self):
        total = cr.build(report_rows(), 'ru')['total']
        auto = acc.summarize(report_rows())['total']
        self.assertAlmostEqual(total['raw_m2'], auto['raw_sum_m2'])
        self.assertAlmostEqual(total['excluded_m2'],
                               auto['confirmed_overstatement_m2'])
        self.assertEqual(total['decided_records'], 0)
        self.assertTrue(total['partition_holds'])

    def test_a_decision_moves_the_effective_totals_only(self):
        decisions = {4: {'id': 1, 'decision_type': dec.CONFIRM_FULL_PHANTOM,
                         'area_algorithm_version':
                             dji_area.AREA_ALGORITHM_VERSION,
                         'calculation_input_hash': 'H1',
                         'performed_by_name': 'SYNTHETIC',
                         'performed_at': datetime(2026, 9, 22, 5, 0)}}
        plain = cr.build(report_rows(), 'ru')
        decided = cr.build(report_rows(), 'ru', decisions=decisions)
        self.assertAlmostEqual(decided['total']['excluded_m2'],
                               plain['total']['excluded_m2'] + 500.0)
        self.assertAlmostEqual(decided['total']['auto_excluded_m2'],
                               plain['total']['auto_excluded_m2'])
        self.assertAlmostEqual(decided['total']['decision_delta_m2'], 500.0)
        self.assertEqual(decided['total']['review_records'], 0)
        item = {i['flight_id']: i for i in decided['items']}[4]
        self.assertEqual(item['accounting_class'], acc.REVIEW)   # автомат
        self.assertEqual(item['state'], dec.S_ADMIN_PHANTOM)       # итог
        self.assertEqual(item['decision']['performed_at_s'],
                         '22.09.2026 10:00')

    def test_the_tree_adds_up_at_every_level(self):
        for list_all in (False, True):
            report = cr.build(report_rows(), 'ru', list_all=list_all)
            for key in ('raw_m2', 'excluded_m2', 'after_m2'):
                self.assertAlmostEqual(
                    sum(n['totals'][key] for n in report['tree']),
                    report['total'][key])
            for node in report['tree']:
                for key in ('raw_m2', 'excluded_m2', 'after_m2'):
                    self.assertAlmostEqual(
                        sum(d['totals'][key] for d in node['days']),
                        node['totals'][key])
                for day in node['days']:
                    listed = sum(i['raw_m2'] for i in day['flights'])
                    rest = day['rest']['raw_m2'] if day['rest'] else 0.0
                    self.assertAlmostEqual(listed + rest,
                                           day['totals']['raw_m2'])
                    accepted = sum(i['accepted_m2'] for i in day['flights'])
                    accepted += day['rest']['accepted_m2'] \
                        if day['rest'] else 0.0
                    self.assertAlmostEqual(accepted,
                                           day['totals']['after_m2'])

    def test_register_only_lists_what_needs_words(self):
        report = cr.build(report_rows(), 'ru')
        listed = {i['flight_id'] for n in report['tree'] for d in n['days']
                  for i in d['flights']}
        self.assertEqual(listed, {3, 4, 5})
        everything = cr.build(report_rows(), 'ru', list_all=True)
        listed = {i['flight_id'] for n in everything['tree']
                  for d in n['days'] for i in d['flights']}
        self.assertEqual(listed, {1, 2, 3, 4, 5, 6})

    def test_views_filter_by_the_effective_state(self):
        decisions = {4: {'id': 1, 'decision_type': dec.CONFIRM_FULL_PHANTOM,
                         'area_algorithm_version':
                             dji_area.AREA_ALGORITHM_VERSION,
                         'calculation_input_hash': 'H1'}}
        confirmed = cr.build(report_rows(), 'ru', cr.VIEW_CONFIRMED,
                             decisions=decisions)
        self.assertEqual({r['flight_id'] for r in confirmed['register']},
                         {3, 4})
        decided = cr.build(report_rows(), 'ru', cr.VIEW_DECIDED,
                           decisions=decisions)
        self.assertEqual({r['flight_id'] for r in decided['register']}, {4})
        review = cr.build(report_rows(), 'ru', cr.VIEW_REVIEW,
                          decisions=decisions)
        self.assertEqual(review['register'], [])

    def test_chain_links_and_times_are_utc_plus_5(self):
        report = cr.build(report_rows(), 'ru')
        item = {i['flight_id']: i for i in report['items']}[3]
        chain = item['chain']
        self.assertEqual([chain['a']['flight_id'],
                          [b['flight_id'] for b in chain['b']],
                          chain['c']['flight_id']], [1, [2], 3])
        self.assertEqual((chain['a']['time_s'], chain['b'][0]['time_s'],
                          chain['c']['time_s']), ('08:00', '08:05', '08:10'))
        for point in [chain['a'], chain['c']] + chain['b']:
            self.assertEqual(point['url'], cr.DJI_RECORD_URL
                             % point['flight_id'])

    def test_a_chain_link_from_the_previous_day_shows_its_date(self):
        rows = report_rows()
        # A начался 20.09 в 23:55 по UTC+5, C -- 21.09 в 08:10.
        rows[0]['start_at_utc'] = datetime(2026, 9, 20, 18, 55)
        chain = {i['flight_id']: i for i in cr.build(rows, 'ru')['items']}[
            3]['chain']
        self.assertEqual(chain['a']['label_s'], '20.09 23:55')
        self.assertEqual(chain['a']['time_s'], '23:55')
        # Звенья того же дня -- только время, как прежде.
        self.assertEqual(chain['b'][0]['label_s'], '08:05')
        self.assertEqual(chain['c']['label_s'], '08:10')

    def test_the_decided_tab_lists_only_decisions_in_force(self):
        lapsed = {4: {'id': 1, 'decision_type': dec.ACCEPT_AUTO_RESULT,
                      'area_algorithm_version':
                          dji_area.AREA_ALGORITHM_VERSION,
                      'calculation_input_hash': 'OLD'}}
        report = cr.build(report_rows(), 'ru', cr.VIEW_DECIDED,
                          decisions=lapsed)
        self.assertEqual(report['register'], [])
        self.assertEqual(report['total']['stale_decision_records'], 1)
        # Отрицательный контроль: то же решение против нынешнего расчёта.
        in_force = {4: dict(lapsed[4], calculation_input_hash='H1')}
        report = cr.build(report_rows(), 'ru', cr.VIEW_DECIDED,
                          decisions=in_force)
        self.assertEqual({r['flight_id'] for r in report['register']}, {4})
        self.assertEqual(report['total']['stale_decision_records'], 0)

    def workbook(self, report, **kwargs):
        from openpyxl import load_workbook
        buffer = io.BytesIO()
        cr.build_workbook(report, 'ru', **kwargs).save(buffer)
        buffer.seek(0)
        return load_workbook(buffer)

    def test_the_workbook_never_writes_a_formula_from_free_text(self):
        from openpyxl import Workbook
        evil = '=HYPERLINK("http://example.invalid","x")'
        rows = report_rows()
        for row in rows:
            if row['flight_id'] == 4:
                row['machine_label'] = '@D1'
        decisions = {4: {'id': 1, 'decision_type': dec.CONFIRM_FULL_PHANTOM,
                         'area_algorithm_version':
                             dji_area.AREA_ALGORITHM_VERSION,
                         'calculation_input_hash': 'H1',
                         'performed_by_name': '+SYNTHETIC',
                         'performed_at': datetime(2026, 9, 22, 5, 0),
                         'comment': evil}}
        report = cr.build(rows, 'ru', decisions=decisions)
        history = [dict(decisions[4], flight_id=4, chain_seq=1,
                        is_current=True, auto_class=acc.REVIEW,
                        raw_area_m2=500.0, auto_accepted_m2=500.0)]
        book = self.workbook(report, period=('=1+1', '-2'),
                             filters_text=[('Машина', '=cmd')],
                             history=history)
        formulas = [(s.title, c.coordinate) for s in book.worksheets
                    for r in s.iter_rows() for c in r
                    if c.data_type == 'f'
                    or (isinstance(c.value, str) and c.value[:1] in '=+@')]
        self.assertEqual(formulas, [])
        sheet = book['Реестр']
        header = [c.value for c in sheet[1]]
        row = {r[header.index('C Flight ID')]: r
               for r in sheet.iter_rows(min_row=2, values_only=True)}[4]
        self.assertEqual(row[header.index('Комментарий решения')],
                         "'" + evil)
        self.assertEqual(row[header.index('Кто решил')], "'+SYNTHETIC")
        # Числа и обычный текст не трогаются.
        self.assertEqual(cr.xlsx_safe(-5.0), -5.0)
        self.assertEqual(cr.xlsx_safe('обычный текст'),
                         'обычный текст')
        # Отрицательный контроль: без защиты та же строка -- формула.
        raw = Workbook()
        raw.active.append([evil])
        self.assertEqual(raw.active['A1'].data_type, 'f')

    def test_history_says_not_in_force_for_a_lapsed_decision(self):
        def history_flag(input_hash):
            decisions = {4: {'id': 1, 'decision_type': dec.ACCEPT_AUTO_RESULT,
                             'area_algorithm_version':
                                 dji_area.AREA_ALGORITHM_VERSION,
                             'calculation_input_hash': input_hash,
                             'performed_by_name': 'SYNTHETIC admin',
                             'performed_at': datetime(2026, 9, 22, 5, 0),
                             'comment': 'визуально'}}
            report = cr.build(report_rows(), 'ru', decisions=decisions)
            history = [dict(decisions[4], flight_id=4, chain_seq=1,
                            is_current=True, auto_class=acc.REVIEW,
                            raw_area_m2=500.0, auto_accepted_m2=500.0)]
            book = self.workbook(report, history=history)
            return list(book['История_решений'].iter_rows(
                min_row=2, values_only=True))[0][3]
        self.assertEqual(history_flag('OLD'), 'Нет')
        self.assertEqual(history_flag('H1'), 'Да')

    def test_string_timestamps_from_raw_sqlite_are_understood(self):
        rows = report_rows()
        rows[2]['start_at_utc'] = '2026-09-21 03:10:00'
        rows[0]['start_at_utc'] = '2026-09-21 03:00:00.000000'
        chain = {i['flight_id']: i for i in cr.build(rows, 'ru')['items']}[
            3]['chain']
        self.assertEqual((chain['a']['time_s'], chain['c']['time_s']),
                         ('08:00', '08:10'))

    def test_the_workbook_carries_times_decisions_and_history(self):
        from openpyxl import load_workbook
        decisions = {4: {'id': 1, 'decision_type': dec.CONFIRM_FULL_PHANTOM,
                         'area_algorithm_version':
                             dji_area.AREA_ALGORITHM_VERSION,
                         'calculation_input_hash': 'H1',
                         'performed_by_name': 'SYNTHETIC admin',
                         'performed_at': datetime(2026, 9, 22, 5, 0),
                         'comment': 'визуально'}}
        report = cr.build(report_rows(), 'ru', decisions=decisions)
        history = [dict(decisions[4], flight_id=4, chain_seq=1,
                        is_current=True, auto_class=acc.REVIEW,
                        raw_area_m2=500.0, auto_accepted_m2=500.0)]
        book = cr.build_workbook(report, 'ru', period=('2026-09-21 00:00',
                                                       '2026-09-22 23:59'),
                                 history=history)
        buffer = io.BytesIO()
        book.save(buffer)
        buffer.seek(0)
        book = load_workbook(buffer)
        sheet = book['Реестр']
        header = [c.value for c in sheet[1]]
        rows = {r[header.index('C Flight ID')]: r
                for r in sheet.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows[3][header.index('A время (UTC+5)')],
                         '21.09.2026 08:00')
        self.assertEqual(rows[4][header.index('Кто решил')], 'SYNTHETIC admin')
        self.assertEqual(rows[4][header.index('Когда решил (UTC+5)')],
                         '22.09.2026 10:00')
        a_link = sheet.cell(row=2, column=header.index('A ссылка') + 1)
        link_rows = [r for r in range(2, sheet.max_row + 1)
                     if sheet.cell(row=r, column=header.index(
                         'C Flight ID') + 1).value == 3]
        a_link = sheet.cell(row=link_rows[0],
                            column=header.index('A ссылка') + 1)
        self.assertEqual(a_link.hyperlink.target, cr.DJI_RECORD_URL % 1)
        history_rows = list(book['История_решений'].iter_rows(
            min_row=2, values_only=True))
        self.assertEqual(history_rows[0][2], 'Полный фантом')
        self.assertEqual(history_rows[0][3], 'Да')
        summary = {r[0]: r for r in book['Сводка'].iter_rows(
            values_only=True) if r[0]}
        self.assertEqual(summary['Период: с'][1], '2026-09-21 00:00')
        self.assertEqual(summary['Спорная площадь входит в принятое'][1],
                         'Да')


if __name__ == '__main__':
    unittest.main()
