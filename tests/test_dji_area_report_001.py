# -*- coding: utf-8 -*-
"""DJI-AREA-REPORT-001: страница «Площадь DJI: техническая оценка».

Отчёт читает `dji_area_calculations` / `dji_field_attributions` и ничего не
пишет. Здесь проверяется то, что ломается молча и стоит дорого:

1. **NULL -- это слово, а не 0.00.** Запись без измеренного начала показывает
   «Недостаточно данных»; запись с ИЗМЕРЕННЫМ нулём (COUNTER_FLAT) показывает
   `0.00`. Обе проверки стоят в одном тесте, и каждая -- отрицательный
   контроль другой: проверка «нигде нет 0.00» запретила бы честный ноль, а
   проверка «где-то есть 0.00» прошла бы и на странице, печатающей ноль
   вместо неизвестного.
2. **Полного итога нет.** Проверенный подытог, предварительная оценка и
   нерешённая RAW-экспозиция показываются раздельно, и ни один элемент
   страницы не равен их арифметической сумме.
3. **Страница не считает сама.** Числа карточек сверяются с корзинами,
   полученными прямым вызовом `dji_area.aggregate.aggregate` на тех же
   строках. Расхождение означает, что арифметика переписана в отчёте.
4. **Устаревшее не суммируется.** Строка с `superseded_at` и строка чужой
   версии алгоритма в итог не входят; рядом стоит контроль -- та же величина
   в ТЕКУЩЕЙ строке итог двигает.
5. **Tier поля не меняет площадь.** TIER5 попадает в корзину «поле не
   определено», но его техническая оценка равна оценке идентичной строки
   TIER1.

ВСЕ ДАННЫЕ СИНТЕТИЧЕСКИЕ. Ни одного настоящего идентификатора вылета, ни
одного hardware_id из кабинета, ни одной координаты. Вылеты нумеруются
с 900001, борта помечены SYNTHETIC-HW-...-NOT-REAL.
"""

import io
import json
import os
import re
import sqlite3
import unittest

from datetime import date, datetime, timedelta

from tests.harness import (app, reset_db, create_admin, create_org, login,
                           CSRF, TEST_DB_PATH)

from models import (db, DjiAreaCalculation, DjiFieldAttribution, DroneFlight,
                    DroneUnit, User, UserModulePermission, ROLE_OPERATOR)

import dji_area
from dji_area import aggregate as dji_aggregate
from dji_area import field as dji_field
from dji_area import resolver as dji_resolver
import drones


REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

PROVIDER = 'SYNTHETIC-ACCOUNT-NOT-REAL'
HW6 = 'SYNTHETIC-HW-6-NOT-REAL'
HW7 = 'SYNTHETIC-HW-7-NOT-REAL'
TOKEN = 'SYNTHETIC-collector-token-NOT-REAL'

# Синтетический день расчётов и окно, которое его покрывает.
DAY = date(2026, 6, 5)
DAY2 = date(2026, 6, 6)
WINDOW = '?date_from=2026-06-01&date_to=2026-06-30'
ALL_TIME = '?date_from=&date_to='

# Шесть классов записей: по одной на каждую корзину модели.
#   (flight_id, статус, право на агрегирование, RAW м2, оценка м2)
BUCKET_ROWS = (
    (900001, dji_resolver.RAW_CORROBORATED, dji_resolver.AGG_CERTIFIED,
     30000.0, 30000.0),
    (900002, dji_resolver.RAW_CORROBORATED_QUALIFIED,
     dji_resolver.AGG_PROVISIONAL, 20000.0, 20000.0),
    (900003, dji_resolver.BASELINE_UNKNOWN, dji_resolver.AGG_UNRESOLVED,
     110000.0, None),
    (900004, dji_resolver.OVERLAP_REVIEW, dji_resolver.AGG_EXCLUDED_OVERLAP,
     70000.0, None),
    (900005, dji_resolver.COUNTER_FLAT_RAW_OVERSTATED,
     dji_resolver.AGG_CERTIFIED, 40000.0, 0.0),
    (900006, dji_resolver.APPLICATION_WITHOUT_MEASURED_AREA,
     dji_resolver.AGG_UNRESOLVED, None, None),
)

# Двадцать строк для сверки с `dji_area.aggregate`. Тот же список кормит и
# базу, и прямой вызов агрегатора: расхождение означает, что страница
# пересчитывает по-своему.
#   (flight_id, статус, право, RAW м2, оценка м2, tier, день, борт, без площади)
PARITY_SPEC = (
    (900101, dji_resolver.RAW_CORROBORATED, dji_resolver.AGG_CERTIFIED,
     12000.0, 12000.0, dji_field.TIER1_EXACT, 0, HW6, False),
    (900102, dji_resolver.RAW_CORROBORATED, dji_resolver.AGG_CERTIFIED,
     15000.0, 14500.0, dji_field.TIER2_STRONG, 0, HW7, False),
    (900103, dji_resolver.RAW_CORROBORATED_QUALIFIED,
     dji_resolver.AGG_PROVISIONAL, 18000.0, 17000.0,
     dji_field.TIER3_SUPPORTED, 0, HW6, False),
    (900104, dji_resolver.RAW_UNVERIFIED, dji_resolver.AGG_PROVISIONAL,
     9000.0, 9000.0, dji_field.TIER4_GEOMETRIC, 0, HW7, False),
    (900105, dji_resolver.BASELINE_UNKNOWN, dji_resolver.AGG_UNRESOLVED,
     21000.0, None, dji_field.TIER5_UNKNOWN, 0, HW6, False),
    (900106, dji_resolver.OVERLAP_REVIEW, dji_resolver.AGG_EXCLUDED_OVERLAP,
     33000.0, None, dji_field.TIER1_EXACT, 0, HW7, False),
    (900107, dji_resolver.COUNTER_FLAT_RAW_OVERSTATED,
     dji_resolver.AGG_CERTIFIED, 27000.0, 0.0, dji_field.TIER2_STRONG,
     0, HW6, False),
    (900108, dji_resolver.APPLICATION_WITHOUT_MEASURED_AREA,
     dji_resolver.AGG_UNRESOLVED, None, None, dji_field.TIER5_UNKNOWN,
     0, HW7, True),
    (900109, dji_resolver.COUNTER_ZERO, dji_resolver.AGG_CERTIFIED,
     0.0, 0.0, dji_field.TIER1_EXACT, 0, HW6, False),
    (900110, dji_resolver.UNKNOWN_SUSPECT, dji_resolver.AGG_UNRESOLVED,
     45000.0, None, dji_field.TIER3_SUPPORTED, 0, HW7, False),
    (900111, dji_resolver.CHANNEL_MISSING, dji_resolver.AGG_UNRESOLVED,
     5000.0, None, dji_field.TIER5_UNKNOWN, 1, HW6, False),
    (900112, dji_resolver.ZERO_RECORDED_UNVERIFIED,
     dji_resolver.AGG_UNRESOLVED, 0.0, None, dji_field.TIER4_GEOMETRIC,
     1, HW7, False),
    (900113, dji_resolver.COUNTER_RELATIONSHIP_OUTLIER,
     dji_resolver.AGG_UNRESOLVED, 61000.0, None, dji_field.TIER1_EXACT,
     1, HW6, False),
    (900114, dji_resolver.COUNTER_NONMONOTONE_REVIEW,
     dji_resolver.AGG_UNRESOLVED, 13000.0, None, dji_field.TIER2_STRONG,
     1, HW7, False),
    (900115, dji_resolver.PARTIAL_RECORDED_OVERSTATEMENT,
     dji_resolver.AGG_CERTIFIED, 24000.0, 2400.0, dji_field.TIER3_SUPPORTED,
     1, HW6, False),
    (900116, dji_resolver.RAW_CORROBORATED, dji_resolver.AGG_CERTIFIED,
     31000.0, 31000.0, dji_field.TIER5_UNKNOWN, 1, HW7, False),
    (900117, dji_resolver.RAW_UNVERIFIED, dji_resolver.AGG_PROVISIONAL,
     7000.0, 7000.0, dji_field.TIER1_EXACT, 1, HW6, False),
    (900118, dji_resolver.OVERLAP_REVIEW, dji_resolver.AGG_EXCLUDED_OVERLAP,
     19000.0, None, dji_field.TIER4_GEOMETRIC, 1, HW7, False),
    (900119, dji_resolver.RAW_CORROBORATED_QUALIFIED,
     dji_resolver.AGG_PROVISIONAL, 26000.0, 25000.0, dji_field.TIER2_STRONG,
     1, HW6, False),
    (900120, dji_resolver.APPLICATION_WITHOUT_MEASURED_AREA,
     dji_resolver.AGG_UNRESOLVED, None, None, dji_field.TIER3_SUPPORTED,
     1, HW7, True),
)


def ha(value_m2):
    """Гектары так, как их печатает страница: два знака, точка, без групп."""
    return None if value_m2 is None else '%.2f' % (value_m2 / 10000.0)


class Base(unittest.TestCase):
    """Синтетическая площадка: две машины, известные борта, один админ."""

    def setUp(self):
        reset_db()
        self.org_id = create_org('SYNTHETIC Org')
        with app.app_context():
            unit6 = DroneUnit(number=6, organization_id=self.org_id,
                              hardware_id=HW6)
            unit7 = DroneUnit(number=7, organization_id=self.org_id,
                              hardware_id=HW7)
            db.session.add_all([unit6, unit7])
            db.session.commit()
            self.unit6_id = unit6.id
            self.unit7_id = unit7.id
        self.admin_id = create_admin('area-admin')
        # Локальный «сегодня» операторов -- UTC+5, как его считает модуль.
        # Вычислен здесь заново, а не взят из drones.py: проверка не должна
        # опираться на ту же функцию, что и проверяемый код.
        self.today = (datetime.utcnow() + timedelta(hours=5)).date()
        self._hash_seq = 0

    # ── посев ────────────────────────────────────────────────────────────

    def _next_hash(self):
        self._hash_seq += 1
        return 'SYNTHETIC-INPUT-HASH-%04d-NOT-REAL' % self._hash_seq

    def calc_object(self, flight_id, status, eligibility, raw_m2=None,
                    corrected_m2=None, day=DAY, hardware_id=HW6, minute=0,
                    application_without_area=None, delta_m2=None,
                    superseded=False, algorithm_version=None, method=None,
                    confidence=None, flags=None, activity=None,
                    channel=None, drone_flight_id=None):
        start = (datetime(day.year, day.month, day.day, 3, 0)
                 + timedelta(minutes=minute))
        return DjiAreaCalculation(
            flight_id=flight_id,
            provider_account_id=PROVIDER,
            drone_flight_id=drone_flight_id,
            hardware_id=hardware_id,
            hardware_id_source='SYNTHETIC',
            area_algorithm_version=(algorithm_version
                                    or dji_area.AREA_ALGORITHM_VERSION),
            calculation_input_hash=self._next_hash(),
            calculated_at=datetime(2026, 9, 8, 0, 0),
            superseded_at=(datetime(2026, 9, 8, 1, 0) if superseded else None),
            supersede_reason='SYNTHETIC' if superseded else None,
            start_at_utc=start,
            end_at_utc=start + timedelta(minutes=20),
            report_timezone=dji_area.REPORT_TIMEZONE,
            report_start_date=day,
            raw_area_m2=raw_m2,
            raw_area_source='CARD' if raw_m2 is not None else None,
            controller_delta_area_m2=delta_m2,
            corrected_recorded_area_m2=corrected_m2,
            area_status=status,
            evidence_status=None,
            area_method=method,
            area_confidence=confidence,
            anomaly_flags_json=json.dumps(list(flags or [])),
            application_activity=activity,
            application_channel_quality=channel,
            application_evidence_kind=None,
            application_without_area=application_without_area,
            aggregation_eligibility=eligibility,
        )

    def add_calc(self, *args, **kwargs):
        with app.app_context():
            db.session.add(self.calc_object(*args, **kwargs))
            db.session.commit()

    def add_calcs(self, objects_spec):
        """Пачкой, одной транзакцией: 500 записей по одной коммитятся долго."""
        with app.app_context():
            db.session.add_all([self.calc_object(**spec)
                                for spec in objects_spec])
            db.session.commit()

    def add_attr(self, flight_id, tier, name=None, superseded=False,
                 version=None, method='SYNTHETIC-FIELD-METHOD'):
        with app.app_context():
            db.session.add(DjiFieldAttribution(
                flight_id=flight_id,
                field_resolver_version=(version
                                        or dji_area.FIELD_RESOLVER_VERSION),
                field_input_hash=self._next_hash(),
                calculated_at=datetime(2026, 9, 8, 0, 0),
                superseded_at=(datetime(2026, 9, 8, 1, 0)
                               if superseded else None),
                historical_geometry_available=False,
                field_attribution_tier=tier,
                field_attribution_method=method,
                field_name_at_snapshot=name,
            ))
            db.session.commit()

    def seed_buckets(self, tier=dji_field.TIER1_EXACT):
        """Шесть классов записей на одном дне и одной машине."""
        for flight_id, status, eligibility, raw_m2, corrected in BUCKET_ROWS:
            self.add_calc(
                flight_id, status, eligibility, raw_m2=raw_m2,
                corrected_m2=corrected, minute=(flight_id - 900001) * 5,
                application_without_area=(
                    status == dji_resolver.APPLICATION_WITHOUT_MEASURED_AREA
                    or None))
            self.add_attr(flight_id, tier, name='SYNTHETIC field A')

    # ── клиент и разбор страницы ─────────────────────────────────────────

    def client_as(self, user_id=None, language='ru'):
        user_id = self.admin_id if user_id is None else user_id
        with app.app_context():
            user = User.query.get(user_id)
            user.language = language
            db.session.commit()
        client = app.test_client()
        login(client, user_id)
        return client

    def page(self, query='', language='ru', user_id=None):
        client = self.client_as(user_id=user_id, language=language)
        response = client.get('/drones/area-evidence' + query)
        self.assertEqual(response.status_code, 200, query)
        return response.get_data(as_text=True)

    def workbook(self, query='', language='ru'):
        from openpyxl import load_workbook
        client = self.client_as(language=language)
        response = client.get('/drones/area-evidence.xlsx' + query)
        self.assertEqual(response.status_code, 200, query)
        return response, load_workbook(io.BytesIO(response.data))

    STAT_RE = re.compile(
        r'<div class="vs-stat">\s*'
        r'<div class="vs-stat-label">(.*?)</div>\s*'
        r'<div class="vs-stat-value">(.*?)</div>\s*'
        r'<div class="vs-muted">(.*?)</div>', re.S)

    def stat_cards(self, html):
        cards = {}
        for label, value, muted in self.STAT_RE.findall(html):
            cards[label.strip()] = {'value': value.strip(),
                                    'muted': ' '.join(muted.split())}
        self.assertEqual(len(cards), 8,
                         'expected eight stat cards, parsed %d' % len(cards))
        return cards

    def record_rows(self, html, title='Записи'):
        """Строки таблицы детализации -- и только они."""
        marker = '<span class="vs-card-title">%s</span>' % title
        self.assertIn(marker, html)
        tail = html.split(marker, 1)[1]
        return [row for row in re.findall(r'<tr>(.*?)</tr>', tail, re.S)
                if '<td' in row]

    def rows_by_flight(self, html, title='Записи'):
        out = {}
        for row in self.record_rows(html, title=title):
            ids = re.findall(r'<td>(\d+)</td>', row)
            self.assertTrue(ids, 'a record row without a DJI flight id')
            out[int(ids[-1])] = row
        return out

    @staticmethod
    def right_cells(row):
        return [cell.strip() for cell in
                re.findall(r'<td class="right">(.*?)</td>', row, re.S)]

    @staticmethod
    def badges(row):
        return [b.strip() for b in
                re.findall(r'<span class="vs-badge[^"]*">(.*?)</span>',
                           row, re.S)]

    def day_rows(self, html):
        marker = '<span class="vs-card-title">По дням и машинам</span>'
        self.assertIn(marker, html)
        tail = html.split(marker, 1)[1].split('<span class="vs-card-title">',
                                              1)[0]
        return [row for row in re.findall(r'<tr>(.*?)</tr>', tail, re.S)
                if '<td' in row]

    # ── SQLite в обход ORM ───────────────────────────────────────────────

    def raw(self, sql, params=()):
        """[REASON]: читается ХРАНИМОЕ значение, а не представление ORM."""
        con = sqlite3.connect(TEST_DB_PATH)
        try:
            return con.execute(sql, params).fetchall()
        finally:
            con.close()


# ─── 1. Доступ ───────────────────────────────────────────────────────────────

class Access(Base):
    """Право `drones` держится на маршруте, а не только в меню."""

    def make_operator(self, username, has_drones):
        with app.app_context():
            user = User(username=username, role=ROLE_OPERATOR,
                        full_name='SYNTHETIC operator', language='ru')
            user.set_password('test-password')
            db.session.add(user)
            db.session.flush()
            if has_drones:
                db.session.add(UserModulePermission(
                    user_id=user.id, module_code='drones', has_access=True))
            db.session.commit()
            return user.id

    def test_admin_opens_the_page_and_the_workbook(self):
        self.seed_buckets()
        client = self.client_as()
        for url in ('/drones/area-evidence',
                    '/drones/area-evidence.xlsx',
                    '/drones/area-evidence' + WINDOW,
                    '/drones/area-evidence.xlsx' + WINDOW):
            self.assertEqual(client.get(url).status_code, 200, url)

    def test_the_module_permission_decides_both_routes(self):
        """Отрицательный контроль стоит рядом: тот же запрос, другое право.

        Без положительного контроля 403 доказывал бы лишь то, что маршрут
        закрыт для всех -- в том числе если бы он был закрыт по ошибке.
        """
        allowed_id = self.make_operator('drones-yes', True)
        denied_id = self.make_operator('drones-no', False)

        allowed = app.test_client()
        login(allowed, allowed_id)
        denied = app.test_client()
        login(denied, denied_id)

        for url in ('/drones/area-evidence', '/drones/area-evidence.xlsx'):
            self.assertEqual(allowed.get(url).status_code, 200, url)
            self.assertEqual(denied.get(url).status_code, 403, url)

        # И право действительно записано в базе, а не подразумевается.
        rows = self.raw('SELECT module_code, has_access '
                        '  FROM user_module_permissions WHERE user_id = ?',
                        (allowed_id,))
        self.assertEqual(rows, [('drones', 1)])
        self.assertEqual(
            self.raw('SELECT COUNT(*) FROM user_module_permissions '
                     ' WHERE user_id = ?', (denied_id,)), [(0,)])

    def test_an_anonymous_visitor_is_not_served_the_report(self):
        anon = app.test_client()
        for url in ('/drones/area-evidence', '/drones/area-evidence.xlsx'):
            self.assertIn(anon.get(url).status_code, (302, 401, 403), url)


# ─── 2. Корзины ──────────────────────────────────────────────────────────────

class Buckets(Base):
    """Шесть классов записей -- шесть разных судеб в карточках."""

    def setUp(self):
        Base.setUp(self)
        self.seed_buckets()

    def test_the_cards_show_each_bucket_separately(self):
        cards = self.stat_cards(self.page(WINDOW))

        # RAW DJI -- всё, что записал DJI, включая пересечение: 27.00 га.
        self.assertEqual(cards['DJI RAW, га']['value'], '27.00')
        self.assertEqual(cards['DJI RAW, га']['muted'], '6 записей')
        # Отрицательный контроль: без RAW пересечения было бы 20.00.
        self.assertNotEqual(cards['DJI RAW, га']['value'], '20.00')

        # Проверено: RAW_CORROBORATED (3 га) + измеренный ноль COUNTER_FLAT.
        self.assertEqual(cards['Проверено, га']['value'], '3.00')
        self.assertEqual(cards['Проверено, га']['muted'], '2 записей')
        # Отрицательный контроль: попади пересечение в проверенное -- 10.00.
        self.assertNotEqual(cards['Проверено, га']['value'], '10.00')

        self.assertEqual(cards['Предварительно, га']['value'], '2.00')
        self.assertEqual(cards['Предварительно, га']['muted'], '1 записей')

        # Известная часть -- ровно проверено + предварительно.
        self.assertEqual(cards['Известная часть, га']['value'], '5.00')
        self.assertIn('не полный итог', cards['Известная часть, га']['muted'])

        # Нерешённое: две записи, и их RAW-экспозиция названа числом.
        self.assertEqual(cards['Недостаточно данных']['value'], '2')
        self.assertIn('11.00', cards['Недостаточно данных']['muted'])
        self.assertIn('пересечений записей: 1',
                      cards['Недостаточно данных']['muted'])

        self.assertEqual(cards['Повторные/перенесённые']['value'], '1')
        self.assertEqual(
            cards['Площадь не измерена при активности']['value'], '1')

    def test_no_element_of_the_page_is_the_complete_total(self):
        """«Полный итог» не вычисляется -- ни в карточке, ни в строке дня.

        Отрицательный контроль -- числа, которые НА странице есть: если бы
        страница вообще ничего не напечатала, все assertNotIn ниже были бы
        зелёными и не проверяли бы ничего.
        """
        html = self.page(WINDOW)
        self.assertIn('27.00', html)   # RAW
        self.assertIn('5.00', html)    # известная часть
        self.assertIn('11.00', html)   # нерешённая RAW-экспозиция

        # 3 + 2 + 11 + 7 -- «всё вместе», которого быть не должно.
        self.assertNotIn('23.00', html)
        # 5 + 11 -- известная часть плюс нерешённое.
        self.assertNotIn('16.00', html)
        # 3 + 2 + 7 -- известная часть плюс пересечение.
        self.assertNotIn('12.00', html)

    def test_the_overlap_row_keeps_its_raw_and_leaves_the_certified_sum(self):
        html = self.page(WINDOW + '&detail=1')
        rows = self.rows_by_flight(html)
        overlap = rows[900004]

        # RAW пересечения показан (7.00 га), оценки нет.
        self.assertEqual(self.right_cells(overlap)[0], '7.00')
        self.assertIn('Недостаточно данных', self.right_cells(overlap)[1])
        self.assertIn('Пересечение записей', self.badges(overlap)[0])

        # И он не попал ни в проверенное, ни в предварительное.
        cards = self.stat_cards(html)
        self.assertEqual(cards['Проверено, га']['value'], '3.00')
        self.assertEqual(cards['Предварительно, га']['value'], '2.00')
        # Отрицательный контроль: 7.00 на странице ЕСТЬ -- в своей строке.
        self.assertIn('7.00', html)

    def test_the_day_row_repeats_the_same_partition(self):
        rows = self.day_rows(self.page(WINDOW))
        self.assertEqual(len(rows), 1)
        cells = [re.sub(r'<[^>]+>', '', cell).strip() for cell in
                 re.findall(r'<td[^>]*>(.*?)</td>', rows[0], re.S)]
        # Дата, машина, записей, RAW, проверено, предварительно, нерешённых,
        # повторных, % поля, поле не определено.
        self.assertEqual(cells[0], '05.06.2026')
        self.assertEqual(cells[1], '№ 6')
        self.assertEqual(cells[2], '6')
        self.assertEqual(cells[3], '27.00')
        self.assertEqual(cells[4], '3.00')
        self.assertEqual(cells[5], '2.00')
        self.assertEqual(cells[6], '2')
        self.assertEqual(cells[7], '1')


# ─── 3. NULL против измеренного нуля ─────────────────────────────────────────

class NullVersusMeasuredZero(Base):
    """Слово и ноль -- разные вещи, и обе должны быть видны на экране."""

    def setUp(self):
        Base.setUp(self)
        self.seed_buckets()

    def test_the_null_is_a_word_and_the_measured_zero_is_a_number(self):
        html = self.page(WINDOW + '&detail=1')
        rows = self.rows_by_flight(html)

        baseline = self.right_cells(rows[900003])   # BASELINE_UNKNOWN
        flat = self.right_cells(rows[900005])       # COUNTER_FLAT

        # Неизмеренное начало: слова, и НИ ОДНОГО нуля в ячейке оценки.
        self.assertEqual(baseline[0], '11.00')
        self.assertEqual(baseline[1], 'Недостаточно данных')
        self.assertNotIn('0.00', baseline[1])

        # Измеренный ноль: именно 0.00, и никаких слов.
        self.assertEqual(flat[0], '4.00')
        self.assertEqual(flat[1], '0.00')
        self.assertNotIn('Недостаточно данных', flat[1])

        # [REASON]: это ВЗАИМНЫЕ отрицательные контроли. Проверка «в ячейке
        # нет 0.00» прошла бы и на странице, которая вообще не умеет печатать
        # ноль; проверка «0.00 есть» прошла бы и на странице, печатающей ноль
        # вместо неизвестного. Различает только пара.
        self.assertNotEqual(baseline[1], flat[1])

    def test_sqlite_agrees_that_one_is_null_and_the_other_is_a_real_zero(self):
        """Хранимое значение, а не то, что нарисовал шаблон."""
        stored = dict((row[0], (row[1], row[2])) for row in self.raw(
            'SELECT flight_id, corrected_recorded_area_m2, '
            '       typeof(corrected_recorded_area_m2) '
            '  FROM dji_area_calculations ORDER BY flight_id'))
        self.assertEqual(stored[900003], (None, 'null'))
        self.assertEqual(stored[900005], (0.0, 'real'))
        # Отрицательный контроль: в базе есть и ненулевая оценка.
        self.assertEqual(stored[900001], (30000.0, 'real'))


# ─── 4. Устаревшие строки ────────────────────────────────────────────────────

class SupersededIsIgnored(Base):
    """Текущая строка -- `superseded_at IS NULL` под текущей версией."""

    def setUp(self):
        Base.setUp(self)
        self.add_calc(900001, dji_resolver.RAW_CORROBORATED,
                      dji_resolver.AGG_CERTIFIED, raw_m2=30000.0,
                      corrected_m2=30000.0)
        self.add_attr(900001, dji_field.TIER1_EXACT)
        self.add_calc(900002, dji_resolver.RAW_CORROBORATED_QUALIFIED,
                      dji_resolver.AGG_PROVISIONAL, raw_m2=20000.0,
                      corrected_m2=20000.0, minute=5)
        self.add_attr(900002, dji_field.TIER1_EXACT)

    def baseline_cards(self):
        cards = self.stat_cards(self.page(WINDOW))
        self.assertEqual(cards['DJI RAW, га']['value'], '5.00')
        self.assertEqual(cards['Проверено, га']['value'], '3.00')
        self.assertEqual(cards['Предварительно, га']['value'], '2.00')
        return cards

    def test_a_superseded_row_and_a_foreign_version_do_not_move_the_total(self):
        before = self.baseline_cards()

        # Устаревшая ревизия того же вылета с чудовищной площадью.
        self.add_calc(900001, dji_resolver.RAW_CORROBORATED,
                      dji_resolver.AGG_CERTIFIED, raw_m2=9990000.0,
                      corrected_m2=9990000.0, superseded=True, minute=10)
        # Текущая строка, но посчитанная ДРУГИМ алгоритмом.
        self.add_calc(900003, dji_resolver.RAW_CORROBORATED,
                      dji_resolver.AGG_CERTIFIED, raw_m2=8880000.0,
                      corrected_m2=8880000.0, minute=15,
                      algorithm_version=dji_area.AREA_ALGORITHM_VERSION
                      + '-SYNTHETIC-OTHER')

        after = self.stat_cards(self.page(WINDOW))
        self.assertEqual(after, before)
        # И обе строки действительно лежат в базе -- иначе тест был бы зелён
        # просто потому, что вставка молча не прошла.
        self.assertEqual(
            self.raw('SELECT COUNT(*) FROM dji_area_calculations'), [(4,)])
        self.assertEqual(
            self.raw('SELECT COUNT(*) FROM dji_area_calculations '
                     ' WHERE superseded_at IS NOT NULL'), [(1,)])
        self.assertEqual(
            self.raw('SELECT COUNT(*) FROM dji_area_calculations '
                     ' WHERE area_algorithm_version <> ?',
                     (dji_area.AREA_ALGORITHM_VERSION,)), [(1,)])

    def test_the_same_number_in_a_current_row_does_move_the_total(self):
        """Отрицательный контроль к тесту выше."""
        self.baseline_cards()
        self.add_calc(900004, dji_resolver.RAW_CORROBORATED,
                      dji_resolver.AGG_CERTIFIED, raw_m2=9990000.0,
                      corrected_m2=9990000.0, minute=20)
        cards = self.stat_cards(self.page(WINDOW))
        self.assertEqual(cards['Проверено, га']['value'], '1002.00')
        self.assertNotEqual(cards['Проверено, га']['value'], '3.00')

    def test_a_superseded_field_attribution_does_not_confirm_the_field(self):
        """Устаревшая привязка не делает поле подтверждённым."""
        self.add_calc(900005, dji_resolver.RAW_CORROBORATED,
                      dji_resolver.AGG_CERTIFIED, raw_m2=10000.0,
                      corrected_m2=10000.0, minute=25)
        self.add_attr(900005, dji_field.TIER1_EXACT, superseded=True)
        rows = self.rows_by_flight(self.page(WINDOW + '&detail=1'))
        self.assertEqual(self.badges(rows[900005])[1], 'Поле не определено')
        # Отрицательный контроль: у вылета с ДЕЙСТВУЮЩЕЙ привязкой -- иначе.
        self.assertEqual(self.badges(rows[900001])[1], 'Поле подтверждено')


# ─── 5. Надёжность привязки к полю ───────────────────────────────────────────

class FieldTiers(Base):
    """TIER меняет корзину поля и НИКОГДА не меняет площадь."""

    TIER_FLIGHTS = (
        (900011, dji_field.TIER1_EXACT, 'Поле подтверждено'),
        (900012, dji_field.TIER2_STRONG, 'Поле подтверждено'),
        (900013, dji_field.TIER3_SUPPORTED, 'Поле предположительно'),
        (900014, dji_field.TIER4_GEOMETRIC, 'Поле предположительно'),
        (900015, dji_field.TIER5_UNKNOWN, 'Поле не определено'),
    )

    def setUp(self):
        Base.setUp(self)
        for index, (flight_id, tier, _label) in enumerate(self.TIER_FLIGHTS):
            # Пять ОДИНАКОВЫХ по площади записей: различается только tier.
            self.add_calc(flight_id, dji_resolver.RAW_CORROBORATED,
                          dji_resolver.AGG_CERTIFIED, raw_m2=10000.0,
                          corrected_m2=10000.0, minute=index * 5)
            self.add_attr(flight_id, tier, name='SYNTHETIC field %d' % index)

    def test_the_three_words_of_field_reliability(self):
        rows = self.rows_by_flight(self.page(WINDOW + '&detail=1'))
        seen = {}
        for flight_id, tier, expected in self.TIER_FLIGHTS:
            seen[tier] = self.badges(rows[flight_id])[1]
            self.assertEqual(seen[tier], expected, tier)
        # Отрицательный контроль: три подписи действительно РАЗНЫЕ, иначе
        # проверка выше прошла бы и на шаблоне, печатающем одно слово всем.
        self.assertEqual(len(set(seen.values())), 3)

    def test_tier5_raw_lands_in_the_unassigned_bucket(self):
        cards = self.stat_cards(self.page(WINDOW))
        self.assertEqual(cards['Поле не определено']['value'], '1.00')
        self.assertIn('1 записей', cards['Поле не определено']['muted'])
        # Отрицательный контроль: это НЕ весь RAW периода.
        self.assertEqual(cards['DJI RAW, га']['value'], '5.00')
        self.assertNotEqual(cards['Поле не определено']['value'], '5.00')

    def test_tier5_does_not_change_the_area(self):
        html = self.page(WINDOW + '&detail=1')
        rows = self.rows_by_flight(html)
        tier1 = self.right_cells(rows[900011])
        tier5 = self.right_cells(rows[900015])
        self.assertEqual(tier5, tier1)
        self.assertEqual(tier5[1], '1.00')
        self.assertNotIn('Недостаточно данных', tier5[1])

        # И проверенный подытог считает все пять, TIER5 включительно.
        cards = self.stat_cards(html)
        self.assertEqual(cards['Проверено, га']['value'], '5.00')
        self.assertEqual(cards['Проверено, га']['muted'], '5 записей')
        # Отрицательный контроль: выброси страница TIER5 -- было бы 4.00.
        self.assertNotEqual(cards['Проверено, га']['value'], '4.00')

    def test_the_missing_attribution_reads_as_tier5(self):
        """Строка без привязки вообще -- тоже «поле не определено»."""
        self.add_calc(900016, dji_resolver.RAW_CORROBORATED,
                      dji_resolver.AGG_CERTIFIED, raw_m2=10000.0,
                      corrected_m2=10000.0, minute=30)
        rows = self.rows_by_flight(self.page(WINDOW + '&detail=1'))
        self.assertEqual(self.badges(rows[900016])[1], 'Поле не определено')
        self.assertEqual(
            self.raw('SELECT COUNT(*) FROM dji_field_attributions '
                     ' WHERE flight_id = ?', (900016,)), [(0,)])


# ─── 6. Фильтры ──────────────────────────────────────────────────────────────

class Filters(Base):
    """Даты, машина, статус, tier -- и что делает неизвестное значение."""

    def setUp(self):
        Base.setUp(self)
        self.recent_day = self.today
        self.old_day = self.today - timedelta(days=200)
        # Свежая запись: борт 6, проверено, поле подтверждено.
        self.add_calc(900021, dji_resolver.RAW_CORROBORATED,
                      dji_resolver.AGG_CERTIFIED, raw_m2=10000.0,
                      corrected_m2=10000.0, day=self.recent_day,
                      hardware_id=HW6)
        self.add_attr(900021, dji_field.TIER1_EXACT)
        # Старая запись: борт 7, предварительно, поле не определено.
        self.add_calc(900022, dji_resolver.RAW_UNVERIFIED,
                      dji_resolver.AGG_PROVISIONAL, raw_m2=20000.0,
                      corrected_m2=20000.0, day=self.old_day,
                      hardware_id=HW7)

    def records_shown(self, html):
        return self.stat_cards(html)['DJI RAW, га']['muted']

    def test_the_default_window_is_bounded_and_recent(self):
        html = self.page()
        cards = self.stat_cards(html)
        self.assertEqual(cards['DJI RAW, га']['value'], '1.00')
        self.assertEqual(cards['DJI RAW, га']['muted'], '1 записей')
        # Границы окна проставлены в форме, а не пусты.
        self.assertRegex(html, r'name="date_from"[^>]*value="\d{4}-\d{2}-\d{2}"')
        self.assertRegex(html, r'name="date_to"[^>]*value="\d{4}-\d{2}-\d{2}"')
        expected_from = (self.today
                         - timedelta(days=drones.DRONE_AREA_DEFAULT_DAYS))
        self.assertIn('value="%s"' % expected_from.isoformat(), html)
        # Отрицательный контроль: старая запись есть в базе, но не в окне.
        self.assertEqual(
            self.raw('SELECT COUNT(*) FROM dji_area_calculations'), [(2,)])

    def test_an_explicit_window_selects_the_old_record_only(self):
        query = '?date_from=%s&date_to=%s' % (self.old_day.isoformat(),
                                              self.old_day.isoformat())
        cards = self.stat_cards(self.page(query))
        self.assertEqual(cards['DJI RAW, га']['value'], '2.00')
        self.assertEqual(cards['Предварительно, га']['value'], '2.00')
        self.assertEqual(cards['Проверено, га']['value'], '0.00')

    def test_the_all_time_link_carries_empty_date_parameters(self):
        html = self.page()
        # Ссылка «За всё время» передаёт даты ПУСТЫМИ, а не опускает их:
        # опущенный параметр вернул бы окно по умолчанию.
        self.assertRegex(html, r'href="[^"]*date_from=&amp;date_to=[^"]*"')

        cards = self.stat_cards(self.page(ALL_TIME))
        self.assertEqual(cards['DJI RAW, га']['value'], '3.00')
        self.assertEqual(cards['DJI RAW, га']['muted'], '2 записей')
        self.assertIn('name="date_from" id="areaDateFrom" '
                      'class="vs-input" value=""', self.page(ALL_TIME))

    def test_the_unit_filter_keeps_one_machine(self):
        cards = self.stat_cards(
            self.page(ALL_TIME + '&unit_id=%d' % self.unit6_id))
        self.assertEqual(cards['DJI RAW, га']['value'], '1.00')
        # Отрицательный контроль: другая машина даёт другое число.
        other = self.stat_cards(
            self.page(ALL_TIME + '&unit_id=%d' % self.unit7_id))
        self.assertEqual(other['DJI RAW, га']['value'], '2.00')

    def test_the_status_filter_keeps_one_status(self):
        cards = self.stat_cards(
            self.page(ALL_TIME + '&status=' + dji_resolver.RAW_UNVERIFIED))
        self.assertEqual(cards['DJI RAW, га']['value'], '2.00')
        self.assertEqual(cards['DJI RAW, га']['muted'], '1 записей')
        other = self.stat_cards(
            self.page(ALL_TIME + '&status=' + dji_resolver.RAW_CORROBORATED))
        self.assertEqual(other['DJI RAW, га']['value'], '1.00')

    def test_the_tier_filter_keeps_one_tier(self):
        cards = self.stat_cards(
            self.page(ALL_TIME + '&tier=' + dji_field.TIER1_EXACT))
        self.assertEqual(cards['DJI RAW, га']['muted'], '1 записей')
        self.assertEqual(cards['DJI RAW, га']['value'], '1.00')
        other = self.stat_cards(
            self.page(ALL_TIME + '&tier=' + dji_field.TIER5_UNKNOWN))
        self.assertEqual(other['DJI RAW, га']['value'], '2.00')

    def test_an_unknown_status_or_tier_is_ignored_not_a_500(self):
        for query in (ALL_TIME + '&status=SYNTHETIC-NO-SUCH-STATUS',
                      ALL_TIME + '&tier=SYNTHETIC-NO-SUCH-TIER',
                      ALL_TIME + '&status=%D0%BC%D1%83%D1%81%D0%BE%D1%80'
                                 '&tier=42',
                      ALL_TIME + '&unit_id=not-a-number',
                      '?date_from=31-02-2026&date_to=nonsense'):
            cards = self.stat_cards(self.page(query))
            self.assertEqual(cards['DJI RAW, га']['muted'], '2 записей',
                             query)
        # Отрицательный контроль: известное значение ФИЛЬТРУЕТ.
        self.assertEqual(
            self.stat_cards(self.page(
                ALL_TIME + '&tier=' + dji_field.TIER1_EXACT)
            )['DJI RAW, га']['muted'], '1 записей')

    def test_a_machine_that_does_not_exist_empties_the_page_without_a_500(self):
        cards = self.stat_cards(self.page(ALL_TIME + '&unit_id=999999'))
        self.assertEqual(cards['DJI RAW, га']['muted'], '0 записей')
        self.assertEqual(cards['DJI RAW, га']['value'], '0.00')
        # Отрицательный контроль: существующая машина строки показывает.
        self.assertEqual(
            self.stat_cards(self.page(
                ALL_TIME + '&unit_id=%d' % self.unit6_id)
            )['DJI RAW, га']['muted'], '1 записей')

    def test_an_inverted_window_is_empty_and_not_an_error(self):
        cards = self.stat_cards(
            self.page('?date_from=2026-06-30&date_to=2026-06-01'))
        self.assertEqual(cards['DJI RAW, га']['muted'], '0 записей')


# ─── 6b. Незнакомые коды ─────────────────────────────────────────────────────

class UnknownCodesSurvive(Base):
    """Код, которого страница не знает, не роняет её и не идёт в проверенное.

    [REASON]: версия модели переживёт эту страницу. Новый статус или новое
    право на агрегирование появятся раньше, чем подпись к ним, и запись с
    таким кодом обязана остаться ВИДИМОЙ и НЕСОСЧИТАННОЙ, а не исчезнуть и
    не попасть в проверенный подытог тихо.
    """

    def setUp(self):
        Base.setUp(self)
        self.add_calc(900031, dji_resolver.RAW_CORROBORATED,
                      dji_resolver.AGG_CERTIFIED, raw_m2=10000.0,
                      corrected_m2=10000.0)
        self.add_attr(900031, dji_field.TIER1_EXACT)
        self.add_calc(900032, 'SYNTHETIC_FUTURE_STATUS',
                      'SYNTHETIC_FUTURE_ELIGIBILITY', raw_m2=50000.0,
                      corrected_m2=50000.0, minute=5)
        self.add_attr(900032, 'SYNTHETIC_FUTURE_TIER')

    def test_the_unknown_row_stays_visible_and_out_of_the_certified_sum(self):
        html = self.page(WINDOW + '&detail=1')
        cards = self.stat_cards(html)

        # RAW показан целиком: незнакомая запись из выборки не выпала.
        self.assertEqual(cards['DJI RAW, га']['value'], '6.00')
        self.assertEqual(cards['DJI RAW, га']['muted'], '2 записей')
        # Но её оценка НЕ проверена и НЕ предварительна.
        self.assertEqual(cards['Проверено, га']['value'], '1.00')
        self.assertEqual(cards['Предварительно, га']['value'], '0.00')
        self.assertEqual(cards['Недостаточно данных']['value'], '1')
        self.assertIn('5.00', cards['Недостаточно данных']['muted'])
        # Отрицательный контроль: 6.00 -- это НЕ проверенный подытог.
        self.assertNotEqual(cards['Проверено, га']['value'], '6.00')

        rows = self.rows_by_flight(html)
        # Незнакомый статус показан своим кодом -- он ASCII и информативнее
        # пустоты; незнакомый tier читается как «поле не определено».
        self.assertIn('SYNTHETIC_FUTURE_STATUS', self.badges(rows[900032])[0])
        self.assertEqual(self.badges(rows[900032])[1], 'Поле не определено')

    def test_a_malformed_flags_json_does_not_break_the_row(self):
        with app.app_context():
            row = DjiAreaCalculation.query.filter_by(flight_id=900032).one()
            row.anomaly_flags_json = 'not a json list'
            db.session.commit()
        rows = self.rows_by_flight(self.page(WINDOW + '&detail=1'))
        self.assertIn(900032, rows)
        self.assertEqual(
            self.raw('SELECT anomaly_flags_json '
                     '  FROM dji_area_calculations WHERE flight_id = ?',
                     (900032,)), [('not a json list',)])


# ─── 7. Детализация: предел, обрезка и приватность ───────────────────────────

class DetailTable(Base):
    """500 строк на экране, остальное -- в Excel; и ничего лишнего в HTML."""

    def test_the_detail_is_capped_and_the_notice_appears(self):
        cap = drones.DRONE_AREA_MAX_DETAIL_ROWS
        self.assertEqual(cap, 500)
        self.add_calcs([
            dict(flight_id=910000 + i,
                 status=dji_resolver.RAW_CORROBORATED,
                 eligibility=dji_resolver.AGG_CERTIFIED,
                 raw_m2=10000.0, corrected_m2=10000.0,
                 minute=i % 600)
            for i in range(cap + 5)])

        html = self.page(WINDOW + '&detail=1')
        self.assertEqual(len(self.record_rows(html)), cap)
        self.assertIn('Показаны первые', html)
        self.assertIn('Сузьте период или выберите машину', html)
        self.assertIn('%d' % (cap + 5), html)
        self.assertEqual(
            self.raw('SELECT COUNT(*) FROM dji_area_calculations'),
            [(cap + 5,)])

    def test_a_short_list_is_shown_whole_without_the_notice(self):
        """Отрицательный контроль к обрезке."""
        self.seed_buckets()
        html = self.page(WINDOW + '&detail=1')
        self.assertEqual(len(self.record_rows(html)), len(BUCKET_ROWS))
        self.assertNotIn('Показаны первые', html)

    def test_the_records_are_hidden_behind_a_link_by_default(self):
        self.seed_buckets()
        html = self.page(WINDOW)
        self.assertEqual(self.record_rows(html), [])
        self.assertIn('Показать записи', html)
        self.assertIn('detail=1', html)

    def test_a_single_machine_opens_its_records_by_itself(self):
        self.seed_buckets()
        html = self.page(WINDOW + '&unit_id=%d' % self.unit6_id)
        self.assertEqual(len(self.record_rows(html)), len(BUCKET_ROWS))

    def test_the_audit_key_is_shown_and_nothing_else_is(self):
        """Идентификатор вылета DJI -- намеренно; всё остальное -- нет."""
        with app.app_context():
            db.session.add(DroneFlight(
                dji_flight_id=900001, drone_unit_id=self.unit6_id,
                nickname_raw='SYNTHETIC-NICK',
                started_at=datetime(2026, 6, 5, 3, 0), area_ha=3.0,
                raw_json=json.dumps({'lat': 39.7, 'lng': 64.4,
                                     'token': TOKEN})))
            db.session.commit()
        self.seed_buckets()
        html = self.page(WINDOW + '&detail=1')

        # Ключ аудита ЕСТЬ -- по нему владелец открывает запись в SmartFarm.
        self.assertIn('900001', html)
        self.assertIn('Вылет DJI', html)

        # Ни координат, ни тела источника, ни токена, ни борта целиком.
        for marker in ('39.7', '64.4', '39,7', '64,4', TOKEN, HW6,
                       'points_json', 'point_shape_census', 'coordinates',
                       'geometry_geojson', 'geometry_md5', 'raw_json',
                       'window_reasons_json', 'anomaly_flags_json',
                       'bridge_flight_ids_json', 'linked_land_uuid',
                       'calculation_input_hash', 'provider_account_id',
                       PROVIDER):
            self.assertNotIn(marker, html, marker)

        # Отрицательный контроль: страница действительно что-то показала.
        self.assertIn('Площадь DJI: техническая оценка', html)
        self.assertIn('27.00', html)


# ─── 8. Два языка ────────────────────────────────────────────────────────────

class Bilingual(Base):
    """Русский и узбекский; узбекский -- только кириллицей."""

    TEMPLATE = 'area_evidence.html'
    # Латиница разрешена только в названиях и в имени файла инструмента,
    # который на странице назван как команда запуска пересчёта.
    ALLOWED_LATIN = ('Excel', 'DJI', 'RAW', 'RFID',
                     'tools/dji_area_recalc.py')

    RU_TITLE = 'Площадь DJI: техническая оценка'
    UZ_TITLE = 'DJI майдони: техник баҳолаш'

    def source(self):
        with open(os.path.join(REPO_ROOT, 'templates', 'drones',
                               self.TEMPLATE), encoding='utf-8') as handle:
            return handle.read()

    def setUp(self):
        Base.setUp(self)
        self.seed_buckets()

    def test_each_language_shows_its_own_heading_and_not_the_other(self):
        ru = self.page(WINDOW + '&detail=1', language='ru')
        self.assertIn(self.RU_TITLE, ru)
        self.assertNotIn(self.UZ_TITLE, ru)
        self.assertIn('Недостаточно данных', ru)

        uz = self.page(WINDOW + '&detail=1', language='uz')
        self.assertIn(self.UZ_TITLE, uz)
        self.assertNotIn(self.RU_TITLE, uz)
        self.assertIn('Маълумот етарли эмас', uz)
        self.assertNotIn('Недостаточно данных', uz)

    def test_the_uzbek_status_and_tier_words_reach_the_page(self):
        uz = self.page(WINDOW + '&detail=1', language='uz')
        self.assertIn('Бошланиши ўлчанмаган', uz)   # BASELINE_UNKNOWN
        self.assertIn('Такрорий/кўчирилган ёзув', uz)
        self.assertIn('Дала тасдиқланган', uz)
        # Отрицательный контроль: русские подписи тех же кодов ушли.
        self.assertNotIn('Начало не измерено', uz)
        self.assertNotIn('Поле подтверждено', uz)

    @staticmethod
    def latin_runs(text, allowed=()):
        for word in allowed:
            text = text.replace(word, ' ')
        return re.findall(r'[A-Za-z]+', text)

    def test_every_uzbek_branch_of_the_template_is_cyrillic(self):
        offenders = []
        halves = re.findall(r"if is_ru else '((?:[^'\\]|\\.)*)'",
                            self.source())
        for text in halves:
            runs = self.latin_runs(text, self.ALLOWED_LATIN)
            if runs:
                offenders.append((text, runs))
        self.assertEqual(offenders, [])
        self.assertGreater(len(halves), 40,
                           'only %d uzbek halves found' % len(halves))

    def test_the_uzbek_scan_fires_on_a_planted_latin_string(self):
        """Сканер, который никогда не срабатывает, -- не сканер."""
        self.assertEqual(self.latin_runs('Дала тасдиқланган',
                                         self.ALLOWED_LATIN), [])
        self.assertEqual(self.latin_runs('Дала tasdiqlangan',
                                         self.ALLOWED_LATIN),
                         ['tasdiqlangan'])
        self.assertEqual(self.latin_runs('Excel DJI RAW',
                                         self.ALLOWED_LATIN), [])

    def test_the_uzbek_letters_really_are_cyrillic_code_points(self):
        self.assertEqual([ord(c) for c in 'Ҳўқғ'],
                         [0x04B2, 0x045E, 0x049B, 0x0493])
        self.assertIn('техник баҳолаш', self.source())

    def test_no_tojson_sits_in_a_double_quoted_attribute(self):
        for match in re.finditer(r'="[^"]*\|\s*tojson', self.source()):
            self.fail(match.group(0))


# ─── 9. Excel ────────────────────────────────────────────────────────────────

class Workbook(Base):
    """Книга -- та копия, которую суммируют. Ноль в ней стоит дороже всего."""

    def setUp(self):
        Base.setUp(self)
        self.seed_buckets()

    def test_the_sheets_and_the_file_name_carry_the_period(self):
        response, wb = self.workbook(WINDOW)
        self.assertEqual(wb.sheetnames,
                         ['Сводка', 'По дням и машинам', 'Записи'])
        disposition = response.headers.get('Content-Disposition', '')
        self.assertIn('drone_area_evidence_2026-06-01_2026-06-30.xlsx',
                      disposition)
        # Отрицательный контроль: без периода имя другое.
        other, _wb = self.workbook(ALL_TIME)
        self.assertIn('drone_area_evidence_all.xlsx',
                      other.headers.get('Content-Disposition', ''))

    def records_sheet(self, wb):
        ws = wb['Записи']
        header = [cell.value for cell in ws[1]]
        rows = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows[row[0]] = row
        return header, rows

    def test_a_null_estimate_is_an_empty_cell_and_a_measured_zero_is_zero(self):
        _response, wb = self.workbook(WINDOW)
        header, rows = self.records_sheet(wb)
        col = header.index('Техническая оценка, м²')

        null_cell = rows[900003][col]      # BASELINE_UNKNOWN
        zero_cell = rows[900005][col]      # COUNTER_FLAT, измеренный ноль
        real_cell = rows[900001][col]      # обычная оценка

        self.assertIsNone(null_cell)
        self.assertNotIsInstance(null_cell, str)
        self.assertNotEqual(null_cell, 0)

        self.assertEqual(zero_cell, 0.0)
        self.assertIsNotNone(zero_cell)
        self.assertNotIsInstance(zero_cell, str)

        # Отрицательный контроль: колонка вообще несёт числа.
        self.assertEqual(real_cell, 30000.0)

    def test_billable_area_exists_and_is_empty_in_every_row(self):
        _response, wb = self.workbook(WINDOW)
        header, rows = self.records_sheet(wb)
        self.assertIn('billable_area_m2', header)
        col = header.index('billable_area_m2')
        self.assertEqual(len(rows), len(BUCKET_ROWS))
        for flight_id, row in rows.items():
            self.assertIsNone(row[col], flight_id)
        # Отрицательный контроль: соседняя колонка версий заполнена.
        version_col = header.index('area_algorithm_version')
        for flight_id, row in rows.items():
            self.assertEqual(row[version_col],
                             dji_area.AREA_ALGORITHM_VERSION, flight_id)

    def test_the_summary_names_the_versions_and_keeps_the_buckets_apart(self):
        _response, wb = self.workbook(WINDOW)
        summary = {}
        for label, value in wb['Сводка'].iter_rows(min_row=2,
                                                   values_only=True):
            summary[label] = value

        self.assertEqual(summary['Версия модели'], dji_area.MODEL_VERSION)
        self.assertEqual(summary['Алгоритм площади'],
                         dji_area.AREA_ALGORITHM_VERSION)
        self.assertEqual(summary['Резолвер поля'],
                         dji_area.FIELD_RESOLVER_VERSION)
        self.assertEqual(summary['Период: с'], '2026-06-01')
        self.assertEqual(summary['Период: по'], '2026-06-30')

        self.assertEqual(summary['Записей'], 6)
        self.assertEqual(summary['DJI RAW, га'], 27.0)
        self.assertEqual(summary['Записей без RAW'], 1)
        self.assertEqual(summary['Проверено, га'], 3.0)
        self.assertEqual(summary['Предварительно, га'], 2.0)
        self.assertEqual(
            summary['Известная часть (проверено + предварительно), га'], 5.0)
        self.assertEqual(summary['Недостаточно данных, записей'], 2)
        self.assertEqual(summary['Недостаточно данных, RAW га'], 11.0)
        self.assertEqual(summary['Пересечение записей, записей'], 1)
        # Пересечение сохраняет свою RAW-экспозицию отдельной строкой.
        self.assertEqual(summary['Пересечение записей, RAW га'], 7.0)
        # Отрицательный контроль: полного итога в книге нет.
        self.assertNotIn(23.0, list(summary.values()))

    def test_the_workbook_carries_no_coordinates_and_no_source_bytes(self):
        _response, wb = self.workbook(WINDOW)
        text = []
        for name in wb.sheetnames:
            for row in wb[name].iter_rows(values_only=True):
                text.extend(str(cell) for cell in row if cell is not None)
        blob = '\n'.join(text)
        for marker in ('points_json', 'coordinates', 'geometry_geojson',
                       'window_reasons_json', TOKEN, '39.7', '64.4'):
            self.assertNotIn(marker, blob, marker)
        # Отрицательный контроль: hardware_id в книге ЕСТЬ намеренно --
        # это колонка аудита, и её отсутствие означало бы пустую книгу.
        self.assertIn(HW6, blob)


# ─── 10. Сверка с агрегатором ────────────────────────────────────────────────

class AggregationParity(Base):
    """Страница обязана показывать то, что даёт `dji_area.aggregate`."""

    def setUp(self):
        Base.setUp(self)
        self.expected_rows = []
        for (flight_id, status, eligibility, raw_m2, corrected, tier,
             day_index, hardware_id, without_area) in PARITY_SPEC:
            day = DAY if day_index == 0 else DAY2
            self.add_calc(flight_id, status, eligibility, raw_m2=raw_m2,
                          corrected_m2=corrected, day=day,
                          hardware_id=hardware_id,
                          minute=(flight_id - 900101) * 3,
                          application_without_area=without_area or None,
                          channel=(dji_resolver.CH_UNRELIABLE
                                   if status == dji_resolver.CHANNEL_MISSING
                                   else None))
            if tier != dji_field.TIER5_UNKNOWN:
                self.add_attr(flight_id, tier,
                              name='SYNTHETIC field %d' % flight_id)
            self.expected_rows.append({
                'flight_id': flight_id,
                'report_start_date': day,
                'hardware_id': hardware_id,
                'raw_area_m2': raw_m2,
                'corrected_recorded_area_m2': corrected,
                'controller_delta_area_m2': None,
                'area_status': status,
                'aggregation_eligibility': eligibility,
                'application_without_area': without_area,
                'application_channel_quality': (
                    dji_resolver.CH_UNRELIABLE
                    if status == dji_resolver.CHANNEL_MISSING else None),
                'field_attribution_tier': tier,
            })

    def expected(self):
        buckets = dji_aggregate.aggregate(
            self.expected_rows,
            lambda r: (r['report_start_date'], r['hardware_id']))
        return buckets, buckets['__total__']

    def test_the_page_totals_equal_the_aggregator_run_on_the_same_rows(self):
        buckets, total = self.expected()
        self.assertEqual(total['records'], len(PARITY_SPEC))
        self.assertTrue(total['eligibility_partition_holds'])
        self.assertTrue(total['tier_partition_holds'])
        # Корзины непустые -- иначе сверка ниже была бы сверкой нулей.
        for key in ('certified_records', 'provisional_records',
                    'unresolved_records', 'overlap_records',
                    'application_without_area_records', 'unassigned_records'):
            self.assertGreater(total[key], 0, key)

        cards = self.stat_cards(self.page(WINDOW))
        self.assertEqual(cards['DJI RAW, га']['value'],
                         ha(total['raw_sum_m2']))
        self.assertEqual(cards['DJI RAW, га']['muted'],
                         '%d записей' % total['records'])
        self.assertEqual(cards['Проверено, га']['value'],
                         ha(total['certified_sum_m2']))
        self.assertEqual(cards['Проверено, га']['muted'],
                         '%d записей' % total['certified_records'])
        self.assertEqual(cards['Предварительно, га']['value'],
                         ha(total['provisional_sum_m2']))
        self.assertEqual(cards['Предварительно, га']['muted'],
                         '%d записей' % total['provisional_records'])
        self.assertEqual(cards['Известная часть, га']['value'],
                         ha(total['known_subtotal_m2']))
        self.assertEqual(cards['Недостаточно данных']['value'],
                         str(total['unresolved_records']))
        self.assertIn(ha(total['unresolved_raw_exposure_m2']),
                      cards['Недостаточно данных']['muted'])
        self.assertIn('пересечений записей: %d' % total['overlap_records'],
                      cards['Недостаточно данных']['muted'])
        self.assertEqual(
            cards['Повторные/перенесённые']['value'],
            str(total['status_counts'][
                dji_resolver.COUNTER_FLAT_RAW_OVERSTATED]))
        self.assertEqual(
            cards['Площадь не измерена при активности']['value'],
            str(total['application_without_area_records']))
        self.assertEqual(cards['Поле не определено']['value'],
                         ha(total['unassigned_raw_m2']))

    def test_the_page_does_not_show_the_arithmetic_sum_of_everything(self):
        _buckets, total = self.expected()
        html = self.page(WINDOW)
        everything = (total['certified_sum_m2'] + total['provisional_sum_m2']
                      + total['unresolved_raw_exposure_m2']
                      + total['overlap_raw_exposure_m2'])
        self.assertGreater(everything, total['known_subtotal_m2'])
        self.assertNotIn(ha(everything), html)
        # Отрицательный контроль: известная часть НА странице есть.
        self.assertIn(ha(total['known_subtotal_m2']), html)

    def test_every_day_and_machine_bucket_matches_too(self):
        buckets, total = self.expected()
        keys = [k for k in buckets if k != '__total__']
        self.assertEqual(len(keys), 4)   # два дня * две машины

        rows = self.day_rows(self.page(WINDOW))
        self.assertEqual(len(rows), 4)
        seen = {}
        for row in rows:
            cells = [re.sub(r'<[^>]+>', '', cell).strip() for cell in
                     re.findall(r'<td[^>]*>(.*?)</td>', row, re.S)]
            seen[(cells[0], cells[1])] = cells

        labels = {HW6: '№ 6', HW7: '№ 7'}
        for key in keys:
            day, hardware_id = key
            bucket = buckets[key]
            cells = seen[(day.strftime('%d.%m.%Y'), labels[hardware_id])]
            self.assertEqual(cells[2], str(bucket['records']))
            self.assertEqual(cells[3], ha(bucket['raw_sum_m2']))
            self.assertEqual(cells[4], ha(bucket['certified_sum_m2']))
            self.assertEqual(cells[5], ha(bucket['provisional_sum_m2']))
            self.assertEqual(cells[6], str(bucket['unresolved_records']))
            self.assertEqual(cells[9], str(bucket['unassigned_records']))

        # Сумма по строкам дней равна итогу -- корзины не потеряны.
        self.assertEqual(sum(int(c[2]) for c in seen.values()),
                         total['records'])


# ─── 11. Лаунчер и полоса разделов ───────────────────────────────────────────

class HubAndNav(Base):

    def test_the_hub_carries_the_tile_and_it_answers(self):
        client = self.client_as()
        html = client.get('/drones/reports').get_data(as_text=True)
        self.assertIn('href="/drones/area-evidence"', html)
        self.assertIn('Площадь DJI: техническая оценка', html)
        self.assertEqual(client.get('/drones/area-evidence').status_code, 200)
        # Плитка объявлена одним источником -- списком в drones.py.
        keys = [tile['key'] for tile in drones.DRONE_REPORT_TILES]
        self.assertIn('area-evidence', keys)
        self.assertEqual(keys.count('area-evidence'), 1)

    def test_the_reports_pill_is_active_on_the_new_page(self):
        html = self.page()
        self.assertRegex(
            html,
            r'<a href="/drones/reports"\s+class="vs-pill is-active"')
        # Отрицательный контроль: на «Вылетах» активна другая пилюля.
        other = self.client_as().get('/drones/').get_data(as_text=True)
        self.assertNotRegex(
            other,
            r'<a href="/drones/reports"\s+class="vs-pill is-active"')

    def test_the_coverage_page_points_at_the_new_report(self):
        html = self.client_as().get('/drones/coverage').get_data(as_text=True)
        self.assertIn('/drones/area-evidence', html)
        self.assertIn('Это геометрическая диагностика покрытия', html)
        self.assertIn('«Площадь DJI: техническая оценка»', html)


if __name__ == '__main__':
    unittest.main()
