# -*- coding: utf-8 -*-
"""FIN-RECON-STORAGE-001: экран «Дроны -> Финансовая сверка» и выгрузка.

Настоящее приложение на выбрасываемой базе (tests/harness.py), снимок
кладётся настоящей миграцией и настоящим импортёром. Держится здесь:

* доступ: аноним -- на вход; оператор, наблюдатель и механик -- 403, В ТОМ
  ЧИСЛЕ с правом на модуль «Дроны»; администратор -- 200. Ссылка в полосе
  модуля -- только у администратора;
* статус «Сверен с исключениями» рядом с «Финансово НЕ утверждён»; ни
  зелёного «касса сходится», ни кнопок утверждения, ни формы загрузки;
* неизвестная зарплата -- словом «неизвестно», доказательный тип у сумм;
* принятая площадь DJI подписана технической, не основание для счёта;
* 38 работ, 7 операторов, 39 аномалий, 2 межмесячные связи синтетики;
* RU и UZ -- каждый на своём языке;
* текст книг экранирован (XSS), в выгрузке не становится формулой, в
  странице и выгрузке нет путей файловой системы;
* без миграции -- предупреждение, а не падение; ответы не кэшируются.

ВСЕ ДАННЫЕ СИНТЕТИЧЕСКИЕ. Запуск:
  python -m unittest tests.test_fin_recon_screen -v
"""

import io
import os
import re
import shutil
import sys
import tempfile
import unittest

from contextlib import redirect_stdout

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from tests.harness import (app, create_admin, login, reset_db,  # noqa: E402
                           TEST_DB_PATH)
from models import (db, User, UserModulePermission, ROLE_MECHANIC,  # noqa: E402
                    ROLE_OPERATOR, ROLE_VIEWER)

import migration_utils  # noqa: E402
import migrate_fin_recon_storage_001 as mig  # noqa: E402
from fin_recon import importer, package  # noqa: E402
from tests import fin_recon_fixture as fx  # noqa: E402

LIST = '/drones/fin-recon'
PAGE = '/drones/fin-recon/%s' % fx.PERIOD
EXPORT = '/drones/fin-recon/%s/export.xlsx' % fx.PERIOD
ROUTES = (LIST, PAGE, EXPORT)
PATH_RE = re.compile(r'[A-Za-z]:\\\\|[A-Za-z]:\\[A-Za-z]|/tmp/|/home/|'
                     r'vehicle_soft_tests_|fin_recon_scr_')


def drop_fin_recon():
    """reset_db() пересоздаёт только таблицы ORM; у fin_recon_* модели нет."""
    import sqlite3
    from fin_recon import schema
    con = sqlite3.connect(TEST_DB_PATH)
    try:
        for table in reversed(schema.TABLE_NAMES):
            con.execute('DROP TABLE IF EXISTS %s' % table)
        con.commit()
    finally:
        con.close()


def migrate():
    saved = migration_utils.DB_PATH
    try:
        with redirect_stdout(io.StringIO()):
            assert mig.run(TEST_DB_PATH) == 0
    finally:
        migration_utils.DB_PATH = saved


class ScreenBase(unittest.TestCase):

    install = True

    def setUp(self):
        reset_db()
        drop_fin_recon()
        self.tmp = tempfile.mkdtemp(prefix='fin_recon_scr_')
        self.admin_id = create_admin()
        if self.install:
            migrate()
            paths = fx.write_package(os.path.join(self.tmp, 'pkg'))
            importer.import_package(TEST_DB_PATH,
                                    package.verify(paths['root']),
                                    imported_by='screen-test')

    def tearDown(self):
        drop_fin_recon()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def user(self, role, drones=False, name='u'):
        with app.app_context():
            u = User(username='%s_%s_%s' % (name, role, drones), role=role,
                     full_name='SYNTHETIC', language='ru')
            u.set_password('test-password')
            db.session.add(u)
            db.session.commit()
            if drones:
                db.session.add(UserModulePermission(
                    user_id=u.id, module_code='drones', has_access=True))
                db.session.commit()
            return u.id

    def client_as(self, user_id=None, language='ru'):
        user_id = self.admin_id if user_id is None else user_id
        with app.app_context():
            u = db.session.get(User, user_id)
            u.language = language
            db.session.commit()
        client = app.test_client()
        login(client, user_id)
        return client

    def page(self, language='ru', url=PAGE):
        r = self.client_as(language=language).get(url)
        self.assertEqual(r.status_code, 200, url)
        return r.get_data(as_text=True)


class Access(ScreenBase):

    def test_anonymous_goes_to_login(self):
        client = app.test_client()
        for url in ROUTES:
            r = client.get(url)
            self.assertEqual(r.status_code, 302, url)
            self.assertIn('/login', r.headers['Location'])

    def test_non_admin_roles_get_403_even_with_drones_access(self):
        for role in (ROLE_OPERATOR, ROLE_VIEWER, ROLE_MECHANIC):
            for drones in (False, True):
                uid = self.user(role, drones)
                client = self.client_as(uid)
                for url in ROUTES:
                    with self.subTest(role=role, drones=drones, url=url):
                        r = client.get(url)
                        self.assertEqual(r.status_code, 403)
                        body = r.get_data(as_text=True)
                        self.assertNotIn('Синтетик Оператор', body)
                        self.assertNotIn('1499900', body)

    def test_admin_gets_200(self):
        client = self.client_as()
        for url in ROUTES:
            self.assertEqual(client.get(url).status_code, 200, url)

    def test_nav_link_only_for_admin(self):
        self.assertIn('Финансовая сверка', self.page(url='/drones/summary'))
        uid = self.user(ROLE_OPERATOR, drones=True)
        r = self.client_as(uid).get('/drones/summary')
        self.assertEqual(r.status_code, 200)
        self.assertNotIn('Финансовая сверка', r.get_data(as_text=True))
        self.assertNotIn('/drones/fin-recon', r.get_data(as_text=True))

    def test_not_cached(self):
        client = self.client_as()
        for url in ROUTES:
            self.assertEqual(client.get(url).headers.get('Cache-Control'),
                             'no-store', url)

    def test_bad_period_and_version_are_404(self):
        client = self.client_as()
        for url in ('/drones/fin-recon/2031-13', '/drones/fin-recon/x',
                    '/drones/fin-recon/2030-01', PAGE + '?version=7'):
            self.assertEqual(client.get(url).status_code, 404, url)

    def test_no_write_routes(self):
        client = self.client_as()
        for url in ROUTES:
            self.assertIn(client.post(url).status_code, (400, 405), url)


class Content(ScreenBase):

    def test_status_is_reconciled_with_exceptions_and_not_approved(self):
        html = self.page()
        self.assertIn('Сверен с исключениями', html)
        self.assertIn('Финансово НЕ утверждён', html)
        self.assertIn('НЕ закрыт', html)
        self.assertIn('не утверждено', html)
        for green in ('касса сходится', 'долгов нет', 'Утвердить</button',
                      'Закрыть месяц</button'):
            self.assertNotIn(green, html)
        # Формы общего шаблона (выход, язык) -- не этого экрана; у самого
        # раздела нет ни одной формы, и ничто не отправляется в его адреса.
        self.assertNotRegex(html, r'action="[^"]*/drones/fin-recon')
        self.assertNotIn('type="file"', html)
        for name in ('fin_recon_period.html', 'fin_recon_periods.html'):
            with open(os.path.join(REPO_ROOT, 'templates', 'drones', name),
                      encoding='utf-8') as fh:
                source = fh.read().lower()
            self.assertNotIn('<form', source, name)
            self.assertNotIn('<button', source, name)

    def test_counts_and_sections(self):
        html = self.page()
        for anchor in ('id="totals"', 'id="operators"', 'id="works"',
                       'id="anomalies"', 'id="sources"', 'id="cross"'):
            self.assertIn(anchor, html)
        self.assertIn('Работы</a>', html.replace(' (38)', ''))
        self.assertIn('(38)', html)
        self.assertIn('(7)', html)
        self.assertIn('(39)', html)
        self.assertEqual(len(re.findall(r'<b>W\d\d</b>', html)), 38)
        self.assertEqual(len(re.findall(r'<td><b>A\d\d</b>', html)), 39)

    def test_unknown_is_a_word_not_zero(self):
        html = self.page()
        self.assertIn('неизвестно', html)
        self.assertIn('зарплата неизвестна', html)
        self.assertIn('га без начисления', html)
        row = re.search(r'Зарплата за|salary for the spravka works(.*?)</tr>',
                        html, re.S)
        self.assertIsNotNone(row)
        self.assertNotRegex(row.group(0), r'>\s*0(\.0+)?\s*<')

    def test_evidence_types_are_shown(self):
        html = self.page()
        for label in ('записано текстом', 'трактовка владельца',
                      'расчёт (формула книги)', 'расчёт (реестр сверки)',
                      'выгрузка системы', 'неизвестно'):
            self.assertIn(label, html)

    def test_dji_accepted_is_technical(self):
        html = self.page()
        self.assertIn('ТЕХНИЧЕСКАЯ: не основание для счёта заказчику', html)
        self.assertIn('880.0034', html)

    def test_open_questions_come_from_open_anomalies(self):
        html = self.page()
        opened = sum(1 for a in fx.build_snapshot()['anomalies']
                     if a['status'] == 'ОТКРЫТО')
        self.assertIn('Открытые хозяйственные вопросы (%d)' % opened, html)

    def test_cross_period_and_no_candidates_yet(self):
        html = self.page()
        self.assertIn(fx.PREV_MONTH, html)
        self.assertIn(fx.NEXT_MONTH, html)
        self.assertIn('Кандидаты в межмесячный дубль (0)', html)

    def test_uzbek_page_is_uzbek(self):
        uz = self.page(language='uz')
        self.assertIn('Истиснолар билан солиштирилган', uz)
        self.assertIn('Молиявий жиҳатдан ТАСДИҚЛАНМАГАН', uz)
        self.assertIn('номаълум', uz)
        self.assertNotIn('Финансово НЕ утверждён', uz)
        self.assertNotIn('Открытые хозяйственные вопросы', uz)
        ru = self.page()
        self.assertNotIn('Очиқ хўжалик саволлари', ru)

    def test_list_page(self):
        html = self.page(url=LIST)
        self.assertIn(fx.PERIOD, html)
        self.assertIn('v1', html)
        self.assertIn('Сверен с исключениями', html)


class Safety(ScreenBase):

    def test_book_text_is_escaped(self):
        html = self.page()
        self.assertNotIn(fx.HOSTILE_FARM, html)
        self.assertNotIn('<script>alert', html)
        self.assertIn('&lt;script&gt;alert', html)
        self.assertNotIn('<b>bold</b> &', html)
        self.assertIn('&lt;b&gt;bold&lt;/b&gt;', html)

    def test_no_filesystem_paths(self):
        client = self.client_as()
        for url in (LIST, PAGE):
            body = client.get(url).get_data(as_text=True)
            self.assertIsNone(PATH_RE.search(body), url)

    def workbook(self):
        from openpyxl import load_workbook
        r = self.client_as().get(EXPORT)
        self.assertEqual(r.status_code, 200)
        self.assertIn('not_source_of_truth', r.headers['Content-Disposition'])
        return load_workbook(io.BytesIO(r.data))

    def test_export_marks_itself_and_has_no_formulas(self):
        wb = self.workbook()
        self.assertEqual(len(wb.sheetnames), 6)
        for ws in wb.worksheets:
            self.assertIn('НЕ ИСТОЧНИК ИСТИНЫ', ws['A1'].value)
            for row in ws.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'f',
                                        '%s!%s' % (ws.title, cell.coordinate))
                    if isinstance(cell.value, str):
                        self.assertIsNone(PATH_RE.search(cell.value))

    def test_export_neutralises_formula_text(self):
        wb = self.workbook()
        texts = [c.value for ws in wb.worksheets for row in ws.iter_rows()
                 for c in row if isinstance(c.value, str)]
        hostile = [t for t in texts if 'HYPERLINK' in t or t.startswith(
            "'=1+1")]
        self.assertTrue(hostile)
        for t in texts:
            self.assertFalse(t.lstrip(' ')[:1] in ('=', '+', '@'), t)

    def test_export_numbers_are_exact(self):
        wb = self.workbook()
        ws = wb['Итоги']
        values = {row[1].value: row[2].value for row in ws.iter_rows(min_row=3)}
        self.assertEqual(str(values['accrued']), '8170000.57')
        self.assertEqual(values['salary for the spravka works'], 'неизвестно')


class NotInstalled(ScreenBase):

    install = False

    def test_list_warns_and_page_404(self):
        html = self.page(url=LIST)
        self.assertIn('не установлено', html)
        self.assertEqual(self.client_as().get(PAGE).status_code, 404)


if __name__ == '__main__':
    unittest.main()
