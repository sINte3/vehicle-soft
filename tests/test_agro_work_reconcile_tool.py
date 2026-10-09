# -*- coding: utf-8 -*-
"""agro-work B3: отчёт сверки файлом -- числа книги равны числам ядра.

Отчёт появляется раньше экрана (экран уедет с релизом), и владелец будет
принимать по нему решения, в том числе о N. Поэтому держится:

  * каждое число листа «Свод» -- то же, что у ядра сверки на тех же данных;
  * строк на листах «Заявка-работа» и «Работа-заявка» -- столько же, сколько
    строк у ядра;
  * база не меняется ни на байт (открыта mode=ro);
  * предпросмотр N помечен в книге как неутверждённый.

Запуск: python -m unittest tests.test_agro_work_reconcile_tool -v
"""

import hashlib
import io
import os
import sqlite3
import sys
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import date
from unittest import mock

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from openpyxl import load_workbook                            # noqa: E402

from agro_work import reconcile as rc                         # noqa: E402
from tests import agro_work_db as dbh                         # noqa: E402
from tests.test_agro_work_reconcile import Fixture            # noqa: E402
from tools import agro_work_reconcile as tool                 # noqa: E402


def sha(path):
    with open(path, 'rb') as fh:
        return hashlib.sha256(fh.read()).hexdigest()


def book_of(path):
    with open(path, 'rb') as fh:
        return load_workbook(fh)


class ReportTool(unittest.TestCase):
    def setUp(self):
        self.fx = Fixture()
        self.fx.app(created=10, completed=11)                  # работа была
        self.fx.app(created=11, completed=13)                  # работы не было
        self.fx.app(transport='T8')                            # не сопоставлена
        self.fx.app(created=12, completed=12, initial='COMPLETED')
        self.fx.con.close()
        self.out = os.path.join(os.path.dirname(self.fx.path), 'report.xlsx')

    def main(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = tool.main(['--db', self.fx.path, '--from', '2026-09-10',
                              '--to', '2026-09-20', '--out', self.out]
                             + list(args))
        return code, out.getvalue(), err.getvalue()

    def test_the_book_repeats_the_core_and_the_database_is_untouched(self):
        before = sha(self.fx.path)
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertTrue(out.isascii())
        self.assertEqual(sha(self.fx.path), before)

        con = sqlite3.connect(self.fx.path)
        try:
            ctx = rc.Reconciliation(con, date(2026, 9, 10), date(2026, 9, 20))
            forward, reverse = ctx.forward_rows(), ctx.reverse_rows()
        finally:
            con.close()
        _, total = ctx.summary(forward, reverse)
        self.assertIn('applications: %d | work confirmed %d | NO WORK %d'
                      % (total['applications'], total['app_rabota_est'],
                         total['app_rabota_net']), out)

        book = book_of(self.out)
        self.assertEqual(book.sheetnames, ['Свод', 'Заявка-работа',
                                           'Работа-заявка', 'Открытые заявки',
                                           'Причины', 'Замер N'])
        summary = [row for row in book['Свод'].iter_rows(values_only=True)]
        totals = [row for row in summary if row[0] == 'Итого / Жами'][0]
        self.assertEqual(totals[2:], (total['machines'], total['applications'],
                                      total['app_rabota_est'],
                                      total['app_rabota_net'],
                                      total['app_bez_verdikta'],
                                      total['work_days'], total['day_pokryta'],
                                      total['day_bez_zayavki'],
                                      total['day_bez_verdikta'],
                                      total[rc.DAY_LATE],
                                      total[rc.DAY_AFTER + rc.A_ORDINARY],
                                      total[rc.DAY_AFTER + rc.A_UNCLEAR],
                                      total[rc.DAY_AFTER + rc.A_NOT_ENTERED],
                                      total[rc.VOL + rc.VOL_OK],
                                      total[rc.VOL + rc.VOL_WARN],
                                      total[rc.VOL + rc.VOL_FAIL],
                                      total[rc.VOL + rc.VOL_NONE]))
        self.assertEqual(book['Заявка-работа'].max_row - 1, len(forward))
        self.assertEqual(book['Работа-заявка'].max_row - 1, len(reverse))
        notes = ' '.join(str(r[0]) for r in summary if r and r[0])
        self.assertIn('N для заявок, заведённых задним числом: 2 '
                      '(утверждён владельцем)', notes)

    def test_the_late_application_and_the_open_list_reach_the_book(self):
        # B4, пункты 4 и 5: заявка задним числом от 18-го для работы 13-го
        # (опоздание 5 суток при N = 2) и открытая заявка от 10-го.
        self.fx.con = sqlite3.connect(self.fx.path)
        self.fx.app(transport='T2', created=18, completed=18,
                    initial='COMPLETED')
        self.fx.app(transport='T1', status='IN_PROGRESS', created=10,
                    completed=None)
        self.fx.con.close()
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertIn('of them with a later backdated application: 1 '
                      '(days late: min 5, max 5)', out)
        self.assertIn('  days late -> machine-days: 5 -> 1\n', out)
        days_open = (date.today() - date(2026, 9, 10)).days
        self.assertIn('open applications: 1 | oldest open %d days since entry'
                      % days_open, out)
        book = book_of(self.out)
        day13 = [row for row in book['Работа-заявка'].iter_rows(values_only=True)
                 if row[0] is not None and str(row[0]).startswith('2026-09-13')
                 and '80 012 EA' in (row[1] or '')][0]
        # Столбцы 9 и 10 -- заявка задним числом и опоздание (B4); столбцы
        # B5 добавлены после них.
        self.assertEqual(day13[8:10], ('N-005', 5))
        opened = list(book['Открытые заявки'].iter_rows(values_only=True))[1:]
        self.assertEqual([(r[2], r[-1]) for r in opened], [('N-006', days_open)])
        totals = [row for row in book['Свод'].iter_rows(values_only=True)
                  if row and row[0] == 'Итого / Жами'][0]
        self.assertEqual(totals[11], 1)

    def test_days_late_are_listed_exactly_not_in_buckets(self):
        # Машина 12: заявка от 18-го опоздала на 5, 4 и 3 суток для 13-го,
        # 14-го и 15-го. Машина 11: заявка от 25-го -- на 9, 8, 7 и 5 суток
        # для 16-го, 17-го, 18-го и 20-го. Значения больше шести нужны
        # затем, что живые заявки опаздывают и на 7-14 суток, а корзина
        # «поздно» -- порог, которого владелец не называл: значения целиком.
        self.fx.con = sqlite3.connect(self.fx.path)
        for day in (14, 15):
            dbh.add_day(self.fx.con, 1002, '2026-09-%02d' % day,
                        sites=[(1.0, None)])
        for day in (16, 17, 18, 20):
            dbh.add_day(self.fx.con, 1001, '2026-09-%02d' % day,
                        sites=[(1.5, None)])
        self.fx.con.commit()
        self.fx.app(transport='T2', created=18, completed=18,
                    initial='COMPLETED')
        self.fx.app(transport='T1', created=25, completed=25,
                    initial='COMPLETED')
        self.fx.con.close()
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertIn('of them with a later backdated application: 7 '
                      '(days late: min 3, max 9)', out)
        self.assertIn('  days late -> machine-days: 3 -> 1 | 4 -> 1 | 5 -> 2 | '
                      '7 -> 1 | 8 -> 1 | 9 -> 1\n', out)

    def test_without_late_applications_there_is_no_distribution_line(self):
        # Строки нет вовсе, а не «0 -> 0»: нечего распределять.
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertIn('of them with a later backdated application: 0\n', out)
        self.assertNotIn('days late ->', out)

    def test_what_came_after_reaches_the_console_and_the_book(self):
        # B5. Машина 12: обычная заявка открыта 15-го -- для работы 13-го и
        # 14-го это 2 и 1 сутки после работы; 16-е она уже может покрыть
        # (без вердикта). Машина 11: работа 20-го, позже заявок нет.
        self.fx.con = sqlite3.connect(self.fx.path)
        for day in (14, 16):
            dbh.add_day(self.fx.con, 1002, '2026-09-%02d' % day,
                        sites=[(1.0, None)])
        dbh.add_day(self.fx.con, 1001, '2026-09-20', sites=[(1.5, None)])
        self.fx.con.commit()
        self.fx.app(transport='T2', status='IN_PROGRESS', created=15,
                    completed=None)
        self.fx.con.close()
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertTrue(out.isascii())
        self.assertIn('WITHOUT APPLICATION 3 ', out)
        self.assertIn('  of them with a later backdated application: 0\n', out)
        self.assertIn('  of them with a later ordinary application: 2 '
                      '(days after: min 1, max 2; still open 2)\n', out)
        self.assertIn('  days after -> machine-days: 1 -> 1 | 2 -> 1\n', out)
        self.assertIn('  of them with no later application at all: 1\n', out)
        self.assertNotIn('unknown entry order', out)
        book = book_of(self.out)
        header = next(book['Работа-заявка'].iter_rows(values_only=True))
        self.assertEqual(header[10], 'Без заявки: что было потом / '
                                     'Буюртмасиз: кейин нима бўлган')
        # Подпись машины -- «название госномер»: ключ -- госномер в ней.
        rows = {(str(r[0])[:10], r[1].split(' ', 1)[1]): r
                for r in book['Работа-заявка'].iter_rows(min_row=2,
                                                         values_only=True)}
        self.assertEqual(rows[('2026-09-13', '80 012 EA')][10:14],
                         ('Позже заведена обычная заявка', 'N-005', 2,
                          'В процессе'))
        self.assertEqual(rows[('2026-09-14', '80 012 EA')][12], 1)
        self.assertEqual(rows[('2026-09-20', '80 011 EA')][10:14],
                         ('Заявка не заведена совсем', None, None, None))
        self.assertEqual(rows[('2026-09-16', '80 012 EA')][10:14],
                         (None, None, None, None))      # без вердикта
        totals = [row for row in book['Свод'].iter_rows(values_only=True)
                  if row and row[0] == 'Итого / Жами'][0]
        # Без заявки -- 3: задним числом 0, обычная позже 2, порядок
        # неизвестен 0, не заведена 1.
        self.assertEqual((totals[9], totals[11:15]), (3, (0, 2, 0, 1)))

    def test_still_open_counts_only_the_open_ones(self):
        # Машина 12: закрытая обычная заявка от 15-го -- для 13-го 2 суток;
        # машина 11: открытая от 22-го -- для 20-го 2 суток. Из двух
        # обычных открыта одна: «все» и «открытые» здесь различимы.
        self.fx.con = sqlite3.connect(self.fx.path)
        dbh.add_day(self.fx.con, 1001, '2026-09-20', sites=[(1.5, None)])
        self.fx.con.commit()
        self.fx.app(transport='T2', created=15, completed=15)
        self.fx.app(transport='T1', status='PENDING', created=22,
                    completed=None)
        self.fx.con.close()
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertIn('  of them with a later ordinary application: 2 '
                      '(days after: min 2, max 2; still open 1)\n', out)

    def test_unknown_entry_order_gets_its_own_console_line(self):
        self.fx.con = sqlite3.connect(self.fx.path)
        self.fx.app(transport='T2', created=18, completed=18, history=False)
        self.fx.con.close()
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertIn('  of them with a later ordinary application: 0\n', out)
        self.assertNotIn('days after ->', out)
        self.assertIn('  of them with a later application of unknown entry '
                      'order: 1\n', out)
        self.assertIn('  of them with no later application at all: 0\n', out)

    def test_the_reasons_in_the_console_are_the_reasons_of_the_book(self):
        # Консоль повторяет лист «Причины» и складывается в итоги, которые
        # напечатаны строкой выше: владелец присылает консоль, не книгу, и
        # по ней я проверяю, что «без вердикта» названо по существу.
        self.fx.con = sqlite3.connect(self.fx.path)
        # Две открытые против одной неопознанной машины: счётчики разные,
        # и порядок «самые частые первыми» отличим от алфавитного и обратного.
        self.fx.app(transport='T1', status='IN_PROGRESS', created=10,
                    completed=None)
        self.fx.app(transport='T1', status='PENDING', created=11,
                    completed=None)
        self.fx.con.close()
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertTrue(out.isascii())

        shown = {}
        for line in out.splitlines():
            for side, title in (('forward', 'no verdict, applications: '),
                                ('reverse', 'no verdict, machine-days: '),
                                ('volume',
                                 '  cannot compare the volume, why: ')):
                if line.startswith(title):
                    shown[side] = {
                        name: int(number) for name, number in (
                            item.rsplit(' ', 1)
                            for item in line[len(title):].split(' | '))}
        sheet = {}
        for row in list(book_of(self.out)['Причины'].iter_rows(
                values_only=True))[1:]:
            side = {'Заявка': 'forward', 'Работа': 'reverse',
                    'Объём': 'volume'}[str(row[0]).split(' ')[0]]
            sheet.setdefault(side, {})[row[2]] = row[3]
        self.assertEqual(shown, sheet)
        # Самые частые причины первыми, при равенстве -- по коду.
        for title in ('no verdict, applications: ',
                      'no verdict, machine-days: '):
            line = [l for l in out.splitlines() if l.startswith(title)][0]
            order = [(item.rsplit(' ', 1)[0], int(item.rsplit(' ', 1)[1]))
                     for item in line[len(title):].split(' | ')]
            self.assertEqual(order, sorted(order,
                                           key=lambda i: (-i[1], i[0])))

        con = sqlite3.connect(self.fx.path)
        try:
            ctx = rc.Reconciliation(con, date(2026, 9, 10), date(2026, 9, 20))
            forward, reverse = ctx.forward_rows(), ctx.reverse_rows()
        finally:
            con.close()
        _, total = ctx.summary(forward, reverse)
        self.assertEqual(sum(shown['forward'].values()),
                         total['app_bez_verdikta'])
        self.assertEqual(sum(shown['reverse'].values()),
                         total['day_bez_verdikta'])
        # Не пусто и не одно значение: иначе равенство ничего не различает.
        self.assertEqual(shown['forward']['otkryta'], 2)
        self.assertIn('mashina_ne_sopostavlena', shown['forward'])
        self.assertGreaterEqual(len(shown['reverse']), 2)
        # Те же сутки, что в первых строках: число открытых заявок списка
        # равно числу заявок с причиной «открыта».
        self.assertIn('open applications: %d' % shown['forward']['otkryta'],
                      out)

    def test_preview_names_itself_and_the_approved_n(self):
        code, out, err = self.main('--lookback-preview', '3')
        self.assertEqual(code, 0, err)
        self.assertIn('PREVIEW N=3', out)
        notes = ' '.join(str(r[0]) for r in book_of(self.out)['Свод']
                         .iter_rows(values_only=True) if r and r[0])
        self.assertIn('ПРЕДПРОСМОТР: N = 3; утверждённое владельцем N = 2',
                      notes)

    def test_without_an_approved_n_the_book_says_so(self):
        with mock.patch.object(rc, 'BACKDATED_LOOKBACK_DAYS', None):
            code, out, err = self.main()
            self.assertEqual(code, 0, err)
            notes = ' '.join(str(r[0]) for r in book_of(self.out)['Свод']
                             .iter_rows(values_only=True) if r and r[0])
            self.assertIn('N для заявок, заведённых задним числом: не '
                          'утверждён', notes)
            code, out, err = self.main('--lookback-preview', '2')
            notes = ' '.join(str(r[0]) for r in book_of(self.out)['Свод']
                             .iter_rows(values_only=True) if r and r[0])
            self.assertIn('ПРЕДПРОСМОТР: N = 2 не утверждён владельцем', notes)

    def test_lag_measurement_is_printed_and_written(self):
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertIn('completion lag', out)
        rows = list(book_of(self.out)['Замер N'].iter_rows(values_only=True))
        self.assertEqual(rows[1][1], 1)                  # одна заявка в замере

    def test_refusals(self):
        bare = os.path.join(os.path.dirname(self.fx.path), 'bare.db')
        sqlite3.connect(bare).close()
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            self.assertEqual(tool.main(['--db', bare]), 2)
            self.assertEqual(tool.main(['--db', self.fx.path, '--from',
                                        '2026-09-20', '--to', '2026-09-10']), 2)
            self.assertEqual(tool.main(['--db', self.fx.path, '--to', '20.09']), 2)
        self.assertFalse(os.path.exists(self.out))



class VolumeInTheReport(unittest.TestCase):
    """U2: объём против GPS -- в консоли, на «Своде», в «Заявка-работа» и
    в «Причинах»; числа те же, что у ядра."""

    def setUp(self):
        self.fx = Fixture()
        # Машина 12: 12..13, по GPS 4,0 га против 5,10 -- вне допуска.
        self.fx.app(transport='T2', created=12, completed=13, volume='5.10')
        # Машина 11: 10..10 и 10..11 пересекаются; общее окно 10..11, по GPS
        # 2,0 га против 1,00 + 1,50.
        self.fx.app(transport='T1', created=10, completed=10, volume='1.00')
        self.fx.app(transport='T1', created=10, completed=11, volume='1.50')
        self.fx.con.close()
        self.out = os.path.join(os.path.dirname(self.fx.path), 'report.xlsx')

    def main(self):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = tool.main(['--db', self.fx.path, '--from', '2026-09-10',
                              '--to', '2026-09-20', '--out', self.out])
        return code, out.getvalue(), err.getvalue()

    def test_console_summary_and_rows(self):
        code, out, err = self.main()
        self.assertEqual(code, 0, err)
        self.assertTrue(out.isascii())
        self.assertIn('applications: 3 | work confirmed 3 | NO WORK 0', out)
        self.assertIn('  work confirmed, volume vs GPS: within tolerance 0 | '
                      'borderline 0 | out of tolerance 1 | cannot compare 2\n',
                      out)
        self.assertIn('  cannot compare the volume, why: '
                      'obem_okna_peresekayutsya 2\n', out)
        book = book_of(self.out)
        summary = list(book['Свод'].iter_rows(values_only=True))
        header = summary[0]
        totals = [row for row in summary if row and row[0] == 'Итого / Жами'][0]
        columns = {title: index for index, title in enumerate(header)}
        volume = [totals[columns[tool.bi(title)]]
                  for _, title in tool.VOLUME_COLUMNS]
        self.assertEqual(volume, [0, 0, 1, 2])
        self.assertEqual(sum(volume),
                         totals[columns[tool.bi(('Работа была', 'Иш бўлган'))]])
        # И в строке группы, а не только в итоге: все три заявки -- машины
        # организации Buxoro, категория «mtz».
        group = [row for row in summary if row and row[0] == 'Buxoro'][0]
        self.assertEqual([group[columns[tool.bi(title)]]
                          for _, title in tool.VOLUME_COLUMNS], [0, 0, 1, 2])

        sheet = list(book['Заявка-работа'].iter_rows(values_only=True))
        header = sheet[0]
        at = {title: index for index, title in enumerate(header)}
        rows = {row[0]: row for row in sheet[1:]}

        def cell(number, title):
            return rows[number][at[tool.bi(title)]]
        self.assertEqual(cell('N-001', ('Объём в заявке, га',
                                        'Буюртмадаги ҳажм, га')), 5.1)
        self.assertEqual(cell('N-001', ('Объём против GPS',
                                        'Ҳажм GPS га қарши')), 'Вне допуска')
        self.assertEqual(cell('N-001', ('Расхождение GPS − заявка, га',
                                        'Фарқ GPS − буюртма, га')), -1.1)
        self.assertEqual(cell('N-001', ('Расхождение, %', 'Фарқ, %')), -21.6)
        self.assertEqual(cell('N-001', ('Почему сверить нельзя',
                                        'Нега солиштириб бўлмайди')), None)
        for number, other in (('N-002', 'N-003'), ('N-003', 'N-002')):
            with self.subTest(application=number):
                self.assertEqual(cell(number, ('Объём против GPS',
                                               'Ҳажм GPS га қарши')),
                                 'Сверить нельзя')
                self.assertEqual(
                    cell(number, ('Почему сверить нельзя',
                                  'Нега солиштириб бўлмайди')),
                    'Окно пересекается с окном другой заявки этой машины')
                self.assertEqual(cell(number, ('Окно пересекается с заявками',
                                               'Ойнаси кесишадиган буюртмалар')),
                                 other)
                self.assertEqual(cell(number, ('Вместе: окно',
                                               'Биргаликда: ойна')),
                                 '10.09 — 11.09')
                self.assertEqual(cell(number, ('Вместе: объём заявок, га',
                                               'Биргаликда: буюртмалар ҳажми, '
                                               'га')), 2.5)
                self.assertEqual(cell(number, ('Вместе: га по GPS',
                                               'Биргаликда: GPS бўйича га')),
                                 2.0)
        reasons = [row for row in book['Причины'].iter_rows(values_only=True)
                   if row[2] == 'obem_okna_peresekayutsya']
        self.assertEqual(len(reasons), 1)
        self.assertEqual(reasons[0][3], 2)
        self.assertTrue(reasons[0][0].startswith('Объём против GPS'))

if __name__ == '__main__':
    unittest.main()
