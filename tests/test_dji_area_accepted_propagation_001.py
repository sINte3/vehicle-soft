# -*- coding: utf-8 -*-
"""DJI-AREA-ACCEPTED-PROPAGATION-001: принятая площадь в рабочих отчётах
модуля Дроны -- на базе приложения.

Чистая семантика (случаи 1-10 этапа F) и пакетное чтение (15) -- в
`tests/test_dji_area_accepted_core.py`: stdlib, идёт в CI. Здесь те же
двенадцать случаев лежат в настоящих таблицах, и проверяется то, что
существует только вместе с Flask:

 1-10. Провайдер на настоящих таблицах даёт ожидаемое по каждому случаю,
       включая решения, записанные штатным `control_store.record_decision`,
       и устаревание после пересчёта.
 11.   Один вылет и один месяц дают одну принятую площадь в списке, его
       Excel, сводке, сверке с ведомостями, календаре, расходе раствора и на
       экране контроля.
 12.   Excel совпадает с экраном (список и сводка).
 13.   RAW не меняется: снимок `drone_flights` и `dji_area_calculations` до
       и после обхода всех страниц и книг на двух языках равен.
 14.   `billable_area_m2` пуст везде.

Июнь 2026 -- двенадцать случаев, июль -- рассчитан частично, ноябрь 2025
-- до запуска контроля площади, май 2026 -- ведомость без вылетов.

ВСЕ ДАННЫЕ СИНТЕТИЧЕСКИЕ: вылеты 930001.., борта SYNTHETIC-HW-...-NOT-REAL.
"""

import io
import re
import sqlite3
import unittest
from datetime import date, datetime, timedelta

from tests.harness import app, TEST_DB_PATH
from tests.test_dji_area_report_001 import Base as ReportBase, HW6

from models import (db, DjiAreaCalculation, DroneCustomer, DroneFlight,
                    DroneWork)

import drones
from dji_area import accepted as acc
from dji_area import control_store
from dji_area import resolver as rs
from dji_area import store as dji_store
from tests.test_dji_area_accepted_core import (CASES, COMMENT, EXPECTED,
                                               JUNE, STALE)

NBSP = '\u00a0'

JULY = date(2026, 7, 10)
NOV = date(2025, 11, 15)
WINDOW_ALL = 'date_from=2025-11-01&date_to=2026-07-31'
WINDOW_JUNE = 'date_from=2026-06-01&date_to=2026-06-30'
JUNE_RAW = sum(c[4] for c in CASES)                       # 59 000 м2
JUNE_ACCEPTED = sum(v[0] for v in EXPECTED.values())      # 34 000 м2
JUNE_EXCLUDED = sum(v[1] for v in EXPECTED.values())      # 25 000 м2

# Июль: один рассчитанный вылет и один без расчёта -> «частично».
JULY_CALC, JULY_NOCALC = 930020, 930021
JULY_CALC_RAW, JULY_NOCALC_HA = 20000.0, 1.2345
# Ноябрь 2025: до запуска контроля площади -- расчётов нет вовсе.
NOV_FLIGHTS = ((930030, 0.5), (930031, 0.7))
LITERS = 10.0
# Август 2026 (REVIEW-FIX): расчёт есть у обоих вылетов, но у одного RAW
# вылета (0.90 га) расходится с RAW расчёта (0.80 га). Полон по расчёту, не
# готов служить базой контроля.
AUG = date(2026, 8, 12)
AUG_OK, AUG_BAD = 930040, 930041
AUG_OK_RAW, AUG_BAD_CALC_RAW, AUG_BAD_FLIGHT_HA = 10000.0, 8000.0, 0.9
AUG_RAW_HA = AUG_OK_RAW / 10000.0 + AUG_BAD_FLIGHT_HA          # 1.90 га
AUG_ACCEPTED_HA = (AUG_OK_RAW + AUG_BAD_CALC_RAW) / 10000.0     # 1.80 га


# ─── 11-15: отчёты на синтетической базе ─────────────────────────────────────

class Seeded(ReportBase):
    """Июнь 2026 -- двенадцать случаев, июль -- частично, ноябрь 2025 --
    до запуска контроля площади."""

    def setUp(self):
        super(Seeded, self).setUp()
        with app.app_context():
            for n, spec in enumerate(CASES):
                fid, _code, status, elig, raw, corr, cand, flags, _k = spec
                obj = self.calc_object(fid, status, elig, raw_m2=raw,
                                       corrected_m2=corr, day=JUNE,
                                       hardware_id=HW6, minute=n * 5,
                                       flags=flags, delta_m2=corr)
                obj.structural_candidate = True if cand else None
                obj.scalar_source_check = True if cand else None
                db.session.add(obj)
                db.session.add(self.flight(fid, JUNE, n * 5, raw / 10000.0))
            db.session.add(self.calc_object(
                JULY_CALC, rs.RAW_CORROBORATED, rs.AGG_CERTIFIED,
                raw_m2=JULY_CALC_RAW, corrected_m2=JULY_CALC_RAW, day=JULY))
            db.session.add(self.flight(JULY_CALC, JULY, 0,
                                       JULY_CALC_RAW / 10000.0))
            db.session.add(self.flight(JULY_NOCALC, JULY, 30, JULY_NOCALC_HA))
            for i, (fid, area) in enumerate(NOV_FLIGHTS):
                db.session.add(self.flight(fid, NOV, i * 10, area))
            for n, (fid, calc_raw, flight_ha) in enumerate((
                    (AUG_OK, AUG_OK_RAW, AUG_OK_RAW / 10000.0),
                    (AUG_BAD, AUG_BAD_CALC_RAW, AUG_BAD_FLIGHT_HA))):
                db.session.add(self.calc_object(
                    fid, rs.RAW_CORROBORATED, rs.AGG_CERTIFIED,
                    raw_m2=calc_raw, corrected_m2=calc_raw, day=AUG,
                    minute=n * 10))
                db.session.add(self.flight(fid, AUG, n * 10, flight_ha))
            customer = DroneCustomer(name='SYNTHETIC customer')
            db.session.add(customer)
            db.session.flush()
            # Ведомость июня ровно на принятую площадь: покрытие 100 % по
            # принятой и 57.6 % по RAW -- два основания различимы.
            db.session.add(DroneWork(
                period_month='2026-06', work_date_from=JUNE,
                work_date_to=JUNE, drone_customer_id=customer.id,
                customer_raw='SYNTHETIC customer',
                area_ha=JUNE_ACCEPTED / 10000.0, amount=0,
                received_amount=0, payment_type='cash'))
            # Август: ведомость ровно на принятую по рассчитанным (1.80) --
            # по принятой было бы 100 %, но база -- RAW (расхождение).
            db.session.add(DroneWork(
                period_month='2026-08', work_date_from=AUG,
                work_date_to=AUG, drone_customer_id=customer.id,
                customer_raw='SYNTHETIC customer', area_ha=AUG_ACCEPTED_HA,
                amount=0, received_amount=0, payment_type='cash'))
            # Май: ведомость есть, вылетов нет -- «—», а не «не рассчитано».
            db.session.add(DroneWork(
                period_month='2026-05', work_date_from=date(2026, 5, 10),
                work_date_to=date(2026, 5, 10),
                drone_customer_id=customer.id,
                customer_raw='SYNTHETIC customer', area_ha=1.0, amount=0,
                received_amount=0, payment_type='cash'))
            db.session.commit()
        con = dji_store.connect(TEST_DB_PATH)
        try:
            for spec in CASES:
                if spec[8]:
                    # Подтверждение переопределения нужно решению по
                    # доказанной корректировке (KR, AA) и безвредно прочим.
                    control_store.record_decision(
                        con, spec[0], spec[8], COMMENT, True, True, None,
                        self.admin_id, 'SYNTHETIC admin')
        finally:
            con.close()
        # Пересчёт после решения: новая строка с тем же статусом и новым
        # отпечатком. Автомат не изменился, но решение принято против
        # прежнего расчёта.
        with app.app_context():
            for fid in STALE:
                spec = [s for s in CASES if s[0] == fid][0]
                for old in DjiAreaCalculation.query.filter_by(
                        flight_id=fid, superseded_at=None):
                    old.superseded_at = datetime(2026, 9, 9, 0, 0)
                obj = self.calc_object(fid, spec[2], spec[3], raw_m2=spec[4],
                                       corrected_m2=spec[5], day=JUNE,
                                       hardware_id=HW6,
                                       minute=CASES.index(spec) * 5,
                                       flags=spec[7], delta_m2=spec[5])
                obj.structural_candidate = True
                obj.scalar_source_check = True
                db.session.add(obj)
            db.session.commit()

    def flight(self, fid, day, minute, area_ha):
        return DroneFlight(
            dji_flight_id=fid, drone_unit_id=self.unit6_id,
            nickname_raw='SYNTHETIC-%d' % fid,
            started_at=datetime(day.year, day.month, day.day, 3, 0)
            + timedelta(minutes=minute),
            work_seconds=600, area_ha=area_ha, spray_liters=LITERS,
            usage_type=0, raw_json='{}')

    def get(self, path, language='ru'):
        response = self.client_as(language=language).get(path)
        self.assertEqual(response.status_code, 200, path)
        return response

    def book(self, path):
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(self.get(path).data))

    @staticmethod
    def num(text):
        text = re.sub(r'<[^>]+>', ' ', text).replace(NBSP, '').strip()
        match = re.match(r'-?[\d ]+\.\d+|-?\d+', text)
        return float(match.group(0).replace(' ', '')) if match else None

    @staticmethod
    def stat(html, label):
        match = re.search(r'<div class="vs-stat-label">%s</div>\s*'
                          r'<div class="vs-stat-value[^"]*">(.*?)</div>'
                          % re.escape(label), html, re.S)
        if not match:
            raise AssertionError('stat %r not found' % label)
        return match.group(1)

    def list_rows(self, html):
        """{flight_id: (RAW, принято-текст, статус)} по нику строки."""
        out = {}
        for row in re.findall(r'<tr>(.*?)</tr>', html, re.S):
            nick = re.search(r'SYNTHETIC-(\d+)', row)
            if not nick:
                continue
            right = re.findall(r'<td class="right"[^>]*>(.*?)</td>', row,
                               re.S)
            status = re.search(r'<span class="vs-badge vs-badge-sm[^"]*">'
                               r'(.*?)</span>', row, re.S)
            out[int(nick.group(1))] = (self.num(right[0]), right[1],
                                       status.group(1).strip())
        return out

    def snapshot(self):
        con = sqlite3.connect(TEST_DB_PATH)
        try:
            cols = [c.name for c in DjiAreaCalculation.__table__.columns]
            return (con.execute('SELECT id, dji_flight_id, area_ha, raw_json '
                                'FROM drone_flights ORDER BY id').fetchall(),
                    con.execute('SELECT %s FROM dji_area_calculations ORDER '
                                'BY id' % ', '.join(cols)).fetchall())
        finally:
            con.close()


class Reports(Seeded):

    def provider(self):
        with app.app_context():
            pairs = [(f.dji_flight_id, f.area_ha)
                     for f in DroneFlight.query.all()]
            with app.test_request_context():
                return drones._drone_accepted_by_dji(pairs)

    def test_the_provider_gives_every_case_on_the_real_tables(self):
        items = self.provider()
        for fid, (accepted, excluded, status, is_open) in EXPECTED.items():
            item = items[fid]
            self.assertAlmostEqual(item['accepted_m2'], accepted, places=4,
                                   msg=fid)
            self.assertAlmostEqual(item['excluded_m2'], excluded, places=4,
                                   msg=fid)
            self.assertEqual(item['status'], status, fid)
            self.assertEqual(item['is_open'], is_open, fid)
        self.assertEqual(items[JULY_NOCALC]['status'], acc.ST_NOT_CALCULATED)
        for fid, _area in NOV_FLIGHTS:
            self.assertEqual(items[fid]['status'], acc.ST_NOT_CALCULATED)
            self.assertIsNone(items[fid]['accepted_m2'])

    def test_the_flight_list_shows_raw_accepted_and_status(self):
        html = self.get('/drones/?' + WINDOW_ALL).get_data(as_text=True)
        rows = self.list_rows(html)
        labels = {s: acc.STATUS_LABELS[s][0] for s in acc.SHORT_STATUSES}
        for spec in CASES:
            fid = spec[0]
            raw_ha, accepted_text, status = rows[fid]
            self.assertAlmostEqual(raw_ha, spec[4] / 10000.0, places=2)
            self.assertAlmostEqual(self.num(accepted_text),
                                   EXPECTED[fid][0] / 10000.0, places=2)
            self.assertEqual(status, labels[EXPECTED[fid][2]], fid)
        raw_ha, accepted_text, status = rows[JULY_NOCALC]
        self.assertAlmostEqual(raw_ha, JULY_NOCALC_HA, places=2)
        self.assertIn('не рассчитано', accepted_text)
        self.assertIsNone(self.num(accepted_text))
        self.assertEqual(status, 'не рассчитано')

    def test_the_flights_book_equals_the_list(self):
        """12: Excel = экран, по каждому вылету; RAW-колонка прежняя."""
        html = self.get('/drones/?' + WINDOW_ALL).get_data(as_text=True)
        rows = self.list_rows(html)
        ws = self.book('/drones/flights.xlsx?' + WINDOW_ALL).worksheets[0]
        head = [c.value for c in ws[1]]
        self.assertEqual(head[5], 'Гектары')
        self.assertEqual(head[14:17], ['Принято, га', 'Исключено, га',
                                       'Статус площади (контроль площади '
                                       'DJI)'])
        seen = 0
        for values in ws.iter_rows(min_row=2, values_only=True):
            fid = values[10]
            raw_ha, accepted_text, status = rows[fid]
            self.assertAlmostEqual(values[5], raw_ha, places=2)
            expected = self.num(accepted_text)
            if expected is None:
                self.assertIsNone(values[14], fid)
                self.assertIsNone(values[15], fid)
            else:
                self.assertAlmostEqual(values[14], expected, places=2)
            self.assertEqual(values[16], status, fid)
            seen += 1
        self.assertEqual(seen, len(CASES) + 2 + len(NOV_FLIGHTS))

    def test_the_summary_and_its_book_agree(self):
        """12: сводка -- экран и книга, итог и месяцы."""
        html = self.get('/drones/summary?' + WINDOW_ALL).get_data(
            as_text=True)
        accepted_ha = (JUNE_ACCEPTED + JULY_CALC_RAW) / 10000.0
        self.assertAlmostEqual(self.num(self.stat(html, 'Принято, га')),
                               accepted_ha, places=2)
        self.assertIn('частично', self.stat(html, 'Принято, га'))
        self.assertAlmostEqual(self.num(self.stat(html, 'Исключено, га')),
                               JUNE_EXCLUDED / 10000.0, places=2)
        self.assertAlmostEqual(
            self.num(self.stat(html, 'Не рассчитано, DJI RAW га')),
            JULY_NOCALC_HA + sum(a for _f, a in NOV_FLIGHTS), places=2)
        self.assertIn('вылетов: 3', html)
        self.assertAlmostEqual(self.num(self.stat(html,
                                                  'Открыто, DJI RAW га')),
                               17500.0 / 10000.0, places=2)
        book = self.book('/drones/summary.xlsx?' + WINDOW_ALL)
        pairs = {row[0]: row[1] for row in
                 book.worksheets[0].iter_rows(min_row=2, values_only=True)}
        self.assertAlmostEqual(pairs['Принято (по рассчитанным вылетам), га'],
                               accepted_ha, places=6)
        self.assertAlmostEqual(pairs['Исключено, га'],
                               JUNE_EXCLUDED / 10000.0, places=6)
        self.assertEqual(pairs['Не рассчитано вылетов'], 3)
        self.assertAlmostEqual(pairs['Не рассчитано, DJI RAW га'],
                               JULY_NOCALC_HA + sum(a for _f, a in
                                                    NOV_FLIGHTS), places=6)
        self.assertEqual(pairs['Открыто (ожидает доказательства или '
                               'решения), вылетов'], 4)
        # Прежние строки -- на месте и с прежним смыслом (DJI RAW).
        self.assertAlmostEqual(pairs['Гектаров'], (JUNE_RAW + JULY_CALC_RAW)
                               / 10000.0 + JULY_NOCALC_HA
                               + sum(a for _f, a in NOV_FLIGHTS), places=6)
        months = {row[0]: row for row in
                  book['По месяцам'].iter_rows(min_row=2, values_only=True)}
        self.assertAlmostEqual(months['2026-06'][4], JUNE_ACCEPTED / 10000.0,
                               places=6)
        self.assertEqual(months['2026-06'][6], 0)
        self.assertEqual(months['2026-06'][9], 'рассчитано полностью')
        self.assertAlmostEqual(months['2026-07'][4], JULY_CALC_RAW / 10000.0,
                               places=6)
        self.assertEqual(months['2026-07'][6], 1)
        self.assertEqual(months['2026-07'][9], 'рассчитано частично')
        self.assertIsNone(months['2025-11'][4],
                          'до запуска контроля -- пусто, а не 0 и не RAW')
        self.assertEqual(months['2025-11'][6], 2)
        self.assertEqual(months['2025-11'][9], 'не рассчитано')

    def test_reconcile_controls_by_accepted_only_in_a_full_month(self):
        with app.app_context():
            with app.test_request_context():
                data = drones._drone_works_flights_reconcile_data()
        rows = {r['month']: r for r in data['rows']}
        june, july, nov = rows['2026-06'], rows['2026-07'], rows['2025-11']
        self.assertEqual(june['control_basis'], 'accepted')
        self.assertAlmostEqual(june['control_area'], JUNE_ACCEPTED / 10000.0)
        self.assertAlmostEqual(june['control_coverage'], 100.0, places=6)
        self.assertEqual(june['control_flag'], '')
        # Отрицательный контроль: по RAW тот же месяц -- 57.6 % и жёлтый.
        self.assertAlmostEqual(june['coverage'],
                               JUNE_ACCEPTED * 100.0 / JUNE_RAW, places=6)
        self.assertEqual(june['flag'], 'is-warning-row')
        # RAW-поля прежние: их читает «Кто не сдал ведомость».
        self.assertAlmostEqual(june['flight_area'], JUNE_RAW / 10000.0)
        self.assertEqual(july['control_basis'], 'raw')
        self.assertEqual(nov['control_basis'], 'raw')
        self.assertEqual(july['acc']['coverage'], acc.COVERAGE_PARTIAL)
        self.assertEqual(nov['acc']['coverage'], acc.COVERAGE_NONE)
        self.assertIsNone(nov['acc']['accepted_ha'])
        may = rows['2026-05']
        self.assertEqual(may['acc']['coverage'], acc.COVERAGE_EMPTY)
        html = self.get('/drones/reports/reconcile').get_data(as_text=True)
        self.assertIn('по принятой', html)
        self.assertIn('по DJI RAW', html)
        may_row = [r for r in re.findall(r'<tr[^>]*>(.*?)</tr>', html, re.S)
                   if '2026-05' in r][0]
        self.assertNotIn('не рассчитано', may_row,
                         'месяц без вылетов -- не «не рассчитано»')
        july_row = [r for r in re.findall(r'<tr[^>]*>(.*?)</tr>', html,
                                          re.S) if '2026-07' in r][0]
        self.assertIn('частично', july_row)
        nov_row = [r for r in re.findall(r'<tr[^>]*>(.*?)</tr>', html, re.S)
                   if '2025-11' in r][0]
        self.assertIn('не рассчитано', nov_row)

    def test_the_calendar_works_on_accepted_and_marks_raw(self):
        with app.app_context():
            with app.test_request_context():
                june = drones._drone_flight_calendar_data('2026-06')
                july = drones._drone_flight_calendar_data('2026-07')
                nov = drones._drone_flight_calendar_data('2025-11')
        row = [r for r in june['rows'] if r['unit_id'] == self.unit6_id][0]
        cell = [c for c in row['cells'] if c['day'] == JUNE][0]
        self.assertEqual(cell['basis'], 'accepted')
        self.assertAlmostEqual(cell['shown'], JUNE_ACCEPTED / 10000.0)
        self.assertAlmostEqual(cell['area'], JUNE_RAW / 10000.0)
        self.assertTrue(cell['open'])
        self.assertEqual(june['total']['basis'], 'accepted')
        jrow = [r for r in july['rows'] if r['unit_id'] == self.unit6_id][0]
        jcell = [c for c in jrow['cells'] if c['day'] == JULY][0]
        self.assertEqual(jcell['basis'], 'raw')
        self.assertAlmostEqual(jcell['shown'], JULY_CALC_RAW / 10000.0
                               + JULY_NOCALC_HA)
        self.assertEqual(nov['total']['basis'], 'raw')
        html = self.get('/drones/reports/calendar?month=2026-07').get_data(
            as_text=True)
        self.assertIn('%.2f*' % (JULY_CALC_RAW / 10000.0 + JULY_NOCALC_HA),
                      html)
        html = self.get('/drones/reports/calendar?month=2026-06').get_data(
            as_text=True)
        self.assertIn('%.2f?' % (JUNE_ACCEPTED / 10000.0), html)

    def test_the_spray_rate_uses_accepted_only_where_calculated(self):
        with app.app_context():
            with app.test_request_context():
                data = drones._drone_spray_usage_data(
                    drones._drone_flight_conditions(drones._drone_filters_from_args(
                        _Args({'date_from': '2025-11-01',
                               'date_to': '2026-07-31'}),
                        default_current_month=False)), [], 30)
        rows = {r['month']: r for r in data['rows']}
        june = rows['2026-06']
        liters = LITERS * len(CASES)
        # Прежний расход -- по RAW, как и был.
        self.assertAlmostEqual(june['rate'], liters / (JUNE_RAW / 10000.0))
        self.assertAlmostEqual(june['rate_accepted'],
                               liters / (JUNE_ACCEPTED / 10000.0))
        self.assertEqual(june['rate_accepted_reason'], 'ok')
        self.assertIsNone(rows['2026-07']['rate_accepted'])
        self.assertEqual(rows['2026-07']['rate_accepted_reason'], 'partial')
        self.assertIsNone(rows['2025-11']['rate_accepted'])
        self.assertEqual(rows['2025-11']['rate_accepted_reason'],
                         'not_calculated')
        self.assertIsNone(data['total']['rate_accepted'])

    def test_one_flight_one_accepted_area_everywhere(self):
        """11: вылет и месяц -- одно число во всех подключённых отчётах."""
        items = self.provider()
        # Экран контроля площади за июнь -- тот же итог принятой.
        with app.app_context():
            filters = drones._drone_area_control_filters(_Args({
                'date_from': '2026-06-01', 'date_to': '2026-06-30'}))
            with app.test_request_context():
                report = drones._drone_area_control_report(filters)
                summary = drones._drone_summary_data(
                    drones._drone_flight_conditions(
                        drones._drone_filters_from_args(_Args({
                            'date_from': '2026-06-01',
                            'date_to': '2026-06-30'}),
                            default_current_month=False)))
                reconcile = drones._drone_works_flights_reconcile_data()
                calendar = drones._drone_flight_calendar_data('2026-06')
        by_control = {i['flight_id']: i for i in report['items']}
        for fid in EXPECTED:
            self.assertEqual(by_control[fid]['accepted_m2'],
                             items[fid]['accepted_m2'], fid)
            self.assertEqual(by_control[fid]['state'], items[fid]['state'])
        june_ha = JUNE_ACCEPTED / 10000.0
        self.assertAlmostEqual(report['total']['accepted_m2'] / 10000.0,
                               june_ha, places=6)
        self.assertAlmostEqual(summary['accepted']['accepted_ha'], june_ha,
                               places=6)
        month = [r for r in reconcile['rows'] if r['month'] == '2026-06'][0]
        self.assertAlmostEqual(month['acc']['accepted_ha'], june_ha, places=6)
        self.assertAlmostEqual(calendar['accepted']['accepted_ha'], june_ha,
                               places=6)
        self.assertAlmostEqual(summary['by_machine']['rows'][0]['acc']
                               ['accepted_ha'], june_ha, places=6)
        # Разрез по операторам -- то же правило привязки: назначений в
        # фикстуре нет, весь июнь -- «оператор не определён».
        self.assertAlmostEqual(summary['by_operator']['undetermined']['acc']
                               ['accepted_ha'], june_ha, places=6)
        self.assertAlmostEqual(summary['by_operator']['total']['acc']
                               ['accepted_ha'], june_ha, places=6)
        book = self.book('/drones/summary.xlsx?' + WINDOW_JUNE)
        rows = list(book['По операторам'].iter_rows(min_row=3,
                                                    values_only=True))
        self.assertEqual(rows[-1][0], 'Итого')
        self.assertAlmostEqual(rows[-1][4], june_ha, places=6)
        self.assertEqual(rows[-1][6], 0)
        self.assertEqual(summary['accepted']['open_records'],
                         report['total']['open_records'])

    def test_raw_is_untouched_and_billable_stays_empty(self):
        """13, 14: обход всех подключённых страниц и книг ничего не пишет."""
        before = self.snapshot()
        for path in ('/drones/?' + WINDOW_ALL,
                     '/drones/flights.xlsx?' + WINDOW_ALL,
                     '/drones/summary?' + WINDOW_ALL,
                     '/drones/summary.xlsx?' + WINDOW_ALL,
                     '/drones/reports/reconcile',
                     '/drones/reports/calendar?month=2026-06',
                     '/drones/reports/spray?' + WINDOW_ALL,
                     '/drones/reports/spray.xlsx?' + WINDOW_ALL,
                     '/drones/area-control?' + WINDOW_JUNE):
            for language in ('ru', 'uz'):
                self.get(path, language=language)
        self.assertEqual(before, self.snapshot())
        con = sqlite3.connect(TEST_DB_PATH)
        try:
            filled = con.execute('SELECT COUNT(*) FROM dji_area_calculations '
                                 'WHERE billable_area_m2 IS NOT NULL'
                                 ).fetchone()[0]
            total = con.execute('SELECT COUNT(*) FROM dji_area_calculations'
                                ).fetchone()[0]
        finally:
            con.close()
        self.assertEqual(filled, 0)
        self.assertGreater(total, len(CASES))

    # ── REVIEW-FIX: расхождение RAW -- fail-closed ────────────────────────

    def test_the_mismatched_month_is_full_but_not_control_ready(self):
        items = self.provider()
        self.assertTrue(items[AUG_BAD]['raw_mismatch'])
        self.assertFalse(items[AUG_OK]['raw_mismatch'])
        with app.app_context():
            with app.test_request_context():
                data = drones._drone_works_flights_reconcile_data()
        aug = [r for r in data['rows'] if r['month'] == '2026-08'][0]
        # (a)/(b): расчёт есть у каждого вылета, но базой месяц не служит.
        self.assertEqual(aug['acc']['coverage'], acc.COVERAGE_FULL)
        self.assertTrue(aug['acc']['complete'])
        self.assertFalse(aug['acc']['control_ready'])
        self.assertIsNone(aug['acc']['accepted_full_ha'])
        self.assertEqual(aug['acc']['raw_mismatch_records'], 1)
        # (c): сверка остаётся на DJI RAW.
        self.assertEqual(aug['control_basis'], 'raw')
        self.assertAlmostEqual(aug['control_area'], AUG_RAW_HA, places=6)
        self.assertAlmostEqual(aug['control_coverage'],
                               AUG_ACCEPTED_HA * 100.0 / AUG_RAW_HA, places=6)
        # Отрицательный контроль: по принятой по рассчитанным покрытие было
        # бы ровно 100 % -- основания различимы, проверка не пустая.
        self.assertAlmostEqual(aug['acc']['accepted_ha'] * 100.0
                               / AUG_ACCEPTED_HA, 100.0, places=6)
        html = self.get('/drones/reports/reconcile').get_data(as_text=True)
        row = [r for r in re.findall(r'<tr[^>]*>(.*?)</tr>', html, re.S)
               if '2026-08' in r][0]
        self.assertIn('по DJI RAW', row)
        self.assertIn('расхождение RAW', row)

    def test_the_calendar_does_not_work_on_a_mismatched_cell(self):
        # (d): число ячейки -- DJI RAW с «*», причина в подсказке.
        with app.app_context():
            with app.test_request_context():
                data = drones._drone_flight_calendar_data('2026-08')
        row = [r for r in data['rows'] if r['unit_id'] == self.unit6_id][0]
        cell = [c for c in row['cells'] if c['day'] == AUG][0]
        self.assertEqual(cell['acc']['coverage'], acc.COVERAGE_FULL)
        self.assertEqual(cell['basis'], 'raw')
        self.assertAlmostEqual(cell['shown'], AUG_RAW_HA, places=6)
        self.assertEqual(data['total']['basis'], 'raw')
        html = self.get('/drones/reports/calendar?month=2026-08').get_data(
            as_text=True)
        self.assertIn('%.2f*' % AUG_RAW_HA, html)
        self.assertNotIn('>%.2f<' % AUG_ACCEPTED_HA, html)
        self.assertIn('RAW расчёта расходится с RAW вылета: 1', html)

    def test_the_spray_rate_names_the_mismatch_instead_of_a_number(self):
        # (e): «л/га по принятой» не считается, причина названа.
        with app.app_context():
            with app.test_request_context():
                data = drones._drone_spray_usage_data(
                    drones._drone_flight_conditions(
                        drones._drone_filters_from_args(_Args({
                            'date_from': '2026-08-01',
                            'date_to': '2026-08-31'}),
                            default_current_month=False)), [], 30)
        aug = data['rows'][0]
        self.assertEqual(aug['month'], '2026-08')
        self.assertIsNotNone(aug['rate'], 'расход по RAW -- как и был')
        self.assertIsNone(aug['rate_accepted'])
        self.assertEqual(aug['rate_accepted_reason'], 'raw_mismatch')
        # Прямой вызов идёт без пользователя -- язык по умолчанию узбекский.
        self.assertIn(aug['rate_accepted_note'],
                      ('расхождение RAW', 'RAW фарқ қилади'))
        self.assertIsNone(data['total']['rate_accepted'])
        book = self.book('/drones/reports/spray.xlsx?date_from=2026-08-01'
                         '&date_to=2026-08-31')
        rows = list(book['Литров на гектар'].iter_rows(min_row=2,
                                                       values_only=True))
        self.assertIsNone(rows[0][11])
        self.assertEqual(rows[0][12], 'расхождение RAW')

    def test_the_summary_shows_the_mismatch_and_no_ready_basis(self):
        query = 'date_from=2026-08-01&date_to=2026-08-31'
        html = self.get('/drones/summary?' + query).get_data(as_text=True)
        self.assertIn('расхождение RAW', self.stat(html, 'Принято, га'))
        self.assertIn('RAW расходится: 1', html)
        self.assertIn('базой контроля не служит', html)
        book = self.book('/drones/summary.xlsx?' + query)
        pairs = {row[0]: row[1] for row in
                 book.worksheets[0].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(pairs['RAW расчёта расходится с RAW вылета, '
                               'вылетов'], 1)
        self.assertIn('RAW расходится: 1',
                      pairs['Полнота контроля площади DJI'])
        # Сумма по рассчитанным -- видна, с подписью; не выдана за базу.
        self.assertAlmostEqual(pairs['Принято (по рассчитанным вылетам), га'],
                               AUG_ACCEPTED_HA, places=6)

    def test_the_uzbek_pages_speak_cyrillic(self):
        html = self.get('/drones/?' + WINDOW_ALL, language='uz').get_data(
            as_text=True)
        self.assertIn('ҳисобланмаган', html)
        self.assertIn('қарор талаб қилинади', html)
        html = self.get('/drones/summary?' + WINDOW_ALL,
                        language='uz').get_data(as_text=True)
        self.assertIn('DJI майдони назорати:', html)
        self.assertNotIn('Area Control', html)


class _Args(dict):
    """Мини-заменитель request.args для прямого вызова построителей."""

    def get(self, key, default=None, type=None):  # noqa: A002
        value = dict.get(self, key, default)
        if type is not None and value is not None:
            try:
                return type(value)
            except (TypeError, ValueError):
                return default
        return value


if __name__ == '__main__':
    unittest.main()
