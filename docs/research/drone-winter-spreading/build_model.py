# -*- coding: utf-8 -*-
"""Генератор параметрической модели «Зимнее разбрасывание гранулированных
удобрений дроном DJI Agras T40» (model.xlsx) и её независимая проверка.

Запуск (из корня репозитория):

    python docs/research/drone-winter-spreading/build_model.py
    python docs/research/drone-winter-spreading/build_model.py --verify

Без аргументов пишет model.xlsx рядом с собой. С ``--verify`` дополнительно
пересчитывает книгу в LibreOffice (headless) во временной папке и сверяет
каждую расчётную ячейку с независимым расчётом на Python (``compute``).
Книга и Python считают одну модель двумя разными способами: формулы
в ячейках и явный код. Если LibreOffice не найден, проверка пропускается
и об этом печатается сообщение.

Вывод в консоль только ASCII (скрипт может выполняться в консоли Windows).

[REASON]: все числовые допущения живут в одном месте -- PARAMS -- и попадают
в лист «Параметры» с пометкой ФАКТ / ОЦЕНКА / ДОГАДКА и источником. Модель
бизнес-правил не выдумывает: нормы внесения, тариф и маржу задаёт владелец,
здесь только арифметика их последствий.
"""
from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DEFAULT = os.path.join(HERE, 'model.xlsx')

# ---------------------------------------------------------------------------
# Параметры: (код, название, значение, единица, тип, источник / комментарий)
# Тип: ФАКТ -- проверяемое с источником; ОЦЕНКА -- вывод из фактов;
#      ДОГАДКА -- правдоподобно, не проверено; РЕШЕНИЕ -- задаёт владелец.
# ---------------------------------------------------------------------------
PARAMS = [
    ('', 'ТЕХНИКА И ФИЗИКА ПОЛЁТА', None, '', '', ''),
    ('payload_kg', 'Полезная нагрузка за вылет', 45, 'кг', 'ОЦЕНКА',
     'Владелец: 40-50 кг. Паспорт DJI: бак 70 л, до 50 кг.'),
    ('swath_eff_m', 'Эффективная ширина захвата', 6.0, 'м', 'ДОГАДКА',
     'Паспорт DJI: 7 м. Минус перекрытие проходов и кромки.'),
    ('speed_ms', 'Рабочая скорость над полем', 6.0, 'м/с', 'ОЦЕНКА',
     'Тестовые параметры DJI для 1,5 т/ч: 7 м/с, высота 3 м.'),
    ('overhead_min', 'Накладное время вылета (взлёт, подлёт, посадка)', 1.5, 'мин', 'ДОГАДКА',
     'Станция у кромки поля, подлёт до 300 м. Проверить пилотом.'),
    ('reload_min', 'Загрузка бака и смена батареи', 1.5, 'мин', 'ДОГАДКА',
     'Два помощника, мешок 50 кг, засыпка по весовому датчику. Главный рычаг.'),
    ('eff_hours_per_day', 'Эффективное лётное время за день', 3.0, 'ч', 'ОЦЕНКА',
     'Сентябрь: ~7 000 га / 15 бортов / ~25 дней при 1,0-1,25 га за вылет -> ~3 ч чистого налёта в день.'),
    ('days_per_season', 'Рабочих дней за зимний сезон на борт', 30, 'дн', 'ДОГАДКА',
     'Окно подкормки ~1.02-20.03 (48 дней) минус выходные и погода.'),
    ('max_flow_kg_min', 'Максимальная подача разбрасывателя', 90, 'кг/мин', 'ОЦЕНКА',
     'Данные ритейлеров DJI; тест DJI 149 кг/га при 17,6 га/ч = 44 кг/мин.'),
    ('cycle_equiv_min', 'Минут полёта на один полный цикл батареи', 9.0, 'мин', 'ОЦЕНКА',
     'DJI: висение 6 мин при 101 кг (полный бак), 18 мин при 50 кг (пустой); вес падает по ходу вылета.'),
    ('max_flight_min', 'Предельная длительность вылета при лёгкой загрузке', 10.0, 'мин', 'ОЦЕНКА',
     'Висение 18 мин без груза; с запасом на возврат ~10-12 мин. Режет вылет при нормах ниже ~25 кг/га (семена).'),
    ('', 'БАТАРЕИ И ЭНЕРГИЯ', None, '', '', ''),
    ('battery_price_usd', 'Цена батареи T40 (DB1560, 30 000 мА·ч)', 1800, 'USD', 'ДОГАДКА',
     'Розница США 2 499 USD; поставки из Китая дешевле. Уточнить у поставщика.'),
    ('battery_cycles', 'Ресурс батареи', 1200, 'циклов', 'ОЦЕНКА',
     'Гарантия DJI по данным ритейлеров: 1 000-1 500 циклов.'),
    ('battery_kwh', 'Ёмкость батареи', 1.567, 'кВт·ч', 'ФАКТ', '30 А·ч x 52,22 В.'),
    ('gen_l_per_kwh', 'Расход бензина генератором D12000iE', 0.5, 'л/кВт·ч', 'ФАКТ',
     'Паспорт DJI: 500 мл/кВт·ч.'),
    ('charge_eff', 'КПД зарядки (доля энергии, дошедшей в батарею)', 0.85, 'доля', 'ДОГАДКА', ''),
    ('petrol_price', 'Цена бензина АИ-92', 13000, 'сум/л', 'ФАКТ',
     'Ташкент, 22.09.2026: 12 600 (местный) / 13 900 (импортный).'),
    ('', 'ЛЮДИ И ТРАНСПОРТ (НА РАБОЧИЙ ДЕНЬ)', None, '', '', ''),
    ('pilot_day', 'Оплата пилота за рабочий день', 400000, 'сум', 'ДОГАДКА',
     'Paylab: операторы сельхозтехники 3,6-7,4 млн/мес; пилот дрона выше. Уточнить по своим ведомостям.'),
    ('helpers_count', 'Помощников на загрузке', 2, 'чел', 'ДОГАДКА',
     'Один человек 60 мешков по 45 кг за день не поднимет.'),
    ('helper_day', 'Оплата помощника за день', 180000, 'сум', 'ДОГАДКА', ''),
    ('vehicle_km', 'Пробег машины сопровождения за день', 100, 'км', 'ДОГАДКА', ''),
    ('diesel_l_100km', 'Расход дизеля машины сопровождения', 12, 'л/100 км', 'ДОГАДКА', ''),
    ('diesel_price', 'Цена дизельного топлива', 15000, 'сум/л', 'ОЦЕНКА',
     'АЗС Ташкента 14 700-15 900; биржа апрель 2026: 15,86 млн сум/т.'),
    ('vehicle_wear_km', 'Износ машины сопровождения', 1500, 'сум/км', 'ДОГАДКА', ''),
    ('per_diem_day', 'Питание и проживание экипажа', 100000, 'сум/день', 'ДОГАДКА', ''),
    ('other_day', 'Прочее за день (износ генератора, связь, промывка)', 50000, 'сум/день', 'ДОГАДКА', ''),
    ('', 'ИЗНОС И АМОРТИЗАЦИЯ ТЕХНИКИ', None, '', '', ''),
    ('drone_wear_usd_h', 'Ремонтный резерв дрона за лётный час', 5, 'USD/ч', 'ДОГАДКА',
     'Винты, моторы, мелкий ремонт: ~2 500 USD на 500 лётных часов.'),
    ('fert_wear_premium', 'Надбавка к износу за удобрения (коррозия, пыль)', 0.5, 'доля', 'ДОГАДКА',
     'Селитра -- самое коррозионно-активное удобрение (РГАУ-МСХА, 2026).'),
    ('spreader_price_usd', 'Цена разбрасывателя T40 Spreading System 3.0', 1300, 'USD', 'ФАКТ',
     'Розница 1 199-1 499 USD (2025).'),
    ('spreader_life_seasons', 'Срок службы разбрасывателя', 2, 'сезонов', 'ДОГАДКА', ''),
    ('drone_price_usd', 'Цена дрона T40 (без разбрасывателя)', 20000, 'USD', 'ОЦЕНКА',
     'Б/у 2025 с 3 АКБ: 19 000 USD; новый без АКБ 13 500-16 000 USD.'),
    ('drone_life_years', 'Срок амортизации дрона', 5, 'лет', 'ДОГАДКА', ''),
    ('drone_dep_winter_share', 'Доля годовой амортизации дрона, относимая на зиму', 0.0, 'доля', 'РЕШЕНИЕ',
     '0 = парк уже окупается летом (маржинальный взгляд). 0,2-0,3 = полная себестоимость.'),
    ('', 'НАКЛАДНЫЕ, МАРЖА, КУРС', None, '', '', ''),
    ('overhead_pct', 'Накладные расходы (диспетчер, учёт, страховка)', 0.10, 'доля выручки', 'ДОГАДКА', ''),
    ('target_margin_pct', 'Целевая маржа', 0.25, 'доля выручки', 'РЕШЕНИЕ', 'Задаёт владелец.'),
    ('usd_rate', 'Курс USD/UZS', 12000, 'сум', 'ФАКТ', 'ЦБ РУз 29.10.2026: 11 965-12 004.'),
    ('', 'РЫНОК И СТОРОНА ФЕРМЕРА', None, '', '', ''),
    ('liquid_tariff_external', 'Текущий тариф опрыскивания, внешние заказчики', 200000, 'сум/га', 'ФАКТ', 'Владелец.'),
    ('liquid_tariff_internal', 'Ставка внутреннего контура (земли Агрокластера)', 80000, 'сум/га', 'ФАКТ',
     'Книги сентября 2025: 75 040 / 76 458 / 85 633 / 86 000.'),
    ('nitrate_price_t', 'Аммиачная селитра (34,4% N)', 3770000, 'сум/т', 'ФАКТ', 'UZEX, август 2026: 3 766,5 тыс. сум/т.'),
    ('urea_price_t', 'Карбамид (46% N)', 3000000, 'сум/т', 'ДОГАДКА',
     'Свежей биржевой цены 2026 не найдено; мировая ~700 USD/т (апр. 2026).'),
    ('wheat_price_t', 'Цена пшеницы для фермера', 3500000, 'сум/т', 'ДОГАДКА',
     'С 2022 закупка по рыночным ценам; цифру 2026 не нашли.'),
    ('wheat_yield_t_ha', 'Урожайность пшеницы фермера', 6.5, 'т/га', 'ДОГАДКА', ''),
    ('ground_spread_cost_ha', 'Альтернатива фермера: трактор/вручную', 100000, 'сум/га', 'ДОГАДКА',
     'Ориентир РФ (Татарстан): ~300 руб/га ~ 45 тыс. сум. Местные МТП -- уточнить.'),
    ('', 'КАЛИБРОВКА ПО ЖИДКОЙ ОБРАБОТКЕ', None, '', '', ''),
    ('liq_ha_per_flight', 'Гектаров за вылет при опрыскивании', 1.2, 'га', 'ФАКТ',
     'Данные вылетов: 1,0-1,25 га за вылет, 27-41 л/га.'),
    ('liq_flight_min', 'Длительность вылета при опрыскивании', 6.5, 'мин', 'ОЦЕНКА', ''),
    ('liq_refill_min', 'Заправка раствора и смена батареи', 2.5, 'мин', 'ДОГАДКА', ''),
    ('liq_cycle_equiv_min', 'Минут полёта на цикл батареи (опрыскивание)', 10.0, 'мин', 'ОЦЕНКА',
     'DJI: 7 мин висения при 90 кг.'),
    ('', 'СЕЗОН', None, '', '', ''),
    ('drones_in_winter', 'Бортов в зимней работе', 5, 'шт', 'РЕШЕНИЕ', 'Из 15 в парке.'),
    ('spreaders_owned', 'Разбрасывателей уже есть', 1, 'шт', 'ФАКТ', 'Один тестовый бак.'),
    ('season_one_time_other', 'Разовые расходы сезона (обучение, демо, оснастка загрузки)', 20000000, 'сум', 'ДОГАДКА', ''),
]

# Множители сценариев: (ключ множителя, подпись, параметр, к которому применяется)
MULTS = [
    ('m_eff', 'эффективные часы', 'eff_hours_per_day'),
    ('m_reload', 'время загрузки', 'reload_min'),
    ('m_overhead', 'накладное время вылета', 'overhead_min'),
    ('m_payload', 'полезная нагрузка', 'payload_kg'),
    ('m_batt', 'цена батареи', 'battery_price_usd'),
    ('m_pilot', 'оплата пилота', 'pilot_day'),
    ('m_petrol', 'цена бензина', 'petrol_price'),
    ('m_cycles', 'ресурс батареи', 'battery_cycles'),
    ('m_days', 'рабочих дней в сезоне', 'days_per_season'),
    ('m_swath', 'ширина захвата', 'swath_eff_m'),
    ('m_speed', 'скорость', 'speed_ms'),
]


def base_scenarios():
    """Сценарии: имя, норма кг/га, словарь множителей (остальные = 1)."""
    sc = [
        ('База 100 кг/га', 100, {}),
        ('База 150 кг/га (карбамид)', 150, {}),
        ('База 200 кг/га', 200, {}),
        ('База 250 кг/га (селитра)', 250, {}),
        ('Пессимистичный 250', 250, {'m_eff': 2.0 / 3.0, 'm_reload': 2.0, 'm_overhead': 4.0 / 3.0, 'm_payload': 40.0 / 45.0}),
        ('Оптимистичный 250', 250, {'m_eff': 1.5, 'm_reload': 2.0 / 3.0, 'm_overhead': 0.8, 'm_payload': 50.0 / 45.0}),
        ('Маржинальный 250 (пилот на окладе)', 250, {'m_pilot': 0.4}),
        ('Пессимистичный 150', 150, {'m_eff': 2.0 / 3.0, 'm_reload': 2.0, 'm_overhead': 4.0 / 3.0, 'm_payload': 40.0 / 45.0}),
        ('Оптимистичный 150', 150, {'m_eff': 1.5, 'm_reload': 2.0 / 3.0, 'm_overhead': 0.8, 'm_payload': 50.0 / 45.0}),
        ('Аэросев саксаула 4 кг/га (бак ~12 кг по объёму)', 4, {'m_payload': 12.0 / 45.0}),
        ('Сев люцерны 15 кг/га', 15, {}),
        ('Подсев озимых 60 кг/га', 60, {}),
    ]
    # Чувствительность при 250 кг/га: по одному параметру, -20% / +20% (нагрузка -10/+10%).
    for key, label, _param in MULTS[:9]:
        lo, hi = (0.9, 1.1) if key == 'm_payload' else (0.8, 1.2)
        sc.append(('Чувств. 250: %s x%.2f' % (label, lo), 250, {key: lo}))
        sc.append(('Чувств. 250: %s x%.2f' % (label, hi), 250, {key: hi}))
    return sc


def params_dict():
    return {k: v for k, _n, v, _u, _t, _s in PARAMS if k}


# ---------------------------------------------------------------------------
# Независимый расчёт на Python (зеркало формул листа «Расчёт»)
# ---------------------------------------------------------------------------
def compute(p, rate, mults):
    m = {k: 1.0 for k, _l, _p in MULTS}
    m.update(mults)
    payload = p['payload_kg'] * m['m_payload']
    eff_hours = p['eff_hours_per_day'] * m['m_eff']
    reload = p['reload_min'] * m['m_reload']
    overhead = p['overhead_min'] * m['m_overhead']
    batt_price = p['battery_price_usd'] * m['m_batt']
    pilot = p['pilot_day'] * m['m_pilot']
    petrol = p['petrol_price'] * m['m_petrol']
    cycles = p['battery_cycles'] * m['m_cycles']
    days = p['days_per_season'] * m['m_days']
    swath = p['swath_eff_m'] * m['m_swath']
    speed = p['speed_ms'] * m['m_speed']
    usd = p['usd_rate']

    r = {}
    r['ha_per_load_weight'] = payload / rate
    r['spread_min_per_ha'] = 10000.0 / (swath * speed) / 60.0
    # [REASON]: при лёгких нормах (семена) 45 кг хватило бы на 15-20 минут
    # разбрасывания, но батареи столько не держат: вылет режется по
    # max_flight_min, и гектаров за вылет становится меньше, чем вмещает бак.
    r['spread_min_per_flight'] = min(r['ha_per_load_weight'] * r['spread_min_per_ha'],
                                     p['max_flight_min'] - overhead)
    r['ha_per_flight'] = r['spread_min_per_flight'] / r['spread_min_per_ha']
    r['loads_per_ha'] = 1.0 / r['ha_per_flight']
    r['kg_per_flight'] = rate * r['ha_per_flight']
    r['flight_min'] = r['spread_min_per_flight'] + overhead
    r['cycle_min'] = r['flight_min'] + reload
    r['flights_per_hour'] = 60.0 / r['cycle_min']
    r['kg_per_hour'] = r['flights_per_hour'] * r['kg_per_flight']
    r['ha_per_hour'] = r['flights_per_hour'] * r['ha_per_flight']
    r['ha_per_day'] = r['ha_per_hour'] * eff_hours
    r['flights_per_day'] = r['flights_per_hour'] * eff_hours
    r['t_per_day'] = r['kg_per_hour'] * eff_hours / 1000.0
    r['required_flow'] = r['kg_per_flight'] / r['spread_min_per_flight']
    r['batt_fraction'] = r['flight_min'] / p['cycle_equiv_min']
    r['charges_per_day'] = r['flights_per_day'] * r['batt_fraction']

    r['battery_cost_per_flight'] = batt_price * usd / cycles * r['batt_fraction']
    r['fuel_per_flight'] = p['battery_kwh'] * r['batt_fraction'] / p['charge_eff'] * p['gen_l_per_kwh'] * petrol
    r['wear_per_flight'] = p['drone_wear_usd_h'] * usd * r['flight_min'] / 60.0 * (1 + p['fert_wear_premium'])
    r['per_flight_total'] = r['battery_cost_per_flight'] + r['fuel_per_flight'] + r['wear_per_flight']

    r['crew_day'] = (pilot + p['helpers_count'] * p['helper_day']
                     + p['vehicle_km'] * (p['diesel_l_100km'] / 100.0 * p['diesel_price'] + p['vehicle_wear_km'])
                     + p['per_diem_day'] + p['other_day'])
    r['flight_costs_per_day'] = r['per_flight_total'] * r['flights_per_day']
    r['spreader_per_day'] = p['spreader_price_usd'] * usd / p['spreader_life_seasons'] / days
    r['drone_dep_per_day'] = p['drone_price_usd'] * usd / p['drone_life_years'] * p['drone_dep_winter_share'] / days
    r['total_cost_per_day'] = r['crew_day'] + r['flight_costs_per_day'] + r['spreader_per_day'] + r['drone_dep_per_day']

    r['flight_costs_per_ha'] = r['per_flight_total'] * r['loads_per_ha']
    r['crew_per_ha'] = r['crew_day'] / r['ha_per_day']
    r['spreader_per_ha'] = r['spreader_per_day'] / r['ha_per_day']
    r['drone_dep_per_ha'] = r['drone_dep_per_day'] / r['ha_per_day']
    r['direct_cost_per_ha'] = (r['flight_costs_per_ha'] + r['crew_per_ha']
                               + r['spreader_per_ha'] + r['drone_dep_per_ha'])
    oh, mg = p['overhead_pct'], p['target_margin_pct']
    r['breakeven_tariff'] = r['direct_cost_per_ha'] / (1 - oh)
    r['target_tariff'] = r['direct_cost_per_ha'] / (1 - oh - mg)
    r['breakeven_per_ton'] = r['breakeven_tariff'] / rate * 1000.0
    r['target_per_ton'] = r['target_tariff'] / rate * 1000.0
    r['revenue_per_day_target'] = r['target_tariff'] * r['ha_per_day']
    r['margin_per_day_target'] = r['revenue_per_day_target'] - r['total_cost_per_day'] - r['revenue_per_day_target'] * oh
    r['ratio_to_liquid'] = r['target_tariff'] / p['liquid_tariff_external']

    # Линейное разложение: себестоимость = C0 (за гектар) + c1 (за кг) x норма.
    r['fixed_per_min'] = (r['crew_day'] + r['spreader_per_day'] + r['drone_dep_per_day']) / (eff_hours * 60.0)
    r['energy_per_flight_min'] = ((batt_price * usd / cycles
                                   + p['battery_kwh'] / p['charge_eff'] * p['gen_l_per_kwh'] * petrol)
                                  / p['cycle_equiv_min']
                                  + p['drone_wear_usd_h'] * usd / 60.0 * (1 + p['fert_wear_premium']))
    r['c0_per_ha'] = r['spread_min_per_ha'] * (r['fixed_per_min'] + r['energy_per_flight_min'])
    r['c1_per_kg'] = ((overhead + reload) * r['fixed_per_min'] + overhead * r['energy_per_flight_min']) / payload
    r['linearity_check'] = r['c0_per_ha'] + r['c1_per_kg'] * rate - r['direct_cost_per_ha']
    r['tariff_base_per_ha'] = r['c0_per_ha'] / (1 - oh - mg)
    r['tariff_per_kg'] = r['c1_per_kg'] / (1 - oh - mg)
    r['tariff_per_ton'] = r['tariff_per_kg'] * 1000.0
    return r


def compute_liquid(p):
    usd = p['usd_rate']
    r = {}
    r['liq_flights_per_hour'] = 60.0 / (p['liq_flight_min'] + p['liq_refill_min'])
    r['liq_ha_per_day'] = r['liq_flights_per_hour'] * p['liq_ha_per_flight'] * p['eff_hours_per_day']
    frac = p['liq_flight_min'] / p['liq_cycle_equiv_min']
    r['liq_batt_fraction'] = frac
    r['liq_battery_per_flight'] = p['battery_price_usd'] * usd / p['battery_cycles'] * frac
    r['liq_fuel_per_flight'] = p['battery_kwh'] * frac / p['charge_eff'] * p['gen_l_per_kwh'] * p['petrol_price']
    r['liq_wear_per_flight'] = p['drone_wear_usd_h'] * usd * p['liq_flight_min'] / 60.0
    r['liq_per_flight_total'] = r['liq_battery_per_flight'] + r['liq_fuel_per_flight'] + r['liq_wear_per_flight']
    r['liq_flight_costs_per_ha'] = r['liq_per_flight_total'] / p['liq_ha_per_flight']
    crew_day = (p['pilot_day'] + p['helpers_count'] * p['helper_day']
                + p['vehicle_km'] * (p['diesel_l_100km'] / 100.0 * p['diesel_price'] + p['vehicle_wear_km'])
                + p['per_diem_day'] + p['other_day'])
    r['liq_crew_per_ha'] = crew_day / r['liq_ha_per_day']
    r['liq_direct_per_ha'] = r['liq_flight_costs_per_ha'] + r['liq_crew_per_ha']
    r['liq_margin_external'] = p['liquid_tariff_external'] - r['liq_direct_per_ha'] - p['liquid_tariff_external'] * p['overhead_pct']
    r['liq_internal_minus_direct'] = p['liquid_tariff_internal'] - r['liq_direct_per_ha']
    return r


def compute_farmer(p, base):
    """Сторона фермера для нормы 200 кг/га селитры по двухставочному тарифу базы 250."""
    rate = 200.0
    r = {}
    r['f_rate'] = rate
    r['f_fert_cost'] = rate / 1000.0 * p['nitrate_price_t']
    r['f_service'] = base['tariff_base_per_ha'] + base['tariff_per_kg'] * rate
    r['f_total_drone'] = r['f_fert_cost'] + r['f_service']
    r['f_total_ground'] = r['f_fert_cost'] + p['ground_spread_cost_ha']
    r['f_premium'] = r['f_service'] - p['ground_spread_cost_ha']
    r['f_breakeven_gain_kg'] = r['f_premium'] / p['wheat_price_t'] * 1000.0
    r['f_breakeven_gain_pct'] = r['f_breakeven_gain_kg'] / 1000.0 / p['wheat_yield_t_ha']
    r['f_service_share_of_fert'] = r['f_service'] / r['f_fert_cost']
    return r


def compute_season(p, scen):
    r = {}
    r['s_ha_per_drone'] = scen['ha_per_day'] * p['days_per_season']
    r['s_ha_total'] = r['s_ha_per_drone'] * p['drones_in_winter']
    r['s_revenue'] = r['s_ha_total'] * scen['target_tariff']
    r['s_direct'] = r['s_ha_total'] * scen['direct_cost_per_ha']
    r['s_overhead'] = r['s_revenue'] * p['overhead_pct']
    r['s_spreaders_to_buy'] = max(0, p['drones_in_winter'] - p['spreaders_owned'])
    r['s_one_time'] = r['s_spreaders_to_buy'] * p['spreader_price_usd'] * p['usd_rate'] + p['season_one_time_other']
    r['s_margin_before_one_time'] = r['s_revenue'] - r['s_direct'] - r['s_overhead']
    r['s_result'] = r['s_margin_before_one_time'] - r['s_one_time']
    r['s_result_usd'] = r['s_result'] / p['usd_rate']
    return r


# ---------------------------------------------------------------------------
# Книга Excel
# ---------------------------------------------------------------------------
HDR = Font(bold=True)
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')   # жёлтый: вход
KEY_FILL = PatternFill('solid', fgColor='E2EFDA')     # зелёный: ключевой итог
SECTION_FILL = PatternFill('solid', fgColor='D9E1F2')
WRAP = Alignment(wrap_text=True, vertical='top')

SHEET_PARAMS = 'Параметры'
SHEET_CALC = 'Расчёт'
SHEET_SENS = 'Чувствительность'
SHEET_SEASON = 'Сезон'
SHEET_CMP = 'Сравнение'
SHEET_KP = 'Тариф для КП'
SHEET_SRC = 'Источники'

# Строки листа «Расчёт»: (ключ, подпись, единица, формула с {c} = колонка, {r:key} = строка ключа)
CALC_ROWS = [
    ('__scenario', 'Сценарий', '', None),
    ('rate', 'Норма внесения', 'кг/га', None),
    ('m_eff', 'Множитель: эффективные часы', 'x', None),
    ('m_reload', 'Множитель: время загрузки', 'x', None),
    ('m_overhead', 'Множитель: накладное время вылета', 'x', None),
    ('m_payload', 'Множитель: полезная нагрузка', 'x', None),
    ('m_batt', 'Множитель: цена батареи', 'x', None),
    ('m_pilot', 'Множитель: оплата пилота', 'x', None),
    ('m_petrol', 'Множитель: цена бензина', 'x', None),
    ('m_cycles', 'Множитель: ресурс батареи', 'x', None),
    ('m_days', 'Множитель: рабочих дней в сезоне', 'x', None),
    ('m_swath', 'Множитель: ширина захвата', 'x', None),
    ('m_speed', 'Множитель: скорость', 'x', None),
    ('__sec1', 'ДЕЙСТВУЮЩИЕ ВХОДЫ СЦЕНАРИЯ', '', None),
    ('payload', 'Полезная нагрузка', 'кг', '=payload_kg*{c}{r:m_payload}'),
    ('eff_hours', 'Эффективное лётное время', 'ч/день', '=eff_hours_per_day*{c}{r:m_eff}'),
    ('reload', 'Загрузка и смена батареи', 'мин', '=reload_min*{c}{r:m_reload}'),
    ('overhead', 'Накладное время вылета', 'мин', '=overhead_min*{c}{r:m_overhead}'),
    ('batt_price', 'Цена батареи', 'USD', '=battery_price_usd*{c}{r:m_batt}'),
    ('pilot', 'Оплата пилота', 'сум/день', '=pilot_day*{c}{r:m_pilot}'),
    ('petrol', 'Цена бензина', 'сум/л', '=petrol_price*{c}{r:m_petrol}'),
    ('cycles', 'Ресурс батареи', 'циклов', '=battery_cycles*{c}{r:m_cycles}'),
    ('days', 'Рабочих дней в сезоне', 'дн', '=days_per_season*{c}{r:m_days}'),
    ('swath', 'Ширина захвата', 'м', '=swath_eff_m*{c}{r:m_swath}'),
    ('speed', 'Скорость', 'м/с', '=speed_ms*{c}{r:m_speed}'),
    ('__sec2', 'ПРОИЗВОДИТЕЛЬНОСТЬ', '', None),
    ('ha_per_load_weight', 'Гектаров на одну загрузку по весу бака', 'га', '={c}{r:payload}/{c}{r:rate}'),
    ('spread_min_per_ha', 'Время разбрасывания на гектар (без накладных)', 'мин', '=10000/({c}{r:swath}*{c}{r:speed})/60'),
    ('spread_min_per_flight', 'Время разбрасывания за вылет (вес бака или предел батареи)', 'мин', '=MIN({c}{r:ha_per_load_weight}*{c}{r:spread_min_per_ha},max_flight_min-{c}{r:overhead})'),
    ('ha_per_flight', 'Гектаров за вылет', 'га', '={c}{r:spread_min_per_flight}/{c}{r:spread_min_per_ha}'),
    ('loads_per_ha', 'Вылетов на гектар', 'шт', '=1/{c}{r:ha_per_flight}'),
    ('kg_per_flight', 'Внесено за вылет', 'кг', '={c}{r:rate}*{c}{r:ha_per_flight}'),
    ('flight_min', 'Длительность вылета', 'мин', '={c}{r:spread_min_per_flight}+{c}{r:overhead}'),
    ('cycle_min', 'Полный цикл (вылет + загрузка)', 'мин', '={c}{r:flight_min}+{c}{r:reload}'),
    ('flights_per_hour', 'Вылетов в час', 'шт', '=60/{c}{r:cycle_min}'),
    ('kg_per_hour', 'Внесено удобрений в час', 'кг', '={c}{r:flights_per_hour}*{c}{r:kg_per_flight}'),
    ('ha_per_hour', 'Гектаров в час', 'га', '={c}{r:flights_per_hour}*{c}{r:ha_per_flight}'),
    ('ha_per_day', 'Гектаров за рабочий день', 'га', '={c}{r:ha_per_hour}*{c}{r:eff_hours}'),
    ('flights_per_day', 'Вылетов за день', 'шт', '={c}{r:flights_per_hour}*{c}{r:eff_hours}'),
    ('t_per_day', 'Тонн за день', 'т', '={c}{r:kg_per_hour}*{c}{r:eff_hours}/1000'),
    ('required_flow', 'Требуемая подача разбрасывателя', 'кг/мин', '={c}{r:kg_per_flight}/{c}{r:spread_min_per_flight}'),
    ('flow_check', 'Подача в пределах паспорта?', '', '=IF({c}{r:required_flow}<=max_flow_kg_min,"да","НЕТ: снизить скорость")'),
    ('batt_fraction', 'Доля цикла батареи за вылет', 'доля', '={c}{r:flight_min}/cycle_equiv_min'),
    ('charges_per_day', 'Полных зарядок батарей за день', 'шт', '={c}{r:flights_per_day}*{c}{r:batt_fraction}'),
    ('__sec3', 'ЗАТРАТЫ НА ОДИН ВЫЛЕТ', '', None),
    ('battery_cost_per_flight', 'Износ батареи', 'сум', '={c}{r:batt_price}*usd_rate/{c}{r:cycles}*{c}{r:batt_fraction}'),
    ('fuel_per_flight', 'Бензин генератора', 'сум', '=battery_kwh*{c}{r:batt_fraction}/charge_eff*gen_l_per_kwh*{c}{r:petrol}'),
    ('wear_per_flight', 'Износ дрона (с надбавкой за удобрения)', 'сум', '=drone_wear_usd_h*usd_rate*{c}{r:flight_min}/60*(1+fert_wear_premium)'),
    ('per_flight_total', 'Итого на вылет', 'сум', '={c}{r:battery_cost_per_flight}+{c}{r:fuel_per_flight}+{c}{r:wear_per_flight}'),
    ('__sec4', 'ЗАТРАТЫ НА РАБОЧИЙ ДЕНЬ', '', None),
    ('crew_day', 'Экипаж, транспорт, прочее', 'сум', '={c}{r:pilot}+helpers_count*helper_day+vehicle_km*(diesel_l_100km/100*diesel_price+vehicle_wear_km)+per_diem_day+other_day'),
    ('flight_costs_per_day', 'Батареи, бензин, износ за все вылеты', 'сум', '={c}{r:per_flight_total}*{c}{r:flights_per_day}'),
    ('spreader_per_day', 'Амортизация разбрасывателя', 'сум', '=spreader_price_usd*usd_rate/spreader_life_seasons/{c}{r:days}'),
    ('drone_dep_per_day', 'Амортизация дрона (зимняя доля)', 'сум', '=drone_price_usd*usd_rate/drone_life_years*drone_dep_winter_share/{c}{r:days}'),
    ('total_cost_per_day', 'Итого за день', 'сум', '={c}{r:crew_day}+{c}{r:flight_costs_per_day}+{c}{r:spreader_per_day}+{c}{r:drone_dep_per_day}'),
    ('__sec5', 'СЕБЕСТОИМОСТЬ И ТАРИФ НА ГЕКТАР', '', None),
    ('flight_costs_per_ha', 'Батареи, бензин, износ', 'сум/га', '={c}{r:per_flight_total}*{c}{r:loads_per_ha}'),
    ('crew_per_ha', 'Экипаж, транспорт, прочее', 'сум/га', '={c}{r:crew_day}/{c}{r:ha_per_day}'),
    ('spreader_per_ha', 'Амортизация разбрасывателя', 'сум/га', '={c}{r:spreader_per_day}/{c}{r:ha_per_day}'),
    ('drone_dep_per_ha', 'Амортизация дрона (зимняя доля)', 'сум/га', '={c}{r:drone_dep_per_day}/{c}{r:ha_per_day}'),
    ('direct_cost_per_ha', 'ПРЯМАЯ СЕБЕСТОИМОСТЬ ГЕКТАРА', 'сум/га', '={c}{r:flight_costs_per_ha}+{c}{r:crew_per_ha}+{c}{r:spreader_per_ha}+{c}{r:drone_dep_per_ha}'),
    ('breakeven_tariff', 'Тариф безубыточности (с накладными, без маржи)', 'сум/га', '={c}{r:direct_cost_per_ha}/(1-overhead_pct)'),
    ('target_tariff', 'ЦЕЛЕВОЙ ТАРИФ (с накладными и маржой)', 'сум/га', '={c}{r:direct_cost_per_ha}/(1-overhead_pct-target_margin_pct)'),
    ('breakeven_per_ton', 'Тариф безубыточности за тонну удобрения', 'сум/т', '={c}{r:breakeven_tariff}/{c}{r:rate}*1000'),
    ('target_per_ton', 'Целевой тариф за тонну удобрения', 'сум/т', '={c}{r:target_tariff}/{c}{r:rate}*1000'),
    ('revenue_per_day_target', 'Выручка за день по целевому тарифу', 'сум', '={c}{r:target_tariff}*{c}{r:ha_per_day}'),
    ('margin_per_day_target', 'Маржа за день по целевому тарифу', 'сум', '={c}{r:revenue_per_day_target}-{c}{r:total_cost_per_day}-{c}{r:revenue_per_day_target}*overhead_pct'),
    ('ratio_to_liquid', 'Целевой тариф к тарифу опрыскивания', 'раз', '={c}{r:target_tariff}/liquid_tariff_external'),
    ('__sec6', 'ДВУХСТАВОЧНЫЙ ТАРИФ: ГЕКТАР + ТОННА', '', None),
    ('fixed_per_min', 'Постоянные затраты на минуту лётного времени', 'сум/мин', '=({c}{r:crew_day}+{c}{r:spreader_per_day}+{c}{r:drone_dep_per_day})/({c}{r:eff_hours}*60)'),
    ('energy_per_flight_min', 'Батарея, бензин, износ на минуту полёта', 'сум/мин', '=({c}{r:batt_price}*usd_rate/{c}{r:cycles}+battery_kwh/charge_eff*gen_l_per_kwh*{c}{r:petrol})/cycle_equiv_min+drone_wear_usd_h*usd_rate/60*(1+fert_wear_premium)'),
    ('c0_per_ha', 'Себестоимость: постоянная часть на гектар (C0)', 'сум/га', '={c}{r:spread_min_per_ha}*({c}{r:fixed_per_min}+{c}{r:energy_per_flight_min})'),
    ('c1_per_kg', 'Себестоимость: переменная часть на килограмм (c1)', 'сум/кг', '=(({c}{r:overhead}+{c}{r:reload})*{c}{r:fixed_per_min}+{c}{r:overhead}*{c}{r:energy_per_flight_min})/{c}{r:payload}'),
    ('linearity_check', 'Контроль: C0 + c1 x норма - прямая себестоимость (0, пока вылет ограничен весом бака, а не батареей)', 'сум', '={c}{r:c0_per_ha}+{c}{r:c1_per_kg}*{c}{r:rate}-{c}{r:direct_cost_per_ha}'),
    ('tariff_base_per_ha', 'ТАРИФ: ставка за гектар', 'сум/га', '={c}{r:c0_per_ha}/(1-overhead_pct-target_margin_pct)'),
    ('tariff_per_kg', 'ТАРИФ: ставка за килограмм удобрения', 'сум/кг', '={c}{r:c1_per_kg}/(1-overhead_pct-target_margin_pct)'),
    ('tariff_per_ton', 'ТАРИФ: ставка за тонну удобрения', 'сум/т', '={c}{r:tariff_per_kg}*1000'),
]

KEY_ROWS = {'direct_cost_per_ha', 'target_tariff', 'ha_per_day', 'tariff_base_per_ha', 'tariff_per_ton', 'target_per_ton'}
MONEY_ROWS = {'battery_cost_per_flight', 'fuel_per_flight', 'wear_per_flight', 'per_flight_total', 'crew_day',
              'flight_costs_per_day', 'spreader_per_day', 'drone_dep_per_day', 'total_cost_per_day',
              'flight_costs_per_ha', 'crew_per_ha', 'spreader_per_ha', 'drone_dep_per_ha', 'direct_cost_per_ha',
              'breakeven_tariff', 'target_tariff', 'breakeven_per_ton', 'target_per_ton', 'revenue_per_day_target',
              'margin_per_day_target', 'fixed_per_min', 'energy_per_flight_min', 'c0_per_ha', 'c1_per_kg',
              'linearity_check', 'tariff_base_per_ha', 'tariff_per_kg', 'tariff_per_ton', 'pilot', 'batt_price', 'petrol'}

SOURCES = [
    ('ФАКТ', 'DJI Agras T40: бак разбрасывателя 70 л / 50 кг, ширина 7 м, гранулы 0,5-5 мм, висение 6 мин при 101 кг, 7 мин при 90 кг, 18 мин при 50 кг; заявка 1,5 т/ч* при 149 кг/га, 7 м/с, 3 м', 'Страницы DJI (dji.com/t40/specs, support) в пересказе поисковой выдачи; ритейлеры DJI (Dronefly, Ferntech, Aerosmart)'),
    ('ФАКТ', 'Батарея DB1560 30 000 мА·ч, 52,22 В; розница США 2 499 USD; гарантия 1 000-1 500 циклов (ритейлеры расходятся)', 'Dronefly, DSLRPros, Reach.dog, Maverick Drone'),
    ('ФАКТ', 'Генератор D12000iE: 500 мл/кВт·ч, заряд батареи T40 9-12 мин', 'Dronefly, Aerosmart, manualslib (пересказ)'),
    ('ФАКТ', 'Разбрасыватель T40 Spreading System 3.0: 1 199-1 499 USD; подача до 90 кг/мин', 'Uniquephoto, Rippingitoutdoors, rcdrone.top'),
    ('ФАКТ', 'Дрон T40 б/у 2025 с 3 АКБ: 19 000 USD (США); новый без АКБ 13 499-15 999 USD', 'Controller.com, Reach.dog'),
    ('ФАКТ', 'АИ-92 Ташкент 22.09.2026: 12 600 / 13 900 сум/л; дизель 14 700-15 900 сум/л; биржа дизель апр. 2026 15,86 млн сум/т', 'Avtoelon.uz, Autostrada.uz, Kun.uz'),
    ('ФАКТ', 'Курс ЦБ РУз 29.10.2026: 11 965-12 004 сум/USD', 'Kursiv.media, Repost.uz'),
    ('ФАКТ', 'Аммиачная селитра UZEX, авг. 2026: 3 813,2 -> 3 766,5 тыс. сум/т', 'Yuz.uz (биржевой обзор)'),
    ('ФАКТ', 'Внутренняя ставка опрыскивания 75 040-86 000 сум/га; внешняя 200 000 сум/га; 1,0-1,25 га за вылет; 27-41 л/га; сентябрь 2026 ~7 071 га RAW по 15 бортам', 'Репозиторий: docs/tracks/drones.md (DRONE-MONEY-001, DRONE-OCT2025-PRICE-001, UAT 30.09.2026)'),
    ('ФАКТ', 'Селитра -- наиболее коррозионно-активное удобрение; следы селитры ускоряют коррозию в 2,5 раза', 'РГАУ-МСХА, Агроинженерия 2026; Тамбов 2018'),
    ('ОЦЕНКА', 'Весенняя подкормка озимой пшеницы: Беларусь 120-200 кг/га селитры (оптимум 180); Украина 30-60 кг д.в./га; РФ 65-75 кг д.в./га первой подкормкой; Узбекистан (Хорезм) оптимум 120-180 кг N/га за сезон в зависимости от N почвы; фон N180P90K60 в Кашкадарье', 'Горки 2006-08; Зернові культури 2020; Новомосковск; UNESCO/ZEF Khorezm 2005-07; AGRIS Кашкадарья'),
    ('ОЦЕНКА', 'Первая весенняя подкормка в Узбекистане -- по талой почве аммиачной селитрой (Карши, 2025); три срока подкормок', 'vestnik.nauka.kz (аннотация)'),
    ('ОЦЕНКА', 'Бухара: январь ~1 °C, февраль ~6 °C; дней с осадками 1-5 в месяц; февраль -- самый облачный месяц', 'Climate-zone.com, Weather Network (аэропорт Бухара), Яндекс Погода'),
    ('ОЦЕНКА', 'Ориентир наземной подкормки: сервисные организации Татарстана ~300 руб/га, по мерзлоталой почве -- «Туманы» на шинах низкого давления', 'Минсельхоз Татарстана (доклад, без даты)'),
    ('ОЦЕНКА', 'США: услуги агродронов 8-25 USD/акр; расчёт Missouri Extension: 12,27 USD/акр собственная обработка, наём ~16 USD/акр (жидкость)', 'Farmonaut, Missouri Extension g01274'),
    ('ДОГАДКА', 'Оплата пилота/помощников, пробег, износ, КПД зарядки, цена карбамида и пшеницы 2026, стоимость наземной альтернативы в районе', 'Не найдено в открытых источниках -- заданы как редактируемые входы'),
]


def build_workbook(p):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PARAMS
    ws.append(['Код', 'Параметр', 'Значение', 'Единица', 'Тип', 'Источник / комментарий'])
    for c in range(1, 7):
        ws.cell(row=1, column=c).font = HDR
    row = 2
    for key, name, value, unit, kind, src in PARAMS:
        if not key:
            ws.cell(row=row, column=2, value=name).font = HDR
            for c in range(1, 7):
                ws.cell(row=row, column=c).fill = SECTION_FILL
            row += 1
            continue
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=name)
        vc = ws.cell(row=row, column=3, value=value)
        vc.fill = INPUT_FILL
        if isinstance(value, float) and value < 1:
            vc.number_format = '0.00'
        elif isinstance(value, (int, float)) and abs(value) >= 1000:
            vc.number_format = '#,##0'
        ws.cell(row=row, column=4, value=unit)
        ws.cell(row=row, column=5, value=kind)
        ws.cell(row=row, column=6, value=src).alignment = WRAP
        wb.defined_names.add(DefinedName(key, attr_text="'%s'!$C$%d" % (SHEET_PARAMS, row)))
        row += 1
    ws.cell(row=row + 1, column=2, value='Жёлтые ячейки -- входы. Всё остальное в книге считается из них. '
            'Тип: ФАКТ -- с источником; ОЦЕНКА -- вывод из фактов; ДОГАДКА -- не проверено; РЕШЕНИЕ -- задаёт владелец.').alignment = WRAP
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 58
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 80
    ws.freeze_panes = 'A2'

    # ---- Расчёт ----
    wc = wb.create_sheet(SHEET_CALC)
    scenarios = base_scenarios()
    rows_index = {key: i + 1 for i, (key, _l, _u, _f) in enumerate(CALC_ROWS)}
    for i, (key, label, unit, _f) in enumerate(CALC_ROWS):
        r = i + 1
        wc.cell(row=r, column=1, value=label)
        wc.cell(row=r, column=2, value=unit)
        if key.startswith('__'):
            wc.cell(row=r, column=1).font = HDR
            for c in range(1, 3 + len(scenarios)):
                wc.cell(row=r, column=c).fill = SECTION_FILL
        if key in KEY_ROWS:
            wc.cell(row=r, column=1).font = HDR
    mult_keys = [k for k, _l, _p in MULTS]
    for j, (name, rate, mults) in enumerate(scenarios):
        col = 3 + j
        cl = get_column_letter(col)
        for key, label, unit, fml in CALC_ROWS:
            r = rows_index[key]
            cell = wc.cell(row=r, column=col)
            if key == '__scenario':
                cell.value = name
                cell.font = HDR
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            elif key == 'rate':
                cell.value = rate
                cell.fill = INPUT_FILL
            elif key in mult_keys:
                cell.value = mults.get(key, 1.0)
                cell.number_format = '0.00'
                cell.fill = INPUT_FILL
            elif fml:
                refs = {}
                for k2 in rows_index:
                    refs[k2] = rows_index[k2]
                text = fml.replace('{c}', cl)
                # подстановка {r:key}
                while '{r:' in text:
                    s = text.index('{r:')
                    e = text.index('}', s)
                    k2 = text[s + 3:e]
                    text = text[:s] + str(rows_index[k2]) + text[e + 1:]
                cell.value = text
                if key in MONEY_ROWS:
                    cell.number_format = '#,##0'
                elif key in ('ratio_to_liquid', 'batt_fraction'):
                    cell.number_format = '0.00'
                elif key != 'flow_check':
                    cell.number_format = '0.00'
                if key in KEY_ROWS:
                    cell.fill = KEY_FILL
        wc.column_dimensions[cl].width = 16
    wc.column_dimensions['A'].width = 60
    wc.column_dimensions['B'].width = 10
    wc.freeze_panes = 'C2'
    wc.row_dimensions[1].height = 48

    # ---- Чувствительность ----
    wsn = wb.create_sheet(SHEET_SENS)
    wsn.append(['Параметр (сценарий 250 кг/га)', 'Множитель', 'Целевой тариф, сум/га', 'Отклонение от базы, %',
                'Га за день', 'Отклонение, %'])
    for c in range(1, 7):
        wsn.cell(row=1, column=c).font = HDR
    base_col = get_column_letter(3 + 3)  # «База 250» -- четвёртый сценарий
    rt, rh = rows_index['target_tariff'], rows_index['ha_per_day']
    wsn.append(['База 250 кг/га', 1.0, "='%s'!%s%d" % (SHEET_CALC, base_col, rt), 0,
                "='%s'!%s%d" % (SHEET_CALC, base_col, rh), 0])
    for j, (name, rate, mults) in enumerate(scenarios):
        if not name.startswith('Чувств.'):
            continue
        cl = get_column_letter(3 + j)
        (mk, mv), = mults.items()
        label = [l for k, l, _p in MULTS if k == mk][0]
        r = wsn.max_row + 1
        wsn.cell(row=r, column=1, value=label)
        wsn.cell(row=r, column=2, value=mv).number_format = '0.00'
        wsn.cell(row=r, column=3, value="='%s'!%s%d" % (SHEET_CALC, cl, rt)).number_format = '#,##0'
        wsn.cell(row=r, column=4, value="=C%d/$C$2-1" % r).number_format = '0.0%'
        wsn.cell(row=r, column=5, value="='%s'!%s%d" % (SHEET_CALC, cl, rh)).number_format = '0.00'
        wsn.cell(row=r, column=6, value="=E%d/$E$2-1" % r).number_format = '0.0%'
    wsn.cell(row=wsn.max_row + 2, column=1,
             value='Чтение: чем больше отклонение при одинаковом множителе, тем сильнее параметр влияет на тариф. '
                   'Это и есть то, что надо измерять в пилоте первым.').alignment = WRAP
    for cl, w in (('A', 34), ('B', 12), ('C', 22), ('D', 20), ('E', 12), ('F', 14)):
        wsn.column_dimensions[cl].width = w

    # ---- Сезон ----
    wss = wb.create_sheet(SHEET_SEASON)
    wss.append(['Показатель', 'Единица', 'База 150 кг/га (карбамид)', 'База 250 кг/га (селитра)'])
    for c in range(1, 5):
        wss.cell(row=1, column=c).font = HDR
    col150 = get_column_letter(3 + 1)
    col250 = get_column_letter(3 + 3)
    season_rows = [
        ('s_ha_per_drone', 'Гектаров за сезон на один борт', 'га', "='%s'!{c}%d*days_per_season" % (SHEET_CALC, rows_index['ha_per_day'])),
        ('s_ha_total', 'Гектаров за сезон, все борта', 'га', '={me}{r:s_ha_per_drone}*drones_in_winter'),
        ('s_revenue', 'Выручка по целевому тарифу', 'сум', "={me}{r:s_ha_total}*'%s'!{c}%d" % (SHEET_CALC, rows_index['target_tariff'])),
        ('s_direct', 'Прямые затраты', 'сум', "={me}{r:s_ha_total}*'%s'!{c}%d" % (SHEET_CALC, rows_index['direct_cost_per_ha'])),
        ('s_overhead', 'Накладные', 'сум', '={me}{r:s_revenue}*overhead_pct'),
        ('s_spreaders_to_buy', 'Разбрасывателей докупить', 'шт', '=MAX(0,drones_in_winter-spreaders_owned)'),
        ('s_one_time', 'Разовые расходы сезона', 'сум', '={me}{r:s_spreaders_to_buy}*spreader_price_usd*usd_rate+season_one_time_other'),
        ('s_margin_before_one_time', 'Маржа до разовых расходов', 'сум', '={me}{r:s_revenue}-{me}{r:s_direct}-{me}{r:s_overhead}'),
        ('s_result', 'РЕЗУЛЬТАТ СЕЗОНА', 'сум', '={me}{r:s_margin_before_one_time}-{me}{r:s_one_time}'),
        ('s_result_usd', 'Результат сезона', 'USD', '={me}{r:s_result}/usd_rate'),
    ]
    s_index = {k: i + 2 for i, (k, _l, _u, _f) in enumerate(season_rows)}
    for i, (key, label, unit, fml) in enumerate(season_rows):
        r = i + 2
        wss.cell(row=r, column=1, value=label)
        wss.cell(row=r, column=2, value=unit)
        for me, c in (('C', col150), ('D', col250)):
            text = fml.replace('{c}', c).replace('{me}', me)
            while '{r:' in text:
                s = text.index('{r:')
                e = text.index('}', s)
                text = text[:s] + str(s_index[text[s + 3:e]]) + text[e + 1:]
            cell = wss.cell(row=r, column=ord(me) - 64, value=text)
            cell.number_format = '#,##0' if unit in ('сум', 'USD') else '0.00'
            if key == 's_result':
                cell.fill = KEY_FILL
                wss.cell(row=r, column=1).font = HDR
    wss.cell(row=wss.max_row + 2, column=1,
             value='Для сравнения: сентябрь 2026 -- около 7 000 га опрыскивания в месяц по 200 000 сум/га, '
                   'то есть ~1,4 млрд сум выручки за месяц пика. Зимнее разбрасывание -- заполнение простоя, не второй сезон.').alignment = WRAP
    for cl, w in (('A', 44), ('B', 10), ('C', 26), ('D', 26)):
        wss.column_dimensions[cl].width = w

    # ---- Сравнение ----
    wcm = wb.create_sheet(SHEET_CMP)
    wcm.append(['КАЛИБРОВКА: та же модель затрат для текущего опрыскивания', 'Единица', 'Значение'])
    wcm.cell(row=1, column=1).font = HDR
    liq_rows = [
        ('liq_flights_per_hour', 'Вылетов в час', 'шт', '=60/(liq_flight_min+liq_refill_min)'),
        ('liq_ha_per_day', 'Гектаров за день', 'га', '={me}{r:liq_flights_per_hour}*liq_ha_per_flight*eff_hours_per_day'),
        ('liq_batt_fraction', 'Доля цикла батареи за вылет', 'доля', '=liq_flight_min/liq_cycle_equiv_min'),
        ('liq_battery_per_flight', 'Износ батареи за вылет', 'сум', '=battery_price_usd*usd_rate/battery_cycles*{me}{r:liq_batt_fraction}'),
        ('liq_fuel_per_flight', 'Бензин за вылет', 'сум', '=battery_kwh*{me}{r:liq_batt_fraction}/charge_eff*gen_l_per_kwh*petrol_price'),
        ('liq_wear_per_flight', 'Износ дрона за вылет', 'сум', '=drone_wear_usd_h*usd_rate*liq_flight_min/60'),
        ('liq_per_flight_total', 'Итого за вылет', 'сум', '={me}{r:liq_battery_per_flight}+{me}{r:liq_fuel_per_flight}+{me}{r:liq_wear_per_flight}'),
        ('liq_flight_costs_per_ha', 'Батареи, бензин, износ на гектар', 'сум/га', '={me}{r:liq_per_flight_total}/liq_ha_per_flight'),
        ('liq_crew_per_ha', 'Экипаж, транспорт, прочее на гектар', 'сум/га', '=(pilot_day+helpers_count*helper_day+vehicle_km*(diesel_l_100km/100*diesel_price+vehicle_wear_km)+per_diem_day+other_day)/{me}{r:liq_ha_per_day}'),
        ('liq_direct_per_ha', 'ПРЯМАЯ СЕБЕСТОИМОСТЬ ГЕКТАРА ОПРЫСКИВАНИЯ', 'сум/га', '={me}{r:liq_flight_costs_per_ha}+{me}{r:liq_crew_per_ha}'),
        ('liq_margin_external', 'Маржа с гектара при 200 000 (после накладных)', 'сум/га', '=liquid_tariff_external-{me}{r:liq_direct_per_ha}-liquid_tariff_external*overhead_pct'),
        ('liq_internal_minus_direct', 'Внутренняя ставка минус прямая себестоимость', 'сум/га', '=liquid_tariff_internal-{me}{r:liq_direct_per_ha}'),
    ]
    l_index = {k: i + 2 for i, (k, _l, _u, _f) in enumerate(liq_rows)}
    for i, (key, label, unit, fml) in enumerate(liq_rows):
        r = i + 2
        wcm.cell(row=r, column=1, value=label)
        wcm.cell(row=r, column=2, value=unit)
        text = fml.replace('{me}', 'C')
        while '{r:' in text:
            s = text.index('{r:')
            e = text.index('}', s)
            text = text[:s] + str(l_index[text[s + 3:e]]) + text[e + 1:]
        cell = wcm.cell(row=r, column=3, value=text)
        cell.number_format = '#,##0' if 'сум' in unit else '0.00'
        if key == 'liq_direct_per_ha':
            cell.fill = KEY_FILL
            wcm.cell(row=r, column=1).font = HDR
    r0 = wcm.max_row + 1
    wcm.cell(row=r0, column=1, value='Чтение: прямая себестоимость опрыскивания по этой модели должна лечь рядом с внутренней '
             'ставкой 75-86 тыс. сум/га. Если нет -- входы блока «Люди и транспорт» завышены или занижены.').alignment = WRAP
    r1 = r0 + 2
    wcm.cell(row=r1, column=1, value='СТОРОНА ФЕРМЕРА: селитра 200 кг/га, двухставочный тариф базы 250').font = HDR
    farmer_rows = [
        ('f_rate', 'Норма внесения', 'кг/га', '=200'),
        ('f_fert_cost', 'Стоимость удобрения', 'сум/га', '={me}{r:f_rate}/1000*nitrate_price_t'),
        ('f_service', 'Услуга дрона: ставка за га + ставка за кг x норма', 'сум/га',
         "='%s'!%s%d+'%s'!%s%d*{me}{r:f_rate}" % (SHEET_CALC, col250, rows_index['tariff_base_per_ha'], SHEET_CALC, col250, rows_index['tariff_per_kg'])),
        ('f_total_drone', 'Итого подкормка дроном', 'сум/га', '={me}{r:f_fert_cost}+{me}{r:f_service}'),
        ('f_total_ground', 'Итого подкормка трактором/вручную', 'сум/га', '={me}{r:f_fert_cost}+ground_spread_cost_ha'),
        ('f_premium', 'Надбавка за дрон против наземной альтернативы', 'сум/га', '={me}{r:f_service}-ground_spread_cost_ha'),
        ('f_breakeven_gain_kg', 'Прибавка урожая, окупающая надбавку', 'кг/га', '={me}{r:f_premium}/wheat_price_t*1000'),
        ('f_breakeven_gain_pct', 'Та же прибавка в долях урожая', 'доля', '={me}{r:f_breakeven_gain_kg}/1000/wheat_yield_t_ha'),
        ('f_service_share_of_fert', 'Услуга к стоимости удобрения', 'доля', '={me}{r:f_service}/{me}{r:f_fert_cost}'),
    ]
    f_index = {k: r1 + 1 + i for i, (k, _l, _u, _f) in enumerate(farmer_rows)}
    for i, (key, label, unit, fml) in enumerate(farmer_rows):
        r = r1 + 1 + i
        wcm.cell(row=r, column=1, value=label)
        wcm.cell(row=r, column=2, value=unit)
        text = fml.replace('{me}', 'C')
        while '{r:' in text:
            s = text.index('{r:')
            e = text.index('}', s)
            text = text[:s] + str(f_index[text[s + 3:e]]) + text[e + 1:]
        cell = wcm.cell(row=r, column=3, value=text)
        cell.number_format = '0.0%' if unit == 'доля' else ('#,##0' if 'сум' in unit else '0.0')
        if key in ('f_premium', 'f_breakeven_gain_pct'):
            cell.fill = KEY_FILL
    wcm.cell(row=wcm.max_row + 2, column=1,
             value='Чтение: фермеру дрон выгоден, если своевременность и равномерность дают прибавку урожая не меньше '
                   'строки «Прибавка урожая, окупающая надбавку». Это главный аргумент КП, и его надо подтвердить агрономом.').alignment = WRAP
    wcm.column_dimensions['A'].width = 62
    wcm.column_dimensions['B'].width = 10
    wcm.column_dimensions['C'].width = 18

    # ---- Тариф для КП ----
    wk = wb.create_sheet(SHEET_KP)
    wk.append(['ТАРИФ ДЛЯ КП (из сценария «База 250», округление до 5 000 сум)', '', ''])
    wk.cell(row=1, column=1).font = HDR
    wk.append(['Ставка за гектар (выезд, облёт поля)', 'сум/га',
               "=ROUND('%s'!%s%d/5000,0)*5000" % (SHEET_CALC, col250, rows_index['tariff_base_per_ha'])])
    wk.append(['Ставка за тонну внесённого удобрения', 'сум/т',
               "=ROUND('%s'!%s%d/5000,0)*5000" % (SHEET_CALC, col250, rows_index['tariff_per_ton'])])
    wk.cell(row=2, column=3).number_format = '#,##0'
    wk.cell(row=3, column=3).number_format = '#,##0'
    wk.cell(row=2, column=3).fill = KEY_FILL
    wk.cell(row=3, column=3).fill = KEY_FILL
    wk.append([])
    wk.append(['Норма, кг/га', 'Цена для фермера, сум/га', 'К тарифу опрыскивания, раз'])
    wk.cell(row=5, column=1).font = HDR
    wk.cell(row=5, column=2).font = HDR
    wk.cell(row=5, column=3).font = HDR
    for i, rate in enumerate((100, 150, 200, 250, 300)):
        r = 6 + i
        wk.cell(row=r, column=1, value=rate)
        wk.cell(row=r, column=2, value='=$C$2+$C$3*A%d/1000' % r).number_format = '#,##0'
        wk.cell(row=r, column=3, value='=B%d/liquid_tariff_external' % r).number_format = '0.00'
    wk.cell(row=12, column=1, value='Почему две ставки: время над гектаром почти не зависит от нормы, а число вылетов '
            'растёт с ней линейно. Один тариф «за гектар» либо разоряет нас при 250 кг/га, либо отпугивает фермера при 100 кг/га. '
            'Проверка в книге: строка «Контроль линейности» на листе «Расчёт» равна нулю.').alignment = WRAP
    wk.column_dimensions['A'].width = 52
    wk.column_dimensions['B'].width = 26
    wk.column_dimensions['C'].width = 26

    # ---- Источники ----
    wsr = wb.create_sheet(SHEET_SRC)
    wsr.append(['Тип', 'Утверждение', 'Источник'])
    for c in range(1, 4):
        wsr.cell(row=1, column=c).font = HDR
    for kind, claim, src in SOURCES:
        wsr.append([kind, claim, src])
        wsr.cell(row=wsr.max_row, column=2).alignment = WRAP
        wsr.cell(row=wsr.max_row, column=3).alignment = WRAP
    wsr.cell(row=wsr.max_row + 2, column=1,
             value='Прямое чтение сайтов из облачной сессии было закрыто сетевой политикой: цифры взяты из поисковых '
                   'сводок и помечены соответственно. Перед печатью КП сверить цены у поставщика и на UZEX.').alignment = WRAP
    wsr.column_dimensions['A'].width = 10
    wsr.column_dimensions['B'].width = 100
    wsr.column_dimensions['C'].width = 60

    return wb, scenarios, rows_index, s_index, l_index, f_index


# ---------------------------------------------------------------------------
# Проверка: пересчёт в LibreOffice и сверка с Python
# ---------------------------------------------------------------------------
def recalc_with_libreoffice(path):
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        return None
    outdir = tempfile.mkdtemp(prefix='spread_model_')
    cmd = [soffice, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', outdir, path]
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180)
    if res.returncode != 0:
        print('libreoffice failed: rc=%d' % res.returncode)
        return None
    out = os.path.join(outdir, os.path.basename(path))
    return out if os.path.exists(out) else None


def verify(path, p, scenarios, rows_index, s_index, l_index, f_index):
    recalced = recalc_with_libreoffice(path)
    if not recalced:
        print('VERIFY: skipped (LibreOffice not available)')
        return 0
    wb = load_workbook(recalced, data_only=True)
    wc = wb[SHEET_CALC]
    checked = mismatches = 0
    errors = []
    tol = 1e-6
    for j, (name, rate, mults) in enumerate(scenarios):
        exp = compute(p, rate, mults)
        col = 3 + j
        for key, val in exp.items():
            if key not in rows_index:
                continue
            got = wc.cell(row=rows_index[key], column=col).value
            checked += 1
            if not isinstance(got, (int, float)) or abs(got - val) > tol * max(1.0, abs(val)):
                mismatches += 1
                errors.append('calc[%d] %s: excel=%r python=%r' % (j, key, got, val))
        got = wc.cell(row=rows_index['flow_check'], column=col).value
        checked += 1
        want = 'да' if exp['required_flow'] <= p['max_flow_kg_min'] else 'НЕТ: снизить скорость'
        if got != want:
            mismatches += 1
            errors.append('calc[%d] flow_check: excel=%r python=%r' % (j, got, want))
    # Сезон
    wss = wb[SHEET_SEASON]
    for me, idx in (('C', 1), ('D', 3)):
        scen = compute(p, scenarios[idx][1], scenarios[idx][2])
        exp = compute_season(p, scen)
        for key, val in exp.items():
            got = wss.cell(row=s_index[key], column=ord(me) - 64).value
            checked += 1
            if not isinstance(got, (int, float)) or abs(got - val) > tol * max(1.0, abs(val)):
                mismatches += 1
                errors.append('season %s %s: excel=%r python=%r' % (me, key, got, val))
    # Сравнение
    wcm = wb[SHEET_CMP]
    expl = compute_liquid(p)
    for key, val in expl.items():
        got = wcm.cell(row=l_index[key], column=3).value
        checked += 1
        if not isinstance(got, (int, float)) or abs(got - val) > tol * max(1.0, abs(val)):
            mismatches += 1
            errors.append('liquid %s: excel=%r python=%r' % (key, got, val))
    base250 = compute(p, scenarios[3][1], scenarios[3][2])
    expf = compute_farmer(p, base250)
    for key, val in expf.items():
        got = wcm.cell(row=f_index[key], column=3).value
        checked += 1
        if not isinstance(got, (int, float)) or abs(got - val) > tol * max(1.0, abs(val)):
            mismatches += 1
            errors.append('farmer %s: excel=%r python=%r' % (key, got, val))
    # Ошибки формул в любой ячейке
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('#'):
                    mismatches += 1
                    errors.append('formula error at %s!%s: %s' % (ws.title, cell.coordinate, cell.value))
    print('VERIFY: cells checked=%d mismatches=%d (recalculated copy: %s)' % (checked, mismatches, recalced))
    for e in errors[:40]:
        print('  ' + e.encode('ascii', 'backslashreplace').decode('ascii'))
    return 1 if mismatches else 0


def main(argv=None):
    ap = argparse.ArgumentParser(description='Build model.xlsx for winter drone spreading economics.')
    ap.add_argument('--out', default=OUT_DEFAULT)
    ap.add_argument('--verify', action='store_true', help='recalculate in LibreOffice and compare with Python')
    ap.add_argument('--set', action='append', default=[], metavar='KEY=VALUE',
                    help='override a parameter (negative control / what-if), e.g. --set payload_kg=40')
    args = ap.parse_args(argv)

    p = params_dict()
    for item in args.set:
        k, v = item.split('=', 1)
        if k not in p:
            print('unknown parameter: %s' % k)
            return 2
        p[k] = float(v)
    # переписать значения в PARAMS для листа
    global PARAMS
    PARAMS = [(k, n, (p[k] if k else v), u, t, s) for k, n, v, u, t, s in PARAMS]

    wb, scenarios, rows_index, s_index, l_index, f_index = build_workbook(p)
    wb.save(args.out)
    print('written: %s' % args.out)
    base = compute(p, 250, {})
    print('base 250 kg/ha: ha/day=%.2f direct=%.0f target=%.0f per_ton=%.0f base_ha=%.0f per_kg=%.0f'
          % (base['ha_per_day'], base['direct_cost_per_ha'], base['target_tariff'], base['target_per_ton'],
             base['tariff_base_per_ha'], base['tariff_per_kg']))
    if args.verify:
        return verify(args.out, p, scenarios, rows_index, s_index, l_index, f_index)
    return 0


if __name__ == '__main__':
    sys.exit(main())
