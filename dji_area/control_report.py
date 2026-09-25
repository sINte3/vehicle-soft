# -*- coding: utf-8 -*-
"""dji_area/control_report.py -- отчёт контроля площади DJI: то, что видит человек.

DJI-AREA-PRODUCTIONIZATION-001, DRONE-AREA-CONTROL-V2-MEGA. Главная мысль
отчёта умещается в строку:

    Площадь по данным DJI  -  доказанное завышение  =  площадь, принятая программой

Модуль НИЧЕГО не решает сам. Класс записи определяет `dji_area.accounting`
поверх решения резолвера, решение администратора -- `dji_area.decisions`;
здесь это только переводится в величины отчёта и в слова. Flask и базы тут
нет -- страницу, книгу и тесты кормит один и тот же код, поэтому числа на
экране и в Excel расходиться не могут.

ТРИ СЛОЯ ОДНОЙ ЗАПИСИ

    DJI RAW          -- что записал DJI; не меняется никогда;
    автоматический   -- класс `accounting` и его принятая площадь:
                        NORMAL = RAW; PHANTOM_PROVEN = проверенный прирост
                        счётчика (0 или малый); PHANTOM_STRUCTURAL и REVIEW =
                        RAW, автоматически НЕ вычитается;
    эффективный      -- автоматический, либо, при действующем решении
                        администратора, результат решения. Итоги отчёта
                        считаются по эффективному слою, а автоматический
                        остаётся рядом, отдельной колонкой.

[REASON]: RAW не переписывается никогда, а мостик B между базой и целью не
исключается: проверка показала, что на нём борт действительно льёт. В цепочке
A -> B -> C корректируется только C, и отчёт говорит это словами.

«Площадь, принятая программой» -- НЕ «точная» и НЕ «финальная» площадь, пока
есть записи, ждущие доказательства или решения человека: они сидят внутри по
своему RAW. Поэтому рядом с числом всегда стоит статус полноты.
"""

import json
from datetime import datetime, timedelta

from dji_area import accounting as acc
from dji_area import decisions as dec
from dji_area import resolver as rs

REPORT_VERSION = 'area-control-report-2'
DJI_RECORD_URL = 'https://www.djiag.com/record/%d'
M2_PER_HA = 10000.0
LOCAL_OFFSET = timedelta(hours=5)

VIEW_ALL = 'all'
VIEW_CONFIRMED = 'confirmed'
VIEW_PENDING = 'pending'
VIEW_REVIEW = 'review'
VIEW_DECIDED = 'decided'
VIEWS = (VIEW_ALL, VIEW_CONFIRMED, VIEW_PENDING, VIEW_REVIEW, VIEW_DECIDED)
# Вкладки фильтруют по ЭФФЕКТИВНОМУ состоянию: запись, которую администратор
# подтвердил фантомом, -- уже корректировка, а не «требует проверки».
_VIEW_STATES = {
    VIEW_CONFIRMED: dec.CORRECTION_STATES,
    VIEW_PENDING: dec.PENDING_STATES,
    VIEW_REVIEW: dec.HUMAN_STATES,
}
# В реестр идут записи, по которым есть что сказать человеку. Обычные записи
# (их тысячи) остаются в итогах и в разрезе по дронам.
REGISTER_CLASSES = (acc.PHANTOM_PROVEN, acc.PHANTOM_STRUCTURAL, acc.REVIEW)

# Причины словами. Код причины -- ключ; (ru, uz). Узбекский -- кириллицей.
EXPLAIN_REPEAT = 'REPEAT_OF_PREVIOUS_AUTO_WORK'
EXPLAIN_PARTIAL = 'PARTIAL_NEW_GROWTH'
EXPLAIN_CONTROL = 'CONTROL_CONFIRMED_OVERSTATEMENT'
EXPLAIN_CONTROL_PARTIAL = 'CONTROL_CONFIRMED_PARTIAL'
EXPLAIN_PENDING = 'PENDING_V4'
EXPLAIN_PENDING_NO_V4 = 'PENDING_NO_V4_AT_SOURCE'

EXPLANATIONS = {
    EXPLAIN_REPEAT: (
        'Повтор площади предыдущей Auto-работы; V4 не подтвердил новый '
        'прирост.',
        'Олдинги Auto-иш майдонининг такрори; V4 янги ўсишни тасдиқламади.'),
    EXPLAIN_PARTIAL: (
        'DJI площадь завышена; V4 подтвердил только %s га нового прироста.',
        'DJI майдони ошириб кўрсатилган; V4 фақат %s га янги ўсишни '
        'тасдиқлади.'),
    EXPLAIN_CONTROL: (
        'Завышение подтверждено V4 контрольной проверкой; нового прироста '
        'нет.',
        'Ошириб кўрсатиш V4 назорат текшируви билан тасдиқланди; янги ўсиш '
        'йўқ.'),
    EXPLAIN_CONTROL_PARTIAL: (
        'Завышение подтверждено V4 контрольной проверкой; подтверждено '
        'только %s га нового прироста.',
        'Ошириб кўрсатиш V4 назорат текшируви билан тасдиқланди; фақат %s га '
        'янги ўсиш тасдиқланди.'),
    EXPLAIN_PENDING: (
        'Сильный кандидат на повтор площади; V4 ещё не получен. Пока учтён '
        'по DJI RAW.',
        'Майдон такрорига кучли номзод; V4 ҳали олинмаган. Ҳозирча DJI RAW '
        'бўйича ҳисобга олинган.'),
    EXPLAIN_PENDING_NO_V4: (
        'Сильный кандидат на повтор площади, но DJI не хранит V4 этой '
        'записи. Учтён по DJI RAW; нужна проверка человеком.',
        'Майдон такрорига кучли номзод, аммо DJI бу ёзувнинг V4 сини '
        'сақламайди. DJI RAW бўйича ҳисобга олинган; инсон текшируви керак.'),
    acc.R_APPLICATION_WITH_FLAT_COUNTER: (
        'Счётчик площади не вырос, но распыление наблюдалось; безопасного '
        'автоматического решения нет.',
        'Майдон ҳисоблагичи ўсмади, аммо пуркаш кузатилди; хавфсиз автоматик '
        'қарор йўқ.'),
    acc.R_OVERSTATEMENT_NOT_CERTIFIED: (
        'Признаки завышения есть, но окно V4 не позволяет подтвердить '
        'поправку.',
        'Ошириб кўрсатиш белгилари бор, аммо V4 ойнаси тузатишни тасдиқлашга '
        'имкон бермайди.'),
    acc.R_INTERVAL_OVERLAP: (
        'Интервал записи пересекается с другой записью того же борта.',
        'Ёзув оралиғи шу бортнинг бошқа ёзуви билан кесишади.'),
    acc.R_COUNTER_RELATIONSHIP: (
        'Счётчик V4 не совпал ни с нулём, ни с площадью DJI.',
        'V4 ҳисоблагичи на нолга, на DJI майдонига мос келмади.'),
    acc.R_COUNTER_NONMONOTONE: (
        'Счётчик V4 сбрасывался внутри записи.',
        'V4 ҳисоблагичи ёзув ичида қайта тикланган.'),
    acc.R_BASELINE_UNKNOWN: (
        'Начало счётчика V4 не измерено.',
        'V4 ҳисоблагичининг бошланиши ўлчанмаган.'),
    acc.R_UNKNOWN_SUSPECT: (
        'Данных для автоматического решения недостаточно.',
        'Автоматик қарор учун маълумот етарли эмас.'),
    acc.R_APPLICATION_WITHOUT_AREA: (
        'Распыление наблюдалось, а площадь не измерена.',
        'Пуркаш кузатилди, майдон эса ўлчанмаган.'),
    acc.R_UNMAPPED_STATUS: (
        'Неизвестный статус расчёта; нужна проверка.',
        'Номаълум ҳисоб ҳолати; текшириш керак.'),
}
# Обычная запись: причина одна на всех, а словами -- коротко.
NORMAL_EXPLANATION = ('Принято по данным DJI; признаков завышения нет.',
                      'DJI маълумотлари бўйича қабул қилинган; ошириб '
                      'кўрсатиш белгилари йўқ.')

CLASS_LABELS = {
    acc.NORMAL: ('Принято по DJI', 'DJI бўйича қабул қилинган'),
    acc.PHANTOM_PROVEN: ('Подтверждённая корректировка',
                         'Тасдиқланган тузатиш'),
    acc.PHANTOM_STRUCTURAL: ('Ожидает V4', 'V4 кутилмоқда'),
    acc.REVIEW: ('Требует проверки', 'Текшириш талаб қилинади'),
}
# Цвет -- семантика дизайн-системы: подтверждённая корректировка нейтральна,
# ожидание -- предупреждение, проверка -- внимание. Красным всё подряд не
# красится: доказанная корректировка -- это порядок, а не тревога.
CLASS_BADGES = {
    acc.NORMAL: '',
    acc.PHANTOM_PROVEN: 'vs-badge-info',
    acc.PHANTOM_STRUCTURAL: 'vs-badge-warning',
    acc.REVIEW: 'vs-badge-danger',
}

EVIDENCE_V4_GOOD = 'V4_VALIDATED'
EVIDENCE_V4_LIMITED = 'V4_LIMITED'
EVIDENCE_V4_MISSING = 'V4_MISSING'
EVIDENCE_NO_V4_AT_SOURCE = 'NO_V4_AT_SOURCE'
EVIDENCE_LABELS = {
    EVIDENCE_V4_GOOD: ('V4 получен, счётчик проверен',
                       'V4 олинган, ҳисоблагич текширилган'),
    EVIDENCE_V4_LIMITED: ('V4 получен, окно неполное',
                          'V4 олинган, ойна тўлиқ эмас'),
    EVIDENCE_V4_MISSING: ('V4 не получен', 'V4 олинмаган'),
    EVIDENCE_NO_V4_AT_SOURCE: ('DJI не хранит V4', 'DJI V4 ни сақламайди'),
}

BRIDGE_NOTE = ('Промежуточный ручной участок; не исключается автоматически.',
               'Оралиқ қўлда бошқарилган қисм; автоматик чиқарилмайди.')

TOOLTIPS = {
    'raw': ('Исходная площадь, полученная от DJI. Не изменяется программой.',
            'DJI дан олинган дастлабки майдон. Дастур уни ўзгартирмайди.'),
    'excluded': ('Площадь, для которой дополнительное доказательство '
                 'или решение администратора подтвердило завышение.',
                 'Қўшимча далил ёки администратор қарори ошириб '
                 'кўрсатилганини тасдиқлаган майдон.'),
    'after': ('DJI RAW минус только доказанные корректировки. Нерешённые '
              'записи пока остаются внутри по RAW.',
              'DJI RAW дан фақат исботланган тузатишлар айирилган. Ҳал '
              'қилинмаган ёзувлар ҳозирча RAW бўйича ичида қолади.'),
    'pending': ('Сильные кандидаты, которые программа пока не корректирует: '
                'ждут доказательства V4.',
                'Дастур ҳозирча тузатмаётган кучли номзодлар: V4 далилини '
                'кутмоқда.'),
    'review': ('Есть доказательство, не позволяющее безопасно принять '
               'автоматическое решение; нужно решение человека.',
               'Автоматик қарорни хавфсиз қабул қилишга имкон бермайдиган '
               'далил бор; инсон қарори керак.'),
}

STATUS_COMPLETE = ('Все записи периода разрешены.',
                   'Даврнинг барча ёзувлари ҳал қилинган.')
STATUS_OPEN = ('Есть нерешённые записи; они пока учтены по DJI RAW.',
               'Ҳал қилинмаган ёзувлар бор; улар ҳозирча DJI RAW бўйича '
               'ҳисобга олинган.')
# Одной фразой: входит ли спорная площадь в принятое.
OPEN_INSIDE = ('Спорная и ожидающая площадь (%s га) сейчас ВХОДИТ в принятую '
               'по DJI RAW, пока администратор не решил иначе.',
               'Баҳсли ва кутилаётган майдон (%s га) ҳозир DJI RAW бўйича '
               'қабул қилинганга КИРАДИ, администратор бошқача ҳал '
               'қилмагунча.')
OPEN_NONE = ('Спорной площади нет: принятое не содержит нерешённых записей.',
             'Баҳсли майдон йўқ: қабул қилинган ҳал қилинмаган ёзувларни '
             'ўз ичига олмайди.')
REST_LABEL = ('Остальные вылеты дня, не показанные поимённо',
              'Куннинг номма-ном кўрсатилмаган қолган парвозлари')


def pick(pair, lang):
    return pair[0] if lang == 'ru' else pair[1]


def xlsx_safe(value):
    """Строка для ячейки книги, которая не станет формулой.

    [REASON]: книга уходит в бухгалтерию и открывается в Excel. Свободный
    текст -- причина решения администратора, имя из профиля, подпись машины,
    строка периода из адреса -- начинающийся с = + - @ Excel прочитал бы как
    формулу (#NAME? в лучшем случае, HYPERLINK наружу -- в худшем). Тот же
    приём, что `drones._drone_xlsx_safe`; модуль Flask не импортирует,
    поэтому своя копия.
    """
    if isinstance(value, str) and value.lstrip()[:1] in ('=', '+', '-', '@'):
        return "'" + value
    return value


def ha(value_m2):
    return None if value_m2 is None else float(value_m2) / M2_PER_HA


def dji_url(flight_id):
    return DJI_RECORD_URL % int(flight_id)


def _as_datetime(value):
    """datetime из datetime либо из строки sqlite (19 или 26 символов).

    [REASON]: строки расчёта приходят двумя путями -- через ORM (datetime)
    и сырым sqlite3 в инструментах (строка, записанная stdlib-писателем без
    дробной части или SQLAlchemy с ней). Неразборчивое -- None, а не
    исключение: время -- подсказка для человека, не величина отчёта.
    """
    if value is None or isinstance(value, datetime):
        return value
    text = str(value).replace('T', ' ')
    for fmt in ('%Y-%m-%d %H:%M:%S.%f', '%Y-%m-%d %H:%M:%S',
                '%Y-%m-%d %H:%M'):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def local_time(utc_dt):
    """UTC -> UTC+5; None (и неразборчивое) остаётся None."""
    utc_dt = _as_datetime(utc_dt)
    return None if utc_dt is None else utc_dt + LOCAL_OFFSET


def _flags(row):
    flags = row.get('anomaly_flags')
    if flags is None:
        try:
            flags = json.loads(row.get('anomaly_flags_json') or '[]')
        except ValueError:
            flags = []
    return [str(flag) for flag in flags if flag]


def _bridges(row):
    if row.get('structural_candidate') not in (True, 1):
        return []
    try:
        ids = json.loads(row.get('bridge_flight_ids_json') or '[]')
    except ValueError:
        return []
    return [int(i) for i in ids if isinstance(i, (int, float))
            or str(i).isdigit()]


def evidence_state(row):
    if 'NO_V4_AT_SOURCE' in _flags(row):
        return EVIDENCE_NO_V4_AT_SOURCE
    if not row.get('v4_revision_id') and not row.get('v4_summary_id'):
        return EVIDENCE_V4_MISSING
    if row.get('aggregation_eligibility') == rs.AGG_CERTIFIED:
        return EVIDENCE_V4_GOOD
    return EVIDENCE_V4_LIMITED


def explanation(decision, row, lang):
    """(код причины, текст) -- человеку, а не enum резолвера."""
    cls = decision['accounting_class']
    if cls == acc.PHANTOM_PROVEN:
        accepted = decision['accounted_area_m2'] or 0.0
        structural = decision['reason'] == acc.R_RETAINED_VALIDATED
        if accepted > 0:
            code = EXPLAIN_PARTIAL if structural else EXPLAIN_CONTROL_PARTIAL
            return code, pick(EXPLANATIONS[code], lang) % (
                '%.4f' % ha(accepted))
        code = EXPLAIN_REPEAT if structural else EXPLAIN_CONTROL
        return code, pick(EXPLANATIONS[code], lang)
    if cls == acc.PHANTOM_STRUCTURAL:
        code = (EXPLAIN_PENDING_NO_V4
                if evidence_state(row) == EVIDENCE_NO_V4_AT_SOURCE
                else EXPLAIN_PENDING)
        return code, pick(EXPLANATIONS[code], lang)
    if cls == acc.NORMAL:
        return decision['reason'], pick(NORMAL_EXPLANATION, lang)
    code = decision['reason']
    pair = EXPLANATIONS.get(code) or EXPLANATIONS[acc.R_UNMAPPED_STATUS]
    return code, pick(pair, lang)


def _point(flight_id, times, day=None):
    """Звено цепочки: id, ссылка DJI и местное время начала (UTC+5).

    ``label_s`` -- то, что стоит в строке дерева: время, а если звено
    началось в другой местный день, чем строка (цепочка пересекла
    полночь), -- дата и время. Иначе «A 23:55» под днём 05.06 читалось бы
    как 05.06 23:55, после C, и искалось бы в DJI не в тот день.
    """
    start = local_time((times or {}).get(int(flight_id)))
    time_s = start.strftime('%H:%M') if start else ''
    other_day = bool(start and hasattr(day, 'year')
                     and start.date() != day)
    return {'flight_id': int(flight_id), 'url': dji_url(flight_id),
            'start_local': start,
            'time_s': time_s,
            'label_s': start.strftime('%d.%m %H:%M') if other_day else time_s,
            'datetime_s': start.strftime('%d.%m.%Y %H:%M') if start else ''}


def _decision_view(decision, row, lang):
    if decision is None:
        return None
    kind = decision.get('decision_type')
    performed_at = decision.get('performed_at')
    at_local = local_time(performed_at) if performed_at else None
    stale = dec.decision_is_stale(decision, row)
    return {
        'id': decision.get('id'),
        'decision_type': kind,
        'label': pick(dec.DECISION_SHORT.get(kind, (kind, kind)), lang),
        'comment': decision.get('comment'),
        'performed_by': decision.get('performed_by_name'),
        'performed_at_local': at_local,
        'performed_at_s': at_local.strftime('%d.%m.%Y %H:%M')
        if at_local else '',
        'is_override': bool(decision.get('is_override')),
        'stale': stale,
        'stale_note': pick(dec.STALE_NOTE, lang) if stale else '',
        'chain_seq': decision.get('chain_seq'),
    }


def record_view(row, lang='ru', decision=None, times=None):
    """Одна строка отчёта. RAW не меняется; исключается только доказанное.

    ``decision`` -- действующее решение администратора (словарь) либо None;
    ``times`` -- {flight_id: start_at_utc} для звеньев цепочки A и B.
    """
    auto = acc.classify(row)
    cls = auto['accounting_class']
    raw = auto['raw_area_m2']
    if cls == acc.PHANTOM_PROVEN:
        auto_accepted = auto['accounted_area_m2']
        auto_excluded = auto['confirmed_overstatement_m2']
    else:
        auto_accepted = raw
        auto_excluded = 0.0 if raw is not None else None
    decision_view = _decision_view(decision, row, lang)
    stale = bool(decision_view and decision_view['stale'])
    state, accepted, excluded, applied = dec.effective(
        cls, raw, auto_accepted, auto_excluded, decision, stale=stale)
    lapsed = decision is not None and not applied
    if decision_view is not None:
        decision_view['applied'] = applied
        decision_view['lapsed_note'] = pick(dec.LAPSED_NOTE, lang) \
            if lapsed else ''
    code, text = explanation(auto, row, lang)
    ev_state = evidence_state(row)
    flight_id = int(row['flight_id'])
    chain = bool(row.get('structural_candidate') in (True, 1)
                 and row.get('candidate_base_flight_id'))
    base_id = row.get('candidate_base_flight_id') if chain else None
    bridges = _bridges(row) if chain else []
    lookup = dict(times or {})
    start_utc = row.get('start_at_utc')
    if start_utc is not None:
        lookup.setdefault(flight_id, start_utc)
    start_local = local_time(start_utc)
    report_day = row.get('report_start_date')
    return {
        'flight_id': flight_id,
        'report_start_date': report_day,
        'start_local': start_local,
        'time_s': start_local.strftime('%H:%M') if start_local else '',
        'datetime_s': start_local.strftime('%d.%m.%Y %H:%M')
        if start_local else (report_day.strftime('%d.%m.%Y')
                             if hasattr(report_day, 'strftime') else ''),
        'machine_key': row.get('machine_key'),
        'machine_label': row.get('machine_label'),
        # Автоматический слой: класс и его величины не зависят от решения.
        'accounting_class': cls,
        'class_label': pick(CLASS_LABELS[cls], lang),
        'class_badge': CLASS_BADGES[cls],
        'auto_accepted_m2': auto_accepted,
        'auto_excluded_m2': auto_excluded,
        # Эффективный слой: то, что входит в итоги отчёта.
        'state': state,
        'state_label': pick(dec.STATE_LABELS[state], lang),
        'state_badge': dec.STATE_BADGES[state],
        'is_open': state in dec.OPEN_STATES,
        'needs_human': state in dec.HUMAN_STATES,
        'raw_m2': raw,
        'accepted_m2': accepted,
        'excluded_m2': excluded,
        'decision': decision_view,
        'decidable': cls in dec.DECIDABLE_CLASSES,
        'is_override_class': dec.is_override(cls),
        'in_register': cls in REGISTER_CLASSES or decision is not None,
        'controller_delta_m2': row.get('controller_delta_area_m2'),
        'reason_code': code,
        'reason_text': text,
        'evidence_state': ev_state,
        'evidence_label': pick(EVIDENCE_LABELS[ev_state], lang),
        'base_flight_id': base_id,
        'bridge_flight_ids': bridges,
        'chain': ({'a': _point(base_id, lookup, report_day),
                   'b': [_point(b, lookup, report_day) for b in bridges],
                   'c': _point(flight_id, lookup, report_day)}
                  if chain else None),
        'dji_url': dji_url(flight_id),
        'structural_rule_version': row.get('structural_rule_version'),
        'area_algorithm_version': row.get('area_algorithm_version'),
    }


# ─── Итоги: раздельные величины эффективного слоя ────────────────────────────

def empty_totals():
    return {
        'records': 0,
        'raw_m2': 0.0,
        'raw_missing_records': 0,
        'excluded_m2': 0.0,
        'excluded_records': 0,
        'auto_excluded_m2': 0.0,
        'auto_excluded_records': 0,
        'pending_m2': 0.0,
        'pending_records': 0,
        'review_m2': 0.0,
        'review_records': 0,
        'decided_records': 0,
        'override_records': 0,
        'stale_decision_records': 0,
        'state_records': {state: 0 for state in dec.STATES},
        'state_raw_m2': {state: 0.0 for state in dec.STATES},
    }


def add_view(totals, item):
    """Добавить строку отчёта в итог. Считается по ЭФФЕКТИВНОМУ слою."""
    totals['records'] += 1
    raw = item['raw_m2']
    exposure = raw if raw is not None else 0.0
    if raw is None:
        totals['raw_missing_records'] += 1
    else:
        totals['raw_m2'] += raw
    state = item['state']
    totals['state_records'][state] += 1
    totals['state_raw_m2'][state] += exposure
    if state in dec.CORRECTION_STATES:
        totals['excluded_records'] += 1
        totals['excluded_m2'] += item['excluded_m2'] or 0.0
    if item['accounting_class'] == acc.PHANTOM_PROVEN:
        totals['auto_excluded_records'] += 1
        totals['auto_excluded_m2'] += item['auto_excluded_m2'] or 0.0
    if state in dec.PENDING_STATES:
        totals['pending_records'] += 1
        totals['pending_m2'] += exposure
    if state in dec.HUMAN_STATES:
        totals['review_records'] += 1
        totals['review_m2'] += exposure
    decision = item.get('decision')
    if decision is not None and decision.get('applied'):
        totals['decided_records'] += 1
        if decision.get('is_override'):
            totals['override_records'] += 1
    if decision is not None and decision.get('stale'):
        # Решение принято против прежнего расчёта: применено оно или
        # потеряло силу -- администратору стоит посмотреть ещё раз.
        totals['stale_decision_records'] += 1
    return totals


def finalize_totals(totals, lang='ru'):
    totals['after_m2'] = totals['raw_m2'] - totals['excluded_m2']
    totals['accepted_m2'] = totals['after_m2']
    totals['decision_delta_m2'] = (totals['excluded_m2']
                                   - totals['auto_excluded_m2'])
    totals['open_m2'] = totals['pending_m2'] + totals['review_m2']
    totals['open_records'] = (totals['pending_records']
                              + totals['review_records'])
    totals['complete'] = totals['open_records'] == 0
    state_raw = sum(totals['state_raw_m2'].values())
    totals['partition_holds'] = (
        sum(totals['state_records'].values()) == totals['records']
        and abs(state_raw - totals['raw_m2']) <= 1e-6 * max(
            1.0, totals['raw_m2']))
    totals['status_text'] = pick(STATUS_COMPLETE if totals['complete']
                                 else STATUS_OPEN, lang)
    if totals['open_records']:
        totals['open_text'] = pick(OPEN_INSIDE, lang) % (
            '%.2f' % ha(totals['open_m2']))
    else:
        totals['open_text'] = pick(OPEN_NONE, lang)
    totals['open_inside_accepted'] = totals['open_records'] > 0
    return totals


def _matches(item, view):
    if view == VIEW_ALL:
        return True
    if view == VIEW_DECIDED:
        return bool(item['decision'] and item['decision'].get('applied'))
    return item['state'] in _VIEW_STATES.get(view, ())


def _rest(items):
    rest = empty_rest()
    for item in items:
        add_to_rest(rest, item)
    return rest if rest['records'] else None


def empty_rest():
    return {'records': 0, 'raw_m2': 0.0, 'auto_accepted_m2': 0.0,
            'accepted_m2': 0.0, 'excluded_m2': 0.0}


def add_to_rest(rest, item):
    """Свернуть вылет в строку «остальные» своего дня."""
    rest['records'] += 1
    rest['raw_m2'] += item['raw_m2'] or 0.0
    rest['auto_accepted_m2'] += item['auto_accepted_m2'] or 0.0
    rest['accepted_m2'] += item['accepted_m2'] or 0.0
    rest['excluded_m2'] += item['excluded_m2'] or 0.0
    return rest


def build(rows, lang='ru', view=VIEW_ALL, decisions=None, times=None,
          list_all=False):
    """Итог периода, разрез по дронам, дерево и реестр -- из одной выборки.

    ``decisions`` -- {flight_id: действующее решение}; ``times`` --
    {flight_id: start_at_utc} звеньев A/B, лежащих вне выборки;
    ``list_all`` -- показывать в дереве каждый вылет, а не только записи
    реестра (обычные вылеты тогда не сворачиваются в строку «остальные»).
    """
    if view not in VIEWS:
        view = VIEW_ALL
    decisions = decisions or {}
    # Время звеньев A/B: сначала из самой выборки (звено часто в ней же),
    # затем переданное снаружи для звеньев вне периода.
    known_times = {int(r['flight_id']): r.get('start_at_utc') for r in rows
                   if r.get('start_at_utc') is not None}
    known_times.update(times or {})
    items = [record_view(row, lang, decisions.get(int(row['flight_id'])),
                         known_times) for row in rows]

    total = empty_totals()
    by_drone = {}
    labels = {}
    for item in items:
        add_view(total, item)
        key = item['machine_key']
        labels.setdefault(key, item['machine_label'])
        drone = by_drone.get(key)
        if drone is None:
            drone = by_drone[key] = {'totals': empty_totals(), 'days': {}}
        add_view(drone['totals'], item)
        day = drone['days'].get(item['report_start_date'])
        if day is None:
            day = drone['days'][item['report_start_date']] = {
                'totals': empty_totals(), 'items': []}
        add_view(day['totals'], item)
        day['items'].append(item)
    finalize_totals(total, lang)

    drones = []
    tree = []
    for key, entry in by_drone.items():
        totals = finalize_totals(entry['totals'], lang)
        flat = dict(totals)
        flat['machine_key'] = key
        flat['machine_label'] = labels.get(key)
        drones.append(flat)
        days = []
        for day_key in sorted(entry['days'], key=lambda d: str(d or '')):
            day = entry['days'][day_key]
            day_totals = finalize_totals(day['totals'], lang)
            day_items = sorted(day['items'], key=lambda i: (
                str(i['start_local'] or ''), i['flight_id']))
            listed = [i for i in day_items
                      if (list_all or i['in_register']) and _matches(i, view)]
            listed_ids = {i['flight_id'] for i in listed}
            rest = _rest([i for i in day_items
                          if i['flight_id'] not in listed_ids])
            days.append({'date': day_key, 'totals': day_totals,
                         'flights': listed, 'rest': rest})
        tree.append({'machine_key': key, 'machine_label': labels.get(key),
                     'totals': totals, 'days': days})
    # Полезный порядок: где больше всего доказанно исключено -- сверху.
    order = lambda d: (-d['excluded_m2'], str(d['machine_label']))  # noqa
    drones.sort(key=order)
    tree.sort(key=lambda node: order(dict(node['totals'],
                                          machine_label=node['machine_label'])))

    register = [i for i in items if i['in_register'] and _matches(i, view)]
    register.sort(key=lambda r: (str(r['report_start_date'] or ''),
                                 r['flight_id']), reverse=True)
    listed_total = sum(len(day['flights']) for node in tree
                       for day in node['days'])
    return {'report_version': REPORT_VERSION, 'view': view, 'total': total,
            'drones': drones, 'tree': tree, 'register': register,
            'listed_flights': listed_total, 'list_all': bool(list_all),
            'items': items}


# ─── Книга «Отчёт контроля площади DJI» ──────────────────────────────────────

SHEET_SUMMARY = ('Сводка', 'Жамланма')
SHEET_DRONES = ('По_дронам', 'Дронлар_бўйича')
SHEET_CORRECTIONS = ('Корректировки', 'Тузатишлар')
SHEET_REVIEW = ('Требует_проверки', 'Текшириш_керак')
SHEET_DAYS = ('По_дням', 'Кунлар_бўйича')
SHEET_REGISTER = ('Реестр', 'Реестр')
SHEET_HISTORY = ('История_решений', 'Қарорлар_тарихи')

YES_NO = {True: ('Да', 'Ҳа'), False: ('Нет', 'Йўқ')}


def build_workbook(report, lang='ru', period=('', ''), versions=None,
                   history=None, filters_text=None):
    """openpyxl.Workbook из результата `build(..., view=VIEW_ALL)`.

    ``history`` -- все строки решений по вылетам выборки (включая
    замещённые) для листа истории; ``filters_text`` -- [(подпись, значение)]
    прочих фильтров экрана, чтобы книга говорила, ЧТО в неё отобрано.

    [REASON]: NULL -- ПУСТАЯ ячейка, никогда не 0: книга -- та копия, которую
    суммируют, и ноль в колонке читался бы как измеренный ноль.

    [REASON]: первые четыре листа -- те же, что до V2, с теми же заголовками
    колонок: их читают бухгалтерия и сверочные скрипты. Новое только
    добавлено -- колонками правее и листами после.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    def tr(ru, uz):
        return ru if lang == 'ru' else uz

    versions = versions or {}
    total = report['total']
    link_font = Font(color='0563C1', underline='single')
    bold = Font(bold=True)

    def header(sheet, titles):
        sheet.append(titles)
        for cell in sheet[1]:
            cell.font = bold
        sheet.freeze_panes = 'A2'

    def link(sheet, column, url):
        if not url:
            return
        cell = sheet.cell(row=sheet.max_row, column=column)
        cell.hyperlink = url
        cell.font = link_font

    book = Workbook()
    sheet = book.active
    sheet.title = pick(SHEET_SUMMARY, lang)
    header(sheet, [tr('Показатель', 'Кўрсаткич'), tr('Значение', 'Қиймат'),
                   tr('Записей', 'Ёзувлар'), tr('Пояснение', 'Изоҳ')])
    sheet.append([tr('Период: с', 'Давр: бошланиши'),
                  xlsx_safe(period[0]) or None])
    sheet.append([tr('Период: по', 'Давр: тугаши'),
                  xlsx_safe(period[1]) or None])
    for label, value in (filters_text or ()):
        sheet.append([label, xlsx_safe(value)])
    sheet.append([tr('DJI RAW, га', 'DJI RAW, га'), ha(total['raw_m2']),
                  total['records'], pick(TOOLTIPS['raw'], lang)])
    sheet.append([tr('Подтверждённо исключено, га',
                     'Тасдиқланган ҳолда чиқарилган, га'),
                  ha(total['excluded_m2']), total['excluded_records'],
                  pick(TOOLTIPS['excluded'], lang)])
    sheet.append([tr('Площадь после подтверждённых корректировок, га',
                     'Тасдиқланган тузатишлардан кейинги майдон, га'),
                  ha(total['after_m2']), None, pick(TOOLTIPS['after'], lang)])
    sheet.append([tr('Ожидает доказательства / V4, га',
                     'Далил / V4 кутилмоқда, га'),
                  ha(total['pending_m2']), total['pending_records'],
                  pick(TOOLTIPS['pending'], lang)])
    sheet.append([tr('Требует проверки, га', 'Текшириш талаб қилинади, га'),
                  ha(total['review_m2']), total['review_records'],
                  pick(TOOLTIPS['review'], lang)])
    sheet.append([tr('Записей без RAW', 'RAW йўқ ёзувлар'),
                  total['raw_missing_records']])
    sheet.append([tr('Статус периода', 'Давр ҳолати'), total['status_text']])
    sheet.append([tr('Автоматически исключено, га',
                     'Автоматик чиқарилган, га'),
                  ha(total['auto_excluded_m2']),
                  total['auto_excluded_records']])
    sheet.append([tr('Изменение решениями администратора, га',
                     'Администратор қарорлари билан ўзгариш, га'),
                  ha(total['decision_delta_m2']), total['decided_records']])
    sheet.append([tr('Спорная площадь входит в принятое',
                     'Баҳсли майдон қабул қилинганга киради'),
                  pick(YES_NO[total['open_inside_accepted']], lang), None,
                  total['open_text']])
    for title, key in ((tr('Алгоритм площади', 'Майдон алгоритми'),
                        'area_algorithm'),
                       (tr('Структурное правило', 'Структура қоидаси'),
                        'structural_rule'),
                       (tr('Учётные классы', 'Ҳисоб синфлари'),
                        'accounting_classes'),
                       (tr('Решения администратора', 'Администратор '
                           'қарорлари'), 'decisions'),
                       (tr('Версия отчёта', 'Ҳисобот версияси'), 'report')):
        sheet.append([title, versions.get(key)])

    sheet = book.create_sheet(pick(SHEET_DRONES, lang))
    header(sheet, [tr('Дрон', 'Дрон'), tr('Вылетов', 'Парвозлар'),
                   tr('DJI RAW, га', 'DJI RAW, га'),
                   tr('Подтверждённо исключено, га',
                      'Тасдиқланган ҳолда чиқарилган, га'),
                   tr('После подтверждённых корректировок, га',
                      'Тасдиқланган тузатишлардан кейин, га'),
                   tr('Ожидает V4, га', 'V4 кутилмоқда, га'),
                   tr('Требует проверки, га', 'Текшириш керак, га'),
                   tr('Решений администратора', 'Администратор қарорлари')])
    for drone in report['drones']:
        sheet.append([xlsx_safe(drone['machine_label']), drone['records'],
                      ha(drone['raw_m2']), ha(drone['excluded_m2']),
                      ha(drone['after_m2']), ha(drone['pending_m2']),
                      ha(drone['review_m2']), drone['decided_records']])

    register_titles = [
        tr('Дата', 'Сана'), tr('Дрон', 'Дрон'), 'C Flight ID',
        'A Flight ID', tr('Bridge IDs (не корректируются)',
                          'Bridge IDs (тузатилмайди)'),
        tr('DJI RAW, га', 'DJI RAW, га'),
        tr('Принято, га', 'Қабул қилинган, га'),
        tr('Исключено, га', 'Чиқарилган, га'), tr('Класс', 'Синф'),
        tr('Причина', 'Сабаб'), tr('Доказательство', 'Далил'),
        tr('Прирост счётчика V4, м²', 'V4 ҳисоблагич ўсиши, м²'),
        tr('Ссылка DJI', 'DJI ҳаволаси'),
        # V2: правее прежних колонок.
        tr('C время (UTC+5)', 'C вақти (UTC+5)'),
        tr('A время (UTC+5)', 'A вақти (UTC+5)'),
        tr('A ссылка', 'A ҳаволаси'),
        tr('B время (UTC+5)', 'B вақти (UTC+5)'),
        tr('B ссылка', 'B ҳаволаси'),
        tr('Автоматически принято, га', 'Автоматик қабул қилинган, га'),
        tr('Автоматически исключено, га', 'Автоматик чиқарилган, га'),
        tr('Решение администратора', 'Администратор қарори'),
        tr('Кто решил', 'Ким ҳал қилди'),
        tr('Когда решил (UTC+5)', 'Қачон ҳал қилди (UTC+5)'),
        tr('Комментарий решения', 'Қарор изоҳи'),
        tr('Итоговый статус', 'Якуний ҳолат'),
        tr('Решение действует', 'Қарор амалда'),
        tr('Пометка к решению', 'Қарорга изоҳ'),
    ]
    c_link_col = register_titles.index(tr('Ссылка DJI', 'DJI ҳаволаси')) + 1
    a_link_col = register_titles.index(tr('A ссылка', 'A ҳаволаси')) + 1
    b_link_col = register_titles.index(tr('B ссылка', 'B ҳаволаси')) + 1

    def register_sheet(title, rows):
        sheet = book.create_sheet(title)
        header(sheet, register_titles)
        for item in rows:
            day = item['report_start_date']
            chain = item['chain'] or {}
            a = chain.get('a') or {}
            bees = chain.get('b') or []
            decision = item['decision'] or {}
            sheet.append([
                day.isoformat() if hasattr(day, 'isoformat') else day,
                xlsx_safe(item['machine_label']), item['flight_id'],
                item['base_flight_id'],
                '; '.join(str(b) for b in item['bridge_flight_ids']) or None,
                ha(item['raw_m2']), ha(item['accepted_m2']),
                ha(item['excluded_m2']), item['class_label'],
                item['reason_text'], item['evidence_label'],
                item['controller_delta_m2'], item['dji_url'],
                item['datetime_s'] or None,
                a.get('datetime_s') or None,
                a.get('url'),
                '; '.join(b['datetime_s'] for b in bees
                          if b['datetime_s']) or None,
                '; '.join(b['url'] for b in bees) or None,
                ha(item['auto_accepted_m2']), ha(item['auto_excluded_m2']),
                decision.get('label'), xlsx_safe(decision.get('performed_by')),
                decision.get('performed_at_s') or None,
                xlsx_safe(decision.get('comment')), item['state_label'],
                pick(YES_NO[bool(decision.get('applied'))], lang)
                if decision else None,
                decision.get('lapsed_note') or decision.get('stale_note')
                or None,
            ])
            link(sheet, c_link_col, item['dji_url'])
            link(sheet, a_link_col, a.get('url'))
            link(sheet, b_link_col, bees[0]['url'] if bees else None)

    everything = [i for i in report.get('items', report['register'])
                  if i['in_register']]
    everything.sort(key=lambda r: (str(r['report_start_date'] or ''),
                                   r['flight_id']), reverse=True)
    register_sheet(pick(SHEET_CORRECTIONS, lang),
                   [r for r in everything
                    if r['state'] in dec.CORRECTION_STATES])
    register_sheet(pick(SHEET_REVIEW, lang),
                   [r for r in everything if r['state'] in dec.OPEN_STATES])

    sheet = book.create_sheet(pick(SHEET_DAYS, lang))
    header(sheet, [tr('Дрон', 'Дрон'), tr('Дата', 'Сана'),
                   tr('Вылетов', 'Парвозлар'),
                   tr('DJI RAW, га', 'DJI RAW, га'),
                   tr('Подтверждённо исключено, га',
                      'Тасдиқланган ҳолда чиқарилган, га'),
                   tr('Принято, га', 'Қабул қилинган, га'),
                   tr('Ожидает V4, га', 'V4 кутилмоқда, га'),
                   tr('Требует проверки, га', 'Текшириш керак, га')])
    for node in report.get('tree', ()):
        for day in node['days']:
            totals = day['totals']
            day_value = day['date']
            sheet.append([xlsx_safe(node['machine_label']),
                          day_value.isoformat()
                          if hasattr(day_value, 'isoformat') else day_value,
                          totals['records'], ha(totals['raw_m2']),
                          ha(totals['excluded_m2']), ha(totals['after_m2']),
                          ha(totals['pending_m2']), ha(totals['review_m2'])])

    register_sheet(pick(SHEET_REGISTER, lang), everything)

    sheet = book.create_sheet(pick(SHEET_HISTORY, lang))
    header(sheet, ['C Flight ID', tr('№ в истории', 'Тарихдаги №'),
                   tr('Действие', 'Амал'), tr('Действует', 'Амалда'),
                   tr('Переопределение доказанного', 'Исботланганни қайта '
                      'белгилаш'),
                   tr('Автомат на момент решения', 'Қарор пайтидаги '
                      'автомат'),
                   tr('DJI RAW на момент решения, га', 'Қарор пайтидаги '
                      'DJI RAW, га'),
                   tr('Автоматически принято на момент решения, га',
                      'Қарор пайтида автоматик қабул қилинган, га'),
                   tr('Кто', 'Ким'), tr('Когда (UTC+5)', 'Қачон (UTC+5)'),
                   tr('Комментарий', 'Изоҳ'), tr('Ссылка DJI', 'DJI ҳаволаси')])
    # «Действует» -- последняя строка цепочки И решение применено к
    # нынешнему расчёту (отменённое «принять автомат» -- не действует).
    applied_by_flight = {
        i['flight_id']: bool((i['decision'] or {}).get('applied'))
        for i in report.get('items', ())}
    for entry in history or ():
        kind = entry.get('decision_type')
        at_local = local_time(entry.get('performed_at'))
        auto_class = entry.get('auto_class')
        in_force = bool(entry.get('is_current')) and applied_by_flight.get(
            entry.get('flight_id'), True)
        sheet.append([
            entry.get('flight_id'), entry.get('chain_seq'),
            pick(dec.DECISION_SHORT.get(kind, (kind, kind)), lang),
            pick(YES_NO[in_force], lang),
            pick(YES_NO[bool(entry.get('is_override'))], lang),
            pick(CLASS_LABELS[auto_class], lang)
            if auto_class in CLASS_LABELS else auto_class,
            ha(entry.get('raw_area_m2')), ha(entry.get('auto_accepted_m2')),
            xlsx_safe(entry.get('performed_by_name')),
            at_local.strftime('%d.%m.%Y %H:%M') if at_local else None,
            xlsx_safe(entry.get('comment')),
            dji_url(entry['flight_id']) if entry.get('flight_id') else None])
        link(sheet, 12, dji_url(entry['flight_id'])
             if entry.get('flight_id') else None)
    return book
