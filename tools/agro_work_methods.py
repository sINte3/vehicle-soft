# -*- coding: utf-8 -*-
"""agro-work B2 -- разметка метода сверки по видам работ agro-work.

ЗАЧЕМ
Единицу работы agro-work отдаёт сам (B0: 8 кодов, HECTARE, MOTOR_HOUR, ...),
а КАК сверять вид работы с GPS -- решение владельца: гектары, время, рейсы
или не сверяется. Код метод не угадывает (план трека, B2): пока вид работы
не размечен, его заявки остаются без вердикта с причиной «метод не
размечен».

КАК
  1. `--export FILE.xlsx` -- книга: все виды работ, их единица, число заявок
     и колонка «Метод сверки» с выпадающим списком из четырёх методов. На
     втором листе -- что сверка делает с заявкой каждого метода.
  2. Владелец выбирает метод в каждой строке и сохраняет файл.
  3. `--import FILE.xlsx` -- план: сколько поставить, сменить, снять.
     Ничего не пишет.
  4. `--import FILE.xlsx --apply` -- записать одной транзакцией, с журналом.

Пустая ячейка у размеченного вида работы -- снятие метода; вид работы,
которого нет в файле, не трогается. Неизвестное значение ячейки -- отказ
всего файла, а не пропуск строки.

Запуск (PowerShell, по одной команде на строку):

  & "C:\\Program Files\\Python314\\python.exe" tools\\agro_work_methods.py --export agro_work_methods.xlsx

  & "C:\\Program Files\\Python314\\python.exe" tools\\agro_work_methods.py --import agro_work_methods.xlsx

  & "C:\\Program Files\\Python314\\python.exe" tools\\agro_work_methods.py --import agro_work_methods.xlsx --apply

Без ключей -- сводка: сколько видов размечено и сколько заявок ждут
разметки. Сеть не нужна. Вывод в консоль -- ASCII.
"""

import argparse
import os
import sqlite3
import sys
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from agro_work import config, methods, store               # noqa: E402

SHEET = 'Виды работ'
LEGEND = 'Методы'
H_ID = 'id'
H_NAME = 'Вид работы / Иш тури'
H_UNIT = 'Единица / Бирлик'
H_UNIT_LABEL = 'Подпись единицы / Бирлик ёзуви'
H_COUNT = 'Заявок / Буюртмалар'
H_FIRST = 'Первая заявка / Биринчи буюртма'
H_LAST = 'Последняя заявка / Охирги буюртма'
H_METHOD = 'Метод сверки / Солиштириш усули'
HEADERS = (H_ID, H_NAME, H_UNIT, H_UNIT_LABEL, H_COUNT, H_FIRST, H_LAST,
           H_METHOD)
WIDTHS = (38, 60, 14, 18, 14, 16, 16, 30)


class Refused(ValueError):
    """Файл разметки нельзя принять. Текст -- ASCII."""


def export(con, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = methods.work_types_for_markup(con)
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for row in rows:
        label = methods.LABELS[row['method']][0] if row['method'] else None
        sheet.append((row['id'], row['name'], row['unit'], row['unit_display'],
                      row['applications'], row['first_day'], row['last_day'],
                      label))
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'B2'
    # [REASON]: выпадающий список, а не свободный текст: «гектар», «га» и
    # «Гектары » -- три опечатки, которые иначе пришлось бы угадывать при
    # загрузке. Загрузка всё равно проверяет каждую ячейку сама.
    choices = ','.join(methods.LABELS[m][0] for m in methods.METHODS)
    validation = DataValidation(type='list', formula1='"%s"' % choices,
                                allow_blank=True)
    validation.error = 'Выберите метод из списка / Рўйхатдан усулни танланг'
    validation.errorTitle = 'Метод сверки'
    sheet.add_data_validation(validation)
    column = get_column_letter(HEADERS.index(H_METHOD) + 1)
    validation.add('%s2:%s%d' % (column, column, max(len(rows) + 1, 2)))

    legend = book.create_sheet(LEGEND)
    legend.append(('Метод', 'Усул', 'Что делает сверка', 'Солиштириш нима '
                                                         'қилади'))
    for cell in legend[1]:
        cell.font = Font(bold=True)
    for slug in methods.METHODS:
        ru, uz = methods.LABELS[slug]
        meaning_ru, meaning_uz = methods.MEANING[slug]
        legend.append((ru, uz, meaning_ru, meaning_uz))
    for index, width in enumerate((16, 18, 70, 70), start=1):
        legend.column_dimensions[get_column_letter(index)].width = width
        for cell in legend[get_column_letter(index)]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    book.save(path)
    return len(rows)


def read_markup(path):
    """{id вида работы: слаг или None} из книги владельца."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[SHEET] if SHEET in book.sheetnames else book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [str(v).strip() if v is not None else '' for v in next(rows, ())]
        for needed in (H_ID, H_METHOD):
            if needed not in header:
                raise Refused('column "%s" is missing in the first row - the '
                              'headers of the exported file must stay as they '
                              'are' % config.ascii_only(needed))
        id_col, method_col = header.index(H_ID), header.index(H_METHOD)
        wanted = {}
        for number, values in enumerate(rows, start=2):
            values = list(values) + [None] * len(header)
            work_type_id = values[id_col]
            if work_type_id is None or not str(work_type_id).strip():
                continue
            work_type_id = str(work_type_id).strip()
            if work_type_id in wanted:
                raise Refused('row %d: work type %s is given twice'
                              % (number, work_type_id))
            try:
                wanted[work_type_id] = methods.parse_method(values[method_col])
            except methods.BadMethod as exc:
                raise Refused('row %d: "%s" is not a method (allowed: %s)'
                              % (number, config.ascii_only(exc.args[0]),
                                 ', '.join(methods.METHODS))) from None
        return wanted
    finally:
        book.close()


def print_status(con):
    rows = methods.work_types_for_markup(con)
    by_method = Counter(row['method'] or 'not marked' for row in rows)
    waiting = sum(row['applications'] for row in rows if not row['method'])
    print('work types: %d' % len(rows))
    for key in methods.METHODS + ('not marked',):
        print('  %-15s %d' % (key, by_method.get(key, 0)))
    print('applications of unmarked work types: %d' % waiting)
    units = Counter(row['unit'] or '-' for row in rows)
    print('units: %s' % ' | '.join('%s %d' % (k, v)
                                   for k, v in sorted(units.items())))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--db', default=config.DB_PATH)
    group = parser.add_mutually_exclusive_group()
    group.add_argument('--export', metavar='FILE.xlsx')
    group.add_argument('--import', dest='import_path', metavar='FILE.xlsx')
    parser.add_argument('--apply', action='store_true',
                        help='with --import: write; without it nothing is written')
    args = parser.parse_args(argv)
    if args.apply and not args.import_path:
        sys.stderr.write('ERROR: --apply works only with --import\n')
        return 2
    if not os.path.exists(args.db):
        sys.stderr.write('ERROR: database not found at %s - refusing to run\n'
                         % args.db)
        return 2
    con = store.connect(args.db)
    con.isolation_level = None
    try:
        missing = store.missing_tables(con)
        if missing:
            sys.stderr.write('ERROR: tables missing: %s - run '
                             'migrate_agro_work_001.py first\n'
                             % ', '.join(missing))
            return 2
        if args.export:
            count = export(con, args.export)
            print('exported %d work type(s) to %s'
                  % (count, config.ascii_only(args.export)))
            if not count:
                print('the dictionary is empty - run tools/agro_work_import.py '
                      'first')
            return 0
        if not args.import_path:
            print_status(con)
            return 0
        try:
            wanted = read_markup(args.import_path)
        except (Refused, OSError, KeyError, ValueError) as exc:
            sys.stderr.write('ERROR: %s\n' % config.ascii_only(exc))
            print('nothing written')
            return 2
        current = methods.current_methods(con)
        changes, unknown = methods.plan(current, wanted)
        if unknown:
            sys.stderr.write('ERROR: %d work type id(s) of the file are not in '
                             'the database, first: %s\n'
                             % (len(unknown), unknown[0]))
            print('nothing written')
            return 2
        kinds = Counter(kind for _, _, _, kind in changes)
        names = {row['id']: row['name'] for row in con.execute(
            'SELECT id, name FROM agro_work_work_types')}
        for work_type_id, old, new, kind in changes:
            if kind != methods.SAME:
                print('%-6s %s: %s -> %s' % (kind, config.ascii_only(
                    names.get(work_type_id, work_type_id))[:60],
                    old or '-', new or '-'))
        print('\nset %d | change %d | clear %d | same %d'
              % (kinds[methods.SET], kinds[methods.CHANGE],
                 kinds[methods.CLEAR], kinds[methods.SAME]))
        if not args.apply:
            print('dry run: nothing was written. Re-run with --apply.')
            return 0
        try:
            written = methods.apply(
                con, changes, 'xlsx:%s' % os.path.basename(args.import_path))
        except sqlite3.OperationalError as exc:
            if not store.is_readonly(exc):
                raise
            sys.stderr.write('ERROR: %s\n' % store.READONLY_HINT)
            return 2
        print('written: %d' % written)
        return 0
    finally:
        con.close()


if __name__ == '__main__':
    sys.exit(main())
