# -*- coding: utf-8 -*-
"""Самопроверка tools/dji_area_bridge_b_validation.py.

Что здесь держится и почему именно так:

* цепочки строит ЗАМОРОЖЕННЫЙ экран, а не копия правила в инструменте;
* мостиков бывает больше одного, и мостик бывает сам целью C -- такой B нельзя
  считать дважды;
* общий для двух цепочек мостик входит в площадь ОДИН раз;
* база открыта только на чтение: SHA-256 до и после, плюс запись, которую
  SQLite обязан отвергнуть;
* категории B -- функция строки резолвера, и каждая ветка проверена парой
  «срабатывает / не срабатывает», иначе проверка ничего не различает.

Stdlib sqlite3 без Flask; xlsx читается обратно openpyxl.

Запуск:  python tools\\test_dji_area_bridge_b_validation.py
"""

import io
import json
import os
import shutil
import sqlite3
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from dji_area import pipeline as pl  # noqa: E402
from dji_area import resolver as rs  # noqa: E402
from dji_area import store  # noqa: E402
from tests.test_dji_area_core import frame, v4_bytes  # noqa: E402
from tests.test_dji_area_identity_001 import _ddl, ts  # noqa: E402
from tools import dji_area_bridge_b_validation as tool  # noqa: E402

DAY = date(2026, 9, 2)
MU = 2000.0 / 3.0
NICK = 'SYNTHETIC-7'

# (id, начало, конец, режим, ширина, RAW м2, расход мл)
A1, B1, C1 = 960001, 960002, 960003
A2, M1, X2, M2, C2 = 960011, 960012, 960013, 960014, 960015
LONE = 960020
FLIGHTS = (
    # Обычная цепочка: Auto -> ручной мостик -> цель.
    (A1, '2026-09-02 08:00:00', '2026-09-02 08:07:00', 4, 6.0, 9000.0, 14000),
    (B1, '2026-09-02 08:07:00', '2026-09-02 08:07:30', 1, None, 600.0, 2400),
    (C1, '2026-09-02 08:07:30', '2026-09-02 08:08:30', 4, None, 9000.0, 16400),
    # Три мостика: ручной, запись mode 4 без ширины (сама цель), ручной.
    (A2, '2026-09-02 10:00:00', '2026-09-02 10:06:00', 4, 6.0, 8000.0, 12000),
    (M1, '2026-09-02 10:06:00', '2026-09-02 10:06:20', 1, None, 400.0, 1500),
    (X2, '2026-09-02 10:06:20', '2026-09-02 10:06:40', 4, None, 8000.0, 13500),
    (M2, '2026-09-02 10:06:40', '2026-09-02 10:07:00', 1, None, 0.0, 0),
    (C2, '2026-09-02 10:07:00', '2026-09-02 10:08:00', 4, None, 8000.0, 13500),
    # Отдельный взлёт: другая точка, цепочки нет.
    (LONE, '2026-09-02 12:00:00', '2026-09-02 12:08:00', 4, 6.0, 7000.0, 11000),
)
BY_ID = {row[0]: row for row in FLIGHTS}
TAKEOFF = {A1: (39.10, 64.10), A2: (39.20, 64.20), LONE: (39.30, 64.30)}


def takeoff_of(flight_id):
    for anchor in (A1, A2, LONE):
        if anchor <= flight_id < anchor + 9:
            return TAKEOFF[anchor]
    raise KeyError(flight_id)


def list_page():
    rows = []
    for fid, start, end, mode, width, raw, usage in FLIGHTS:
        lat, lng = takeoff_of(fid)
        rows.append({'id': fid, 'new_work_area': raw, 'mode_name': mode,
                     'manual_mode': mode != 4, 'spray_width': width,
                     'start_timestamp': ts(start), 'end_timestamp': ts(end),
                     'nickname': NICK, 'spray_usage': usage, 'lat': lat,
                     'lng': lng, 'work_time_seconds': ts(end) - ts(start),
                     'serial_number': 'R%d' % fid})
    return json.dumps({'code': 0, 'data': rows}).encode('utf-8')


class Fixture(object):

    def __init__(self):
        self.tmp = tempfile.mkdtemp(prefix='bridge_b_')
        self.db = os.path.join(self.tmp, 'instance', 'test.db')
        self.out = os.path.join(self.tmp, 'out')
        os.makedirs(os.path.dirname(self.db))
        con = sqlite3.connect(self.db)
        con.execute('CREATE TABLE drone_units (id INTEGER PRIMARY KEY, '
                    'hardware_id TEXT)')
        con.execute('CREATE TABLE drone_flights (id INTEGER PRIMARY KEY, '
                    'dji_flight_id BIGINT UNIQUE, started_at TEXT, '
                    'finished_at TEXT, raw_json TEXT, drone_unit_id INTEGER, '
                    'nickname_raw TEXT)')
        for stmt in _ddl():
            con.execute(stmt)
        con.execute("INSERT INTO drone_units VALUES (1, 'BODY-CODE')")
        for fid, start, end, _m, _w, _r, _u in FLIGHTS:
            con.execute(
                'INSERT INTO drone_flights (dji_flight_id, started_at, '
                'finished_at, raw_json, drone_unit_id, nickname_raw) '
                'VALUES (?,?,?,?,1,?)', (fid, start, end, None, NICK))
        con.commit()
        con.close()
        body = list_page()
        con = store.connect(self.db)
        root = store.source_root(os.path.abspath(self.db))
        store.begin_immediate(con)
        for fid in BY_ID:
            store.upsert_source_revision(con, root, 'list', body,
                                         flight_id=fid, inline_max_bytes=0)
            store.refresh_flight_evidence(con, root, fid)
        con.execute('COMMIT')
        con.close()

    def add_v4(self, flight_id, first_mu, last_mu, spray):
        row = BY_ID[flight_id]
        t0, t1 = ts(row[1]), ts(row[2])
        n = t1 - t0
        frames = [frame((t0 + i) * 1000,
                        area=first_mu + (last_mu - first_mu) * i / float(n),
                        spray_flag=1 if spray else None,
                        flow=100 if spray else None)
                  for i in range(n + 1)]
        con = store.connect(self.db)
        root = store.source_root(os.path.abspath(self.db))
        store.begin_immediate(con)
        store.upsert_source_revision(con, root, 'v4', v4_bytes(frames),
                                     flight_id=flight_id)
        store.refresh_flight_evidence(con, root, flight_id)
        con.execute('COMMIT')
        con.close()

    def recalculate(self):
        pl.recalculate(self.db, DAY, DAY, apply=True)

    def run(self, *extra):
        out = io.StringIO()
        with redirect_stdout(out):
            code = tool.main(['--db', self.db, '--from', '2026-09-02', '--to',
                              '2026-09-02', '--out', self.out] + list(extra))
        return code, out.getvalue()

    def summary(self):
        with io.open(os.path.join(self.out, 'bridge_b_summary.json'),
                     encoding='utf-8') as fh:
            return json.load(fh)

    def close(self):
        shutil.rmtree(self.tmp, ignore_errors=True)


class Base(unittest.TestCase):

    def setUp(self):
        self.fx = Fixture()
        self.addCleanup(self.fx.close)


class ChainsComeFromTheFrozenScreen(Base):

    def chains(self):
        con = tool.open_readonly(self.fx.db)
        try:
            _items, by_key = tool.grouped_records(con, DAY, DAY)
            return tool.reconstruct_chains(by_key)
        finally:
            con.close()

    def test_a_single_bridge_chain_is_reconstructed(self):
        chain = [c for c in self.chains() if c['c_id'] == C1][0]
        self.assertEqual(chain['a_id'], A1)
        self.assertEqual(chain['b_ids'], [B1])
        self.assertIs(chain['scalar_source_check'], True)

    def test_a_chain_may_carry_several_bridges(self):
        chain = [c for c in self.chains() if c['c_id'] == C2][0]
        self.assertEqual(chain['a_id'], A2)
        self.assertEqual(chain['b_ids'], [M1, X2, M2])

    def test_a_bridge_can_itself_be_the_target_of_another_chain(self):
        found = {c['c_id']: c for c in self.chains()}
        self.assertEqual(set(found), {C1, X2, C2})
        self.assertEqual(found[X2]['b_ids'], [M1])
        # M1 -- общий мостик двух цепочек, X2 -- и мостик, и цель.
        self.assertIn(M1, found[C2]['b_ids'])
        self.assertIn(X2, found[C2]['b_ids'])

    def test_the_tool_does_not_carry_its_own_copy_of_the_rule(self):
        # [REASON]: скопированное правило разошлось бы с замороженным при
        # первой же правке одного из двух. Экран подменяется -- и инструмент
        # обязан пойти за ним, а не за собственной логикой.
        from dji_area import structural
        real = structural.screen_all
        structural.screen_all = lambda records: [
            structural.not_applicable('TEST') for _ in records]
        self.addCleanup(setattr, structural, 'screen_all', real)
        self.assertEqual(self.chains(), [])

    def test_the_lone_flight_is_not_a_chain(self):
        ids = {c['c_id'] for c in self.chains()}
        self.assertNotIn(LONE, ids)


class TheRunIsReadOnly(Base):

    def test_the_database_is_byte_identical_after_a_run(self):
        self.fx.recalculate()
        before = tool.file_sha256(self.fx.db)
        code, text = self.fx.run()
        self.assertEqual(code, tool.EXIT_OK, text)
        self.assertEqual(tool.file_sha256(self.fx.db), before)
        self.assertIn('database sha256 unchanged: yes', text)

    def test_sqlite_itself_refuses_a_write_on_this_handle(self):
        con = tool.open_readonly(self.fx.db)
        self.addCleanup(con.close)
        with self.assertRaises(sqlite3.OperationalError):
            con.execute('DELETE FROM drone_flights')

    def test_the_same_delete_succeeds_on_a_normal_handle(self):
        # Отрицательный контроль: отказ выше -- свойство `mode=ro`, а не
        # того, что DELETE на этой таблице невозможен вообще.
        con = sqlite3.connect(self.fx.db)
        self.addCleanup(con.close)
        con.execute('DELETE FROM drone_flights WHERE dji_flight_id = ?',
                    (LONE,))
        self.assertEqual(con.total_changes, 1)

    def test_a_missing_database_is_code_2_and_no_file_appears(self):
        ghost = os.path.join(self.fx.tmp, 'nowhere', 'absent.db')
        out = io.StringIO()
        with redirect_stdout(out):
            code = tool.main(['--db', ghost, '--from', '2026-09-02', '--to',
                              '2026-09-02', '--out', self.fx.out])
        self.assertEqual(code, tool.EXIT_NO_DATABASE)
        self.assertFalse(os.path.exists(ghost))

    def test_console_output_is_ascii_only(self):
        self.fx.recalculate()
        _code, text = self.fx.run()
        self.assertTrue(all(ord(ch) < 128 for ch in text), text)


class TheReferenceSetStopsTheRun(Base):

    def test_a_plan_that_disagrees_is_exit_3(self):
        self.fx.recalculate()
        plan = os.path.join(self.fx.tmp, 'plan.json')
        with io.open(plan, 'w', encoding='utf-8') as fh:
            json.dump({'locked': {'candidates': [
                {'flight_id': C1}, {'flight_id': 999}]}}, fh)
        code, text = self.fx.run('--plan', plan)
        self.assertEqual(code, tool.EXIT_MISMATCH)
        self.assertIn('STOP', text)

    def test_a_plan_that_agrees_lets_the_run_through(self):
        self.fx.recalculate()
        plan = os.path.join(self.fx.tmp, 'plan.json')
        with io.open(plan, 'w', encoding='utf-8') as fh:
            json.dump({'locked': {'candidates': [
                {'flight_id': C1}, {'flight_id': X2}, {'flight_id': C2}]}}, fh)
        code, text = self.fx.run('--plan', plan)
        self.assertEqual(code, tool.EXIT_OK, text)


def profile(**over):
    base = {'is_itself_candidate': False, 'raw_area_m2': 600.0,
            'v4_present': True, 'area_status': rs.RAW_CORROBORATED,
            'application_activity': rs.ACT_PRESENT,
            'application_channel_quality': rs.CH_INFORMATIVE}
    base.update(over)
    return base


class ResearchCategories(unittest.TestCase):
    """Каждая ветка -- парой: срабатывает и НЕ срабатывает."""

    def category(self, **over):
        return tool.research_category(profile(**over))[0]

    def test_real_work_needs_both_counter_and_application(self):
        self.assertEqual(self.category(), tool.B_REAL_WORK)
        self.assertNotEqual(
            self.category(application_activity=rs.ACT_NOT_OBSERVED),
            tool.B_REAL_WORK)
        self.assertNotEqual(
            self.category(application_activity=rs.ACT_UNKNOWN),
            tool.B_REAL_WORK)

    def test_counter_growth_without_application_is_a_conflict(self):
        self.assertEqual(
            self.category(application_activity=rs.ACT_NOT_OBSERVED),
            tool.B_EVIDENCE_CONFLICT)

    def test_zero_proven_needs_a_flat_counter_and_no_application(self):
        flat = rs.COUNTER_FLAT_RAW_OVERSTATED
        self.assertEqual(self.category(
            area_status=flat, application_activity=rs.ACT_NOT_OBSERVED),
            tool.B_ZERO_PROVEN)
        # Применение при плоском счётчике -- конфликт, а не ноль.
        self.assertEqual(self.category(
            area_status=flat, application_activity=rs.ACT_PRESENT),
            tool.B_EVIDENCE_CONFLICT)
        # Канал молчит -- ноль не доказан.
        self.assertEqual(self.category(
            area_status=flat, application_activity=rs.ACT_UNKNOWN),
            tool.B_VISUAL_REVIEW_REQUIRED)

    def test_partial_follows_the_resolver_status(self):
        self.assertEqual(self.category(
            area_status=rs.PARTIAL_RECORDED_OVERSTATEMENT), tool.B_PARTIAL)

    def test_without_v4_nothing_is_proven_either_way(self):
        self.assertEqual(self.category(v4_present=False),
                         tool.B_VISUAL_REVIEW_REQUIRED)
        self.assertEqual(
            tool.research_category(profile(v4_present=False))[1],
            'V4_NOT_CAPTURED')

    def test_no_area_claimed_is_not_a_correction(self):
        self.assertEqual(self.category(raw_area_m2=0.0),
                         tool.B_NO_AREA_CLAIMED)
        self.assertNotEqual(self.category(raw_area_m2=1.0),
                            tool.B_NO_AREA_CLAIMED)

    def test_a_bridge_that_is_itself_c_is_labelled_before_anything_else(self):
        self.assertEqual(self.category(
            is_itself_candidate=True,
            area_status=rs.PARTIAL_RECORDED_OVERSTATEMENT), tool.B_IS_ITSELF_C)

    def test_an_unlisted_status_goes_to_a_person(self):
        self.assertEqual(self.category(area_status=rs.BASELINE_UNKNOWN),
                         tool.B_VISUAL_REVIEW_REQUIRED)

    def test_production_class_names_are_not_reused(self):
        from dji_area import accounting
        for name in tool.CATEGORIES:
            self.assertNotIn(name, accounting.ACCOUNTING_CLASSES)
            self.assertTrue(name.startswith('B_'))


class EndToEnd(Base):

    def setUp(self):
        super(EndToEnd, self).setUp()
        # База цепочки доказывает, что канал борта информативен.
        self.fx.add_v4(A1, 0.0, 9000.0 / MU, spray=True)
        # B1: счётчик вырос на RAW, насос работал -- настоящая обработка.
        self.fx.add_v4(B1, 0.0, 600.0 / MU, spray=True)
        # M1: счётчик вырос на RAW, но применения не наблюдалось.
        self.fx.add_v4(M1, 0.0, 400.0 / MU, spray=False)
        self.fx.recalculate()
        code, self.text = self.fx.run()
        self.assertEqual(code, tool.EXIT_OK, self.text)
        self.summary = self.fx.summary()

    def test_counts_match_the_fixture(self):
        self.assertEqual(self.summary['chains'], 3)
        self.assertEqual(self.summary['bridges_per_chain'], {'1': 2, '3': 1})
        self.assertEqual(self.summary['bridge_slots'], 5)
        self.assertEqual(self.summary['unique_bridges'], 4)
        self.assertEqual(self.summary['bridges_shared_by_several_chains'], 1)
        self.assertEqual(self.summary['bridges_that_are_themselves_c'], 1)

    def test_a_shared_bridge_enters_the_area_once(self):
        # 600 (B1) + 400 (M1, один раз) + 8000 (X2) + 0 (M2).
        self.assertAlmostEqual(self.summary['bridge_raw_total_ha'], 0.9)

    def test_categories_follow_the_evidence(self):
        cats = self.summary['categories']
        self.assertEqual(cats[tool.B_REAL_WORK]['bridges'], 1)
        self.assertEqual(cats[tool.B_EVIDENCE_CONFLICT]['bridges'], 1)
        self.assertEqual(cats[tool.B_IS_ITSELF_C]['bridges'], 1)
        self.assertEqual(cats[tool.B_NO_AREA_CLAIMED]['bridges'], 1)
        self.assertEqual(cats[tool.B_ZERO_PROVEN]['bridges'], 0)
        self.assertEqual(
            self.summary['B_EVIDENCE_CONFIRMED_OVERSTATEMENT_ha'], 0.0)

    def test_the_takeoff_point_is_shared_inside_a_chain_only(self):
        self.assertEqual(self.summary['same_takeoff_point_chains'], 3)
        control = self.summary['takeoff_point_control']
        # Отрицательный контроль: у записей РАЗНЫХ взлётов точка разная,
        # иначе совпадение внутри цепочки ничего бы не значило.
        self.assertEqual(control['neighbours_gap_over_60s'],
                         {'different': 2})
        self.assertNotIn('different', control['neighbours_gap_within_rule'])

    def test_the_cumulative_bridge_is_not_added_to_the_usage_sum(self):
        # C2 = 13500 = A2 12000 + M1 1500 + M2 0. X2 несёт накопительные
        # 13500, и сложить его значило бы получить отрицательный остаток.
        self.assertEqual(self.summary['c_usage_equals_a_plus_b'], 3)
        self.assertEqual(self.summary['c_usage_residual_ml']['min'], 0)

    def test_sortie_signatures_describe_without_classifying(self):
        self.assertEqual(self.summary['sortie_signatures'],
                         {'AMN': 1, 'AMNMN': 1, 'A': 1})

    def test_the_workbook_is_usable_by_a_person(self):
        from openpyxl import load_workbook
        book = load_workbook(os.path.join(
            self.fx.out, 'DJI_SEPTEMBER_BRIDGE_B_VERIFICATION.xlsx'))
        self.assertEqual(book.sheetnames, [
            'Сводка', 'Цепочки_ABC', 'Bridge_B_проверка',
            'B_для_выборочной_проверки', 'Необычные_B'])
        chains = book['Цепочки_ABC']
        self.assertEqual(chains.max_row, 1 + 3)
        bridges = book['Bridge_B_проверка']
        self.assertEqual(bridges.max_row, 1 + 5)
        header = [c.value for c in bridges[1]]
        for title in ('ОТКРЫТЬ A', 'ОТКРЫТЬ B', 'ОТКРЫТЬ C',
                      'Визуальный вердикт владельца', 'Комментарий владельца'):
            self.assertIn(title, header)
        column = header.index('ОТКРЫТЬ B') + 1
        ids = header.index('B ID') + 1
        for row in range(2, bridges.max_row + 1):
            cell = bridges.cell(row=row, column=column)
            self.assertEqual(
                cell.hyperlink.target,
                'https://www.djiag.com/record/%d'
                % bridges.cell(row=row, column=ids).value)
            # Поля владельца обязаны быть пустыми: их заполняет человек.
            self.assertIsNone(bridges.cell(
                row=row, column=header.index(
                    'Визуальный вердикт владельца') + 1).value)

    def test_the_sample_is_deterministic(self):
        con = tool.open_readonly(self.fx.db)
        try:
            root = store.source_root(os.path.abspath(self.fx.db))
            items, by_key = tool.grouped_records(con, DAY, DAY)
            chains = tool.reconstruct_chains(by_key)
            slots = tool.build_profiles(con, root, items, chains,
                                        with_geometry=False)
        finally:
            con.close()
        first = [s['b_id'] for s, _why in tool.select_sample(slots, seed=7)]
        again = [s['b_id'] for s, _why in tool.select_sample(slots, seed=7)]
        self.assertEqual(first, again)
        # Необычные попадают в выборку всегда, при любом посеве.
        for seed in (1, 2, 3):
            picked = {s['b_id'] for s, _w in tool.select_sample(slots,
                                                                seed=seed)}
            self.assertIn(M1, picked)       # конфликт доказательств
            self.assertIn(X2, picked)       # мостик, который сам цель

    def test_every_drone_reaches_the_sample(self):
        """Борт без необычных записей не выпадает из проверки человеком."""
        def slot(b_id, nickname, category):
            return {'b_id': b_id, 'nickname': nickname, 'category': category,
                    'bridges_in_chain': 1, 'raw_area_m2': 100.0 + b_id,
                    'duration_s': 10 + b_id, 'list_spray_usage_ml': 500}
        slots = [slot(i, 'DRONE-X', tool.B_EVIDENCE_CONFLICT)
                 for i in range(1, 9)]
        small = slot(99, 'DRONE-Y', tool.B_VISUAL_REVIEW_REQUIRED)
        # Маленькая и короткая: иначе она попала бы как выброс по площади, и
        # проход покрытия бортов остался бы непроверенным.
        small['raw_area_m2'] = 1.0
        small['duration_s'] = 1
        slots.append(small)
        # typical=0: типичных пулов нет, и DRONE-Y может попасть в выборку
        # ТОЛЬКО проходом покрытия бортов.
        picked = tool.select_sample(slots, seed=1, typical=0)
        names = {s['nickname'] for s, _why in picked}
        self.assertEqual(names, {'DRONE-X', 'DRONE-Y'})
        reasons = {s['b_id']: why for s, why in picked}
        self.assertEqual(reasons[99], 'DRONE_COVERAGE')

    def test_the_capture_manifest_is_written_only_on_request(self):
        path = os.path.join(self.fx.out, 'bridge_b_v4_capture_ids.txt')
        self.assertFalse(os.path.exists(path))
        code, _text = self.fx.run('--capture-manifest')
        self.assertEqual(code, tool.EXIT_OK)
        self.assertTrue(os.path.exists(path))


if __name__ == '__main__':
    unittest.main()
