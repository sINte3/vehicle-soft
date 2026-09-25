# -*- coding: utf-8 -*-
"""Самопроверка tools/dji_area_control_acceptance.py.

Приёмка, которая не умеет падать, ничего не принимает. Поэтому каждая
проверка здесь стоит парой: верный оракул -- PASS, тот же оракул с одним
испорченным числом -- FAIL с названием именно этого числа.

Данные строит НАСТОЯЩИЙ конвейер на синтетической базе: цепочка
база -> мостик -> цель, у цели V4 с неподвижным счётчиком. Отдельно лежит
сентябрьский оракул репозитория: его арифметика обязана сходиться сама с
собой, а якоря -- совпадать с числами постановки.

Stdlib sqlite3, без Flask.

Запуск:  python tools\\test_dji_area_control_acceptance.py
"""

import copy
import hashlib
import io
import json
import os
import sys
import unittest
from contextlib import redirect_stdout

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from dji_area import pipeline as pl  # noqa: E402
from tests.test_dji_area_identity_001 import (  # noqa: E402
    BASE, BRIDGE, DAY, Fixture, MU, TARGET)
from tools import dji_area_control_acceptance as tool  # noqa: E402

ORACLE = os.path.join(REPO_ROOT, 'docs', 'DJI_AREA_SEPTEMBER_2026_ORACLE.json')


class Base(unittest.TestCase):

    def setUp(self):
        self.fx = Fixture()
        self.addCleanup(self.fx.close)
        # База доказывает, что канал борта информативен; мостик льёт; у цели
        # счётчик стоит на значении базы -- retained-запись.
        self.fx.add_v4(BASE, 0.0, 10000.0 / MU, spray=True)
        self.fx.add_v4(BRIDGE, 0.0, 500.0 / MU, spray=True)
        self.fx.add_v4(TARGET, 10000.0 / MU, 10000.0 / MU, spray=False)
        pl.recalculate(self.fx.db, DAY, DAY, apply=True)
        self.oracle_path = os.path.join(self.fx.tmp, 'oracle.json')
        self.expected = self.observed()
        self.write_oracle(self.expected)

    def observed(self):
        import sqlite3
        con = sqlite3.connect(self.fx.db)
        con.row_factory = sqlite3.Row
        try:
            return tool.observe(tool.load_rows(con, DAY, DAY))[0]
        finally:
            con.close()

    def write_oracle(self, expected):
        with io.open(self.oracle_path, 'w', encoding='utf-8') as fh:
            json.dump({'period': [DAY.isoformat(), DAY.isoformat()],
                       'expected': expected}, fh)

    def run_tool(self, *extra):
        out = io.StringIO()
        with redirect_stdout(out):
            code = tool.main(['--db', self.fx.db, '--oracle',
                              self.oracle_path] + list(extra))
        return code, out.getvalue()


class TheFixtureIsARealCase(Base):

    def test_the_target_is_a_proven_phantom_and_the_bridge_is_not(self):
        self.assertEqual(self.expected['structural_candidates'], 1)
        self.assertEqual(self.expected['excluded_records'], 1)
        self.assertAlmostEqual(self.expected['excluded_ha'], 1.0)
        self.assertEqual(self.expected['chains_shown'], 1)
        self.assertEqual(self.expected['bridges_excluded'], 0)
        # RAW: база 1.0 + мостик 0.05 + цель 1.0 + одиночный 0.9 га.
        self.assertAlmostEqual(self.expected['raw_ha'], 2.95)
        self.assertAlmostEqual(self.expected['after_ha'], 1.95)


class PassAndFail(Base):

    def test_the_right_oracle_passes(self):
        code, text = self.run_tool()
        self.assertEqual(code, tool.EXIT_PASS, text)
        self.assertIn('VERDICT: PASS', text)

    def test_one_wrong_total_fails_and_is_named(self):
        wrong = copy.deepcopy(self.expected)
        wrong['excluded_ha'] = wrong['excluded_ha'] + 0.5
        self.write_oracle(wrong)
        code, text = self.run_tool()
        self.assertEqual(code, tool.EXIT_FAIL)
        self.assertIn('MISMATCH: excluded_ha', text)
        self.assertIn('VERDICT: FAIL', text)

    def test_one_wrong_drone_figure_fails_even_when_the_total_is_right(self):
        wrong = copy.deepcopy(self.expected)
        name = sorted(wrong['by_drone'])[0]
        wrong['by_drone'][name]['raw_ha'] += 0.25
        self.write_oracle(wrong)
        code, text = self.run_tool()
        self.assertEqual(code, tool.EXIT_FAIL)
        self.assertIn('by_drone[', text)

    def test_a_missing_drone_is_a_mismatch(self):
        wrong = copy.deepcopy(self.expected)
        wrong['by_drone']['DRONE-THAT-IS-NOT-THERE'] = {
            'records': 1, 'raw_ha': 1.0, 'excluded_ha': 0.0, 'after_ha': 1.0,
            'pending_ha': 0.0, 'review_ha': 0.0}
        self.write_oracle(wrong)
        self.assertEqual(self.run_tool()[0], tool.EXIT_FAIL)

    def test_a_wrong_count_fails(self):
        wrong = copy.deepcopy(self.expected)
        wrong['review_application_with_flat_counter'] = 4
        self.write_oracle(wrong)
        code, text = self.run_tool()
        self.assertEqual(code, tool.EXIT_FAIL)
        self.assertIn('review_application_with_flat_counter', text)

    def test_a_tiny_rounding_difference_is_tolerated(self):
        # Отрицательный контроль к допуску: он узкий, а не «что угодно».
        near = copy.deepcopy(self.expected)
        near['raw_ha'] += 0.00004
        self.write_oracle(near)
        self.assertEqual(self.run_tool()[0], tool.EXIT_PASS)
        near['raw_ha'] += 0.001
        self.write_oracle(near)
        self.assertEqual(self.run_tool()[0], tool.EXIT_FAIL)


class ReadOnly(Base):

    def digest(self):
        with open(self.fx.db, 'rb') as fh:
            return hashlib.sha256(fh.read()).hexdigest()

    def test_the_database_is_byte_identical_after_the_run(self):
        before = self.digest()
        code, text = self.run_tool()
        self.assertEqual(code, tool.EXIT_PASS)
        self.assertEqual(self.digest(), before)
        self.assertIn('database sha256 unchanged: yes', text)

    def test_a_missing_database_is_code_2_and_no_file_appears(self):
        ghost = os.path.join(self.fx.tmp, 'nowhere', 'absent.db')
        out = io.StringIO()
        with redirect_stdout(out):
            code = tool.main(['--db', ghost, '--oracle', self.oracle_path])
        self.assertEqual(code, tool.EXIT_NO_DATABASE)
        self.assertFalse(os.path.exists(ghost))

    def test_console_output_is_ascii_only(self):
        _code, text = self.run_tool()
        self.assertTrue(all(ord(ch) < 128 for ch in text), text)

    def test_the_out_directory_gets_the_verdict_and_the_workbook(self):
        from openpyxl import load_workbook
        out_dir = os.path.join(self.fx.tmp, 'out')
        self.run_tool('--out', out_dir)
        with io.open(os.path.join(out_dir, 'area_control_acceptance.json'),
                     encoding='utf-8') as fh:
            self.assertEqual(json.load(fh)['verdict'], 'PASS')
        # Через открытый дескриптор: иначе openpyxl оставляет файл на сборщик
        # мусора, и прогон шумит ResourceWarning.
        with open(os.path.join(out_dir, 'area_control_report.xlsx'),
                  'rb') as fh:
            book = load_workbook(fh)
        # DRONE-AREA-CONTROL-V2-MEGA: книга выросла с четырёх листов до
        # семи; прежние четыре -- первыми и в прежнем порядке.
        self.assertEqual(book.sheetnames[:4], ['Сводка', 'По_дронам',
                                               'Корректировки',
                                               'Требует_проверки'])
        self.assertEqual(len(book.sheetnames), 7)


class TheDryRunRecalculation(Base):
    """Нынешний код не хочет переписать ни одной строки периода."""

    def summary(self, document):
        path = os.path.join(self.fx.tmp, 'dry.json')
        with io.open(path, 'w', encoding='utf-8') as fh:
            if isinstance(document, str):
                fh.write(document)
            else:
                json.dump(document, fh)
        return path

    def records(self):
        return self.expected['records']

    def test_a_real_dry_run_over_the_same_evidence_passes(self):
        # Не подставная сводка: её пишет сам пересчёт на той же базе.
        summary = pl.recalculate(self.fx.db, DAY, DAY, apply=False)
        summary.pop('flights', None)
        code, text = self.run_tool('--recalc-summary', self.summary(summary))
        self.assertEqual(code, tool.EXIT_PASS, text)
        self.assertIn('recalc dry-run calc_writes', text)

    def test_a_summary_that_wants_to_write_fails_and_says_what(self):
        n = self.records()
        code, text = self.run_tool('--recalc-summary', self.summary(
            {'flights_in_period': n,
             'calc_writes': {'unchanged': n - 1, 'new': 1}}))
        self.assertEqual(code, tool.EXIT_FAIL)
        self.assertIn('would write {"new": 1}', text)

    def test_another_period_is_not_accepted_as_this_one(self):
        n = self.records()
        code, text = self.run_tool('--recalc-summary', self.summary(
            {'flights_in_period': n + 5, 'calc_writes': {'unchanged': n + 5}}))
        self.assertEqual(code, tool.EXIT_FAIL)
        self.assertIn('flights_in_period', text)

    def test_an_empty_summary_proves_nothing(self):
        code, text = self.run_tool('--recalc-summary', self.summary(
            {'flights_in_period': self.records()}))
        self.assertEqual(code, tool.EXIT_FAIL)
        self.assertIn('calc_writes is missing', text)

    def test_unchanged_must_cover_every_record(self):
        n = self.records()
        code, _text = self.run_tool('--recalc-summary', self.summary(
            {'flights_in_period': n, 'calc_writes': {'unchanged': n - 1}}))
        self.assertEqual(code, tool.EXIT_FAIL)

    def test_an_unreadable_summary_is_a_usage_error_not_a_pass(self):
        code, _text = self.run_tool('--recalc-summary',
                                    self.summary('{not json'))
        self.assertEqual(code, tool.EXIT_USAGE)
        code, _text = self.run_tool('--recalc-summary', os.path.join(
            self.fx.tmp, 'absent.json'))
        self.assertEqual(code, tool.EXIT_USAGE)

    def test_without_the_flag_nothing_about_the_recalc_is_claimed(self):
        _code, text = self.run_tool()
        self.assertNotIn('recalc dry-run', text)


class TheContractWithTheOwnerBlocks(Base):

    def test_exit_codes_are_the_literal_numbers_the_runbook_reads(self):
        # [REASON]: блоки PowerShell сравнивают $LASTEXITCODE с ЧИСЛОМ. Пока
        # тесты сверяли код только с именем константы, FAIL, ставший нулём,
        # проходил бы их все -- и блок принял бы провал за приёмку.
        self.assertEqual((tool.EXIT_PASS, tool.EXIT_USAGE,
                          tool.EXIT_NO_DATABASE, tool.EXIT_FAIL), (0, 1, 2, 3))

    def test_a_database_that_changed_under_the_reader_fails(self):
        digests = iter(['a' * 64, 'b' * 64])
        original = tool.file_sha256
        tool.file_sha256 = lambda _path: next(digests)
        self.addCleanup(setattr, tool, 'file_sha256', original)
        code, text = self.run_tool()
        self.assertEqual(code, 3)
        self.assertIn('read-only: the database changed', text)
        self.assertIn('database sha256 unchanged: NO', text)


class TheSeptemberOracle(unittest.TestCase):
    """Оракул репозитория: якоря постановки и внутренняя арифметика."""

    def setUp(self):
        with io.open(ORACLE, encoding='utf-8') as fh:
            self.oracle = json.load(fh)
        self.expected = self.oracle['expected']

    def test_the_anchors_are_the_ones_the_owner_named(self):
        self.assertEqual(self.oracle['period'], ['2026-09-01', '2026-09-18'])
        self.assertEqual(self.expected['records'], 4623)
        self.assertEqual(self.expected['raw_missing_records'], 0)
        self.assertEqual(self.expected['structural_candidates'], 233)
        self.assertAlmostEqual(self.expected['raw_ha'], 4413.2825, places=4)
        self.assertAlmostEqual(self.expected['excluded_ha'], 167.8088,
                               places=4)
        self.assertAlmostEqual(self.expected['after_ha'], 4245.4737, places=4)

    def test_the_four_flat_counter_cases_stay_in_review(self):
        self.assertEqual(
            self.expected['review_application_with_flat_counter'], 4)
        self.assertEqual(self.expected['proven_structural'], 229)
        self.assertEqual(self.expected['proven_by_control_only'], 1)
        self.assertEqual(self.expected['pending_records'], 0)
        self.assertFalse(self.expected['complete'])

    def test_the_oracle_adds_up_by_itself(self):
        self.assertEqual(tool.formula_problems(self.expected), [])

    def test_no_flight_id_is_hard_coded_in_the_product_code(self):
        # Сентябрьские идентификаторы живут в документах, не в коде.
        for rel in ('dji_area/control_report.py', 'dji_area/accounting.py',
                    'dji_area/capture_manifest.py',
                    'tools/dji_area_control_acceptance.py'):
            with io.open(os.path.join(REPO_ROOT, rel),
                         encoding='utf-8') as fh:
                text = fh.read()
            for flight_id in ('701661028', '702797709', '703830892',
                              '703847599'):
                self.assertNotIn(flight_id, text, rel)


if __name__ == '__main__':
    unittest.main()
