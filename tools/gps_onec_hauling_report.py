# -*- coding: utf-8 -*-
"""GPS-35 -- недельная сверка перевозок хлопка по рейсам 1С (Excel). Только чтение.

ЗАЧЕМ
Решение владельца 27.10 (трек §0.7): «Сверка подходит, раз в неделю, список
отдавать мне». GPS сам перевозку хлопка не подтверждает (контроль 27.10 --
NOT MET), а связка «госномер рейса 1С -- наша машина» проверена по трекам и
надёжна (GPS-34, 27.10 -- MET). Значит, доказательство перевозки -- рейс в
1С: вывоз ПТМ -> завод («Вывоз по дате взвешивания брутто») или приём на
ПТМ (подробный реестр ПК-17).

ЧТО ДЕЛАЕТ
  1. Читает выгрузки 1С за неделю (как GPS-34: xlsx стандартной
     библиотекой) -- пункт, госномер, время и массу нетто: у вывоза «Вес
     нетто, кг», у реестра «Масса физическая,кг».
  2. Госномер -> наша машина по правилу GPS-34: свёртка, у реестра серия
     переносится в конец, машина -- ровно одна строка справочника техники.
     Объект Wialon здесь не нужен: заявка и рейс сходятся на машине.
  3. Сутки сверки -- сутки, которые покрывает реестр ПК-17 (его выгружают
     ровно по «Дате брутто»); без реестра -- сутки рейсов.
  4. Заявки agro-work, чьё окно задевает эти сутки, -- как «Сверка»
     (`agro_work.reconcile`): машина заявки, окно, объём.
  5. Книга Excel, три листа:
     «Заявки» -- каждые сутки сверки в окне каждой заявки «пахта ташиш» /
         «тола ташиш»: сколько у её машины в эти сутки рейсов 1С (вывоз,
         приём), нетто по 1С и итог: есть рейс в 1С, нет рейса, машина не
         сопоставлена;
     «Без заявки» -- машино-сутки наших машин с рейсами 1С, которые не
         покрывает ни одна заявка agro-work (любой вид работы): рейсы,
         нетто, пункты и вердикт покрытия суток с экрана «Сверка»
         (`Reconciliation.coverage`): «работа без заявки» или «без
         вердикта» с причиной -- например, открытая заявка машины может
         покрыть сутки после закрытия;
     «Сводка» -- период, числа, заявки без окна по причинам.
  6. В консоль -- только числа, ASCII.

ЧЕГО НЕ ДЕЛАЕТ
  НИЧЕГО НЕ ПИШЕТ, кроме книги: база открывается `mode=ro`.
  Не судит заявки на экране «Сверка» и в базу вердикт не кладёт.
  Не берёт из выгрузок фермеров, товароведов и машины не из справочника.
  Тонны заявки и тонны 1С не сравнивает и правила о расхождении не вводит:
      тонны в заявках пишут примерно (владелец 24.10), правила нет.

КОДЫ ВЫХОДА
  0  готово, книга записана;
  2  неверный ввод: нет базы, файла, колонки, таблиц agro-work; ничего не
     записано;
  1  инструмент упал.

Запуск -- системный Python (ему нужен openpyxl из requirements.txt), из
C:\\gps-tools, по одной команде на строку:

  & "C:\\Program Files\\Python314\\python.exe" C:\\gps-tools\\tools\\gps_onec_hauling_report.py --db C:\\transport-report\\instance\\transport.db --vyvoz C:\\gps-tools\\check\\onec_vyvoz_po_brutto_2026-10-20_26.xlsx --reestr C:\\gps-tools\\check\\onec_pk17_reestr_2026-10-20_26.xlsx --out C:\\gps-tools\\check\\perevozki_1c_2026-10-20_26.xlsx

Ключи --vyvoz и --reestr повторяются.
"""

import argparse
import os
import sys
from collections import Counter, defaultdict
from datetime import timedelta

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from agro_work import labels                                        # noqa: E402
from agro_work import reconcile as rc                               # noqa: E402
from agro_work import store                                         # noqa: E402
from gps.hauling import is_cotton, is_hauling                       # noqa: E402
from gps_collector import config as collector_config                # noqa: E402
from plate_norm import normalize_plate                              # noqa: E402
from tools.gps_onec_trips_check import (REESTR, VYVOZ, BadInput,    # noqa: E402
                                        covered_days, read_trips)
from tools.gps_track_only_days import open_readonly                 # noqa: E402

EXIT_OK = 0
EXIT_CRASHED = 1
EXIT_BAD_INPUT = 2

# [REASON]: нетто рейса -- то, что машина везла: у вывоза «Вес нетто»
# (брутто минус тара на весах ПТМ), у приёма «Масса физическая» (то же до
# скидок на засор и влажность). Кондиционная масса -- расчёт для оплаты
# фермеру, а не груз машины.
MASS = {VYVOZ: 'Вес нетто, кг', REESTR: 'Масса физическая,кг'}

HAS_TRIP = ('Есть рейс в 1С', '1С да рейс бор')
NO_TRIP = ('Нет рейса в 1С', '1С да рейс йўқ')
NO_MACHINE = ('Машина не сопоставлена', 'Машина боғланмаган')

APP_HEADER = (('№ заявки', 'Буюртма №'), ('Сутки', 'Кун'), ('Техника', 'Техника'),
              ('Госномер', 'Давлат рақами'), ('Организация', 'Ташкилот'),
              ('Вид работы', 'Иш тури'), ('Объём заявки', 'Буюртма ҳажми'),
              ('Единица', 'Бирлик'),
              ('Рейсов вывоза на завод (1С)', 'Заводга олиб кетиш рейслари (1С)'),
              ('Рейсов приёма на ПТМ (1С)', 'ПТМда қабул рейслари (1С)'),
              ('Нетто по 1С, т', '1С бўйича нетто, т'), ('Итог', 'Натижа'))
APP_WIDTHS = (22, 12, 22, 14, 24, 28, 12, 10, 16, 16, 14, 26)
FREE_HEADER = (('Сутки', 'Кун'), ('Техника', 'Техника'),
               ('Госномер', 'Давлат рақами'), ('Организация', 'Ташкилот'),
               ('Рейсов вывоза на завод (1С)', 'Заводга олиб кетиш рейслари (1С)'),
               ('Рейсов приёма на ПТМ (1С)', 'ПТМда қабул рейслари (1С)'),
               ('Нетто по 1С, т', '1С бўйича нетто, т'),
               ('Пункты приёма', 'Қабул пунктлари'),
               ('Первый рейс', 'Биринчи рейс'), ('Последний рейс', 'Охирги рейс'),
               ('Покрытие (как на «Сверке»)', 'Қопланиш («Солиштириш»даги каби)'),
               ('Причина', 'Сабаб'),
               ('Заявки, которые могут покрыть', 'Қоплаши мумкин бўлган буюртмалар'))
FREE_WIDTHS = (12, 22, 14, 24, 16, 16, 14, 40, 10, 10, 22, 40, 24)


def bi(pair):
    return '%s / %s' % pair


def parse_mass(value):
    """Килограммы из ячейки: число или текст с пробелами и запятой; иначе 0."""
    text = str(value or '').replace('\xa0', '').replace(' ', '').replace(',', '.')
    try:
        return float(text)
    except ValueError:
        return 0.0


def equipment_of(con):
    """{номер: {id техники}} и {id: (название, госномер, организация)}."""
    by_plate = defaultdict(set)
    info = {}
    for equipment_id, name, plate, org in con.execute(
            'SELECT e.id, e.name, e.plate, o.name FROM equipment e '
            'LEFT JOIN organizations o ON o.id = e.organization_id'):
        info[equipment_id] = (name or '', plate or '', org or '')
        key = normalize_plate(plate)
        if key:
            by_plate[key].add(equipment_id)
    return by_plate, info


def machine_days(trips, by_plate):
    """{(машина, сутки): [рейсы]} и счёт рейсов по причине."""
    found = defaultdict(list)
    reasons = Counter()
    for trip in trips:
        if trip['time'] is None:
            reasons['no time'] += 1
            continue
        machines = by_plate.get(trip['plate'], set()) if trip['plate'] else set()
        if len(machines) != 1:
            reasons['not ours' if not machines else 'several machines'] += 1
            continue
        reasons['ours'] += 1
        found[(next(iter(machines)), trip['time'].date())].append(trip)
    return found, reasons


def summary(trips):
    """(рейсов вывоза, рейсов приёма, нетто в тоннах) списка рейсов."""
    removal = sum(1 for trip in trips if trip['kind'] == VYVOZ)
    tonnes = sum(parse_mass(trip['extra'][MASS[trip['kind']]]) for trip in trips) / 1000
    return removal, len(trips) - removal, round(tonnes, 3)


def application_lines(ctx, days, found, info):
    """Строки листа «Заявки» и {причина: заявок} без окна (их нечем судить)."""
    lines, no_window = [], Counter()
    for row in ctx.forward_rows():
        app = row['app']
        name = (ctx.work_type_names.get(app.work_type_id)
                or app.row.get('work_type_name') or '')
        if not (is_hauling(name) and is_cotton(name)):
            continue
        if not row['window']:
            no_window[app.window_reason] += 1
            continue
        machine = row['equipment_id']
        title, plate, org = info.get(machine, ('', app.row.get('plate_number') or '', ''))
        day, end = row['window']
        while day <= end:
            if day in days:
                trips = found.get((machine, day), [])
                removal, receiving, tonnes = summary(trips)
                verdict = (NO_MACHINE if machine is None
                           else HAS_TRIP if trips else NO_TRIP)
                lines.append((app.number, day, title, plate, org, name,
                              row['volume'], app.row.get('unit') or '',
                              removal, receiving, tonnes if trips else None,
                              verdict))
            day += timedelta(days=1)
    lines.sort(key=lambda line: (line[1], line[4], line[2], line[3], line[0]))
    return lines, no_window


def free_lines(ctx, found, info):
    """Строки листа «Без заявки» (машино-сутки с рейсами, не покрытые заявкой)
    и их счёт {(покрытие, причина): машино-суток}."""
    lines, counts = [], Counter()
    for (machine, day), trips in found.items():
        # [REASON]: покрыты ли сутки -- решает `Reconciliation.coverage`, как
        # сутки «работа -> заявка» на экране «Сверка»: те же отменённые и
        # удалённые заявки (не покрывают), та же открытая заявка (может
        # покрыть -- вердикт после закрытия). Своё правило здесь дало бы
        # одним суткам два разных ответа на двух выходах.
        state, reason, apps = ctx.coverage(machine, day)
        if state == rc.C_COVERED:
            continue
        counts[(state, reason)] += 1
        title, plate, org = info.get(machine, ('', '', ''))
        removal, receiving, tonnes = summary(trips)
        times = sorted(trip['time'] for trip in trips)
        points = ', '.join(sorted({trip['point'] for trip in trips if trip['point']}))
        lines.append((day, title, plate, org, removal, receiving, tonnes, points,
                      times[0].strftime('%H:%M'), times[-1].strftime('%H:%M'),
                      labels.COVERAGE[state], labels.REASONS[reason] if reason else '',
                      ', '.join(sorted(str(app.number) for app in apps))))
    lines.sort(key=lambda line: (line[0], line[3], line[1], line[2]))
    return lines, counts


def write_book(path, app_lines, free, summary_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)

    def sheet_with(title, header, rows, widths):
        sheet = book.create_sheet(title)
        sheet.append([bi(pair) for pair in header])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        for row in rows:
            sheet.append([bi(value) if isinstance(value, tuple) else value
                          for value in row])
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A2'
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if hasattr(cell.value, 'year'):
                    cell.number_format = 'DD.MM.YYYY'
        return sheet

    sheet_with('Сводка — Жамланма', (('Показатель', 'Кўрсаткич'),
                                       ('Значение', 'Қиймат')),
               summary_rows, (70, 30))
    sheet_with('Заявки — Буюртмалар', APP_HEADER, app_lines, APP_WIDTHS)
    sheet_with('Без заявки — Буюртмасиз', FREE_HEADER, free, FREE_WIDTHS)
    book.save(path)


def check_inputs(args):
    if not os.path.isfile(args.db):
        raise BadInput('database not found: %s' % args.db)
    files = [(path, VYVOZ) for path in args.vyvoz] + [(path, REESTR) for path in args.reestr]
    if not files:
        raise BadInput('give at least one --vyvoz or --reestr file')
    for path, _kind in files:
        if not os.path.isfile(path):
            raise BadInput('file not found: %s' % path)
    out = os.path.abspath(args.out)
    if (not out.lower().endswith('.xlsx') or not os.path.isdir(os.path.dirname(out))
            or os.path.isdir(out)
            or out in {os.path.abspath(p) for p, _k in files}
            or out == os.path.abspath(args.db)):
        raise BadInput('cannot write %s (an .xlsx path in an existing folder)' % args.out)
    return files


def run(args, out=print):
    files = check_inputs(args)
    trips = []
    for path, kind in files:
        trips.extend(read_trips(path, kind, extra=(MASS[kind],)))
    days = covered_days(trips)
    if not days:
        raise BadInput('the exports have no dated trips')
    con = open_readonly(args.db)
    try:
        if store.missing_tables(con):
            raise BadInput('agro-work tables are missing in %s' % args.db)
        by_plate, info = equipment_of(con)
        ctx = rc.Reconciliation(con, min(days), max(days))
        found, reasons = machine_days(trips, by_plate)
        app_lines, no_window = application_lines(ctx, days, found, info)
        free, free_counts = free_lines(ctx, found, info)
    finally:
        con.close()

    verdicts = Counter(line[-1] for line in app_lines)
    free_states = Counter()
    free_reasons = Counter()
    for (state, reason), count in free_counts.items():
        free_states[state] += count
        if reason:
            free_reasons[reason] += count
    kinds = Counter(trip['kind'] for trip in trips)
    first, last = min(days).strftime('%d.%m.%Y'), max(days).strftime('%d.%m.%Y')
    summary_rows = [
        (('Сутки сверки (сутки реестра ПК-17)', 'Солиштириш кунлари (ПК-17 реестри кунлари)'),
         '%s — %s, %d' % (first, last, len(days))),
        (('Рейсов в выгрузках 1С: вывоз на завод', '1С юкламаларидаги рейслар: заводга олиб кетиш'),
         kinds[VYVOZ]),
        (('Рейсов в выгрузках 1С: приём на ПТМ', '1С юкламаларидаги рейслар: ПТМда қабул'),
         kinds[REESTR]),
        (('Из них — наших машин', 'Шулардан — бизнинг машиналар'), reasons['ours']),
        (('Сутки заявок «пахта/тола ташиш»', '«пахта/тола ташиш» буюртмалари кунлари'),
         len(app_lines)),
        (HAS_TRIP, verdicts[HAS_TRIP]),
        (NO_TRIP, verdicts[NO_TRIP]),
        (NO_MACHINE, verdicts[NO_MACHINE]),
        (('Заявок «пахта/тола ташиш» без окна — сутки не судятся',
          '«пахта/тола ташиш» ойнасиз буюртмалари — кунлари баҳоланмайди'),
         sum(no_window.values())),
    ] + [
        (('— ' + labels.REASONS[reason][0], '— ' + labels.REASONS[reason][1]), count)
        for reason, count in sorted(no_window.items())
    ] + [
        (('Машино-суток рейсов 1С, не покрытых заявкой',
          'Буюртма билан қопланмаган 1С рейслари машина-кунлари'), len(free)),
        (labels.COVERAGE[rc.C_UNCOVERED], free_states[rc.C_UNCOVERED]),
        (labels.COVERAGE[rc.C_NONE], free_states[rc.C_NONE]),
    ] + [
        (('— ' + labels.REASONS[reason][0], '— ' + labels.REASONS[reason][1]), count)
        for reason, count in sorted(free_reasons.items())
    ] + [
        (('Тонны — нетто по 1С: вывоз — «Вес нетто», приём — «Масса физическая»',
          'Тонна — 1С бўйича нетто: олиб кетиш — «Вес нетто», қабул — «Масса физическая»'), ''),
    ]
    write_book(args.out, app_lines, free, summary_rows)

    out('days %s .. %s: %d' % (min(days), max(days), len(days)))
    out('trips read: %d (removal %d, receiving %d) | ours %d | not ours %d | '
        'several machines %d | no time %d'
        % (len(trips), kinds[VYVOZ], kinds[REESTR], reasons['ours'],
           reasons['not ours'], reasons['several machines'], reasons['no time']))
    out('cotton hauling application-days %d | with a 1C trip %d | without %d | '
        'machine not matched %d | applications without a window %d (%s)'
        % (len(app_lines), verdicts[HAS_TRIP], verdicts[NO_TRIP],
           verdicts[NO_MACHINE], sum(no_window.values()),
           ', '.join('%s %d' % item for item in sorted(no_window.items())) or '-'))
    out('machine-days with 1C trips not covered by an application: %d | '
        'without an application %d | no verdict %d (%s)'
        % (len(free), free_states[rc.C_UNCOVERED], free_states[rc.C_NONE],
           ', '.join('%s %d' % item for item in sorted(free_reasons.items())) or '-'))
    out('xlsx: %s' % args.out)
    out('READ ONLY: nothing was written but the workbook')
    return EXIT_OK


def main(argv=None, out=print):
    parser = argparse.ArgumentParser(
        description=collector_config.ascii_only(__doc__.splitlines()[0]))
    parser.add_argument('--db', required=True)
    parser.add_argument('--vyvoz', action='append', default=[],
                        help='1C report of removal trips by gross weighing date')
    parser.add_argument('--reestr', action='append', default=[],
                        help='1C detailed PK-17 register')
    parser.add_argument('--out', required=True, help='the workbook (.xlsx)')
    args = parser.parse_args(argv)
    try:
        return run(args, out)
    except BadInput as problem:
        sys.stderr.write('ERROR: %s\n' % collector_config.ascii_only(str(problem)))
        return EXIT_BAD_INPUT
    except Exception as problem:                                  # noqa: BLE001
        sys.stderr.write('ERROR: %s: %s\n' % (
            type(problem).__name__, collector_config.ascii_only(str(problem))))
        return EXIT_CRASHED


if __name__ == '__main__':
    sys.exit(main())
