# -*- coding: utf-8 -*-
"""agro-work B3 -- отчёт сверки заявок с GPS в обе стороны и замер N.

ЗАЧЕМ
Экран сверки (`/agro-work/`) живёт в приложении и появится на сервере с
релизом. Этот инструмент отдаёт те же числа файлом уже сейчас -- на копии
базы из docs/AGRO_WORK_B1_RUNBOOK.md -- и меряет то, без чего владелец не
утвердит окно заявок, заведённых задним числом: сколько суток проходит от
последнего дня работы по GPS до отметки «Выполнено» (вопрос 4).

ЧТО В КНИГЕ
  Свод           -- организация x категория: машины, заявки (работа была /
                    не было / без вердикта), машино-сутки с работой по GPS
                    (покрыты / без заявки / без вердикта);
  Заявка-работа  -- каждая заявка периода: окно, вердикт или причина, сутки
                    с работой, гектары GPS в окне;
  Работа-заявка  -- каждые машино-сутки с работой по GPS и их покрытие;
  Причины        -- сколько строк без вердикта и почему;
  Замер N        -- распределение лага по трём выборкам и доля заявок,
                    которую покрыло бы окно для N = 0..14.

Ядро одно с экраном: `agro_work/reconcile.py`. Правила -- там.

Только чтение: база открывается `mode=ro`, SQLite сам откажет в записи.
Сеть не нужна.

Запуск (PowerShell, по одной команде на строку):

  & "C:\\Program Files\\Python314\\python.exe" tools\\agro_work_reconcile.py --from 2026-09-01 --to 2026-09-27

`--lookback-preview N` -- то же, как если бы N был утверждён; лист «Свод»
тогда помечен «ПРЕДПРОСМОТР». Утверждённое N ставится в коде, не ключом.

Вывод в консоль -- ASCII.
"""

import argparse
import os
import sqlite3
import sys
from datetime import date, datetime, timedelta

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from agro_work import config, labels, methods, reconcile as rc   # noqa: E402
from agro_work import store                                      # noqa: E402

SAMPLE_TITLES = {
    'all': ('Все обычные закрытые заявки', 'Барча оддий ёпилган буюртмалар'),
    'unit_hectare': ('Только единица HECTARE из API (фильтр замера)',
                     'Фақат API даги HECTARE бирлиги (ўлчов филтри)'),
    'method_ga': ('Только размеченные методом «гектары»',
                  'Фақат «гектарлар» усули билан белгиланганлар'),
}
EXCLUDED_TITLES = {
    'net_raboty_v_okne': ('В окне нет ни одних суток с работой',
                          'Ойнада иш бўлган бирорта кун йўқ'),
    'neizvestnye_sutki_posle_raboty': (
        'После последней работы в окне есть неизвестные сутки',
        'Охирги ишдан кейин ойнада номаълум кунлар бор'),
}


def bi(pair):
    return '%s / %s' % pair


def _open_ro(path):
    # [REASON]: `mode=ro` -- запрет на уровне SQLite, а не обещание
    # скрипта: отчёт, случайно ставший писателем, SQLite остановит сам.
    return sqlite3.connect('file:%s?mode=ro' % os.path.abspath(path)
                           .replace('\\', '/'), uri=True, timeout=30)


def _machine(ctx, equipment_id):
    item = ctx.equipment.get(equipment_id)
    if item is None:
        return ''
    return ('%s %s' % (item['name'], item['plate'])).strip()


def _org(ctx, equipment_id):
    org_id, _ = ctx.group_of(equipment_id)
    return ctx.orgs.get(org_id, '') if org_id else ''


def build_book(ctx, forward, reverse, lags, excluded, preview):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    book = Workbook()

    def sheet_with(title, header, rows, widths):
        sheet = book.create_sheet(title)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        for row in rows:
            sheet.append(row)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A2'
        return sheet

    book.remove(book.active)
    groups, total = ctx.summary(forward, reverse)
    summary_rows = []
    for (org_id, category), counter in sorted(
            groups.items(), key=lambda kv: (kv[0][0] is None, kv[0])):
        org = ctx.orgs.get(org_id, '') if org_id else bi((
            'Машина не сопоставлена', 'Машина боғланмаган'))
        summary_rows.append((
            org, rc_category(category), counter['machines'],
            counter['applications'], counter['app_' + rc.V_WORK],
            counter['app_' + rc.V_NO_WORK], counter['app_' + rc.V_NONE],
            counter['work_days'], counter['day_' + rc.C_COVERED],
            counter['day_' + rc.C_UNCOVERED], counter['day_' + rc.C_NONE]))
    summary_rows.append((bi(('Итого', 'Жами')), '', total['machines'],
                         total['applications'], total['app_' + rc.V_WORK],
                         total['app_' + rc.V_NO_WORK], total['app_' + rc.V_NONE],
                         total['work_days'], total['day_' + rc.C_COVERED],
                         total['day_' + rc.C_UNCOVERED], total['day_' + rc.C_NONE]))
    sheet = sheet_with(
        'Свод',
        (bi(('Организация', 'Ташкилот')), bi(('Категория', 'Тоифа')),
         bi(('Машин', 'Машиналар')), bi(('Заявок', 'Буюртмалар')),
         bi(labels.VERDICTS[rc.V_WORK]), bi(labels.VERDICTS[rc.V_NO_WORK]),
         bi(labels.VERDICTS[rc.V_NONE]),
         bi(('Машино-суток с работой по GPS', 'GPS бўйича иш бўлган машина-кунлар')),
         bi(labels.COVERAGE[rc.C_COVERED]), bi(labels.COVERAGE[rc.C_UNCOVERED]),
         bi(labels.COVERAGE[rc.C_NONE])),
        summary_rows, (28, 28, 10, 10, 14, 16, 14, 18, 16, 16, 14))
    notes = [
        bi(('Период: %s — %s' % (ctx.date_from.strftime('%d.%m.%Y'),
                                 ctx.date_to.strftime('%d.%m.%Y')),
            'Давр: %s — %s' % (ctx.date_from.strftime('%d.%m.%Y'),
                               ctx.date_to.strftime('%d.%m.%Y')))),
        bi(('Работа по GPS у объектов без машины в справочнике, объекто-суток: %d'
            % ctx.orphan_work_days(),
            'Маълумотномада машинаси йўқ объектларнинг GPS бўйича иши, '
            'объект-кунлар: %d' % ctx.orphan_work_days())),
    ]
    if preview is not None:
        notes.insert(0, bi(('ПРЕДПРОСМОТР: N = %d не утверждён владельцем'
                            % preview,
                            'ОЛДИНДАН КЎРИШ: N = %d эгаси томонидан '
                            'тасдиқланмаган' % preview)))
    else:
        notes.append(bi(('N для заявок, заведённых задним числом: не утверждён',
                         'Орқа сана билан киритилган буюртмалар учун N: '
                         'тасдиқланмаган')))
    sheet.append(())
    for note in notes:
        sheet.append((note,))

    forward_rows = []
    for row in forward:
        app = row['app']
        window = row['window']
        forward_rows.append((
            app.number, app.created_day, app.completed_day,
            labels.pick(labels.STATUSES, app.status, True),
            app.row.get('company_name') or '', app.row.get('plate_number') or '',
            _machine(ctx, row['equipment_id']), _org(ctx, row['equipment_id']),
            app.row.get('work_type_name') or '',
            labels.pick(methods.LABELS, row['method'], True)
            if row['method'] else '',
            app.row.get('volume') or '',
            window[0] if window else None, window[1] if window else None,
            labels.pick(labels.VERDICTS, row['verdict'], True),
            labels.pick(labels.REASONS, row['reason'], True)
            if row['reason'] else '',
            ', '.join(d.strftime('%d.%m') for d in row['work_days']),
            row['gps_ha'] if row['work_days'] else None,
            row['unknown_days'] or None,
            'да' if app.row.get('gone_at') else ''))
    sheet_with(
        'Заявка-работа',
        (bi(('Номер заявки', 'Буюртма рақами')), bi(('Создана', 'Яратилган')),
         bi(('Выполнена', 'Бажарилган')), bi(('Статус', 'Ҳолат')),
         bi(('Предприятие', 'Корхона')), bi(('Госномер agro-work',
                                              'agro-work давлат рақами')),
         bi(('Наша машина', 'Бизнинг машина')), bi(('Организация', 'Ташкилот')),
         bi(('Вид работы', 'Иш тури')), bi(('Метод', 'Усул')),
         bi(('Объём', 'Ҳажм')), bi(('Окно с', 'Ойна дан')),
         bi(('Окно по', 'Ойна гача')), bi(('Вердикт', 'Ҳукм')),
         bi(('Причина', 'Сабаб')), bi(('Сутки с работой', 'Иш бўлган кунлар')),
         bi(('Га по GPS в окне', 'Ойнада GPS бўйича га')),
         bi(('Неизвестных суток', 'Номаълум кунлар')),
         bi(('Нет в последней выгрузке', 'Охирги юкламада йўқ'))),
        forward_rows,
        (24, 12, 12, 14, 26, 16, 24, 20, 36, 14, 10, 12, 12, 22, 44, 22, 12, 12, 12))

    reverse_rows = []
    for row in reverse:
        reverse_rows.append((
            row['day'], _machine(ctx, row['equipment_id']),
            _org(ctx, row['equipment_id']), rc_category(
                ctx.group_of(row['equipment_id'])[1]),
            row['gps_ha'], labels.pick(labels.COVERAGE, row['coverage'], True),
            labels.pick(labels.REASONS, row['reason'], True)
            if row['reason'] else '',
            ', '.join(app.number for app in row['apps'][:5])))
    sheet_with(
        'Работа-заявка',
        (bi(('Сутки', 'Кун')), bi(('Машина', 'Машина')),
         bi(('Организация', 'Ташкилот')), bi(('Категория', 'Тоифа')),
         bi(('Га по GPS', 'GPS бўйича га')), bi(('Покрытие', 'Қопланиш')),
         bi(('Причина', 'Сабаб')), bi(('Заявки', 'Буюртмалар'))),
        reverse_rows, (12, 26, 20, 26, 10, 22, 48, 40))

    reason_rows = []
    counts = {}
    for row in forward:
        if row['reason']:
            counts[('forward', row['reason'])] = counts.get(
                ('forward', row['reason']), 0) + 1
    for row in reverse:
        if row['reason']:
            counts[('reverse', row['reason'])] = counts.get(
                ('reverse', row['reason']), 0) + 1
    for (side, reason), count in sorted(counts.items(), key=lambda kv: -kv[1]):
        reason_rows.append((bi(('Заявка → работа', 'Буюртма → иш'))
                            if side == 'forward'
                            else bi(('Работа → заявка', 'Иш → буюртма')),
                            bi(labels.REASONS[reason]), reason, count))
    sheet_with('Причины',
               (bi(('Сторона', 'Томон')), bi(('Причина', 'Сабаб')), 'code',
                bi(('Строк', 'Қаторлар'))),
               reason_rows, (24, 70, 30, 10))

    lag_sheet = sheet_with(
        'Замер N',
        (bi(('Выборка', 'Танлама')), bi(('Заявок', 'Буюртмалар')),
         bi(('Медиана, сут', 'Медиана, кун')), 'p75', 'p90', 'p95',
         bi(('Максимум', 'Энг кўпи'))),
        [(bi(SAMPLE_TITLES[name]), stats['n'], stats['median'], stats['p75'],
          stats['p90'], stats['p95'], stats['max'])
         for name, stats in lags.items()],
        (48, 10, 12, 8, 8, 8, 12))
    lag_sheet.append(())
    lag_sheet.append((bi(('Доля заявок, работа которых попала бы в окно «день '
                          'ввода и N суток до него»',
                          'Иши «киритилган кун ва ундан олдинги N кун» '
                          'ойнасига тушадиган буюртмалар улуши')),))
    lag_sheet.append(('N',) + tuple(bi(SAMPLE_TITLES[name]) for name in lags))
    for n in range(0, 15):
        lag_sheet.append((n,) + tuple(
            round(stats['covered_by_n'][n], 3)
            if stats['covered_by_n'][n] is not None else None
            for stats in lags.values()))
    lag_sheet.append(())
    lag_sheet.append((bi(('Не вошли в замер', 'Ўлчовга кирмаганлар')),))
    for reason, count in sorted(excluded.items()):
        title = EXCLUDED_TITLES.get(reason) or labels.REASONS.get(reason)
        lag_sheet.append((bi(title) if title else reason, count))
    return book


def rc_category(slug):
    return labels.CATEGORIES.get(slug, slug or '')


def print_lags(lags, excluded, log=print):
    log('')
    log('completion lag, days from the last GPS work day to "Vypolneno":')
    for name, stats in lags.items():
        log('  %-13s n=%-5d median=%s p75=%s p90=%s p95=%s max=%s'
            % (name, stats['n'], stats['median'], stats['p75'], stats['p90'],
               stats['p95'], stats['max']))
    all_stats = lags.get('all') or {}
    if all_stats.get('n'):
        cover = all_stats['covered_by_n']
        log('  share of ordinary applications a window of N days would cover:')
        log('  ' + ' | '.join('N=%d %.0f%%' % (n, 100 * cover[n])
                              for n in range(0, 8)))
    for reason, count in sorted(excluded.items()):
        log('  not measured, %s: %d' % (reason, count))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument('--db', default=config.DB_PATH)
    parser.add_argument('--from', dest='date_from', default=None)
    parser.add_argument('--to', dest='date_to', default=None)
    parser.add_argument('--out', default=None,
                        help='xlsx path (default agro_work_reconcile_<to>.xlsx)')
    parser.add_argument('--lookback-preview', type=int, default=None,
                        metavar='N')
    args = parser.parse_args(argv)
    today = date.today()
    try:
        date_to = (datetime.strptime(args.date_to, '%Y-%m-%d').date()
                   if args.date_to else today - timedelta(days=1))
        date_from = (datetime.strptime(args.date_from, '%Y-%m-%d').date()
                     if args.date_from else date_to - timedelta(days=13))
    except ValueError:
        sys.stderr.write('ERROR: dates must look like YYYY-MM-DD\n')
        return 2
    if date_from > date_to:
        sys.stderr.write('ERROR: --from is after --to\n')
        return 2
    if args.lookback_preview is not None and not 0 <= args.lookback_preview <= 60:
        sys.stderr.write('ERROR: --lookback-preview must be 0..60\n')
        return 2
    if not os.path.exists(args.db):
        sys.stderr.write('ERROR: database not found at %s\n' % args.db)
        return 2
    con = _open_ro(args.db)
    try:
        missing = store.missing_tables(con)
        if missing:
            sys.stderr.write('ERROR: tables missing: %s - run '
                             'migrate_agro_work_001.py first\n'
                             % ', '.join(missing))
            return 2
        kwargs = {}
        if args.lookback_preview is not None:
            kwargs['lookback'] = args.lookback_preview
        ctx = rc.Reconciliation(con, date_from, date_to, today=today, **kwargs)
        forward = ctx.forward_rows()
        reverse = ctx.reverse_rows()
        lags, excluded = rc.measure_lags(con, today=today)
    finally:
        con.close()
    groups, total = ctx.summary(forward, reverse)
    print('period: %s .. %s%s' % (date_from, date_to,
                                  ' (PREVIEW N=%d)' % args.lookback_preview
                                  if args.lookback_preview is not None else ''))
    print('applications: %d | work confirmed %d | NO WORK %d | no verdict %d'
          % (total['applications'], total['app_' + rc.V_WORK],
             total['app_' + rc.V_NO_WORK], total['app_' + rc.V_NONE]))
    print('machine-days with GPS work: %d | covered %d | WITHOUT APPLICATION %d '
          '| no verdict %d' % (total['work_days'], total['day_' + rc.C_COVERED],
                               total['day_' + rc.C_UNCOVERED],
                               total['day_' + rc.C_NONE]))
    print('object-days with work and no machine in our registry: %d'
          % ctx.orphan_work_days())
    print_lags(lags, excluded)
    out = args.out or 'agro_work_reconcile_%s.xlsx' % date_to.isoformat()
    book = build_book(ctx, forward, reverse, lags, excluded,
                      args.lookback_preview)
    book.save(out)
    print('\nwritten %s' % config.ascii_only(out))
    return 0


if __name__ == '__main__':
    sys.exit(main())
